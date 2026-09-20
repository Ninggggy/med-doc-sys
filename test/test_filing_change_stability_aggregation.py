import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT_DIR = Path(__file__).resolve().parents[1]
PACKAGE_PARENT = ROOT_DIR.parent
if str(PACKAGE_PARENT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_PARENT))

from agent.agent_backend.services.filing_change_stability_service import FilingChangeStabilityService
from agent.agent_backend.services.filing_change_submission_parser import parse_filing_change_docx


def stability_record(
    batch_no,
    month,
    result,
    *,
    orientation="正置",
    role="自制制剂",
    indicator="杂质A",
    standard="≤0.10",
):
    time_point = f"{month}月{orientation}" if orientation else f"{month}月"
    return {
        "product_role": role,
        "batch_no": batch_no,
        "time_point": time_point,
        "month": month,
        "orientation": orientation,
        "indicator": indicator,
        "standard_text": standard,
        "result_text": str(result),
    }


class FilingChangeStabilityAggregationTests(unittest.TestCase):
    def setUp(self):
        self.service = FilingChangeStabilityService()

    @staticmethod
    def run_service(submissions):
        return FilingChangeStabilityService().run(
            {
                "form_json": {
                    "item_15_validity_period": {
                        "sub_fields": {"proposed_validity_period": "24个月"}
                    }
                }
            },
            submissions,
            "aggregation_test",
        )

    def test_extrema_use_all_self_batches_and_exclude_reference(self):
        records = [
            stability_record("B1", 0, "0.010", orientation=""),
            stability_record("B1", 24, "0.020", orientation="正置"),
            stability_record("B1", 24, "0.018", orientation="倒置"),
            stability_record("B2", 0, "0.008", orientation=""),
            stability_record("B2", 24, "0.025", orientation="正置"),
            stability_record("B3", 0, "0.009", orientation=""),
            stability_record("B3", 24, "0.025", orientation="倒置"),
            stability_record("REF1", 24, "0.099", role="参比制剂"),
        ]
        result = self.run_service(
            [
                {
                    "doc_id": "direct_doc",
                    "file_name": "稳定性数据.docx",
                    "quality_tags": ["稳定性研究资料"],
                    "extracted_json": {"stability_records": records},
                }
            ]
        )
        indicator = result["key_indicators"][0]
        self.assertEqual(result["data_scope"]["self_batches"], ["B1", "B2", "B3"])
        self.assertEqual(result["data_scope"]["reference_batches"], ["REF1"])
        self.assertEqual(indicator["min_value"], 0.008)
        self.assertIn("B2批0月", indicator["min_display"])
        self.assertEqual(indicator["max_value"], 0.025)
        self.assertIn("B2批24月正置", indicator["max_display"])
        self.assertIn("B3批24月倒置", indicator["max_display"])
        self.assertNotIn("REF1", indicator["max_display"])
        self.assertEqual(indicator["latest_time_point"], "全部自制批次汇总")

    def test_each_document_prefers_structured_records_then_falls_back_to_table(self):
        direct = stability_record("B1", 24, "0.020")
        ignored_markdown_table = {
            "headers": ["时间", "批号", "指标", "结果", "限度"],
            "rows": [["24月", "SHOULD_NOT_EXIST", "杂质A", "0.099", "≤0.10"]],
        }
        fallback_table = {
            "headers": ["时间", "批号", "指标", "结果", "限度"],
            "rows": [["24月倒置", "B2", "杂质A", "0.030", "≤0.10"]],
        }
        result = self.run_service(
            [
                {
                    "doc_id": "doc1",
                    "file_name": "稳定性直接记录.docx",
                    "extracted_json": {
                        "stability_records": [direct],
                        "markdown_tables": [ignored_markdown_table],
                    },
                },
                {
                    "doc_id": "doc2",
                    "file_name": "稳定性旧版表格.docx",
                    "extracted_json": {"markdown_tables": [fallback_table]},
                },
            ]
        )
        batches = {row["batch_no"] for row in result["records"]}
        source_types = {row["data_source_type"] for row in result["records"]}
        self.assertEqual(batches, {"B1", "B2"})
        self.assertNotIn("SHOULD_NOT_EXIST", batches)
        self.assertEqual(source_types, {"structured_stability_table", "markdown_table_fallback"})
        self.assertFalse(result["data_scope"]["narrative_text_used_for_numeric_analysis"])

    def test_censored_and_composite_text_are_not_exact_numeric_trend_points(self):
        self.assertIsNone(self.service._extract_numeric("<10cfu/ml"))
        self.assertIsNone(self.service._extract_numeric("0.1/0.2"))
        self.assertIsNone(self.service._extract_numeric("符合规定"))
        self.assertEqual(self.service._extract_numeric("40cfu/ml"), 40.0)
        self.assertTrue(self.service._judge_within_standard("<10cfu/ml", "需氧菌总数≤100cfu/ml"))

    def test_limit_check_reports_exceeded_result_with_context(self):
        records = [
            stability_record("B1", 0, "0.05"),
            stability_record("B1", 24, "0.12"),
        ]
        result = self.run_service(
            [
                {
                    "doc_id": "oos_doc",
                    "file_name": "稳定性超限数据.docx",
                    "quality_tags": ["稳定性研究资料"],
                    "extracted_json": {"stability_records": records},
                }
            ]
        )

        self.assertEqual(result["limit_check"]["status"], "已超限")
        self.assertEqual(result["limit_check"]["status_code"], "out_of_spec")
        self.assertEqual(result["limit_check"]["severity"], "error")
        self.assertEqual(result["limit_check"]["out_of_spec_count"], 1)
        self.assertEqual(result["limit_check"]["out_of_spec_details"][0]["result_text"], "0.12")
        self.assertEqual(result["limit_check"]["out_of_spec_details"][0]["limit_text"], "≤0.10")
        self.assertEqual(
            result["limit_check"]["out_of_spec_details"][0]["reason_text"],
            "数值明确超出可接受范围",
        )
        indicator = result["key_indicators"][0]
        self.assertEqual(indicator["limit_status"], "已超限")
        self.assertEqual(indicator["risk_level"], "高")
        self.assertEqual(result["result"], "不支持延长")

    def test_limit_check_marks_uncomparable_result_for_manual_confirmation(self):
        records = [
            stability_record("B1", 24, "无异常", indicator="性状", standard="应符合规定"),
        ]
        result = self.run_service(
            [
                {
                    "doc_id": "pending_doc",
                    "file_name": "稳定性定性数据.docx",
                    "quality_tags": ["稳定性研究资料"],
                    "extracted_json": {"stability_records": records},
                }
            ]
        )

        self.assertEqual(result["limit_check"]["status"], "待确认")
        self.assertEqual(result["limit_check"]["status_code"], "manual_review")
        self.assertEqual(result["limit_check"]["undecidable_count"], 1)
        self.assertEqual(len(result["limit_check"]["undecidable_details"]), 1)
        self.assertIn("无法自动比较", result["limit_check"]["undecidable_details"][0]["reason_text"])
        self.assertEqual(result["key_indicators"][0]["limit_status"], "待确认")

    def test_common_pharmaceutical_limit_wording_is_compared(self):
        self.assertTrue(self.service._judge_within_standard("0.08", "不得过0.10"))
        self.assertFalse(self.service._judge_within_standard("0.12", "不得过0.10"))
        self.assertTrue(self.service._judge_within_standard("95.0", "不小于90.0"))
        self.assertFalse(self.service._judge_within_standard("85.0", "不小于90.0"))
        self.assertTrue(self.service._judge_within_standard("95.0", "90.0～110.0"))
        self.assertIsNone(self.service._judge_within_standard("101.1", "95.0%~105.0%"))
        self.assertTrue(self.service._judge_within_standard("101.1%", "95.0%~105.0%"))
        self.assertFalse(self.service._judge_within_standard("106.0", "95.0%~105.0%"))
        self.assertTrue(self.service._judge_within_standard("符合规定", "应符合规定"))
        self.assertFalse(self.service._judge_within_standard("检出", "不得检出"))

    def test_strict_scientific_censored_and_unit_rules(self):
        self.assertTrue(self.service._judge_within_standard("9.9e-2", "<0.10"))
        self.assertFalse(self.service._judge_within_standard("0.10", "<0.10"))
        self.assertTrue(self.service._judge_within_standard("90", "≥90"))
        self.assertTrue(self.service._judge_within_standard("<10cfu/ml", "≤100cfu/ml"))
        self.assertIsNone(self.service._judge_within_standard("0.10mg/g", "≤0.20%"))
        mismatch = self.service._evaluate_limit("0.10mg/g", "≤0.20%")
        self.assertEqual(mismatch["reason"], "unit_mismatch")
        self.assertIsNone(self.service._judge_within_standard("约0.10", "≤0.20"))

    def test_stability_table_header_is_detected_dynamically(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "dynamic-stability.docx"
            doc = Document()
            doc.add_paragraph("表1：稳定性数据")
            table = doc.add_table(rows=4, cols=5)
            values = [
                ["批号：B100", "规格：10mg", "考察条件：25℃", "", ""],
                ["说明", "下表为检测结果", "", "", ""],
                ["考察项目", "明细项目", "可接受标准", "0月", "6月正置"],
                ["有关物质", "杂质A", "≤0.10%", "0.05%", "0.12%"],
            ]
            for row_idx, row in enumerate(values):
                for col_idx, value in enumerate(row):
                    table.cell(row_idx, col_idx).text = value
            doc.save(path)

            chunks = parse_filing_change_docx(str(path))
            stability_tables = [
                table_item
                for chunk in chunks
                for table_item in (chunk.get("tables", []) or [])
                if table_item.get("table_type") == "stability_result"
            ]
            self.assertEqual(len(stability_tables), 1)
            structured = stability_tables[0]["structured_data"]
            self.assertEqual(structured["header_row_index"], 2)
            self.assertEqual(len(structured["records"]), 2)
            self.assertEqual(structured["records"][1]["time_point"], "6月正置")
            self.assertEqual(structured["records"][1]["result_text"], "0.12%")

    def test_real_handoff_word_document_aggregates_three_self_batches(self):
        path = ROOT_DIR / "task" / "change_review" / "T1延长有效期资料24月-发送IT(1).docx"
        if not path.exists():
            self.skipTest(f"真实交接材料不存在: {path}")
        chunks = parse_filing_change_docx(str(path))
        records = [
            record
            for chunk in chunks
            for table in (chunk.get("tables", []) or [])
            if table.get("table_type") == "stability_result"
            for record in (table.get("structured_data", {}).get("records", []) or [])
        ]
        result = self.run_service(
            [
                {
                    "doc_id": "real_handoff",
                    "file_name": path.name,
                    "quality_tags": ["稳定性研究资料"],
                    "extracted_json": {"stability_records": records},
                }
            ]
        )
        by_name = {item["indicator"]: item for item in result["key_indicators"]}
        impurity = by_name["杂质A"]
        self.assertEqual(result["data_scope"]["self_batches"], ["230501", "230502", "230503"])
        self.assertEqual(result["data_scope"]["reference_batches"], ["0196"])
        self.assertEqual(impurity["point_count"], 30)
        self.assertEqual(impurity["min_display"], "0.004（230503批6月倒置；230503批6月正置）")
        self.assertEqual(impurity["max_display"], "0.053（230502批24月倒置）")
        impurity_chart = next(item for item in result["charts"] if item["indicator"] == "杂质A")
        self.assertEqual(len(impurity_chart["series_groups"]), 6)
        self.assertIn('height="293"', impurity_chart["svg_content"])
        self.assertIn('y="246"', impurity_chart["svg_content"])
        self.assertIn("乙醇（mg/g）", by_name)
        self.assertIn("1,2-丙二醇（mg/g）", by_name)
        self.assertNotIn("乙醇和1,2-丙二醇（mg/g）", by_name)
        chart_names = {item["indicator"] for item in result["charts"]}
        self.assertNotIn("微生物限度-霉菌和酵母菌总数", chart_names)
        self.assertNotIn("微生物限度-控制菌", chart_names)


if __name__ == "__main__":
    unittest.main()
