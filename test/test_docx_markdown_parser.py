import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch
import sys

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT_DIR = Path(__file__).resolve().parents[1]
PACKAGE_PARENT = ROOT_DIR.parent
if str(PACKAGE_PARENT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_PARENT))

from agent.agent_backend.utils.parser.docx_markdown_parser import (
    _build_sections_from_markdown_text,
    parse_docx_to_markdown_json,
)
from agent.agent_backend.utils.parser.docx_parser import parse_docx


class DocxMarkdownParserTests(unittest.TestCase):
    @staticmethod
    def _add_ctd_auto_numbering(doc, paragraphs, template="3.2.P.2.%1") -> None:
        numbering = doc.part.numbering_part.element
        abstract_id = 91
        num_id = 91
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), template)
        level.extend([start, num_fmt, level_text])
        abstract.append(level)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        for paragraph in paragraphs:
            p_pr = paragraph._p.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            num_id_element = OxmlElement("w:numId")
            num_id_element.set(qn("w:val"), str(num_id))
            num_pr.extend([ilvl, num_id_element])
            p_pr.append(num_pr)

    @staticmethod
    def _flatten_sections(sections):
        flattened = []
        for section in sections or []:
            flattened.append(section)
            flattened.extend(DocxMarkdownParserTests._flatten_sections(section.get("children_sections", [])))
        return flattened

    @staticmethod
    def _case_dir(name: str) -> Path:
        return Path(tempfile.mkdtemp(prefix=f"docx-parser-{name}-"))

    def test_markitdown_path_builds_sections_and_preserves_tables(self) -> None:
        case_dir = self._case_dir("markitdown")
        docx_path = case_dir / "sample.docx"
        docx_path.write_bytes(b"placeholder")

        markdown_text = (
            "# Title\n\n"
            "Paragraph line one\n"
            "line two\n\n"
            "| A | B |\n"
            "| --- | --- |\n"
            "| 1 | 2 |\n\n"
            "## 3.2.s.1.1 Details\n\n"
            "Body text\n"
        )

        with patch(
            "agent.agent_backend.utils.parser.docx_markdown_parser._try_convert_docx_with_markitdown",
            return_value=(markdown_text, "Converted title", None),
        ):
            payload = parse_docx_to_markdown_json(
                input_path=docx_path,
                title=None,
                embed_images=False,
                skip_toc=True,
                filter_header_footer=True,
                merge_continuous_tables=True,
            )

        self.assertEqual(payload["parser"]["markdown_source"], "markitdown")
        self.assertEqual(payload["title"], "Converted title")
        self.assertEqual(len(payload["sections"]), 1)

        root_section = payload["sections"][0]
        self.assertEqual(root_section["section_title"], "Title")
        self.assertIn("Paragraph line one\nline two", root_section["content"])
        self.assertIn("| A | B |", root_section["content"])
        self.assertEqual(len(root_section["tables"]), 1)
        self.assertEqual(root_section["children_sections"][0]["section_title"], "3.2.s.1.1 Details")

    def test_markitdown_failure_falls_back_to_python_docx(self) -> None:
        case_dir = self._case_dir("fallback")
        docx_path = case_dir / "fallback.docx"
        doc = Document()
        doc.add_heading("3.2.s.1 Basic info", level=1)
        doc.add_paragraph("Fallback paragraph.")
        doc.save(docx_path)

        with patch(
            "agent.agent_backend.utils.parser.docx_markdown_parser._try_convert_docx_with_markitdown",
            return_value=(None, None, "boom"),
        ):
            payload = parse_docx_to_markdown_json(
                input_path=docx_path,
                title=None,
                embed_images=False,
                skip_toc=True,
                filter_header_footer=True,
                merge_continuous_tables=True,
            )

        self.assertEqual(payload["parser"]["markdown_source"], "word_ooxml_native")
        self.assertIn("boom", payload["parser"]["markitdown_error"])
        self.assertEqual(len(payload["sections"]), 1)
        section = payload["sections"][0]
        self.assertIn("3.2.s.1 Basic info", section["section_title"])
        self.assertIn("Fallback paragraph.", section["content"])

    def test_markdown_section_builder_keeps_paragraph_line_breaks(self) -> None:
        sections, _ = _build_sections_from_markdown_text(
            "# Heading\n\nLine one\nline two\n",
            title=None,
            skip_toc=True,
            filter_header_footer=True,
            merge_continuous_tables=True,
        )

        self.assertEqual(len(sections), 1)
        self.assertIn("Line one\nline two", sections[0]["content"])

    def test_docx_parser_keeps_parent_section_content_when_child_exists(self) -> None:
        case_dir = self._case_dir("hierarchy")
        docx_path = case_dir / "hierarchy.docx"
        doc = Document()
        doc.add_heading("一、总述", level=1)
        doc.add_paragraph("这是一级标题自己的正文，不应被丢失。")
        doc.add_heading("（一）子章节", level=2)
        doc.add_paragraph("这是子章节内容。")
        doc.save(docx_path)

        chunks = parse_docx(str(docx_path))

        joined = "\n".join(str(item.get("text", "")) for item in chunks)
        self.assertIn("这是一级标题自己的正文，不应被丢失。", joined)
        self.assertIn("这是子章节内容。", joined)

    def test_native_ooxml_reconstructs_automatic_ctd_numbering(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            docx_path = Path(temp_dir) / "automatic-numbering.docx"
            doc = Document()
            first = doc.add_paragraph("处方组成", style="Heading 4")
            doc.add_paragraph("处方组成正文。")
            second = doc.add_paragraph("制剂", style="Heading 4")
            doc.add_paragraph("制剂正文。")
            self._add_ctd_auto_numbering(doc, [first, second])
            doc.save(docx_path)

            with patch(
                "agent.agent_backend.utils.parser.docx_markdown_parser._try_convert_docx_with_markitdown",
                return_value=("# 3.2.P.2 \u4ea7品开发\n\n所有内容被合并", None, None),
            ):
                payload = parse_docx_to_markdown_json(docx_path)

            self.assertEqual(payload["parser"]["markdown_source"], "word_ooxml_native")
            flattened = self._flatten_sections(payload["sections"])
            by_number = {
                (item.get("heading_evidence", {}) or {}).get("number"): item
                for item in flattened
            }
            self.assertIn("3.2.p.2.1", by_number)
            self.assertIn("处方组成正文", by_number["3.2.p.2.1"]["content"])
            self.assertIn("3.2.p.2.2", by_number)
            self.assertIn("制剂正文", by_number["3.2.p.2.2"]["content"])

    def test_fullwidth_no_space_table_and_textbox_headings_are_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            docx_path = Path(temp_dir) / "mixed-layout.docx"
            doc = Document()
            doc.add_paragraph("3．2．P．2．1处方组成", style="Heading 4")
            doc.add_paragraph("全角编号正文。")
            table = doc.add_table(rows=4, cols=1)
            table.cell(0, 0).text = "3.2.P.2.2 制剂"
            table.cell(1, 0).text = "表格内的制剂正文。"
            table.cell(2, 0).text = "3.2.P.2.3生产工艺的开发"
            table.cell(3, 0).text = "表格内的工艺正文。"
            textbox_host = doc.add_paragraph()
            textbox_host._p.append(
                parse_xml(
                    f'<w:txbxContent {nsdecls("w")}><w:p><w:r><w:t>3.2.P.2.4 包装系统</w:t></w:r></w:p></w:txbxContent>'
                )
            )
            doc.add_paragraph("文本框标题后的正文。")
            doc.save(docx_path)

            with patch(
                "agent.agent_backend.utils.parser.docx_markdown_parser._try_convert_docx_with_markitdown",
                return_value=(None, None, "disabled"),
            ):
                payload = parse_docx_to_markdown_json(docx_path)

            flattened = self._flatten_sections(payload["sections"])
            joined_titles = "\n".join(item.get("section_title", "") for item in flattened)
            joined_content = "\n".join(item.get("content", "") for item in flattened)
            for section_id in ["3.2.p.2.1", "3.2.p.2.2", "3.2.p.2.3", "3.2.p.2.4"]:
                self.assertIn(section_id, joined_titles.lower())
            self.assertIn("表格内的制剂正文", joined_content)
            self.assertIn("表格内的工艺正文", joined_content)
            self.assertIn("文本框标题后的正文", joined_content)
            self.assertGreaterEqual(payload["parse_quality"]["detected_ctd_heading_count"], 4)


if __name__ == "__main__":
    unittest.main()
