"""真实服务与隔离SQLite验证；模型执行被监测，未就绪时必须零调用零任务。"""
import unittest
from unittest.mock import patch, Mock
from datetime import datetime
import json
from flask import Flask
from test_filing_change_review_service_regression import (
    FilingChangeReviewServiceRegressionTest as Fixture,
    FilingChangeReviewRun, FilingChangeProject, FilingChangeReviewAppService,
    filing_change_review_controller as controller,
)


class ReadinessIntegrationTests(unittest.TestCase):
    setUp = Fixture.setUp
    tearDown = Fixture.tearDown
    _add_project = Fixture._add_project
    _add_submission = Fixture._add_submission

    def test_catalog_completeness_uses_effective_application_fields(self):
        from test_filing_change_review_service_regression import FilingChangeApplicationForm
        self._add_project()
        original = {'item_6_generic_name': {'value': ''}, 'item_5_application_matter_category': {'selected_values': []}}
        with self.connection.get_session() as session:
            session.add(FilingChangeApplicationForm(project_id='project-main', original_file_id='synthetic-source',
                form_json=json.dumps(original), raw_text='原始识别', parse_status='partial', confidence=.5,
                created_at=datetime.now(), updated_at=datetime.now()))
            session.commit()
        effective = {'item_6_generic_name': {'value': '合成甲片'}, 'item_5_application_matter_category': {'selected_values': ['延长有效期']}}
        with patch.object(self.service, '_effective_form_revision', return_value=(effective, '已修订', {})) as read:
            result = self.service.check_submission_completeness('project-main')
        self.assertIn('application_info', result['uploaded_codes'])
        read.assert_called_once()
        with self.connection.get_session() as session:
            row = session.query(FilingChangeApplicationForm).filter_by(project_id='project-main').one()
            self.assertEqual(json.loads(row.form_json), original)
        with patch.object(self.service, '_effective_form_revision', side_effect=ValueError('invalid revision')):
            invalid = self.service.check_submission_completeness('project-main')
        self.assertNotIn('application_info', invalid['uploaded_codes'])
        self.assertFalse(invalid['pass'])

    def test_real_application_all_pages_joint_revision_reaches_report(self):
        import os
        from pathlib import Path
        from copy import deepcopy
        keys = ['INTEGRITY_PAGE1_TEXT', 'INTEGRITY_PAGE2_REVIEW', 'INTEGRITY_PAGE3_REVIEW', 'INTEGRITY_PAGE4_TEXT']
        if not all(os.getenv(key) for key in keys):
            self.skipTest('需四页独立原图核对样本；私有内容仅本地，不输出原值')
        self.test_recorded_real_application_pages_survive_refresh_and_block_review()
        samples = [Path(os.environ[key]).read_text(encoding='utf8').strip() for key in keys]
        samples[1:3] = [json.loads(value) for value in samples[1:3]]
        doc_id = self.service.get_application_form('project-main')[2]['original_file_id']
        _, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        original = deepcopy(context['original_chunks'])
        items = []
        for index, sample in enumerate(samples):
            issues = [i for i in context['issues'] if i.get('chunk_index') == index]
            primary = next(i for i in issues if i['code'] == 'ocr_coverage')
            item = {'issue_key': primary['issue_key'], 'reason': '逐页独立原图核对，完整保留正文与表格',
                    'numeric_text_verified': True,
                    'related_issues': [{'issue_key': i['issue_key'], 'reason': '逐项核对本页完整区域'}
                        for i in issues if i['issue_key'] != primary['issue_key'] and i['code'] in ('ocr_quality', 'ocr_coverage')]}
            if isinstance(sample, dict):
                rows = sample['table_rows']
                item.update(action='correct_table', outside_text=sample['outside_text'], outside_text_verified=True,
                            table={'row_count': len(rows), 'column_count': len(rows[0]), 'cells': [
                                {'row': r, 'column': c, 'text': value} for r, row in enumerate(rows) for c, value in enumerate(row)]})
                if index == 2:
                    item['continuation_item_no'] = 23
            else:
                item.update(action='correct_text', text=sample)
            if index == 3:
                fields = sorted([i for i in issues if i['code'] == 'field_region_crossing'], key=lambda i: i['bbox_pdf'][1])
                excerpts = ['电子信箱：    手机：', '《药品生产许可证》编号：浙', 'GMP证书 编号（如有）：']
                self.assertEqual(len(fields), len(excerpts))
                item['related_issues'].extend({'issue_key': i['issue_key'], 'reason': '对照原页生产企业对应完整字段',
                    'target_item_no': 31, 'reviewed_text': excerpt, 'field_text_verified': True} for i, excerpt in zip(fields, excerpts))
            items.append(item)
            ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, {
                'source_identity': context['source_identity'], 'expected_revision': context['revision'], 'items': items})
            self.assertTrue(ok, message)
            _, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
            self.assertTrue(context['original_chunks'] == original, '原始四页OCR改变')
            self.assertTrue(all(i['resolved'] for i in context['issues'] if i.get('chunk_index', 99) <= index), '已核对页面重新出现阻塞')
            self.assertTrue(all(not i['resolved'] for i in context['issues'] if i.get('chunk_index', -1) > index), '未核对页面被放行')
        for index, sample in enumerate(samples):
            chunk = context['effective_chunks'][index]
            expected_text = sample['outside_text'] if isinstance(sample, dict) else sample
            self.assertTrue(expected_text in chunk['raw_text'], '整份核对丢失正文')
            if isinstance(sample, dict):
                self.assertTrue(chunk['tables'][0]['rows'] == sample['table_rows'], '整份核对改变表格')
        def check_form(form):
            self.assertTrue(form['parse_resolution']['manually_reviewed'])
            self.assertEqual(form['parse_resolution']['original_status'], 'partial', '不得伪造自动解析成功')
            for sample in samples[1:3]:
                for key, expected in sample['expected_fields'].items():
                    self.assertTrue(form['form_json'][key]['value'] == expected, '联合修订字段不一致：' + key)
            for key, expected in samples[2]['expected_applicant'].items():
                self.assertTrue(form['form_json']['item_30_applicant_info']['sub_fields'].get(key) == expected, '主体子字段不一致：' + key)
            self.assertTrue(samples[2]['table_rows'][1][0] in form['form_json']['item_29_other_related_info']['value'], '历史表缺失')
        for _ in range(2):
            check_form(self.service.get_application_form('project-main')[2])
        readiness = self.service.get_parse_readiness('project-main')
        self.assertTrue(all(i['code'] == 'required_material_missing' for i in readiness['blocking_issues']),
                        '四页核对后仍有非资料缺项阻塞')
        # 申请表核对完成不等于资料齐全。补入明确标注的合成必需资料，
        # 不绕过目录完整性检查，不声称这些替身具备真实审评依据。
        for category in ('1', '2', '4', '5'):
            doc = 'joint-required-' + category
            self._add_submission('project-main', doc, created_at=datetime.now(), category=category, parse_status='success')
            path = self.service._project_path('project-main') / 'parsed' / (doc + '.json')
            path.write_text(json.dumps([{'page': 1, 'status': 'success', 'text': 'Synthetic required material ' + category,
                                         'errors': [], 'tables': []}]))
        joint_readiness = self.service.get_parse_readiness('project-main')
        diagnostic = {
            'blocking_codes': [i['code'] for i in joint_readiness['blocking_issues']],
            'blocking_kinds': [i.get('source_kind') for i in joint_readiness['blocking_issues']],
            'matter_selection_count': len(self.service.get_application_form('project-main')[2]['form_json']
                                          ['item_5_application_matter_category']['selected_values'])}
        # 原页是管理类别说明加未勾选的文字事项，不自动把候选清单当选择。
        # 通过既有人工字段保存明确选择原页1.7项，再检验就绪与执行链路。
        self.assertFalse(joint_readiness['ready'])
        self.assertEqual(diagnostic['matter_selection_count'], 0, diagnostic)
        self.assertTrue(all(code == 'required_material_missing' for code in diagnostic['blocking_codes']), diagnostic)
        current_form = self.service.get_application_form('project-main')[2]['form_json']
        filled = deepcopy(current_form)
        filled['item_5_application_matter_category']['selected_values'] = ['1.7']
        ok, message, _ = self.service.save_application_form('project-main', {
            'form_json': filled, 'edited_paths': {'item_5_application_matter_category': ['selected_values']},
            'expected_values': current_form})
        self.assertTrue(ok, message)
        self.assertTrue(self.service.get_parse_readiness('project-main')['ready'])
        for _ in range(2):
            refreshed = self.service.get_application_form('project-main')[2]
            check_form(refreshed)
            self.assertEqual(refreshed['form_json']['item_5_application_matter_category']['selected_values'], ['1.7'])
        _, _, reread = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertTrue(reread['original_chunks'] == original, '人工选项确认改变原始OCR')
        # 仅验证执行输入、持久化与报告同源；模型为确定性检查器，不代表真实审评质量。
        from test_filing_change_review_service_regression import _RuleList, FilingChangeReviewResult
        from agent.agent_backend.services.filing_change_report_service import FilingChangeReportService
        from docx import Document
        self.service.rule_service = _RuleList()
        self.service.report_service = FilingChangeReportService(self.service.root_dir)
        marker = '四页人工核对输入一致性验证完成'
        def inspect_input(**inputs):
            check_form(inputs['application_form'])
            self.assertEqual(inputs['application_form']['form_json']['item_5_application_matter_category']['selected_values'], ['1.7'])
            return {'overall_conclusion': {'result': marker}}
        self.service.orchestrator = Mock()
        self.service.orchestrator.run.side_effect = inspect_input
        import time
        from test_filing_change_review_service_regression import RuntimeTaskStore
        app = Flask(__name__)
        app.register_blueprint(controller.filing_change_review_bp)
        app_service = object.__new__(FilingChangeReviewAppService)
        app_service.service = self.service
        store = RuntimeTaskStore(connection=self.connection, ensure_schema=False)
        with patch.object(controller, 'get_service', return_value=app_service), patch.object(
                controller, '_prune_tasks'), patch.object(controller, 'get_runtime_task_store', return_value=store):
            response = app.test_client().post('/filing-change-review/projects/project-main/review/start')
            self.assertEqual(response.status_code, 200)
            ids = response.get_json(force=True)['data']
            # QEMU内四页完整重放较慢；仅为测试等待上限，不修改业务/模型超时。
            deadline = time.monotonic() + 300
            task = store.get_task(ids['task_id'])
            while task['status'] in ('pending', 'running') and time.monotonic() < deadline:
                time.sleep(.1)
                task = store.get_task(ids['task_id'])
            self.assertEqual(task['status'], 'completed', '四页公开入口后台任务未完成')
        run_id = ids['run_id']
        result = task['result']
        self.service.orchestrator.run.assert_called_once()
        self.assertIn(marker, result['review_report_markdown'])
        with self.connection.get_session() as session:
            stored = session.query(FilingChangeReviewResult).filter_by(run_id=run_id).one()
            self.assertEqual(json.loads(stored.conclusion_json)['overall_conclusion']['result'], marker)
        report = Document(self.service.root_dir / 'projects' / 'project-main' / 'reports' / (run_id + '.docx'))
        self.assertIn(marker, '\n'.join(p.text for p in report.paragraphs))

    def test_real_application_fourth_page_complete_manual_revision(self):
        import os
        from pathlib import Path
        from copy import deepcopy
        sample_path = os.getenv('INTEGRITY_PAGE4_TEXT')
        if not sample_path:
            self.skipTest('需独立原图核对的第四页全文，私有样本仅本地提供')
        self.test_recorded_real_application_pages_survive_refresh_and_block_review()
        text = Path(sample_path).read_text(encoding='utf8').strip()
        doc_id = self.service.get_application_form('project-main')[2]['original_file_id']
        _, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        original = deepcopy(context['original_chunks'])
        issues = [i for i in context['issues'] if i.get('chunk_index') == 3]
        primary = next(i for i in issues if i['code'] == 'ocr_coverage')
        item = {'issue_key': primary['issue_key'], 'action': 'correct_text', 'text': text,
                'reason': '独立原图逐字核对第四页主体、签名栏、审查机关和打印说明',
                'numeric_text_verified': True,
                'related_issues': [{'issue_key': i['issue_key'], 'reason': '完整核对本页正文并保留原空白字段'}
                    for i in issues if i['issue_key'] != primary['issue_key'] and i['code'] in ('ocr_quality', 'ocr_coverage')]}
        # 独立原图：这三个区域依次是生产企业的邮箱/手机、生产许可和GMP栏。
        # 不把所有字段诊断自动标为解决；逐项提交完整对应文字及明确归属。
        field_issues = sorted([i for i in issues if i['code'] == 'field_region_crossing'], key=lambda i: i['bbox_pdf'][1])
        excerpts = ['电子信箱：    手机：', '《药品生产许可证》编号：浙', 'GMP证书 编号（如有）：']
        self.assertEqual(len(field_issues), len(excerpts))
        item['related_issues'].extend({'issue_key': issue['issue_key'], 'reason': '按独立原图核对生产企业区域并对应完整修订文字',
                                      'target_item_no': 31, 'reviewed_text': excerpt, 'field_text_verified': True}
                                     for issue, excerpt in zip(field_issues, excerpts))
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, {
            'source_identity': context['source_identity'], 'expected_revision': context['revision'], 'items': [item]})
        self.assertTrue(ok, message)
        _, _, after = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertTrue(after['original_chunks'] == original, '原始OCR证据改变')
        self.assertTrue(after['effective_chunks'][3]['raw_text'] == text, '第四页完整修订正文改变')
        for index in (0, 1, 2):
            self.assertTrue(after['effective_chunks'][index] == original[index], '第四页修订影响其他页')
        for _ in range(2):
            form = self.service.get_application_form('project-main')[2]
            self.assertFalse(form['parse_resolution']['manually_reviewed'])
        self.assertFalse(self.service.get_parse_readiness('project-main')['ready'])
        remaining = [i['code'] for i in after['issues'] if i.get('chunk_index') == 3 and not i['resolved']]
        wrong = deepcopy(item)
        wrong['related_issues'][-1]['target_item_no'] = 32
        accepted, _, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, {
            'source_identity': after['source_identity'], 'expected_revision': after['revision'], 'items': [wrong]})
        self.assertFalse(accepted, '错误归属不应放行')
        _, _, retained = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertTrue(retained['effective_chunks'] == after['effective_chunks'], '拒绝错误归属后修改了有效内容')
        # 当前死路的安全边界也须保留：小区域不能覆盖整页修订造成二次丢字。
        # 真正修复须提供显式关联字段核对，而不是放开重叠覆盖。
        if remaining:
            field_issue = next(i for i in after['issues'] if i.get('chunk_index') == 3 and i['code'] == 'field_region_crossing')
            rejected, _, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, {
                'source_identity': after['source_identity'], 'expected_revision': after['revision'],
                'items': [item, {'issue_key': field_issue['issue_key'], 'action': 'correct_text',
                                 'target_item_no': 31, 'text': '已核对生产企业字段',
                                 'reason': '验证原区域不能再次覆盖整页修订'}]})
            self.assertFalse(rejected, '重叠修订不应覆盖已有整页正文')
            _, _, retained = self.service.get_parse_review('project-main', 'application_form', doc_id)
            self.assertTrue(retained['effective_chunks'] == after['effective_chunks'], '失败保存改变此前修订')
        self.assertEqual(remaining, [], '已完整修订第四页仍有未完成核对；仅输出原因码')

    def test_real_application_third_page_complete_manual_revision(self):
        import os
        from pathlib import Path
        from copy import deepcopy
        sample_path = os.getenv('INTEGRITY_PAGE3_REVIEW')
        if not sample_path:
            self.skipTest('需独立原页核对的第三页全文及表格，包含私有信息的样本仅本地提供')
        self.test_recorded_real_application_pages_survive_refresh_and_block_review()
        sample = json.loads(Path(sample_path).read_text(encoding='utf8'))
        doc_id = self.service.get_application_form('project-main')[2]['original_file_id']
        _, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        original = deepcopy(context['original_chunks'])
        page_issues = [i for i in context['issues'] if i.get('chunk_index') == 2]
        primary = next(i for i in page_issues if i['code'] == 'ocr_coverage')
        rows = sample['table_rows']
        item = {'issue_key': primary['issue_key'], 'action': 'correct_table', 'continuation_item_no': 23,
                'reason': '独立原页核对第三页全部表外正文和历次申请表，并对照前页23项确认续文',
                'outside_text': sample['outside_text'], 'outside_text_verified': True, 'numeric_text_verified': True,
                'table': {'row_count': len(rows), 'column_count': len(rows[0]), 'cells': [
                    {'row': r, 'column': c, 'text': text} for r, row in enumerate(rows) for c, text in enumerate(row)]},
                'related_issues': [{'issue_key': i['issue_key'], 'reason': '整页逐字逐格核对，空白字段与签名栏亦保留'}
                    for i in page_issues if i['issue_key'] != primary['issue_key'] and i['code'] in ('ocr_quality', 'ocr_coverage')]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, {
            'source_identity': context['source_identity'], 'expected_revision': context['revision'], 'items': [item]})
        self.assertTrue(ok, message)
        _, _, after = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertTrue(after['original_chunks'] == original, '原始OCR证据改变')
        self.assertTrue(after['effective_chunks'][2]['tables'][0]['rows'] == rows, '已核对表格改变')
        self.assertTrue(sample['outside_text'] in after['effective_chunks'][2]['raw_text'], '表外正文未完整保留')
        for index in (0, 1, 3):
            self.assertTrue(after['effective_chunks'][index] == original[index], '修订波及其他页')
        self.assertTrue(all(i['resolved'] for i in after['issues'] if i.get('chunk_index') == 2))
        for _ in range(2):
            form = self.service.get_application_form('project-main')[2]
            for key, value in sample['expected_fields'].items():
                self.assertTrue(form['form_json'][key]['value'] == value, '已核对字段不一致：' + key)
            self.assertTrue(rows[1][0] in form['form_json']['item_29_other_related_info']['value'], '历次申请表未进入相应字段')
            mismatches = [key for key, value in sample.get('expected_applicant', {}).items()
                          if form['form_json']['item_30_applicant_info']['sub_fields'].get(key) != value]
            self.assertEqual(mismatches, [], '申请人字段未正确保留；只输出字段名，不输出私有值')
            self.assertFalse(form['parse_resolution']['manually_reviewed'])
        self.assertFalse(self.service.get_parse_readiness('project-main')['ready'])

    def test_real_application_second_page_full_mixed_revision(self):
        import os
        from pathlib import Path
        from copy import deepcopy
        sample_path = os.getenv('INTEGRITY_PAGE2_REVIEW')
        if not sample_path:
            self.skipTest('需独立原页核对后的第二页表格及表外全文')
        self.test_recorded_real_application_pages_survive_refresh_and_block_review()
        sample = json.loads(Path(sample_path).read_text(encoding='utf8'))
        doc_id = self.service.get_application_form('project-main')[2]['original_file_id']
        _, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        original = deepcopy(context['original_chunks'])
        page_issues = [i for i in context['issues'] if i.get('chunk_index') == 1]
        primary = next(i for i in page_issues if i['code'] == 'ocr_coverage')
        rows = sample['table_rows']
        payload = {'source_identity': context['source_identity'], 'expected_revision': context['revision'], 'items': [{
            'issue_key': primary['issue_key'], 'action': 'correct_table',
            'reason': '对照独立原页核对整页字段与原辅包来源表，不以OCR候选作为真值',
            'outside_text': sample['outside_text'], 'outside_text_verified': True, 'numeric_text_verified': True,
            'table': {'row_count': len(rows), 'column_count': len(rows[0]), 'cells': [
                {'row': r, 'column': c, 'text': text} for r, row in enumerate(rows) for c, text in enumerate(row)]},
            'related_issues': [{'issue_key': i['issue_key'], 'reason': '完整区域逐字逐格核对并保留空字段'}
                for i in page_issues if i['issue_key'] != primary['issue_key'] and i['code'] in ('ocr_quality', 'ocr_coverage')]}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, payload)
        self.assertTrue(ok, message)
        _, _, after = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertTrue(after['original_chunks'] == original, '原始页证据被改变')
        effective = after['effective_chunks'][1]
        self.assertEqual(effective['tables'][0]['rows'], rows)
        self.assertIn(sample['outside_text'], effective['raw_text'])
        self.assertTrue(all(i['resolved'] for i in after['issues'] if i.get('chunk_index') == 1))
        self.assertTrue(any(not i['resolved'] for i in after['issues'] if i.get('chunk_index') != 1))
        for index in (0, 2, 3):
            self.assertTrue(after['effective_chunks'][index] == original[index], '修订波及其他页')
        for _ in range(2):
            form = self.service.get_application_form('project-main')[2]
            for key, expected in sample['expected_fields'].items():
                self.assertTrue(form['form_json'][key]['value'] == expected, '人工核对字段不一致：' + key)
            material = form['form_json']['item_17_material_source']['table_rows']
            expected_material = [dict(zip(('material_name', 'register_no', 'accept_no', 'manufacturer'), row))
                                 for row in rows[1:]]
            self.assertTrue(material == expected_material, '原辅包来源表未正确进入申请字段')
            self.assertFalse(form['parse_resolution']['manually_reviewed'])
        self.assertFalse(self.service.get_parse_readiness('project-main')['ready'])

    def test_real_application_first_page_manual_revision_keeps_other_pages_blocked(self):
        import os
        from pathlib import Path
        from copy import deepcopy
        text_path = os.getenv('INTEGRITY_PAGE1_TEXT')
        if not text_path:
            self.skipTest('需独立原页核对后的第一页全文')
        self.test_recorded_real_application_pages_survive_refresh_and_block_review()
        form = self.service.get_application_form('project-main')[2]
        doc_id = form['original_file_id']
        _, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        original = deepcopy(context['original_chunks'])
        page_issues = [i for i in context['issues'] if i.get('chunk_index') == 0]
        primary = next(i for i in page_issues if i['code'] == 'ocr_coverage')
        related = [{'issue_key': i['issue_key'], 'reason': '独立渲染原页逐句核对，该区域正文已完整补录'}
                   for i in page_issues if i['issue_key'] != primary['issue_key'] and i['code'] in ('ocr_quality', 'ocr_coverage')]
        text = Path(text_path).read_text(encoding='utf8').strip()
        payload = {'source_identity': context['source_identity'], 'expected_revision': context['revision'], 'items': [{
            'issue_key': primary['issue_key'], 'action': 'correct_text', 'text': text,
            'reason': '对照独立Poppler整页及条码文字放大图核对公开声明页；不以OCR候选作为真值',
            'numeric_text_verified': True, 'related_issues': related}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, payload)
        self.assertTrue(ok, message)
        _, _, after = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertTrue(original == after['original_chunks'], '原始OCR证据改变')
        self.assertEqual(after['effective_chunks'][0]['text'], text)
        self.assertTrue(all(i['resolved'] for i in after['issues'] if i.get('chunk_index') == 0))
        self.assertTrue(any(not i['resolved'] for i in after['issues'] if i.get('chunk_index') in (1, 2, 3)))
        self.assertTrue(original[1:] == after['effective_chunks'][1:], '修订波及其他页面')
        for _ in range(2):
            refreshed = self.service.get_application_form('project-main')[2]
            self.assertFalse(refreshed['parse_resolution']['manually_reviewed'])
            self.assertIn('本申请一并提交的电子文件与打印文件内容完全一致', refreshed['raw_text'])
            for key, value in form['form_json'].items():
                if key.startswith('item_') and isinstance(value, dict) and 'value' in value:
                    self.assertTrue(value['value'] == refreshed['form_json'][key]['value'],
                                    '声明页修订改变了其他申请字段：' + key)
        self.assertFalse(self.service.get_parse_readiness('project-main')['ready'])

    def test_application_numeric_text_review_persists_and_revokes_with_correction(self):
        import io
        import fitz
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import build_pdf_form_from_pages
        self._add_project()
        with fitz.open() as doc:
            doc.new_page(width=300, height=200)
            upload = io.BytesIO(doc.tobytes())
        upload.filename = 'synthetic-numeric-text.pdf'
        word = {'text': '规格：8mg', 'bbox': [20, 40, 160, 60], 'source': 'native'}
        pages = [{'page': 1, 'page_bbox': [0, 0, 300, 200], 'status': 'partial', 'text': word['text'],
                  'raw_text': word['text'], 'words': [word], 'lines': [word], 'tables': [],
                  'errors': [{'code': 'ocr_quality', 'bbox_pdf': [18, 38, 162, 62]},
                             {'code': 'numeric_uncertain', 'bbox_pdf': [0, 0, 300, 200],
                              'numeric_verification': [{'status': 'uncertain', 'bbox_pdf': [100, 40, 160, 60]}]}]}]
        self.service.form_parser = Mock()
        self.service.form_parser.parse_form_file.return_value = FilingChangeFormParserService().form_from_pdf_result(
            build_pdf_form_from_pages(pages))
        ok, message, form = self.service.import_application_form('project-main', upload)
        self.assertTrue(ok, message)
        doc_id = form['original_file_id']
        _, _, initial = self.service.get_parse_review('project-main', 'application_form', doc_id)
        quality = next(i for i in initial['issues'] if i['code'] == 'ocr_quality')
        numeric = next(i for i in initial['issues'] if i['code'] == 'numeric_uncertain')
        item = {'issue_key': quality['issue_key'], 'action': 'correct_text', 'text': '规格：3mg',
                'reason': '原页数值为3mg，完整核对规格区域'}
        def save(context, items):
            ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, {
                'source_identity': context['source_identity'], 'expected_revision': context['revision'], 'items': items})
            self.assertTrue(ok, message)
            return self.service.get_parse_review('project-main', 'application_form', doc_id)[2]
        pending = save(initial, [item])
        self.assertFalse(next(i for i in pending['issues'] if i['issue_key'] == numeric['issue_key'])['resolved'])
        verified = save(pending, [{**item, 'numeric_text_verified': True}])
        self.assertTrue(next(i for i in verified['issues'] if i['issue_key'] == numeric['issue_key'])['resolved'])
        refreshed = self.service.get_application_form('project-main')[2]
        self.assertEqual(refreshed['form_json']['item_12_specification']['value'], '3mg')
        self.assertTrue(refreshed['parse_resolution']['manually_reviewed'])
        self.assertEqual(verified['original_chunks'], initial['original_chunks'])
        revoked = save(verified, [])
        self.assertFalse(next(i for i in revoked['issues'] if i['issue_key'] == numeric['issue_key'])['resolved'])
        self.assertFalse(self.service.get_application_form('project-main')[2]['parse_resolution']['manually_reviewed'])
        self.assertEqual(revoked['original_chunks'], initial['original_chunks'])
        direct = save(revoked, [{**item, 'issue_key': numeric['issue_key'], 'numeric_text_verified': True}])
        self.assertTrue(next(i for i in direct['issues'] if i['issue_key'] == numeric['issue_key'])['resolved'])
        self.assertFalse(next(i for i in direct['issues'] if i['issue_key'] == quality['issue_key'])['resolved'])
        direct_form = self.service.get_application_form('project-main')[2]
        self.assertEqual(direct_form['form_json']['item_12_specification']['value'], '3mg')
        self.assertFalse(direct_form['parse_resolution']['manually_reviewed'])
        self.assertEqual(direct['original_chunks'], initial['original_chunks'])

    def test_recorded_real_application_pages_survive_refresh_and_block_review(self):
        """重放真实OCR页证据；不声称重新执行OCR或全文正确。"""
        import os
        import io
        from pathlib import Path
        from copy import deepcopy
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import build_pdf_form_from_pages
        evidence_path = os.getenv('INTEGRITY_REPLAY_PAGES')
        if not evidence_path:
            self.skipTest('需指定已保存的真实OCR页证据')
        pages = json.loads(Path(evidence_path).read_text(encoding='utf8'))
        before = deepcopy(pages)
        self.assertEqual([p['page'] for p in pages], [1, 2, 3, 4])
        self.assertTrue(all(p['status'] == 'partial' for p in pages))
        self._add_project()
        source = Path(__file__).resolve().parents[1] / 'task' / 'change_review' / '申请表模板.pdf'
        content = source.read_bytes()
        upload = io.BytesIO(content)
        upload.filename = source.name
        self.service.form_parser = Mock()
        self.service.form_parser.parse_form_file.return_value = FilingChangeFormParserService().form_from_pdf_result(
            build_pdf_form_from_pages(deepcopy(pages)))
        ok, message, imported = self.service.import_application_form('project-main', upload)
        self.assertTrue(ok, message)
        self.assertEqual(imported['content_status'], 'partial')
        doc_id = imported['original_file_id']
        for _ in range(2):
            refreshed = self.service.get_application_form('project-main')[2]
            self.assertTrue(imported['raw_text'] == refreshed['raw_text'], '刷新改变了保存正文')
            self.assertTrue(imported['form_json'] == refreshed['form_json'], '刷新改变了申请字段')
        ok, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertTrue(ok)
        stored_pages = context['original_chunks']
        self.assertEqual(len(stored_pages), 4)
        for original, stored in zip(before, stored_pages):
            self.assertTrue(original['words'] == stored['words'], '保存改变了原始识别词')
            self.assertEqual(len(stored['errors']), len(original['errors']), '同一证据重建后增加了重复问题')
            for index, error in enumerate(original['errors']):
                matches = [i for i in context['issues'] if i.get('chunk_index') == original['page'] - 1
                           and i.get('error_index') == index and i['code'] == error['code']]
                self.assertEqual(len(matches), 1, '原始问题未唯一保留')
                self.assertFalse(matches[0]['resolved'])
        self.assertFalse(any('diagnostic_index' in i for i in context['issues']), '同一页证据出现额外汇总阻塞')
        self.assertFalse(self.service.get_parse_readiness('project-main')['ready'])
        app = Flask(__name__)
        app.register_blueprint(controller.filing_change_review_bp)
        app_service = object.__new__(FilingChangeReviewAppService)
        app_service.service = self.service
        with patch.object(controller, 'get_service', return_value=app_service), patch.object(
                controller, '_prune_tasks'), patch.object(controller, 'get_runtime_task_store') as store, patch.object(
                self.service, '_run_review_once') as review:
            response = app.test_client().post('/filing-change-review/projects/project-main/review/start')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()['data']['code'], 'parse_review_required')
        store.assert_not_called()
        review.assert_not_called()
        with self.connection.get_session() as session:
            self.assertEqual(session.query(FilingChangeReviewRun).count(), 0)
        stored_file = self.service._project_path('project-main') / 'application_form' / (doc_id + '.pdf')
        self.assertTrue(content == stored_file.read_bytes(), '原件被改变')
        self.assertTrue(before == pages, '重放改变了输入证据')

    def test_mixed_manual_multifield_save_refresh_and_failed_retry(self):
        import io
        import fitz
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import build_pdf_form_from_pages
        self._add_project()
        with fitz.open() as doc:
            doc.new_page(width=300, height=400)
            upload = io.BytesIO(doc.tobytes())
        upload.filename = 'synthetic-multifield.pdf'
        words = [{'text': '药品通用名称：旧名称', 'bbox': [10, 10, 180, 30], 'source': 'native'},
                 {'text': '规格：8mg', 'bbox': [10, 50, 180, 70], 'source': 'native'}]
        pages = [{'page': 1, 'page_bbox': [0, 0, 300, 400], 'status': 'partial',
                  'content_available': True, 'text': '旧正文', 'raw_text': '旧正文', 'words': words, 'lines': words,
                  'tables': [{'id': 't', 'page': 1, 'bbox_pdf': [10, 100, 200, 160], 'markdown': '|旧表|', 'cells': []}],
                  'errors': [{'code': 'ocr_quality', 'bbox_pdf': [0, 0, 210, 170]}]}]
        self.service.form_parser = Mock()
        self.service.form_parser.parse_form_file.return_value = FilingChangeFormParserService().form_from_pdf_result(
            build_pdf_form_from_pages(pages))
        ok, message, form = self.service.import_application_form('project-main', upload)
        self.assertTrue(ok, message)
        doc_id = form['original_file_id']
        _, _, before = self.service.get_parse_review('project-main', 'application_form', doc_id)
        issue = next(i for i in before['issues'] if i['code'] == 'ocr_quality')
        item = {'issue_key': issue['issue_key'], 'action': 'correct_table', 'reason': '合成原页逐字段核对',
                'outside_text': '药品通用名称：合成甲片\n规格：3mg', 'outside_text_verified': True,
                'table': {'row_count': 1, 'column_count': 2, 'cells': [
                    {'row': 0, 'column': 0, 'text': '英文名称'},
                    {'row': 0, 'column': 1, 'text': 'Synthetic'}]}}
        payload = {'source_identity': before['source_identity'], 'expected_revision': before['revision'], 'items': [item]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, payload)
        self.assertTrue(ok, message)
        for _ in range(2):
            fields = self.service.get_application_form('project-main')[2]['form_json']
            self.assertEqual(fields['item_6_generic_name']['value'], '合成甲片')
            self.assertEqual(fields['item_12_specification']['value'], '3mg')
            self.assertEqual(fields['item_7_english_or_latin_name']['value'], 'Synthetic')
        _, _, after = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertEqual(after['original_chunks'], before['original_chunks'])
        invalid = {**payload, 'expected_revision': after['revision'], 'items': [
            {**item, 'outside_text': '药品通用名称：甲\n规格：3mg\n通用名称：乙'}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, invalid)
        self.assertFalse(ok, message)
        _, _, retained = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertEqual(retained['revision'], after['revision'])
        self.assertEqual(retained['original_chunks'], before['original_chunks'])
        fields = self.service.get_application_form('project-main')[2]['form_json']
        self.assertEqual(fields['item_6_generic_name']['value'], '合成甲片')
        self.assertEqual(fields['item_12_specification']['value'], '3mg')

    def test_data_table_revision_reaches_persisted_parent_field(self):
        import io
        import fitz
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import build_pdf_form_from_pages
        self._add_project()
        with fitz.open() as doc:
            page=doc.new_page(width=300,height=400)
            page.insert_text((10,25),'21. 补充申请的内容：',fontname='china-s',fontsize=10)
            upload=io.BytesIO(doc.tobytes())
        upload.filename='synthetic-parent-table.pdf'
        word={'text':'21. 补充申请的内容：旧说明','bbox':[10,10,290,30],'source':'native'}
        bounds=[10,50,290,120]
        pages=[{'page':1,'page_bbox':[0,0,300,400],'status':'partial','content_available':True,
            'text':word['text'],'raw_text':word['text'],'words':[word],'lines':[word],
            'tables':[{'id':'original','page':1,'bbox_pdf':bounds,'markdown':'待核对','cells':[]}],
            'errors':[{'code':'table_structure_unresolved','bbox_pdf':bounds}]}]
        self.service.form_parser=Mock()
        self.service.form_parser.parse_form_file.return_value=FilingChangeFormParserService().form_from_pdf_result(build_pdf_form_from_pages(pages))
        ok,message,form=self.service.import_application_form('project-main',upload)
        self.assertTrue(ok,message)
        doc_id=form['original_file_id']
        before=self.service.get_parse_review('project-main','application_form',doc_id)[2]
        issue=next(i for i in before['issues'] if i['code']=='table_structure_unresolved')
        rows=[['项目','数值'],['合成甲','3 mg']]
        payload={'source_identity':before['source_identity'],'expected_revision':before['revision'],
            'items':[{'issue_key':issue['issue_key'],'action':'correct_table','reason':'逐格核对合成原表',
                'table':{'row_count':2,'column_count':2,'cells':[
                    {'row':r,'column':c,'text':value} for r,row in enumerate(rows) for c,value in enumerate(row)]}}]}
        ok,message,_=self.service.save_parse_review('project-main','application_form',doc_id,payload)
        self.assertTrue(ok,message)
        for _ in range(2):
            current=self.service.get_application_form('project-main')[2]
            field=current['form_json']['item_21_change_content']
            self.assertIn('3 mg',field['value']);self.assertIn('3 mg',field['source_text'])
            self.assertIn('旧说明',field['value'])
            self.assertTrue(any(r.get('bbox_pdf')==bounds for r in field['source_regions']))
        after=self.service.get_parse_review('project-main','application_form',doc_id)[2]
        self.assertEqual(after['original_chunks'],before['original_chunks'])

        # 模型替身只验证实际审评输入；公开启动、任务持久化和报告生成使用真实服务。
        from test_filing_change_review_service_regression import RuntimeTaskStore, _RuleList
        from agent.agent_backend.services.filing_change_report_service import FilingChangeReportService
        from docx import Document
        from copy import deepcopy
        import time
        for category in ('1','2','4','5'):
            material_id='parent-table-required-'+category
            self._add_submission('project-main',material_id,created_at=datetime.now(),category=category,parse_status='success')
            path=self.service._project_path('project-main')/'parsed'/(material_id+'.json')
            path.write_text(json.dumps([{'page':1,'status':'success','text':'Synthetic material','errors':[],'tables':[]}]))
        self.service.rule_service=_RuleList()
        self.service.report_service=FilingChangeReportService(self.service.root_dir)
        current=self.service.get_application_form('project-main')[2]['form_json']
        filled=deepcopy(current)
        filled['item_5_application_matter_category']['selected_values']=['延长有效期']
        filled['item_6_generic_name']['value']='合成甲片'
        ok,message,_=self.service.save_application_form('project-main',{'form_json':filled,
            'edited_paths':{'item_5_application_matter_category':['selected_values'],'item_6_generic_name':['value']},
            'expected_values':current})
        self.assertTrue(ok,message)
        def inspect_input(**inputs):
            field=inputs['application_form']['form_json']['item_21_change_content']
            self.assertIn('3 mg',field['value']);self.assertIn('旧说明',field['value'])
            self.assertIn('3 mg',field['source_text'])
            return {'overall_conclusion':{'result':field['value']}}
        self.service.orchestrator=Mock();self.service.orchestrator.run.side_effect=inspect_input
        app=Flask(__name__);app.register_blueprint(controller.filing_change_review_bp)
        app_service=object.__new__(FilingChangeReviewAppService);app_service.service=self.service
        store=RuntimeTaskStore(connection=self.connection,ensure_schema=False)
        with patch.object(controller,'get_service',return_value=app_service),patch.object(controller,'_prune_tasks'),patch.object(controller,'get_runtime_task_store',return_value=store):
            response=app.test_client().post('/filing-change-review/projects/project-main/review/start')
            self.assertEqual(response.status_code,200,response.get_json())
            ids=response.get_json(force=True)['data'];deadline=time.monotonic()+20
            task=store.get_task(ids['task_id'])
            while task['status'] in ('pending','running') and time.monotonic()<deadline:
                time.sleep(.05);task=store.get_task(ids['task_id'])
            self.assertEqual(task['status'],'completed',task.get('error_message'))
        self.service.orchestrator.run.assert_called_once()
        report=self.service.root_dir/'projects'/'project-main'/'reports'/(ids['run_id']+'.docx')
        exported=Document(report)
        text='\n'.join([p.text for p in exported.paragraphs]+[c.text for t in exported.tables for r in t.rows for c in r.cells])
        for output in (task['result']['review_report_markdown'],text):
            self.assertIn('3 mg',output);self.assertIn('旧说明',output)

    def test_manual_form_table_save_read_and_ambiguous_retry_preserves_revision(self):
        import io
        import fitz
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import build_pdf_form_from_pages
        self._add_project()
        with fitz.open() as doc:
            doc.new_page(width=600, height=400)
            upload = io.BytesIO(doc.tobytes())
        upload.filename = 'synthetic-table-form.pdf'
        bounds = [30, 40, 400, 160]
        existing = {'text': '药品英文名称：Synthetic', 'bbox': [30, 200, 350, 220], 'source': 'native'}
        pages = [{'page': 1, 'page_bbox': [0, 0, 600, 400], 'status': 'partial',
                  'content_available': True, 'text': '原始表格待核对\n' + existing['text'],
                  'raw_text': '原始表格待核对\n' + existing['text'],
                  'lines': [existing], 'words': [existing],
                  'tables': [{'id': 'original', 'page': 1, 'bbox_pdf': bounds, 'markdown': '原始表格待核对', 'cells': []}],
                  'errors': [{'code': 'table_structure_unresolved', 'bbox_pdf': bounds}]}]
        self.service.form_parser = Mock()
        self.service.form_parser.parse_form_file.return_value = FilingChangeFormParserService().form_from_pdf_result(
            build_pdf_form_from_pages(pages))
        ok, message, form = self.service.import_application_form('project-main', upload)
        self.assertTrue(ok, message)
        doc_id = form['original_file_id']
        _, _, before = self.service.get_parse_review('project-main', 'application_form', doc_id)
        issue = next(i for i in before['issues'] if i['code'] == 'table_structure_unresolved')
        corrected = {'row_count': 3, 'column_count': 2, 'cells': [
            {'row': 0, 'column': 0, 'rowspan': 2, 'text': '药品通用名称'},
            {'row': 0, 'column': 1, 'rowspan': 2, 'text': '合成甲片'},
            {'row': 2, 'column': 0, 'text': '规格'},
            {'row': 2, 'column': 1, 'text': '3mg'}]}
        payload = {'source_identity': before['source_identity'], 'expected_revision': before['revision'],
                   'items': [{'issue_key': issue['issue_key'], 'action': 'correct_table',
                              'table': corrected, 'reason': '合成原表逐格核对'}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, payload)
        self.assertTrue(ok, message)
        for _ in range(2):
            current = self.service.get_application_form('project-main')[2]
            self.assertEqual(current['form_json']['item_6_generic_name']['value'], '合成甲片')
            self.assertEqual(current['form_json']['item_12_specification']['value'], '3mg')
            self.assertTrue(current['parse_resolution']['manually_reviewed'])
        _, _, after = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertEqual(after['original_chunks'], before['original_chunks'])
        ambiguous = {'row_count': 2, 'column_count': 2, 'cells': [
            {'row': 0, 'column': 0, 'text': '药品通用名称'},
            {'row': 1, 'column': 0, 'text': '规格'},
            {'row': 0, 'column': 1, 'rowspan': 2, 'text': '合成甲片 3mg'}]}
        invalid = {**payload, 'expected_revision': after['revision'],
                   'items': [{**payload['items'][0], 'table': ambiguous}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, invalid)
        self.assertFalse(ok, message)
        self.assertIn('新的字段归属问题', message)
        _, _, retained = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertEqual(retained['revision'], after['revision'])
        self.assertEqual(retained['original_chunks'], before['original_chunks'])
        current = self.service.get_application_form('project-main')[2]
        self.assertEqual(current['form_json']['item_6_generic_name']['value'], '合成甲片')
        self.assertEqual(current['form_json']['item_12_specification']['value'], '3mg')

        # 正式启动与执行服务重新读取修订字段；仅模型编排为无外发替身。
        from test_filing_change_review_service_regression import RuntimeTaskStore, _RuleList
        from agent.agent_backend.services.filing_change_report_service import FilingChangeReportService
        from docx import Document
        import time
        for category in ('1', '2', '4', '5'):
            doc = 'form-required-' + category
            self._add_submission('project-main', doc, created_at=datetime.now(), category=category, parse_status='success')
            path = self.service._project_path('project-main') / 'parsed' / (doc + '.json')
            path.write_text(json.dumps([{'page': 1, 'status': 'success', 'text': 'Synthetic material',
                                         'errors': [], 'tables': []}]))
        self.service.rule_service = _RuleList()
        self.service.report_service = FilingChangeReportService(self.service.root_dir)
        current_form = self.service.get_application_form('project-main')[2]['form_json']
        import copy
        filled = copy.deepcopy(current_form)
        filled['item_5_application_matter_category']['selected_values'] = ['延长有效期']
        ok, message, _ = self.service.save_application_form('project-main', {
            'form_json': filled, 'edited_paths': {'item_5_application_matter_category': ['selected_values']},
            'expected_values': current_form})
        self.assertTrue(ok, message)
        def inspect_form(**inputs):
            actual = inputs['application_form']['form_json']
            self.assertEqual(actual['item_6_generic_name']['value'], '合成甲片')
            self.assertEqual(actual['item_12_specification']['value'], '3mg')
            return {'overall_conclusion': {'result': actual['item_6_generic_name']['value'] + ' / ' +
                                           actual['item_12_specification']['value']}}
        self.service.orchestrator = Mock()
        self.service.orchestrator.run.side_effect = inspect_form
        app = Flask(__name__)
        app.register_blueprint(controller.filing_change_review_bp)
        app_service = object.__new__(FilingChangeReviewAppService)
        app_service.service = self.service
        store = RuntimeTaskStore(connection=self.connection, ensure_schema=False)
        with patch.object(controller, 'get_service', return_value=app_service), patch.object(
                controller, '_prune_tasks'), patch.object(controller, 'get_runtime_task_store', return_value=store):
            response = app.test_client().post('/filing-change-review/projects/project-main/review/start')
            self.assertEqual(response.status_code, 200, response.get_json())
            ids = response.get_json(force=True)['data']
            deadline = time.monotonic() + 20
            task = store.get_task(ids['task_id'])
            while task['status'] in ('pending', 'running') and time.monotonic() < deadline:
                time.sleep(.05)
                task = store.get_task(ids['task_id'])
            self.assertEqual(task['status'], 'completed', task.get('error_message'))
        self.service.orchestrator.run.assert_called_once()
        report = self.service.root_dir / 'projects' / 'project-main' / 'reports' / (ids['run_id'] + '.docx')
        exported = Document(report)
        text = '\n'.join([p.text for p in exported.paragraphs] +
                         [c.text for t in exported.tables for r in t.rows for c in r.cells])
        for output in (task['result']['review_report_markdown'], text):
            self.assertIn('合成甲片', output)
            self.assertIn('3mg', output)

    def test_form_partial_correction_can_save_without_dismissing_other_field_issue(self):
        import io
        import fitz
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import build_pdf_form_from_pages
        self._add_project()
        with fitz.open() as doc:
            doc.new_page(width=600, height=400)
            upload = io.BytesIO(doc.tobytes())
        upload.filename = 'partial-form.pdf'
        lines = [
            {'text': '药品通用名称：旧名称', 'bbox': [30, 40, 190, 60], 'source': 'native'},
            {'text': '规格：3mg', 'bbox': [300, 40, 390, 60], 'source': 'native'},
            {'text': '跨界文字需要核对', 'bbox': [260, 80, 390, 100], 'source': 'native'},
        ]
        pages = [{'page': 1, 'page_bbox': [0, 0, 600, 400], 'status': 'partial',
                  'content_available': True, 'text': '\n'.join(x['text'] for x in lines),
                  'raw_text': '\n'.join(x['text'] for x in lines), 'lines': lines, 'words': lines,
                  'tables': [], 'errors': [{'code': 'ocr_coverage', 'bbox_pdf': [28, 38, 192, 62]}]}]
        parsed_pdf = build_pdf_form_from_pages(pages)
        self.assertTrue(parsed_pdf['unassigned_regions'])
        self.service.form_parser = Mock()
        self.service.form_parser.parse_form_file.return_value = FilingChangeFormParserService().form_from_pdf_result(parsed_pdf)
        ok, message, form = self.service.import_application_form('project-main', upload)
        self.assertTrue(ok, message)
        doc_id = form['original_file_id']
        _, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        issue = next(i for i in context['issues'] if i['code'] == 'ocr_coverage')
        payload = {'source_identity': context['source_identity'], 'expected_revision': context['revision'],
                   'items': [{'issue_key': issue['issue_key'], 'action': 'correct_text',
                              'text': '药品通用名称：修订名称', 'reason': '只核对本区域，其他跨界区域未处理'}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, payload)
        self.assertTrue(ok, message)
        current = self.service.get_application_form('project-main')[2]
        self.assertEqual(current['form_json']['item_6_generic_name']['value'], '修订名称')
        self.assertFalse(current['parse_resolution']['manually_reviewed'])
        _, _, after = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertEqual(after['original_chunks'], context['original_chunks'])
        remaining = [i for i in after['issues'] if not i['resolved']]
        self.assertTrue(any(i['code'] == 'field_region_crossing' for i in remaining))
        self.assertFalse(self.service.get_parse_readiness('project-main')['ready'])
        # 改变左侧字段身份会改变跨界文字的候选归属，不能借分步保存放行。
        changed = {**payload, 'expected_revision': after['revision'], 'items': [
            {**payload['items'][0], 'text': '规格：3mg'}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, changed)
        self.assertFalse(ok, message)
        self.assertIn('新的字段归属问题', message)
        _, _, preserved = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertEqual(preserved['revision'], after['revision'])
        self.assertEqual(self.service.get_application_form('project-main')[2]['form_json']['item_6_generic_name']['value'], '修订名称')
        crossing = next(i for i in remaining if i['code'] == 'field_region_crossing')
        assigned = {'issue_key': crossing['issue_key'], 'action': 'correct_text',
                    'text': '补充名称', 'target_item_no': 6, 'reason': '对照原页确认此段属于药品通用名称'}
        next_payload = {'source_identity': preserved['source_identity'], 'expected_revision': preserved['revision'],
                        'items': preserved['items'] + [assigned]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, next_payload)
        self.assertTrue(ok, message)
        final = self.service.get_application_form('project-main')[2]
        self.assertIn('补充名称', final['form_json']['item_6_generic_name']['value'])
        self.assertNotIn('补充名称', final['form_json']['item_12_specification']['value'])
        self.assertEqual(final['form_json']['item_12_specification']['value'], '3mg')
        self.assertTrue(final['parse_resolution']['manually_reviewed'],
                        self.service.get_parse_review('project-main', 'application_form', doc_id)[2]['issues'])
        _, _, final_context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertEqual(final_context['original_chunks'], context['original_chunks'])
        split = {**assigned, 'text': '补充名称\n每盒十片', 'all_text_verified': True,
                 'field_assignments': [{'target_item_no': 6, 'text': '补充名称'},
                                       {'target_item_no': 12, 'text': '每盒十片'}]}
        payload_split = {'source_identity': final_context['source_identity'], 'expected_revision': final_context['revision'],
                         'items': [i for i in final_context['items'] if i['issue_key'] != crossing['issue_key']] + [split]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, payload_split)
        self.assertTrue(ok, message)
        split_form = self.service.get_application_form('project-main')[2]
        self.assertIn('补充名称', split_form['form_json']['item_6_generic_name']['value'])
        self.assertNotIn('每盒十片', split_form['form_json']['item_6_generic_name']['value'])
        self.assertIn('每盒十片', split_form['form_json']['item_12_specification']['value'])
        self.assertNotIn('补充名称', split_form['form_json']['item_12_specification']['value'])
        self.assertTrue(split_form['parse_resolution']['manually_reviewed'])

    def test_identity_timestamp_transport_preserves_nanoseconds(self):
        from types import SimpleNamespace
        path = Mock()
        path.is_file.return_value = True
        path.stat.return_value = SimpleNamespace(st_mtime_ns=1700000000000000001, st_size=12)
        identity = self.service._parse_source_identity({'doc_id': 'synthetic'}, path)
        self.assertEqual(identity['parsed_mtime_ns'], '1700000000000000001')
        legacy = {**identity, 'parsed_mtime_ns': 1700000000000000001}
        self.assertTrue(self.service._parse_identity_matches(legacy, identity))
        rounded = {**identity, 'parsed_mtime_ns': int(float(legacy['parsed_mtime_ns']))}
        self.assertFalse(self.service._parse_identity_matches(rounded, identity))

    def test_form_region_save_reextracts_and_preserves_manual_fields(self):
        import io
        import fitz
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import build_pdf_form_from_pages
        from test_filing_change_review_service_regression import FilingChangeApplicationForm
        self._add_project()
        with fitz.open() as doc:
            doc.new_page(width=300, height=400)
            upload = io.BytesIO(doc.tobytes())
        upload.filename = 'form.pdf'
        lines = [{'text': '6. 药品通用名称', 'bbox': [20, 40, 160, 60], 'source': 'native'},
                 {'text': '旧药品名称', 'bbox': [20, 65, 180, 85], 'source': 'native'}]
        pages = [{'page': 1, 'page_bbox': [0, 0, 300, 400], 'status': 'partial', 'content_available': True,
                  'text': '\n'.join(l['text'] for l in lines), 'raw_text': '\n'.join(l['text'] for l in lines),
                  'lines': lines, 'words': lines, 'tables': [],
                  'errors': [{'code': 'ocr_coverage', 'bbox_pdf': [18, 63, 195, 87]}]}]
        parsed = FilingChangeFormParserService().form_from_pdf_result(build_pdf_form_from_pages(pages))
        self.service.form_parser = Mock()
        self.service.form_parser.parse_form_file.return_value = parsed
        ok, message, form = self.service.import_application_form('project-main', upload)
        self.assertTrue(ok, message)
        with self.connection.get_session() as session:
            row = session.query(FilingChangeApplicationForm).first()
            original = json.loads(row.form_json)
            original['item_7_english_or_latin_name'].update(value='manual retained', manual_modified=True, manual_paths=['value'])
            row.form_json = json.dumps(original, ensure_ascii=False)
            session.commit()
        doc_id = form['original_file_id']
        ok, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertTrue(ok)
        payload = {'source_identity': context['source_identity'], 'expected_revision': context['revision'],
                   'items': [{'issue_key': context['issues'][0]['issue_key'], 'action': 'correct_text',
                              'text': '修订药品名称', 'reason': '核对合成原页'}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, payload)
        self.assertTrue(ok, message)
        effective = self.service.get_application_form('project-main')[2]
        self.assertEqual(effective['form_json']['item_6_generic_name']['value'], '修订药品名称')
        self.assertEqual(effective['form_json']['item_7_english_or_latin_name']['value'], 'manual retained')
        self.assertEqual(effective['parse_status'], 'partial')
        self.assertTrue(effective['parse_resolution']['manually_reviewed'], self.service.get_parse_review('project-main', 'application_form', doc_id)[2]['issues'])
        with self.connection.get_session() as session:
            row = session.query(FilingChangeApplicationForm).first()
            self.assertEqual(json.loads(row.form_json), original)
        self.assertFalse(any(i['source_kind'] == 'application_form'
                             for i in self.service.get_parse_readiness('project-main')['blocking_issues']))
        self.assertFalse(self.service.save_parse_review('project-main', 'application_form', doc_id, payload)[0])
        _, _, current = self.service.get_parse_review('project-main', 'application_form', doc_id)
        ok, message, _ = self.service.save_parse_review('project-main', 'application_form', doc_id, {
            'source_identity': current['source_identity'], 'expected_revision': current['revision'], 'items': []})
        self.assertTrue(ok, message)
        restored = self.service.get_application_form('project-main')[2]
        self.assertEqual(restored['form_json']['item_6_generic_name']['value'], '旧药品名称')
        self.assertEqual(restored['form_json']['item_7_english_or_latin_name']['value'], 'manual retained')
        self.assertTrue(any(i['source_kind'] == 'application_form'
                            for i in self.service.get_parse_readiness('project-main')['blocking_issues']))
        with self.connection.get_session() as session:
            self.assertEqual(json.loads(session.query(FilingChangeApplicationForm).first().form_json), original)

    def test_conflicting_form_revision_keeps_original_readable_and_blocks_review(self):
        import io
        import fitz
        from test_filing_change_review_service_regression import FilingChangeApplicationForm
        self._add_project()
        with fitz.open() as doc:
            doc.new_page(width=300, height=200)
            upload = io.BytesIO(doc.tobytes())
        upload.filename = 'conflict.pdf'
        pages = [{'page': 1, 'status': 'partial', 'text': 'ORIGINAL', 'raw_text': 'ORIGINAL',
                  'words': [{'text': 'ORIGINAL', 'bbox': [1, 1, 4, 4]}], 'tables': [],
                  'errors': [{'code': 'ocr_coverage', 'bbox_pdf': [0, 0, 10, 10]},
                             {'code': 'ocr_coverage', 'bbox_pdf': [0, 0, 12, 12]}]}]
        self.service.form_parser = Mock()
        self.service.form_parser.parse_form_file.return_value = {
            'form_json': {}, 'raw_text': 'ORIGINAL', 'confidence': .8,
            'pdf_parse_result': {'pages': pages}}
        ok, message, form = self.service.import_application_form('project-main', upload)
        self.assertTrue(ok, message)
        doc_id = form['original_file_id']
        manifest, meta, source, path, chunks = self.service._parse_review_context('project-main', 'application_form', doc_id)
        _, _, context = self.service.get_parse_review('project-main', 'application_form', doc_id)
        self.assertGreaterEqual(len(context['issues']), 2)
        meta['parse_revision'] = {'source_identity': context['source_identity'], 'revision': 1,
            'form_json': {}, 'raw_text': 'CONFLICT_CANDIDATE',
            'items': [{'issue_key': issue['issue_key'], 'action': 'correct_text', 'text': 'candidate',
                       'reason': 'historical synthetic correction'} for issue in context['issues']]}
        self.service._save_submission_manifest('project-main', manifest)
        before = path.read_bytes()
        ok, message, recovered = self.service.get_application_form('project-main')
        self.assertTrue(ok, message)
        self.assertEqual(recovered['raw_text'], 'ORIGINAL')
        self.assertEqual(recovered['parse_resolution']['revision_error']['code'], 'parse_revision_invalid')
        self.assertFalse(recovered['parse_resolution']['manually_reviewed'])
        self.assertEqual(path.read_bytes(), before)
        self.assertFalse(self.service.get_parse_readiness('project-main')['ready'])
        self.service.orchestrator = Mock()
        ok, _, blocked = self.service.prepare_review_run('project-main')
        self.assertFalse(ok)
        self.assertEqual(blocked['code'], 'parse_review_required')
        self.assertTrue(any(i['code'] == 'parse_revision_conflict' for i in blocked['blocking_issues']))
        self.service.orchestrator.run.assert_not_called()
        with self.connection.get_session() as session:
            self.assertEqual(session.query(FilingChangeReviewRun).count(), 0)
            row = session.query(FilingChangeApplicationForm).first()
            with self.assertRaises(ValueError):
                self.service._effective_form_revision('project-main', row)

    def test_application_form_page_evidence_commits_with_original(self):
        import io
        import fitz
        self._add_project()
        with fitz.open() as doc:
            doc.new_page(width=300, height=200).insert_text((20, 30), 'Synthetic form evidence')
            content = doc.tobytes()
        upload = io.BytesIO(content)
        upload.filename = 'synthetic-form.pdf'
        pages = [{'page': 1, 'status': 'partial', 'text': 'Synthetic form evidence',
                  'raw_text': 'Synthetic form evidence', 'content_available': True,
                  'words': [{'text': 'Synthetic', 'bbox': [20, 20, 60, 30]}], 'tables': [],
                  'errors': [{'code': 'ocr_quality', 'bbox_pdf': [0, 0, 100, 50], 'reason': 'quality check'}]}]
        self.service.form_parser = Mock()
        self.service.form_parser.parse_form_file.return_value = {
            'form_json': {}, 'raw_text': 'Synthetic form evidence', 'confidence': .8,
            'pdf_parse_result': {'pages': pages}}
        ok, message, result = self.service.import_application_form('project-main', upload)
        self.assertTrue(ok, message)
        source = self.service._load_submission_manifest('project-main')['application_form_source']
        self.assertEqual(source['parsed_pages'], pages)
        self.assertNotIn('parsed_pages', self.service._load_submission_manifest('project-main')['application_form_attempt'])
        ok, _, context = self.service.get_parse_review('project-main', 'application_form', result['original_file_id'])
        self.assertTrue(ok)
        self.assertEqual(context['original_chunks'], pages)
        ok, message, image = self.service.get_parse_review_page('project-main', 'application_form', result['original_file_id'], 1,
                                                              {'source_identity': context['source_identity']})
        self.assertTrue(ok, message)
        self.assertEqual(image['page_bbox'], [0, 0, 300, 200])
        readiness = self.service.get_parse_readiness('project-main')
        self.assertTrue(any(i.get('chunk_index') == 0 and i['source_kind'] == 'application_form'
                            for i in readiness['blocking_issues']))
        self.assertFalse(self.service.get_parse_review('project-main', 'application_form', 'old-file')[0])

    def test_original_page_render_coordinates_and_stale_source(self):
        import base64
        import fitz
        self._add_project()
        self._add_submission('project-main', 'page-doc', created_at=datetime.now(), parse_status='partial')
        path = self.service._project_dir('project-main') / 'submissions' / 'page-doc_material.txt'
        with fitz.open() as pdf:
            page = pdf.new_page(width=300, height=200)
            page.insert_text((30, 40), 'original evidence')
            page.set_rotation(90)
            path.write_bytes(pdf.tobytes())
        original = path.read_bytes()
        ok, _, context = self.service.get_parse_review('project-main', 'submission', 'page-doc')
        self.assertTrue(ok)
        payload = {'source_identity': context['source_identity']}
        app = Flask(__name__)
        app.register_blueprint(controller.filing_change_review_bp)
        app_service = object.__new__(FilingChangeReviewAppService)
        app_service.service = self.service
        url = '/filing-change-review/projects/project-main/parse-review/submission/page-doc/page/1'
        with patch.object(controller, 'get_service', return_value=app_service):
            response = app.test_client().post(url, json=payload)
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertEqual(response.headers['Cache-Control'], 'no-store')
            data = response.get_json()['data']
            self.assertEqual(data['page_bbox'], [0, 0, 300, 200])
            self.assertEqual(data['source_rotation'], 90)
            png = base64.b64decode(data['image_data_url'].split(',', 1)[1])
            self.assertTrue(png.startswith(b'\x89PNG'))
            self.assertEqual((data['image_width'], data['image_height']), (450, 300))
            response = app.test_client().post(url, json={'source_identity': {}})
            self.assertEqual(response.status_code, 409)
            self.assertEqual(response.get_json()['data']['code'], 'parse_revision_conflict')
        self.assertEqual(path.read_bytes(), original)
        self.assertFalse(self.service.get_parse_review_page('project-main', 'submission', 'page-doc', 2, payload)[0])
        self.assertFalse(self.service.get_parse_review_page('another-project', 'submission', 'page-doc', 1, payload)[0])

    def test_unparsed_form_prevents_run_creation_and_keeps_project_state(self):
        self._add_project()
        self.service.orchestrator = Mock()
        ok, _, result = self.service.prepare_review_run('project-main')
        self.assertFalse(ok)
        self.assertEqual(result['code'], 'parse_review_required')
        self.assertTrue(any(x['source_kind'] == 'application_form' for x in result['blocking_issues']))
        with self.connection.get_session() as session:
            self.assertEqual(session.query(FilingChangeReviewRun).count(), 0)
            self.assertEqual(session.query(FilingChangeProject).first().review_status, 'not_started')
        self.service.orchestrator.run.assert_not_called()

    def test_http_start_is_409_and_creates_no_runtime_task(self):
        self._add_project()
        app = Flask(__name__)
        app.register_blueprint(controller.filing_change_review_bp)
        app_service = object.__new__(FilingChangeReviewAppService)
        app_service.service = self.service
        with patch.object(controller, 'get_service', return_value=app_service), patch.object(
                controller, '_prune_tasks'), patch.object(controller, 'get_runtime_task_store') as store:
            response = app.test_client().post('/filing-change-review/projects/project-main/review/start')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()['data']['code'], 'parse_review_required')
        store.assert_not_called()

    def test_changed_inputs_stop_execution_even_if_current_status_is_ready(self):
        self._add_project()
        first = {'ready': True, '_input_identity': [{'doc_id': 'd1', 'attempt': 'old'}]}
        second = {'ready': True, '_input_identity': [{'doc_id': 'd1', 'attempt': 'new'}]}
        with patch.object(self.service, 'get_parse_readiness', return_value=first):
            ok, _, data = self.service.prepare_review_run('project-main')
        self.assertTrue(ok)
        with patch.object(self.service, 'get_parse_readiness', return_value=second), patch.object(
                self.service, '_run_review_once') as run:
            ok, _, result = self.service.execute_review_run('project-main', data['run_id'])
        self.assertFalse(ok)
        self.assertEqual(result['code'], 'parse_review_required')
        run.assert_not_called()
        with self.connection.get_session() as session:
            self.assertEqual(session.query(FilingChangeReviewRun).first().status, 'failed')

    def test_ready_unchanged_inputs_reach_execution(self):
        self._add_project()
        ready = {'ready': True, '_input_identity': [{'doc_id': 'd1', 'attempt': 'same'}]}
        with patch.object(self.service, 'get_parse_readiness', return_value=ready):
            ok, _, data = self.service.prepare_review_run('project-main')
            self.assertTrue(ok)
            with patch.object(self.service, '_run_review_once', return_value=(True, 'success', {})) as run:
                self.assertTrue(self.service.execute_review_run('project-main', data['run_id'])[0])
                run.assert_called_once()

    def test_success_status_without_parsed_file_is_blocked(self):
        self._add_project()
        self._add_submission('project-main', 'd1', created_at=datetime.now(), parse_status='success')
        result = self.service.get_parse_readiness('project-main')
        self.assertTrue(any(x.get('doc_id') == 'd1' and x['code'] == 'parsed_content_missing'
                            for x in result['blocking_issues']))

    def test_excluded_submission_does_not_hide_required_category(self):
        self._add_project()
        self._add_submission('project-main', 'd1', created_at=datetime.now(), category='5')
        manifest = self.service._load_submission_manifest('project-main')
        manifest.setdefault('doc_meta', {})['d1'] = {'review_enabled': False}
        self.service._save_submission_manifest('project-main', manifest)
        result = self.service.get_parse_readiness('project-main')
        self.assertFalse(any(x.get('doc_id') == 'd1' for x in result['blocking_issues']))
        self.assertTrue(any(x['code'] == 'required_material_missing' and '药学' in x['message']
                            for x in result['blocking_issues']))

    def test_readiness_http_does_not_expose_internal_inputs(self):
        self._add_project()
        app = Flask(__name__)
        app.register_blueprint(controller.filing_change_review_bp)
        app_service = object.__new__(FilingChangeReviewAppService)
        app_service.service = self.service
        with patch.object(controller, 'get_service', return_value=app_service):
            response = app.test_client().get('/filing-change-review/projects/project-main/review/readiness')
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()['data']
        self.assertFalse(payload['ready'])
        self.assertNotIn('_input_identity', payload)

    def test_manual_text_roundtrip_source_unchanged_and_reparse_invalidates(self):
        self._add_project()
        self._add_submission('project-main', 'd1', created_at=datetime.now(), parse_status='partial')
        chunks = [{'page': 1, 'status': 'partial', 'text': 'old', 'raw_text': 'old',
                   'words': [{'text': 'old', 'bbox': [1, 1, 4, 4]}], 'tables': [],
                   'errors': [{'code': 'ocr_coverage', 'bbox_pdf': [0, 0, 10, 10]}]}]
        path = self.service._project_dir('project-main') / 'parsed' / 'd1.json'
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(chunks))
        original = path.read_bytes()
        ok, _, context = self.service.get_parse_review('project-main', 'submission', 'd1')
        self.assertTrue(ok)
        payload = {'source_identity': context['source_identity'], 'expected_revision': context['revision'],
                   'items': [{'issue_key': context['issues'][0]['issue_key'], 'action': 'correct_text',
                              'text': 'corrected content', 'reason': '逐字核对原件'}]}
        ok, message, _ = self.service.save_parse_review('project-main', 'submission', 'd1', payload)
        self.assertTrue(ok, message)
        self.assertEqual(path.read_bytes(), original)
        effective = self.service._load_parsed_submission_map('project-main', ['d1'])['d1']
        self.assertIn('corrected content', effective[0]['text'])
        self.assertNotIn('old', effective[0]['text'])
        # 真实读取/持久化/报告生成，模型边界使用确定性检查器；不证明模型判断质量。
        from agent.agent_backend.services.filing_change_report_service import FilingChangeReportService
        from test_filing_change_review_service_regression import _RuleList, FilingChangeReviewResult
        from docx import Document
        self.service.rule_service = _RuleList()
        self.service.report_service = FilingChangeReportService(self.service.root_dir)
        def review_input(**inputs):
            self.assertEqual(inputs['parsed_submission_map']['d1'][0]['text'], effective[0]['text'])
            self.assertTrue(inputs['submissions'][0]['parse_resolution']['manually_reviewed'])
            self.assertEqual(inputs['submissions'][0]['parse_status'], 'partial')
            return {'overall_conclusion': {'result': inputs['parsed_submission_map']['d1'][0]['text']}}
        self.service.orchestrator = Mock()
        self.service.orchestrator.run.side_effect = review_input
        with self.connection.get_session() as session:
            session.add(FilingChangeReviewRun(run_id='synthetic-effective-report', project_id='project-main',
                task_type='extend_validity_period', status='running', started_at=datetime.now(), result_json='{}'))
            session.commit()
        ok, message, reviewed = self.service._run_review_once('project-main', 'synthetic-effective-report')
        self.assertTrue(ok, message)
        self.service.orchestrator.run.assert_called_once()
        self.assertIn('corrected content', reviewed['review_report_markdown'])
        with self.connection.get_session() as session:
            persisted = session.query(FilingChangeReviewResult).filter_by(run_id='synthetic-effective-report').one()
            self.assertEqual(json.loads(persisted.conclusion_json)['overall_conclusion']['result'], effective[0]['text'])
        report_path = self.service.root_dir / 'projects' / 'project-main' / 'reports' / 'synthetic-effective-report.docx'
        self.assertIn('corrected content', '\n'.join(p.text for p in Document(report_path).paragraphs))
        self.assertEqual(path.read_bytes(), original)
        self.assertFalse(any(i.get('doc_id') == 'd1' for i in self.service.get_parse_readiness('project-main')['blocking_issues']))
        ok, _, result = self.service.save_parse_review('project-main', 'submission', 'd1', payload)
        self.assertFalse(ok)
        self.assertEqual(result['code'], 'parse_revision_conflict')
        path.write_text(json.dumps(chunks) + '\n')
        _, _, reread = self.service.get_parse_review('project-main', 'submission', 'd1')
        self.assertTrue(reread['stale'])
        self.assertEqual(reread['effective_chunks'][0]['text'], 'old')
        self.assertTrue(any(i.get('doc_id') == 'd1' for i in self.service.get_parse_readiness('project-main')['blocking_issues']))

    def test_corrected_table_reaches_stability_without_overwriting_original(self):
        from copy import deepcopy
        from agent.agent_backend.services.filing_parse_resolution import manual_table
        from agent.agent_backend.services.filing_change_stability_service import FilingChangeStabilityService
        from test_filing_change_review_service_regression import FilingChangeSubmissionFile
        self._add_project()
        self._add_submission('project-main', 'table-doc', created_at=datetime.now(), parse_status='partial')
        with self.connection.get_session() as session:
            row = session.query(FilingChangeSubmissionFile).filter_by(doc_id='table-doc').one()
            row.file_name = '稳定性研究.pdf'
            row.material_category = '5.6'
            session.commit()
        rows = [['批号', '时间点', '指标', '限度', '结果'],
                ['SYNTHETIC-B1', '0月', '杂质A', '≤0.20%', '0.10%'],
                ['SYNTHETIC-B1', '3月', '杂质A', '≤0.20%', '0.10%']]
        def spec(values):
            return {'row_count': 3, 'column_count': 5, 'cells': [
                {'row': r, 'column': c, 'text': value} for r, row in enumerate(values) for c, value in enumerate(row)]}
        table = manual_table(spec(rows), [0, 0, 500, 150], 1, 'synthetic')
        table['needs_review'] = True
        words = [{'text': value, 'bbox': [c*100+2, r*50+2, (c+1)*100-2, (r+1)*50-2],
                  'source': 'ocr'} for r, row in enumerate(rows) for c, value in enumerate(row)]
        chunks = [{'page': 1, 'status': 'partial', 'text': table['markdown'], 'raw_text': table['markdown'],
                   'tables': [table], 'tables_in_text': True, 'words': words,
                   'errors': [{'code': 'table_structure_unresolved', 'bbox_pdf': [0, 0, 500, 150]}]}]
        path = self.service._project_path('project-main') / 'parsed' / 'table-doc.json'
        path.write_text(json.dumps(chunks))
        original = path.read_bytes()
        before_analysis = FilingChangeStabilityService().run({}, [{'doc_id': 'table-doc',
            'file_name': '稳定性研究.pdf', 'material_category': '5.6', 'parse_status': 'partial',
            'extracted_json': self.service._extract_structured_payload('稳定性研究.pdf', chunks)}])
        self.assertEqual(before_analysis['limit_check']['status_code'], 'within_spec')
        # 完整合成资料：就绪检查及HTTP启动不做替身；本例不验收文件识别效果。
        from test_filing_change_review_service_regression import FilingChangeApplicationForm, RuntimeTaskStore
        for category in ('1', '2', '4'):
            doc_id = 'required-' + category
            self._add_submission('project-main', doc_id, created_at=datetime.now(), category=category, parse_status='success')
            (path.parent / (doc_id + '.json')).write_text(json.dumps([
                {'page': 1, 'status': 'success', 'text': 'Synthetic required material ' + category,
                 'errors': [], 'tables': []}]))
        with self.connection.get_session() as session:
            session.add(FilingChangeApplicationForm(project_id='project-main', original_file_id='synthetic-form',
                form_json=json.dumps({'item_6_generic_name': {'value': '合成药品'},
                                      'item_5_application_matter_category': {'selected_values': ['延长有效期']}}),
                raw_text='Synthetic application form', parse_status='success', confidence=1,
                created_at=datetime.now(), updated_at=datetime.now()))
            session.commit()
        app = Flask(__name__)
        app.register_blueprint(controller.filing_change_review_bp)
        app_service = object.__new__(FilingChangeReviewAppService)
        app_service.service = self.service
        store = RuntimeTaskStore(connection=self.connection, ensure_schema=False)
        start_url = '/filing-change-review/projects/project-main/review/start'
        with patch.object(controller, 'get_service', return_value=app_service), patch.object(
                controller, '_prune_tasks'), patch.object(controller, 'get_runtime_task_store') as blocked_store:
            blocked = app.test_client().post(start_url)
            self.assertEqual(blocked.status_code, 409, blocked.get_json())
            self.assertEqual(blocked.get_json()['data']['code'], 'parse_review_required')
            blocked_store.assert_not_called()
        with self.connection.get_session() as session:
            self.assertEqual(session.query(FilingChangeReviewRun).count(), 0)
        _, _, context = self.service.get_parse_review('project-main', 'submission', 'table-doc')
        issue = next(i for i in context['issues'] if i['code'] == 'table_review_required')
        corrected = deepcopy(rows)
        corrected[2][4] = '0.30%'
        ok, message, _ = self.service.save_parse_review('project-main', 'submission', 'table-doc', {
            'source_identity': context['source_identity'], 'expected_revision': context['revision'],
            'items': [{'issue_key': issue['issue_key'], 'action': 'correct_table', 'reason': '逐格核对合成原表',
                       'table': spec(corrected)}]})
        self.assertTrue(ok, message)
        displayed = self.service.get_submission_parsed_markdown('project-main', 'table-doc')[2]['parsed_chunks']
        execution = self.service._load_parsed_submission_map('project-main', ['table-doc'])['table-doc']
        self.assertEqual(displayed, execution)
        self.assertEqual(execution[0]['tables'][0]['rows'], corrected)
        extracted = self.service._extract_structured_payload('稳定性研究.pdf', execution)
        analyzed = FilingChangeStabilityService().run({}, [{'doc_id': 'table-doc', 'file_name': '稳定性研究.pdf',
            'material_category': '5.6', 'parse_status': 'partial', 'extracted_json': extracted}])
        self.assertEqual(analyzed['record_count'], 2)
        self.assertEqual(analyzed['limit_check']['status_code'], 'out_of_spec')
        self.assertEqual(path.read_bytes(), original)

        # 通过执行服务重新提取覆盖层，不能把上面预先算好的结果注入报告。
        from test_filing_change_review_service_regression import (
            _RuleList, FilingChangeReviewResult, FilingChangeSubmissionFile)
        from agent.agent_backend.services.filing_change_report_service import FilingChangeReportService
        from docx import Document
        self.service.rule_service = _RuleList()
        self.service.report_service = FilingChangeReportService(self.service.root_dir)
        def review_corrected_table(**inputs):
            self.assertEqual(inputs['parsed_submission_map']['table-doc'][0]['tables'][0]['rows'], corrected)
            submission = next(s for s in inputs['submissions'] if s['doc_id'] == 'table-doc')
            self.assertTrue(submission['parse_resolution']['manually_reviewed'], submission['parse_resolution'])
            self.assertEqual(submission['parse_status'], 'partial')
            actual = FilingChangeStabilityService().run(inputs['application_form'], inputs['submissions'])
            self.assertEqual(actual['limit_check']['status_code'], 'out_of_spec')
            return {'stability_trend_analysis': actual, 'overall_conclusion': {'result': '需核对超限结果'}}
        self.service.orchestrator = Mock()
        self.service.orchestrator.run.side_effect = review_corrected_table
        readiness = self.service.get_parse_readiness('project-main')
        self.assertTrue(readiness['ready'], readiness['blocking_issues'])
        import time
        with patch.object(controller, 'get_service', return_value=app_service), patch.object(
                controller, '_prune_tasks'), patch.object(controller, 'get_runtime_task_store', return_value=store):
            response = app.test_client().post(start_url)
            self.assertEqual(response.status_code, 200, response.get_json())
            ids = response.get_json(force=True)['data']
            run_id = ids['run_id']
            deadline = time.monotonic() + 20
            task = store.get_task(ids['task_id'])
            while task['status'] in ('pending', 'running') and time.monotonic() < deadline:
                time.sleep(.05)
                task = store.get_task(ids['task_id'])
            self.assertEqual(task['status'], 'completed', task.get('error_message'))
            reviewed = task['result']
        self.service.orchestrator.run.assert_called_once()
        with self.connection.get_session() as session:
            persisted = session.query(FilingChangeReviewResult).filter_by(run_id=run_id).one()
            self.assertEqual(json.loads(persisted.stability_analysis_json), reviewed['stability_trend_analysis'])
        markdown = reviewed['review_report_markdown']
        report_path = self.service.root_dir / 'projects' / 'project-main' / 'reports' / (run_id + '.docx')
        doc = Document(report_path)
        exported = '\n'.join([p.text for p in doc.paragraphs] +
                             [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
        for output in (markdown, exported):
            self.assertIn('0.30', output)
            self.assertIn('超限', output)
            self.assertIn('杂质A', output)
        self.assertEqual(path.read_bytes(), original)
        # 读取历史结果不能重写报告；资料变化后必须明确标旧。
        report_bytes = report_path.read_bytes()
        ok, message, current = self.service.get_run_result(run_id)
        self.assertTrue(ok,message)
        self.assertEqual(current['input_state']['status'],'current')
        with self.connection.get_session() as session:
            saved_result = session.query(FilingChangeReviewResult).filter_by(run_id=run_id).one().conclusion_json
        path.write_bytes(original + b'\n')  # 仅本测试临时解析文件，模拟新的解析尝试文件身份。
        ok, message, stale = self.service.get_run_result(run_id)
        self.assertTrue(ok,message)
        self.assertEqual(stale['input_state']['status'],'stale')
        self.assertEqual(report_path.read_bytes(),report_bytes)
        with self.connection.get_session() as session:
            self.assertEqual(session.query(FilingChangeReviewResult).filter_by(run_id=run_id).one().conclusion_json,saved_result)

    def test_conflicting_saved_revision_remains_inspectable(self):
        self._add_project()
        self._add_submission('project-main', 'd1', created_at=datetime.now(), parse_status='partial')
        chunks = [{'page': 1, 'status': 'partial', 'text': 'old', 'raw_text': 'old',
                   'words': [{'text': 'old', 'bbox': [1, 1, 4, 4]}], 'tables': [],
                   'errors': [{'code': 'ocr_coverage', 'bbox_pdf': [0, 0, 10, 10]},
                              {'code': 'ocr_coverage', 'bbox_pdf': [0, 0, 12, 12]}]}]
        path = self.service._project_dir('project-main') / 'parsed' / 'd1.json'
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(chunks))
        original = path.read_bytes()
        _, _, context = self.service.get_parse_review('project-main', 'submission', 'd1')
        manifest, meta, _, _, _ = self.service._parse_review_context('project-main', 'submission', 'd1')
        items = [{'issue_key': issue['issue_key'], 'action': 'correct_text',
                  'text': 'candidate', 'reason': 'historical synthetic correction'} for issue in context['issues']]
        meta['parse_revision'] = {'source_identity': context['source_identity'], 'revision': 1,
                                  'source_status': 'partial', 'items': items}
        self.service._save_submission_manifest('project-main', manifest)
        ok, message, recovered = self.service.get_parse_review('project-main', 'submission', 'd1')
        self.assertTrue(ok, message)
        self.assertEqual(recovered['revision_error']['code'], 'parse_revision_invalid')
        self.assertEqual(recovered['effective_chunks'], chunks)
        self.assertEqual(recovered['items'], items)
        self.assertFalse(any(i['resolved'] for i in recovered['issues']))
        self.assertEqual(path.read_bytes(), original)
        ok, message, result = self.service.save_parse_review('project-main', 'submission', 'd1', {
            'source_identity': recovered['source_identity'], 'expected_revision': 1, 'items': []})
        self.assertTrue(ok, message)
        self.assertEqual(result['resolved_issue_keys'], [])
        _, meta, _, _, _ = self.service._parse_review_context('project-main', 'submission', 'd1')
        self.assertEqual(meta['parse_revision']['items'], [])
        self.assertEqual(meta['parse_revision_history'][-1]['items'], items)
        self.assertTrue(any(i.get('doc_id') == 'd1' for i in self.service.get_parse_readiness('project-main')['blocking_issues']))
        self.assertEqual(path.read_bytes(), original)

    def test_cannot_read_another_projects_parse_review(self):
        self._add_project('p1'); self._add_project('p2')
        self._add_submission('p1', 'd1', created_at=datetime.now())
        self.assertFalse(self.service.get_parse_review('p2', 'submission', 'd1')[0])


if __name__ == '__main__':
    unittest.main()
