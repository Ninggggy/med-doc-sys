import io
import importlib.util
import json
import multiprocessing
import os
import tempfile
import threading
import time
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest import mock

from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker

from agent.agent_backend.application.filing_change_review_app_service import FilingChangeReviewAppService
from agent.agent_backend.database.mysql.db_model import (
    FilingChangeApplicationForm,
    FilingChangeProject,
    FilingChangeReferenceMaterial,
    FilingChangeReviewReport,
    FilingChangeReviewResult,
    FilingChangeReviewRun,
    FilingChangeSubmissionFile,
    RuntimeTask,
)
from agent.agent_backend.database.mysql.mysql_conn import Base
from agent.agent_backend.services.filing_change_review_service import FilingChangeReviewService
from agent.agent_backend.services.runtime_task_store import RuntimeTaskStore


_CONTROLLER_PATH = Path(__file__).resolve().parents[1] / "agent_backend" / "controller" / "filing_change_review_controller.py"
_CONTROLLER_SPEC = importlib.util.spec_from_file_location("_filing_change_review_controller_regression", _CONTROLLER_PATH)
if _CONTROLLER_SPEC is None or _CONTROLLER_SPEC.loader is None:
    raise RuntimeError("cannot load filing change review controller")
filing_change_review_controller = importlib.util.module_from_spec(_CONTROLLER_SPEC)
_CONTROLLER_SPEC.loader.exec_module(filing_change_review_controller)


class _TestConnection:
    def __init__(self, database_path: Path) -> None:
        self.engine = create_engine(
            f"sqlite:///{database_path}",
            connect_args={"check_same_thread": False},
        )

        @event.listens_for(self.engine, "connect")
        def _enable_foreign_keys(dbapi_connection, _connection_record):
            cursor = dbapi_connection.cursor()
            cursor.execute("PRAGMA foreign_keys=ON")
            cursor.close()

        self._session_factory = sessionmaker(bind=self.engine, expire_on_commit=False)
        self._fail_guard = threading.Lock()
        self.fail_next_commit = False
        Base.metadata.create_all(
            self.engine,
            tables=[
                FilingChangeProject.__table__,
                FilingChangeApplicationForm.__table__,
                FilingChangeSubmissionFile.__table__,
                FilingChangeReferenceMaterial.__table__,
                FilingChangeReviewRun.__table__,
                FilingChangeReviewResult.__table__,
                FilingChangeReviewReport.__table__,
                RuntimeTask.__table__,
            ],
        )

    def get_session(self):
        session = self._session_factory()
        original_commit = session.commit

        def _commit():
            with self._fail_guard:
                should_fail = self.fail_next_commit
                if should_fail:
                    self.fail_next_commit = False
            if should_fail:
                raise RuntimeError("injected database commit failure")
            return original_commit()

        session.commit = _commit
        return session


class _Upload(io.BytesIO):
    def __init__(self, filename: str, content: bytes) -> None:
        super().__init__(content)
        self.filename = filename


class _MaterialParser:
    def __init__(self, parsed=None) -> None:
        self.parsed = parsed or [{"chunk_id": "chunk-1", "text": "new parsed content"}]

    def parse_file(self, _file_path: str):
        return self.parsed


class _FormParser:
    def parse_form_file(self, _file_path: str):
        return {"form_json": {"field": "value"}, "raw_text": "form text", "confidence": 0.9}


class _ReportCapture:
    def __init__(self, root_dir: Path) -> None:
        self.root_dir = root_dir
        self.payload = None

    def generate_report(self, project_id: str, run_id: str, result_payload):
        self.payload = result_payload
        report_dir = self.root_dir / "projects" / project_id / "reports"
        report_dir.mkdir(parents=True, exist_ok=True)
        word_path = report_dir / f"{run_id}.docx"
        word_path.write_bytes(b"report")
        return {
            "markdown": "generated report",
            "word_path": str(word_path),
            "markdown_path": str(report_dir / f"{run_id}.md"),
        }


class _ReportFileWriter:
    def __init__(self, root_dir: Path, content_prefix: bytes = b"replacement") -> None:
        self.root_dir = root_dir
        self.content_prefix = content_prefix

    def generate_report(self, project_id: str, run_id: str, result_payload):
        report_dir = self.root_dir / "projects" / project_id / "reports"
        report_dir.mkdir(parents=True, exist_ok=True)
        paths = {
            suffix: report_dir / f"{run_id}{suffix}"
            for suffix in (".md", ".docx", ".txt")
        }
        for suffix, path in paths.items():
            path.write_bytes(self.content_prefix + suffix.encode("ascii"))
        return {
            "markdown": "replacement report",
            "word_path": str(paths[".docx"]),
            "markdown_path": str(paths[".md"]),
        }


class _RuleList:
    @staticmethod
    def list_rules(_payload):
        return {"list": []}


class _ReviewOrchestrator:
    @staticmethod
    def run(**_kwargs):
        return {
            "overall_conclusion": {"result": "continue"},
            "consistency_check": {"source": "review"},
            "conclusion": {"result": "review"},
        }


class _BlockingRuntimeTaskStore:
    def __init__(self, delegate: RuntimeTaskStore, entered: threading.Event, release: threading.Event) -> None:
        self.delegate = delegate
        self.entered = entered
        self.release = release

    def create_task(self, **kwargs):
        self.entered.set()
        if not self.release.wait(timeout=5):
            raise TimeoutError("task creation boundary was not released")
        return self.delegate.create_task(**kwargs)

    def __getattr__(self, name: str):
        return getattr(self.delegate, name)


def _increment_manifest(root_dir: str, project_id: str, loops: int) -> None:
    service = object.__new__(FilingChangeReviewService)
    service.root_dir = Path(root_dir)
    for _ in range(loops):
        with service._submission_manifest_lock(project_id):
            manifest = service._load_submission_manifest(project_id)
            manifest["counter"] = int(manifest.get("counter", 0)) + 1
            service._save_submission_manifest(project_id, manifest)


class FilingChangeReviewServiceRegressionTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="filing-change-regression-")
        self.base_dir = Path(self.temp_dir.name)
        self.root_dir = self.base_dir / "filing_change_review"
        self.root_dir.mkdir(parents=True, exist_ok=True)
        self.connection = _TestConnection(self.base_dir / "regression.sqlite")
        self.service = object.__new__(FilingChangeReviewService)
        self.service.root_dir = self.root_dir
        self.service.db_conn = self.connection
        self.service.material_service = _MaterialParser()

    def tearDown(self) -> None:
        self.connection.engine.dispose()
        self.temp_dir.cleanup()

    def _add_project(self, project_id: str = "project-main", *, deleted: bool = False) -> FilingChangeProject:
        now = datetime(2026, 1, 1, 9, 0, 0)
        row = FilingChangeProject(
            project_id=project_id,
            project_name="回归验证项目",
            registration_category="化学药品",
            registration_classification="二类",
            task_type="extend_validity_period",
            review_status="not_started",
            report_status="not_generated",
            remark="",
            created_at=now,
            updated_at=now,
            deleted=deleted,
        )
        session = self.connection.get_session()
        try:
            session.add(row)
            session.commit()
        finally:
            session.close()
        return row

    def _add_submission(
        self,
        project_id: str,
        doc_id: str,
        *,
        created_at: datetime,
        category: str = "5",
        parse_status: str = "pending",
    ) -> Path:
        project_dir = self.service._project_dir(project_id)
        storage_path = project_dir / "submissions" / f"{doc_id}_material.txt"
        storage_path.write_text(f"source-{doc_id}", encoding="utf-8")
        session = self.connection.get_session()
        try:
            session.add(
                FilingChangeSubmissionFile(
                    project_id=project_id,
                    doc_id=doc_id,
                    file_name=f"{doc_id}.txt",
                    file_type="txt",
                    material_category=category,
                    storage_path=str(storage_path),
                    parse_status=parse_status,
                    chunk_status=parse_status,
                    index_status=parse_status,
                    created_at=created_at,
                )
            )
            session.commit()
        finally:
            session.close()
        return storage_path

    def _add_run_and_result(self, project_id: str, run_id: str = "run-main") -> None:
        now = datetime(2026, 1, 2, 10, 0, 0)
        conclusion = {
            "overall_conclusion": {"result": "continue"},
            "field_check": {"source": "field"},
            "consistency_check": {"source": "consistency"},
            "manual_review_items": [],
            "conclusion": {"result": "review"},
        }
        session = self.connection.get_session()
        try:
            session.add(
                FilingChangeReviewRun(
                    run_id=run_id,
                    project_id=project_id,
                    task_type="extend_validity_period",
                    status="completed",
                    started_at=now,
                    finished_at=now,
                    result_json=json.dumps({"run_id": run_id}),
                    error_message="",
                )
            )
            session.flush()
            session.add(
                FilingChangeReviewResult(
                    run_id=run_id,
                    project_id=project_id,
                    formal_review_json="{}",
                    category_suggestion_json="{}",
                    quality_check_json="{}",
                    stability_analysis_json="{}",
                    risk_points_json="[]",
                    evidence_json="[]",
                    conclusion_json=json.dumps(conclusion),
                    created_at=now,
                )
            )
            session.commit()
        finally:
            session.close()

    def test_delete_project_hard_deletes_all_domain_data_and_files(self):
        project_id = "project-delete"
        self._add_project(project_id, deleted=True)
        storage_path = self._add_submission(project_id, "doc-delete", created_at=datetime(2026, 1, 1))
        self._add_run_and_result(project_id, "run-delete")
        project_dir = self.service._project_dir(project_id)
        (project_dir / "parsed" / "doc-delete.json").write_text("[]", encoding="utf-8")
        (project_dir / "charts" / "trend.svg").write_text("<svg/>", encoding="utf-8")
        (project_dir / "application_form" / "form.docx").write_bytes(b"form")
        self.service._save_submission_manifest(project_id, {"doc_meta": {"doc-delete": {}}, "category_meta": {}})
        report_path = project_dir / "reports" / "run-delete.docx"
        report_path.write_bytes(b"report")

        session = self.connection.get_session()
        try:
            now = datetime(2026, 1, 1)
            session.add(
                FilingChangeApplicationForm(
                    project_id=project_id,
                    original_file_id="form",
                    form_json="{}",
                    raw_text="form",
                    parse_status="success",
                    confidence=1.0,
                    created_at=now,
                    updated_at=now,
                )
            )
            session.add(
                FilingChangeReviewReport(
                    report_id="report-delete",
                    project_id=project_id,
                    run_id="run-delete",
                    report_type="review_report",
                    report_content="report",
                    report_file_path=str(report_path),
                    created_at=now,
                    updated_at=now,
                )
            )
            session.add(
                RuntimeTask(
                    task_id="task-delete",
                    domain="filing_change_review",
                    task_type="review",
                    project_id=project_id,
                    run_id="run-delete",
                    status="completed",
                    log_offset=0,
                    created_at=1.0,
                    updated_at=1.0,
                )
            )
            session.add(
                RuntimeTask(
                    task_id="task-keep",
                    domain="pre_review",
                    task_type="review",
                    project_id=project_id,
                    status="completed",
                    log_offset=0,
                    created_at=1.0,
                    updated_at=1.0,
                )
            )
            session.commit()
        finally:
            session.close()

        ok, message = self.service.delete_project(project_id)

        self.assertTrue(ok, message)
        self.assertFalse(project_dir.exists())
        self.assertFalse(storage_path.exists())
        session = self.connection.get_session()
        try:
            for model in [
                FilingChangeProject,
                FilingChangeApplicationForm,
                FilingChangeSubmissionFile,
                FilingChangeReviewRun,
                FilingChangeReviewResult,
                FilingChangeReviewReport,
            ]:
                self.assertEqual(session.query(model).filter_by(project_id=project_id).count(), 0)
            self.assertEqual(session.query(RuntimeTask).filter_by(task_id="task-delete").count(), 0)
            self.assertEqual(session.query(RuntimeTask).filter_by(task_id="task-keep").count(), 1)
        finally:
            session.close()
        self.assertFalse(self.service.get_application_form(project_id)[0])
        self.assertFalse(self.service.get_run_result("run-delete")[0])
        self.assertFalse(self.service.get_project_latest_report(project_id)[0])
        self.assertFalse(self.service.get_submission_parsed_markdown(project_id, "doc-delete")[0])
        self.assertFalse(self.service.get_review_history(project_id)["project_exists"])

    def test_delete_project_commit_failure_restores_database_and_directory(self):
        project_id = "project-delete-rollback"
        self._add_project(project_id)
        storage_path = self._add_submission(project_id, "doc-rollback", created_at=datetime(2026, 1, 1))
        manifest = {"doc_meta": {"doc-rollback": {"material_sub_category": "5.1"}}, "category_meta": {}}
        self.service._save_submission_manifest(project_id, manifest)
        manifest_before = self.service._submission_manifest_path(project_id).read_bytes()
        self.connection.fail_next_commit = True

        ok, message = self.service.delete_project(project_id)

        self.assertFalse(ok)
        self.assertIn("injected database commit failure", message)
        self.assertTrue(storage_path.exists())
        self.assertEqual(self.service._submission_manifest_path(project_id).read_bytes(), manifest_before)
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeProject).filter_by(project_id=project_id).count(), 1)
            self.assertEqual(session.query(FilingChangeSubmissionFile).filter_by(project_id=project_id).count(), 1)
        finally:
            session.close()
        trash_dir = self.root_dir / ".trash"
        self.assertFalse(trash_dir.exists() and any(trash_dir.iterdir()))

    def test_upload_manifest_failure_rolls_back_database_and_files(self):
        project_id = "project-upload-rollback"
        self._add_project(project_id)
        with mock.patch.object(self.service, "_save_submission_manifest", side_effect=OSError("manifest write failed")):
            ok, message, _ = self.service.upload_submission_files(
                project_id,
                [_Upload("material.txt", b"payload")],
                "5",
                "5.1",
            )

        self.assertFalse(ok)
        self.assertIn("manifest write failed", message)
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeSubmissionFile).filter_by(project_id=project_id).count(), 0)
        finally:
            session.close()
        submissions_dir = self.service._project_path(project_id) / "submissions"
        self.assertEqual([path for path in submissions_dir.iterdir() if path.is_file()], [])

    def test_application_form_commit_failure_rolls_back_file_and_database(self):
        project_id = "project-form-rollback"
        self._add_project(project_id)
        self.service.form_parser = _FormParser()
        self.connection.fail_next_commit = True

        ok, message, _ = self.service.import_application_form(
            project_id,
            _Upload("application.docx", b"form-content"),
        )

        self.assertFalse(ok)
        self.assertIn("保存失败", message)
        attempt = self.service.get_application_form(project_id)[2]['latest_attempt']
        self.assertEqual(('failed', 'persist', 'RuntimeError'), (attempt['content_status'], attempt['stage'], attempt['exception_type']))
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeApplicationForm).filter_by(project_id=project_id).count(), 0)
        finally:
            session.close()
        form_dir = self.service._project_path(project_id) / "application_form"
        self.assertEqual([path for path in form_dir.iterdir() if path.is_file()], [])

    def test_application_form_replacement_failure_restores_previous_version(self):
        project_id = "project-form-replacement-rollback"
        self._add_project(project_id)
        form_dir = self.service._project_dir(project_id) / "application_form"
        old_path = form_dir / "app_form_previous.docx"
        old_path.write_bytes(b"previous-form")
        now = datetime(2026, 1, 1)
        session = self.connection.get_session()
        try:
            session.add(
                FilingChangeApplicationForm(
                    project_id=project_id,
                    original_file_id="app_form_previous",
                    form_json=json.dumps({"field": "previous"}),
                    raw_text="previous text",
                    parse_status="success",
                    confidence=0.8,
                    created_at=now,
                    updated_at=now,
                )
            )
            session.commit()
        finally:
            session.close()
        # 持久化故障必须发生在真实可用的填写内容通过判定之后；
        # 不用无填写证据的解析 stub 绕过原件替换保护。
        from docx import Document
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        replacement = Document()
        replacement.add_paragraph('药品通用名称：替换药品')
        replacement_stream = io.BytesIO()
        replacement.save(replacement_stream)
        self.service.form_parser = FilingChangeFormParserService()
        self.connection.fail_next_commit = True

        ok, message, _ = self.service.import_application_form(
            project_id,
            _Upload("replacement.docx", replacement_stream.getvalue()),
        )

        self.assertFalse(ok)
        self.assertIn("保存失败", message)
        self.assertEqual('failed', self.service.get_application_form(project_id)[2]['latest_attempt']['content_status'])
        self.assertEqual(old_path.read_bytes(), b"previous-form")
        self.assertEqual([path.name for path in form_dir.iterdir() if path.is_file()], [old_path.name])
        session = self.connection.get_session()
        try:
            row = session.query(FilingChangeApplicationForm).filter_by(project_id=project_id).one()
            self.assertEqual(row.original_file_id, "app_form_previous")
            self.assertEqual(json.loads(row.form_json), {"field": "previous"})
        finally:
            session.close()

    def test_parse_commit_failure_restores_previous_outputs_manifest_and_status(self):
        project_id = "project-parse-rollback"
        doc_id = "doc-parse-rollback"
        self._add_project(project_id)
        self._add_submission(project_id, doc_id, created_at=datetime(2026, 1, 1))
        manifest = {"doc_meta": {doc_id: {"parse_status": "pending", "extracted_json": {}}}, "category_meta": {}}
        self.service._save_submission_manifest(project_id, manifest)
        manifest_path = self.service._submission_manifest_path(project_id)
        manifest_before = manifest_path.read_bytes()
        parsed_dir = self.service._project_path(project_id) / "parsed"
        old_json = parsed_dir / f"{doc_id}.json"
        old_markdown = parsed_dir / f"{doc_id}.md"
        old_json.write_bytes(b"old-json")
        old_markdown.write_bytes(b"old-markdown")
        self.connection.fail_next_commit = True

        ok, message, _ = self.service.parse_submission(project_id, doc_id)

        self.assertFalse(ok)
        self.assertIn("保存失败", message)
        self.assertEqual(old_json.read_bytes(), b"old-json")
        self.assertEqual(old_markdown.read_bytes(), b"old-markdown")
        # 业务结果必须完全回滚，但新尝试的失败记录不能被一并回滚。
        restored = json.loads(manifest_path.read_text())
        meta = restored['doc_meta'][doc_id]
        self.assertEqual('failed', meta.pop('latest_attempt')['content_status'])
        self.assertEqual('persist', meta.pop('parse_attempts')[0]['stage'])
        self.assertEqual(json.loads(manifest_before), restored)
        session = self.connection.get_session()
        try:
            row = session.query(FilingChangeSubmissionFile).filter_by(doc_id=doc_id).one()
            self.assertEqual(row.parse_status, "pending")
            self.assertEqual(row.chunk_status, "pending")
            self.assertEqual(row.index_status, "pending")
        finally:
            session.close()
        self.assertFalse(any(path.suffix in {".pending", ".backup"} for path in parsed_dir.iterdir()))

    def test_real_pdf_parse_save_refresh_keeps_table_and_partial_diagnostics(self):
        import fitz
        from agent.test.test_pdf_page_extraction import gradient_page
        from agent.agent_backend.services.filing_change_material_service import FilingChangeMaterialService
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import OCRServiceClient
        self._add_project('pdf-project')
        self._add_submission('pdf-project', 'pdf-doc', created_at=datetime(2026, 1, 1))
        path = self.base_dir / 'gradient.pdf'
        doc = fitz.open()
        gradient_page(doc)
        source = fitz.open()
        gradient_page(source)
        image = source[0].get_pixmap()
        page = doc.new_page()
        page.insert_image(page.rect, stream=image.tobytes('png'))
        doc.save(path)
        doc.close()
        source.close()
        session = self.connection.get_session()
        row = session.query(FilingChangeSubmissionFile).filter_by(doc_id='pdf-doc').one()
        row.storage_path, row.file_type, row.file_name = str(path), 'pdf', 'gradient.pdf'
        session.commit()
        session.close()
        self.service.material_service = FilingChangeMaterialService()
        with mock.patch.object(OCRServiceClient, 'image_to_data', side_effect=TimeoutError('page 2 timeout')):
            ok, status, payload = self.service.parse_submission('pdf-project', 'pdf-doc')
        self.assertTrue(ok)
        self.assertEqual(payload['content_status'], 'partial')
        self.assertIn('部分成功', status)
        self.assertEqual(payload['parse_diagnostics']['failed_pages'], [2])
        ok, _, saved = self.service.get_submission_parsed_markdown('pdf-project', 'pdf-doc')
        self.assertTrue(ok)
        self.assertEqual(saved['parsed_chunks'][0]['tables'][0]['rows'][2], ['0.5', '50.0', '50.0'])
        self.assertEqual(saved['markdown'].count('| 0.5 | 50.0 | 50.0 |'), 1)
        self.assertIn('第 2 页', saved['markdown'])
        self.assertEqual(saved['parse_diagnostics']['status'], 'partial')
        session = self.connection.get_session()
        self.assertEqual(session.query(FilingChangeSubmissionFile).filter_by(doc_id='pdf-doc').one().parse_status, 'partial')
        session.close()

    def test_empty_reparse_retains_previous_content_and_retry_replaces_only_on_success(self):
        project_id, doc_id = 'p-retry', 'd-retry'
        self._add_project(project_id)
        source = self._add_submission(project_id, doc_id, created_at=datetime(2026, 1, 1))
        self.assertTrue(self.service.parse_submission(project_id, doc_id)[0])
        parsed_path = self.service._project_path(project_id) / 'parsed' / f'{doc_id}.json'
        previous = parsed_path.read_bytes()
        with mock.patch.object(self.service.material_service, 'parse_file', return_value=[]):
            ok, _, attempt = self.service.parse_submission(project_id, doc_id)
        self.assertFalse(ok)
        self.assertEqual('empty_result', attempt['code'])
        self.assertEqual(previous, parsed_path.read_bytes())
        self.assertEqual(f'source-{doc_id}', source.read_text())
        refreshed = self.service.list_submissions(project_id, {})['list'][0]
        self.assertEqual('success', refreshed['parse_status'])
        self.assertEqual('failed', refreshed['latest_attempt']['content_status'])
        with mock.patch.object(self.service.material_service, 'parse_file', return_value=[{'text': 'fresh actual body'}]):
            self.assertTrue(self.service.parse_submission(project_id, doc_id)[0])
        self.assertIn('fresh actual body', parsed_path.read_text())
        history = self.service._load_submission_manifest(project_id)['doc_meta'][doc_id]['parse_attempts']
        self.assertEqual(['success', 'failed', 'success'], [a['content_status'] for a in history])

    def test_first_empty_parse_and_batch_failures_do_not_create_fake_outputs(self):
        self._add_project('p-batch')
        for doc_id in ('empty', 'good'):
            self._add_submission('p-batch', doc_id, created_at=datetime(2026, 1, 1))
        with mock.patch.object(self.service.material_service, 'parse_file', side_effect=[[], [{'text': 'actual body'}]]):
            ok, _, result = self.service.batch_parse_submissions('p-batch', {'doc_ids': ['empty', 'good']})
        self.assertTrue(ok)
        self.assertEqual('partial', result['content_status'])
        self.assertEqual('empty_result', result['failed'][0]['code'])
        self.assertFalse((self.service._project_path('p-batch') / 'parsed' / 'empty.json').exists())
        self.assertEqual('success', result['success'][0]['content_status'])

    def test_parser_exception_preserves_source_and_logs_only_safe_diagnostics(self):
        self._add_project('p-parser-error')
        source = self._add_submission('p-parser-error', 'a', created_at=datetime(2026, 1, 1))
        self.assertTrue(self.service.parse_submission('p-parser-error', 'a')[0])
        artifact = self.service._project_path('p-parser-error') / 'parsed' / 'a.json'
        previous = artifact.read_bytes()
        with self.assertLogs('agent.agent_backend.services.filing_parse_outcome', level='ERROR') as logs:
            with mock.patch.object(self.service.material_service, 'parse_file', side_effect=RuntimeError('SECRET_FULL_CONTENT /private/server')):
                ok, message, attempt = self.service.parse_submission('p-parser-error', 'a')
        self.assertFalse(ok)
        self.assertEqual('parser_failed', attempt['code'])
        self.assertEqual(previous, artifact.read_bytes())
        self.assertEqual('source-a', source.read_text())
        self.assertEqual('failed', self.service.list_submissions('p-parser-error', {})['list'][0]['latest_attempt']['content_status'])
        self.assertNotIn('SECRET_FULL_CONTENT', message + str(logs.output))
        self.assertNotIn('/private/server', message + str(logs.output))
        self.assertIn('p-parser-error', str(logs.output))
        self.assertIn('RuntimeError', str(logs.output))

    def test_review_blocks_selected_runtime_interruption_and_ignores_excluded(self):
        project_id = 'p-selected'
        self._add_project(project_id)
        for doc_id in ('selected', 'excluded'):
            self._add_submission(project_id, doc_id, created_at=datetime(2026, 1, 1))
        self.service._save_submission_manifest(project_id, {'doc_meta': {'selected': {'review_enabled': True}, 'excluded': {'review_enabled': False, 'latest_attempt': {'content_status': 'failed'}}}})
        store = RuntimeTaskStore(connection=self.connection, ensure_schema=False)
        store.create_task(task_id='interrupted', domain='filing_change_review', project_id=project_id, source_doc_id='selected', task_type='parse_submission', now=100)
        store.claim_task('interrupted', now=101)
        store.recover_stale_active_tasks(domain='filing_change_review', stale_after_seconds=10, now=120)
        self.service.rule_service = _RuleList()
        self.service.orchestrator = mock.Mock()
        self.service.orchestrator.run.return_value = _ReviewOrchestrator.run()
        self.service.report_service = _ReportCapture(self.root_dir)
        ok, message, run = self.service.prepare_review_run(project_id)
        self.assertFalse(ok, message)
        self.assertEqual(run['code'], 'parse_review_required')
        self.assertTrue(any(x.get('doc_id') == 'selected' for x in run['blocking_issues']))
        self.assertFalse(any(x.get('doc_id') == 'excluded' for x in run['blocking_issues']))
        self.service.orchestrator.run.assert_not_called()
        selected = next(x for x in self.service.list_submissions(project_id, {})['list'] if x['doc_id'] == 'selected')
        self.assertEqual('failed', selected['latest_attempt']['content_status'])
        self.assertEqual('interrupted', selected['latest_attempt']['task_id'])
        with self.connection.get_session() as session:
            self.assertEqual(session.query(FilingChangeReviewRun).count(), 0)

    def test_real_docx_import_then_empty_replacement_preserves_file_ownership(self):
        from docx import Document
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        self._add_project('p-real-form')
        self.service.form_parser = FilingChangeFormParserService()
        doc = Document()
        doc.add_paragraph('药品通用名称：米诺地尔；申请内容：拟将有效期延长至24个月。')
        stream = io.BytesIO()
        doc.save(stream)
        ok, _, original = self.service.import_application_form('p-real-form', _Upload('申请表.docx', stream.getvalue()))
        self.assertTrue(ok)
        self.assertIn('米诺地尔', original['raw_text'])
        empty = io.BytesIO()
        Document().save(empty)
        ok, _, failed = self.service.import_application_form('p-real-form', _Upload('替换.docx', empty.getvalue()))
        self.assertFalse(ok)
        self.assertEqual('empty_result', failed['code'])
        restored = self.service.get_application_form('p-real-form')[2]
        self.assertEqual(original['original_file_id'], restored['original_file_id'])
        self.assertEqual(original['raw_text'], restored['raw_text'])
        files = list((self.service._project_path('p-real-form') / 'application_form').glob('*.docx'))
        self.assertEqual([original['original_file_id'] + '.docx'], [p.name for p in files])
        read_bytes = Path.read_bytes
        def unreadable_source(path):
            if path == files[0]:
                raise PermissionError()
            return read_bytes(path)
        with mock.patch.object(Path, 'read_bytes', unreadable_source):
            ok, _, attempt = self.service.parse_application_form('p-real-form')
        self.assertFalse(ok)
        self.assertEqual('file_access', attempt['code'])
        self.assertEqual(original['original_file_id'], self.service.get_application_form('p-real-form')[2]['original_file_id'])
        ok, _, retry = self.service.parse_application_form('p-real-form')
        self.assertTrue(ok)
        # 同一原件重解析不再伪造一次文件替换，原件关联和字节保持不变。
        self.assertEqual(original['original_file_id'], retry['original_file_id'])
        self.assertEqual(original['raw_text'], retry['raw_text'])

    def test_interrupted_task_cannot_commit_late_parser_output(self):
        from agent.agent_backend.services.filing_parse_outcome import parse_task_id
        self._add_project('p-interrupt')
        self._add_submission('p-interrupt', 'a', created_at=datetime(2026, 1, 1))
        self.assertTrue(self.service.parse_submission('p-interrupt', 'a')[0])
        artifact = self.service._project_path('p-interrupt') / 'parsed' / 'a.json'
        previous = artifact.read_bytes()
        store = RuntimeTaskStore(connection=self.connection, ensure_schema=False)
        store.create_task(task_id='late', domain='filing_change_review', project_id='p-interrupt', source_doc_id='a', task_type='parse_submission', now=100)
        store.claim_task('late', now=101)
        def interrupted_parser(_):
            store.recover_stale_active_tasks(domain='filing_change_review', stale_after_seconds=10, now=120)
            return [{'text': 'late new result must not overwrite previous result'}]
        token = parse_task_id.set('late')
        try:
            with mock.patch.object(self.service.material_service, 'parse_file', side_effect=interrupted_parser):
                ok, _, attempt = self.service.parse_submission('p-interrupt', 'a')
        finally:
            parse_task_id.reset(token)
        self.assertFalse(ok)
        self.assertEqual('task_interrupted', attempt['code'])
        self.assertEqual(previous, artifact.read_bytes())
        self.assertEqual('failed', store.get_task('late')['status'])

    def test_batch_progress_records_each_file_in_existing_runtime_json(self):
        from agent.agent_backend.services.filing_parse_outcome import parse_task_id
        self._add_project('p-progress')
        for doc_id in ('a', 'b'):
            self._add_submission('p-progress', doc_id, created_at=datetime(2026, 1, 1))
        store = RuntimeTaskStore(connection=self.connection, ensure_schema=False)
        store.create_task(task_id='batch-progress', domain='filing_change_review', project_id='p-progress', task_type='parse_submissions_batch')
        store.claim_task('batch-progress')
        token = parse_task_id.set('batch-progress')
        try:
            ok, _, result = self.service.batch_parse_submissions('p-progress', {'all': True})
        finally:
            parse_task_id.reset(token)
        self.assertTrue(ok)
        task = store.get_task('batch-progress')
        self.assertEqual(['a', 'b'], task['payload']['doc_ids'])
        self.assertEqual([], task['result']['data']['pending'])
        self.assertEqual(result['success'], task['result']['data']['success'])
        self.assertEqual('success', task['result']['content_status'])

    def test_real_invalid_pdf_first_import_has_no_form_or_claimed_source(self):
        import fitz
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        self._add_project('p-invalid-form')
        self.service.form_parser = FilingChangeFormParserService()
        doc = fitz.open()
        doc.new_page().insert_text((50, 50), 'Protected form')
        encrypted = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='reader')
        doc.close()
        for content, code in [(b'', 'empty_file'), (b'%PDF-1.7 broken', 'pdf_damaged'), (encrypted, 'pdf_password')]:
            with self.subTest(code=code):
                ok, _, attempt = self.service.import_application_form('p-invalid-form', _Upload('申请表.pdf', content))
                self.assertFalse(ok)
                self.assertEqual(code, attempt['code'])
                refreshed = self.service.get_application_form('p-invalid-form')[2]
                # 空模板供纯人工填写，不包含识别值、人工标记或文件来源。
                from agent.agent_backend.services.filing_form_revision import values, present
                for field in refreshed['form_json'].values():
                    if isinstance(field, dict) and 'field_type' in field:
                        self.assertFalse(any(present(v) for k, v in values(field).items() if k != 'sub_fields.validity_unit'))
                        self.assertFalse(field.get('source_file'))
                        self.assertFalse(field.get('manual_modified'))
                self.assertEqual('not_parsed', refreshed['parse_status'])
                self.assertEqual(code, refreshed['latest_attempt']['code'])
                self.assertFalse(list((self.service._project_path('p-invalid-form') / 'application_form').glob('*.pdf')))

    def test_real_pdf_upload_and_parse_keeps_uploaded_source_and_saved_content(self):
        import fitz
        from agent.agent_backend.services.filing_change_material_service import FilingChangeMaterialService
        self._add_project('p-upload-pdf')
        self.service.material_service = FilingChangeMaterialService()
        doc = fitz.open()
        doc.new_page().insert_text((50, 50), 'Drug registration study\nReal material body and results for shelf life 24 months.')
        content = doc.tobytes()
        doc.close()
        ok, message, uploaded = self.service.upload_submission_files('p-upload-pdf', [_Upload('研究.pdf', content)], '5')
        self.assertTrue(ok, message)
        doc_id = uploaded['created'][0]['doc_id']
        ok, message, parsed = self.service.parse_submission('p-upload-pdf', doc_id)
        self.assertTrue(ok, message)
        self.assertEqual('success', parsed['content_status'])
        view = self.service.get_submission_parsed_markdown('p-upload-pdf', doc_id)[2]
        self.assertIn('shelf life 24 months', view['markdown'])
        session = self.connection.get_session()
        try:
            row = session.query(FilingChangeSubmissionFile).filter_by(doc_id=doc_id).one()
            self.assertEqual(content, Path(row.storage_path).read_bytes())
            self.assertEqual('p-upload-pdf', row.project_id)
            self.assertEqual('success', row.parse_status)
        finally:
            session.close()

    @unittest.skipUnless(os.getenv('T2_OCR_URL'), '需要实际 OCR HTTP 服务')
    def test_repository_application_pdf_import_with_real_ocr_and_refresh(self):
        from agent.agent_backend.config.settings import settings
        from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
        source = Path(__file__).resolve().parents[1] / 'task' / 'change_review' / '申请表模板.pdf'
        self._add_project('p-real-application-pdf')
        self.service.form_parser = FilingChangeFormParserService()
        content = source.read_bytes()
        with mock.patch.object(settings, 'ocr_service_url', os.environ['T2_OCR_URL']):
            ok, message, result = self.service.import_application_form('p-real-application-pdf', _Upload(source.name, content))
        if os.getenv('INTEGRITY_APPLICATION_EVIDENCE'):
            # 仅授权的本地验收输出，便于逐页核对；不打印申报正文。
            evidence = Path(os.environ['INTEGRITY_APPLICATION_EVIDENCE'])
            evidence.mkdir(parents=True, exist_ok=True)
            (evidence / 'application-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf8')
            manifest = self.service._load_submission_manifest('p-real-application-pdf')
            (evidence / 'application-pages.json').write_text(json.dumps(
                (manifest.get('application_form_source') or {}).get('parsed_pages', []),
                ensure_ascii=False, indent=2), encoding='utf8')
        self.assertTrue(ok, message)
        self.assertIn(result['content_status'], ('success', 'partial'))
        self.assertGreater(len(result['raw_text']), 100)
        self.assertTrue(result['parse_diagnostics']['available_pages'])
        refreshed = self.service.get_application_form('p-real-application-pdf')[2]
        self.assertEqual(result['raw_text'], refreshed['raw_text'])
        self.assertEqual(result['form_json'], refreshed['form_json'])
        stored = self.service._project_path('p-real-application-pdf') / 'application_form' / (refreshed['original_file_id'] + '.pdf')
        self.assertEqual(content, stored.read_bytes())

    def test_delete_submission_commit_failure_restores_files_manifest_and_row(self):
        project_id = "project-submission-delete-rollback"
        doc_id = "doc-delete-rollback"
        self._add_project(project_id)
        storage_path = self._add_submission(project_id, doc_id, created_at=datetime(2026, 1, 1))
        manifest = {"doc_meta": {doc_id: {"material_sub_category": "5.1"}}, "category_meta": {}}
        self.service._save_submission_manifest(project_id, manifest)
        manifest_path = self.service._submission_manifest_path(project_id)
        manifest_before = manifest_path.read_bytes()
        parsed_path = self.service._project_path(project_id) / "parsed" / f"{doc_id}.json"
        parsed_path.write_text("[]", encoding="utf-8")
        self.connection.fail_next_commit = True

        ok, message = self.service.delete_submission(project_id, doc_id)

        self.assertFalse(ok)
        self.assertIn("injected database commit failure", message)
        self.assertTrue(storage_path.exists())
        self.assertTrue(parsed_path.exists())
        self.assertEqual(manifest_path.read_bytes(), manifest_before)
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeSubmissionFile).filter_by(doc_id=doc_id).count(), 1)
        finally:
            session.close()
        trash_dir = self.root_dir / ".trash"
        self.assertFalse(trash_dir.exists() and any(trash_dir.iterdir()))

    def test_reference_upload_parse_and_delete_failures_restore_consistent_state(self):
        reference_dir = self.root_dir / "reference_materials"
        with mock.patch.object(self.service, "_save_reference_manifest", side_effect=OSError("reference manifest failed")):
            ok, message, _ = self.service.upload_reference_materials(
                [_Upload("failed-reference.txt", b"failed")],
                "guideline",
                "extend_validity_period",
                {},
            )
        self.assertFalse(ok)
        self.assertIn("reference manifest failed", message)
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeReferenceMaterial).count(), 0)
        finally:
            session.close()
        self.assertFalse(reference_dir.exists() and any(path.is_file() for path in reference_dir.iterdir()))

        ok, message, uploaded = self.service.upload_reference_materials(
            [_Upload("reference.txt", b"source")],
            "guideline",
            "extend_validity_period",
            {"drug_category": "化学药品"},
        )
        self.assertTrue(ok, message)
        doc_id = uploaded["created"][0]["doc_id"]
        session = self.connection.get_session()
        try:
            row = session.query(FilingChangeReferenceMaterial).filter_by(doc_id=doc_id).one()
            storage_path = Path(row.storage_path)
        finally:
            session.close()
        parsed_dir = reference_dir / "parsed"
        parsed_dir.mkdir(parents=True, exist_ok=True)
        parsed_json = parsed_dir / f"{doc_id}.json"
        parsed_markdown = parsed_dir / f"{doc_id}.md"
        parsed_json.write_bytes(b"old-reference-json")
        parsed_markdown.write_bytes(b"old-reference-markdown")
        manifest_path = self.service._reference_manifest_path()
        manifest_before_parse = manifest_path.read_bytes()
        self.connection.fail_next_commit = True

        ok, message, _ = self.service.parse_reference_material(doc_id)

        self.assertFalse(ok)
        self.assertIn("解析结果保存失败", message)
        self.assertEqual(parsed_json.read_bytes(), b"old-reference-json")
        self.assertEqual(parsed_markdown.read_bytes(), b"old-reference-markdown")
        latest_manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        attempt = latest_manifest['doc_meta'][doc_id].pop('latest_attempt')
        self.assertEqual(attempt['content_status'], 'failed')
        self.assertEqual(attempt['code'], 'persist_failed')
        self.assertEqual(latest_manifest, json.loads(manifest_before_parse))
        manifest_before_parse = manifest_path.read_bytes()
        session = self.connection.get_session()
        try:
            row = session.query(FilingChangeReferenceMaterial).filter_by(doc_id=doc_id).one()
            self.assertEqual(row.parse_status, "pending")
        finally:
            session.close()

        self.connection.fail_next_commit = True
        ok, message = self.service.delete_reference_material(doc_id)
        self.assertFalse(ok)
        self.assertIn("injected database commit failure", message)
        self.assertTrue(storage_path.exists())
        self.assertTrue(parsed_json.exists())
        self.assertEqual(manifest_path.read_bytes(), manifest_before_parse)
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeReferenceMaterial).filter_by(doc_id=doc_id).count(), 1)
        finally:
            session.close()

    def test_submission_and_reference_filters_apply_before_pagination(self):
        project_id = "project-pagination"
        self._add_project(project_id)
        base_time = datetime(2026, 1, 1)
        for index in range(3):
            self._add_submission(
                project_id,
                f"doc-page-{index}",
                created_at=base_time + timedelta(days=index),
                category="5",
            )
        submission_manifest = {
            "doc_meta": {
                "doc-page-0": {"material_category": "4", "material_sub_category": "target"},
                "doc-page-1": {"material_category": "5", "material_sub_category": "other"},
                "doc-page-2": {"material_category": "5", "material_sub_category": "other"},
            },
            "category_meta": {},
        }
        self.service._save_submission_manifest(project_id, submission_manifest)
        submission_page = self.service.list_submissions(
            project_id,
            {"page": 1, "page_size": 1, "material_category": "4", "material_sub_category": "target"},
        )
        self.assertEqual(submission_page["total"], 1)
        self.assertEqual([item["doc_id"] for item in submission_page["list"]], ["doc-page-0"])

        session = self.connection.get_session()
        try:
            for index in range(3):
                session.add(
                    FilingChangeReferenceMaterial(
                        doc_id=f"ref-page-{index}",
                        title=f"reference-{index}",
                        material_type="database-type",
                        task_type="extend_validity_period",
                        storage_path=str(self.root_dir / "reference_materials" / f"ref-page-{index}.txt"),
                        parse_status="pending",
                        index_status="pending",
                        created_at=base_time + timedelta(days=index),
                    )
                )
            session.commit()
        finally:
            session.close()
        reference_manifest = {
            "doc_meta": {
                "ref-page-0": {
                    "drug_category": "中药",
                    "material_type": "target-type",
                    "applicable_change_item": "target-change",
                    "enabled": False,
                },
                "ref-page-1": {"drug_category": "化学药品", "material_type": "other", "enabled": True},
                "ref-page-2": {"drug_category": "化学药品", "material_type": "other", "enabled": True},
            }
        }
        self.service._save_reference_manifest(reference_manifest)
        reference_page = self.service.list_reference_materials(
            {
                "page": 1,
                "page_size": 1,
                "task_type": "extend_validity_period",
                "drug_category": "中药",
                "material_type": "target-type",
                "applicable_change_item": "target-change",
                "enabled": "false",
            }
        )
        self.assertEqual(reference_page["total"], 1)
        self.assertEqual([item["doc_id"] for item in reference_page["list"]], ["ref-page-0"])

    def test_soft_deleted_project_submission_reads_return_no_legacy_data(self):
        project_id = "project-legacy-hidden"
        initial_doc_id = "legacy-initial-document"
        supplement_doc_id = "legacy-supplement-document"
        self._add_project(project_id, deleted=True)
        self._add_submission(project_id, initial_doc_id, created_at=datetime(2026, 1, 1), category="5")
        self._add_submission(project_id, supplement_doc_id, created_at=datetime(2026, 1, 2), category="4")
        parsed_dir = self.service._project_path(project_id) / "parsed"
        (parsed_dir / f"{initial_doc_id}.md").write_text("legacy initial", encoding="utf-8")
        (parsed_dir / f"{supplement_doc_id}.md").write_text("legacy supplement", encoding="utf-8")
        self.service._save_submission_manifest(
            project_id,
            {
                "doc_meta": {
                    initial_doc_id: {"material_sub_category": "legacy-category"},
                    supplement_doc_id: {"material_sub_category": "legacy-category"},
                },
                "category_meta": {
                    "legacy-category": {
                        "applicable_flag": False,
                        "not_applicable_reason": "legacy private reason",
                    }
                },
            },
        )
        session = self.connection.get_session()
        try:
            session.add(
                FilingChangeReviewRun(
                    run_id="legacy-running-review",
                    project_id=project_id,
                    task_type="extend_validity_period",
                    status="running",
                    started_at=datetime(2026, 1, 3),
                    result_json="{}",
                    error_message="",
                )
            )
            session.commit()
        finally:
            session.close()

        with mock.patch.object(
            self.service,
            "_load_submission_manifest",
            wraps=self.service._load_submission_manifest,
        ) as manifest_loader:
            catalog = self.service.get_submission_catalog(project_id)
            completeness = self.service.check_submission_completeness(project_id)
            batch = self.service.batch_parse_submissions(project_id, {"all": True})
            comparison = self.service.compare_submission_materials(
                project_id,
                {"initial_doc_id": initial_doc_id, "supplement_doc_id": supplement_doc_id},
            )
            submission_list = self.service.list_submissions(project_id, {})

        self.assertEqual(manifest_loader.call_count, 0)
        self.assertEqual(catalog, {"catalog": [], "category_meta": {}, "project_exists": False})
        self.assertEqual(
            completeness,
            {
                "uploaded_codes": [],
                "missing_items": [],
                "not_applicable_issues": [],
                "pass": False,
                "project_exists": False,
            },
        )
        self.assertEqual(batch, (False, "project not found", {"total": 0, "success": [], "failed": []}))
        self.assertEqual(comparison, (False, "project not found", None))
        self.assertEqual(submission_list["list"], [])
        serialized = json.dumps([catalog, completeness, batch, comparison, submission_list], ensure_ascii=False)
        self.assertNotIn(initial_doc_id, serialized)
        self.assertNotIn(supplement_doc_id, serialized)
        self.assertNotIn("legacy private reason", serialized)

        ok, message, changed = self.service.mark_review_run_interrupted(
            project_id,
            "legacy-running-review",
            "worker stopped",
        )
        self.assertFalse(ok)
        self.assertEqual(message, "project not found")
        self.assertFalse(changed["run_updated"])
        session = self.connection.get_session()
        try:
            self.assertEqual(
                session.query(FilingChangeReviewRun).filter_by(run_id="legacy-running-review").one().status,
                "running",
            )
        finally:
            session.close()

    def test_delete_ignores_empty_and_out_of_root_storage_paths(self):
        project_id = "project-safe-delete-paths"
        self._add_project(project_id)
        empty_doc_id = "submission-empty-path"
        outside_doc_id = "submission-outside-path"
        self._add_submission(project_id, empty_doc_id, created_at=datetime(2026, 1, 1))
        self._add_submission(project_id, outside_doc_id, created_at=datetime(2026, 1, 2))
        outside_file = self.base_dir / "must-remain-submission.txt"
        outside_file.write_bytes(b"outside submission")
        session = self.connection.get_session()
        try:
            session.query(FilingChangeSubmissionFile).filter_by(doc_id=empty_doc_id).one().storage_path = ""
            session.query(FilingChangeSubmissionFile).filter_by(doc_id=outside_doc_id).one().storage_path = str(outside_file)
            session.commit()
        finally:
            session.close()
        self.service._save_submission_manifest(
            project_id,
            {"doc_meta": {empty_doc_id: {}, outside_doc_id: {}}, "category_meta": {}},
        )

        working_dir = self.base_dir / "working-directory-must-remain"
        working_dir.mkdir()
        working_sentinel = working_dir / "sentinel.txt"
        working_sentinel.write_bytes(b"cwd sentinel")
        original_cwd = Path.cwd()
        try:
            os.chdir(working_dir)
            ok, message = self.service.delete_submission(project_id, empty_doc_id)
        finally:
            os.chdir(original_cwd)
        self.assertTrue(ok, message)
        self.assertTrue(working_dir.is_dir())
        self.assertEqual(working_sentinel.read_bytes(), b"cwd sentinel")

        ok, message = self.service.delete_submission(project_id, outside_doc_id)
        self.assertTrue(ok, message)
        self.assertEqual(outside_file.read_bytes(), b"outside submission")

        reference_outside = self.base_dir / "must-remain-reference.txt"
        reference_outside.write_bytes(b"outside reference")
        session = self.connection.get_session()
        try:
            for doc_id, storage_path in [
                ("reference-empty-path", ""),
                ("reference-outside-path", str(reference_outside)),
            ]:
                session.add(
                    FilingChangeReferenceMaterial(
                        doc_id=doc_id,
                        title=doc_id,
                        material_type="guideline",
                        task_type="extend_validity_period",
                        storage_path=storage_path,
                        parse_status="pending",
                        index_status="pending",
                        created_at=datetime(2026, 1, 3),
                    )
                )
            session.commit()
        finally:
            session.close()
        self.service._save_reference_manifest(
            {"doc_meta": {"reference-empty-path": {}, "reference-outside-path": {}}}
        )

        original_cwd = Path.cwd()
        try:
            os.chdir(working_dir)
            ok, message = self.service.delete_reference_material("reference-empty-path")
        finally:
            os.chdir(original_cwd)
        self.assertTrue(ok, message)
        self.assertEqual(working_sentinel.read_bytes(), b"cwd sentinel")
        ok, message = self.service.delete_reference_material("reference-outside-path")
        self.assertTrue(ok, message)
        self.assertEqual(reference_outside.read_bytes(), b"outside reference")

    def test_review_commit_failure_removes_first_generated_report_files(self):
        project_id = "project-review-report-rollback"
        run_id = "run-review-report-rollback"
        self._add_project(project_id)
        session = self.connection.get_session()
        try:
            project = session.query(FilingChangeProject).filter_by(project_id=project_id).one()
            project.review_status = "running"
            session.add(
                FilingChangeReviewRun(
                    run_id=run_id,
                    project_id=project_id,
                    task_type="extend_validity_period",
                    status="running",
                    started_at=datetime(2026, 1, 1),
                    result_json=json.dumps({'input_identity': [{'kind': 'synthetic-lock-fixture'}]}),
                    error_message="",
                )
            )
            session.commit()
        finally:
            session.close()
        self.service.rule_service = _RuleList()
        self.service.orchestrator = _ReviewOrchestrator()
        self.service.report_service = _ReportFileWriter(self.root_dir)
        self._assume_ready_for_lock_or_transaction_test()
        self.connection.fail_next_commit = True

        ok, message, _ = self.service.execute_review_run(project_id, run_id)

        self.assertFalse(ok)
        self.assertIn("injected database commit failure", message)
        report_dir = self.service._project_path(project_id) / "reports"
        for suffix in (".md", ".docx", ".txt"):
            self.assertFalse((report_dir / f"{run_id}{suffix}").exists())
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeReviewResult).filter_by(run_id=run_id).count(), 0)
            self.assertEqual(session.query(FilingChangeReviewReport).filter_by(run_id=run_id).count(), 0)
            self.assertEqual(session.query(FilingChangeReviewRun).filter_by(run_id=run_id).one().status, "failed")
            project = session.query(FilingChangeProject).filter_by(project_id=project_id).one()
            self.assertEqual(project.review_status, "failed")
            self.assertEqual(project.report_status, "not_generated")
        finally:
            session.close()

    def test_report_rebuild_commit_failure_restores_all_previous_files_and_row(self):
        project_id = "project-report-rebuild-rollback"
        run_id = "run-report-rebuild-rollback"
        self._add_project(project_id)
        self._add_run_and_result(project_id, run_id)
        report_dir = self.service._project_dir(project_id) / "reports"
        previous_content = {
            ".md": b"previous markdown",
            ".docx": b"previous word",
            ".txt": b"previous fallback",
        }
        for suffix, content in previous_content.items():
            (report_dir / f"{run_id}{suffix}").write_bytes(content)
        previous_updated_at = datetime(2026, 1, 2, 11, 0, 0)
        session = self.connection.get_session()
        try:
            session.add(
                FilingChangeReviewReport(
                    report_id="report-rebuild-rollback",
                    project_id=project_id,
                    run_id=run_id,
                    report_type="review_report",
                    report_content="previous database report",
                    report_file_path=str(report_dir / f"{run_id}.docx"),
                    created_at=previous_updated_at,
                    updated_at=previous_updated_at,
                )
            )
            session.commit()
        finally:
            session.close()
        self.service.report_service = _ReportFileWriter(self.root_dir)
        self.connection.fail_next_commit = True

        ok, message, _ = self.service.generate_report(run_id)

        self.assertFalse(ok)
        self.assertIn("injected database commit failure", message)
        for suffix, content in previous_content.items():
            self.assertEqual((report_dir / f"{run_id}{suffix}").read_bytes(), content)
        session = self.connection.get_session()
        try:
            report = session.query(FilingChangeReviewReport).filter_by(report_id="report-rebuild-rollback").one()
            self.assertEqual(report.report_content, "previous database report")
            self.assertEqual(report.report_file_path, str(report_dir / f"{run_id}.docx"))
            self.assertEqual(report.updated_at, previous_updated_at)
        finally:
            session.close()

    def test_report_uses_consistency_result_and_manual_confirmation_round_trips(self):
        project_id = "project-result"
        run_id = "run-result"
        self._add_project(project_id)
        self._add_run_and_result(project_id, run_id)
        report_capture = _ReportCapture(self.root_dir)
        self.service.report_service = report_capture

        ok, message, _ = self.service.generate_report(run_id)
        self.assertTrue(ok, message)
        self.assertEqual(report_capture.payload["consistency_check"], {"source": "consistency"})
        self.assertNotEqual(report_capture.payload["consistency_check"], {"source": "field"})

        ok, message, confirmed = self.service.manual_confirm_run(
            run_id,
            {"confirmed": True, "comment": "reviewed", "reviewer": "reviewer-a"},
        )
        self.assertTrue(ok, message)
        ok, message, result = self.service.get_run_result(run_id)
        self.assertTrue(ok, message)
        self.assertEqual(result["manual_confirmation"], confirmed["manual_confirmation"])
        self.assertTrue(result["manual_confirmation"]["confirmed"])

    def test_evidence_and_correction_reads_do_not_build_full_run_result(self):
        project_id = "project-read-subset"
        run_id = "run-read-subset"
        self._add_project(project_id)
        self._add_run_and_result(project_id, run_id)
        session = self.connection.get_session()
        try:
            row = session.query(FilingChangeReviewResult).filter_by(run_id=run_id).one()
            payload = json.loads(row.conclusion_json or "{}")
            payload.update(
                {
                    "matched_rules": [{"rule_id": "rule-1"}],
                    "module_evidence": {"stability": ["evidence-1"]},
                    "llm_calls": [{"model": "local"}],
                    "correction_notice_draft": {"title": "已持久化草稿"},
                }
            )
            row.conclusion_json = json.dumps(payload, ensure_ascii=False)
            row.evidence_json = json.dumps([{"source": "source-1"}], ensure_ascii=False)
            session.commit()
        finally:
            session.close()

        self.service.get_run_result = mock.Mock(
            side_effect=AssertionError("子集读取不得构造完整运行结果")
        )

        ok, message, evidence = self.service.get_run_evidence(run_id)
        self.assertTrue(ok, message)
        self.assertEqual(evidence["matched_rules"], [{"rule_id": "rule-1"}])
        self.assertEqual(evidence["evidence_refs"], [{"source": "source-1"}])
        ok, message, correction = self.service.generate_correction_notice(run_id)
        self.assertTrue(ok, message)
        self.assertEqual(correction["correction_notice_draft"]["title"], "已持久化草稿")
        self.service.get_run_result.assert_not_called()

    def test_project_list_query_count_does_not_grow_per_project(self):
        for index in range(12):
            self._add_project(f"project-list-{index:02d}")
        statements = []

        def capture_statement(_conn, _cursor, statement, _parameters, _context, _executemany):
            statements.append(statement)

        event.listen(self.connection.engine, "before_cursor_execute", capture_statement)
        try:
            result = self.service.list_projects({"page": 1, "page_size": 20})
        finally:
            event.remove(self.connection.engine, "before_cursor_execute", capture_statement)

        self.assertEqual(result["total"], 12)
        self.assertEqual(len(result["list"]), 12)
        # 总数、分页项目、材料数统计、最新申请表：查询数与项目数无关。
        self.assertLessEqual(len(statements), 5)

    def test_interrupted_run_compare_and_set_does_not_overwrite_completed_state(self):
        project_id = "project-interrupted"
        run_id = "run-interrupted"
        self._add_project(project_id)
        session = self.connection.get_session()
        try:
            project_row = session.query(FilingChangeProject).filter_by(project_id=project_id).one()
            project_row.review_status = "running"
            session.add(
                FilingChangeReviewRun(
                    run_id=run_id,
                    project_id=project_id,
                    task_type="extend_validity_period",
                    status="running",
                    started_at=datetime(2026, 1, 1),
                    result_json="{}",
                    error_message="",
                )
            )
            session.add(
                FilingChangeReviewRun(
                    run_id="run-interrupted-other",
                    project_id=project_id,
                    task_type="extend_validity_period",
                    status="running",
                    started_at=datetime(2026, 1, 1, 0, 1),
                    result_json="{}",
                    error_message="",
                )
            )
            session.commit()
        finally:
            session.close()

        ok, message, changed = self.service.mark_review_run_interrupted(project_id, run_id, "worker stopped")
        self.assertTrue(ok, message)
        self.assertTrue(changed["run_updated"])
        self.assertFalse(changed["project_updated"])
        session = self.connection.get_session()
        try:
            run_row = session.query(FilingChangeReviewRun).filter_by(run_id=run_id).one()
            project_row = session.query(FilingChangeProject).filter_by(project_id=project_id).one()
            self.assertEqual(run_row.status, "failed")
            self.assertEqual(run_row.error_message, "worker stopped")
            self.assertIsNotNone(run_row.finished_at)
            self.assertEqual(project_row.review_status, "running")
            self.assertEqual(
                session.query(FilingChangeReviewRun).filter_by(run_id="run-interrupted-other").one().status,
                "running",
            )
        finally:
            session.close()

        ok, message, changed = self.service.mark_review_run_interrupted(
            project_id,
            "run-interrupted-other",
            "other worker stopped",
        )
        self.assertTrue(ok, message)
        self.assertTrue(changed["run_updated"])
        self.assertTrue(changed["project_updated"])
        session = self.connection.get_session()
        try:
            run_row = session.query(FilingChangeReviewRun).filter_by(run_id=run_id).one()
            project_row = session.query(FilingChangeProject).filter_by(project_id=project_id).one()
            self.assertEqual(project_row.review_status, "failed")
            run_row.status = "completed"
            run_row.error_message = ""
            project_row.review_status = "completed"
            session.commit()
        finally:
            session.close()

        ok, message, changed = self.service.mark_review_run_interrupted(project_id, run_id, "late recovery")
        self.assertTrue(ok, message)
        self.assertFalse(changed["run_updated"])
        self.assertFalse(changed["project_updated"])
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeReviewRun).filter_by(run_id=run_id).one().status, "completed")
            self.assertEqual(session.query(FilingChangeProject).filter_by(project_id=project_id).one().review_status, "completed")
        finally:
            session.close()

    def test_soft_deleted_project_cannot_generate_or_export_report(self):
        project_id = "project-soft-deleted"
        run_id = "run-soft-deleted"
        self._add_project(project_id, deleted=True)
        self._add_run_and_result(project_id, run_id)
        report_capture = _ReportCapture(self.root_dir)
        self.service.report_service = report_capture
        report_path = self.root_dir / "projects" / project_id / "reports" / "existing.docx"
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_bytes(b"existing")
        session = self.connection.get_session()
        try:
            now = datetime(2026, 1, 1)
            session.add(
                FilingChangeReviewReport(
                    report_id="report-soft-deleted",
                    project_id=project_id,
                    run_id=run_id,
                    report_type="review_report",
                    report_content="existing",
                    report_file_path=str(report_path),
                    created_at=now,
                    updated_at=now,
                )
            )
            session.commit()
        finally:
            session.close()

        self.assertFalse(self.service.generate_report(run_id)[0])
        self.assertIsNone(report_capture.payload)
        self.assertFalse(self.service.export_report_word("report-soft-deleted")[0])

    def test_manifest_update_is_safe_across_processes(self):
        project_id = "project-process-lock"
        loops = 25
        self.service._save_submission_manifest(project_id, {"doc_meta": {}, "category_meta": {}, "counter": 0})
        available_methods = multiprocessing.get_all_start_methods()
        if "fork" not in available_methods:
            self.skipTest("process lock test requires fork start method")
        context = multiprocessing.get_context("fork")
        processes = [
            context.Process(target=_increment_manifest, args=(str(self.root_dir), project_id, loops))
            for _ in range(4)
        ]
        for process in processes:
            process.start()
        for process in processes:
            process.join(timeout=15)
            self.assertEqual(process.exitcode, 0)
        manifest = self.service._load_submission_manifest(project_id)
        self.assertEqual(manifest["counter"], loops * len(processes))

    def _assume_ready_for_lock_or_transaction_test(self):
        # 本夹具只隔离解析前置条件；并发、事务与删除仍使用真实数据库/锁。
        # 解析拒绝及零任务由独立readiness集成测试覆盖，不在此模拟通过验收。
        def ready(project_id):
            with self.connection.get_session() as session:
                exists = session.query(FilingChangeProject).filter_by(project_id=project_id, deleted=False).first() is not None
            return {'project_exists': exists, 'ready': exists, 'blocking_issues': [],
                    '_input_identity': [{'kind': 'synthetic-lock-fixture'}]}
        patched = mock.patch.object(self.service, 'get_parse_readiness', side_effect=ready)
        patched.start()
        self.addCleanup(patched.stop)

    def test_delete_waits_for_running_review_then_removes_late_outputs(self):
        project_id = "project-running-boundary"
        self._add_project(project_id)
        self._assume_ready_for_lock_or_transaction_test()
        ok, message, run = self.service.prepare_review_run(project_id)
        self.assertTrue(ok, message)
        started = threading.Event()
        release = threading.Event()
        review_result = []
        delete_result = []

        def _slow_review(_project_id: str, _run_id: str):
            report_path = self.service._project_dir(project_id) / "reports" / "late.txt"
            report_path.write_text("late", encoding="utf-8")
            started.set()
            release.wait(timeout=5)
            return True, "success", {}

        self.service._run_review_once = _slow_review
        review_thread = threading.Thread(
            target=lambda: review_result.append(self.service.execute_review_run(project_id, run['run_id']))
        )
        delete_thread = threading.Thread(target=lambda: delete_result.append(self.service.delete_project(project_id)))
        review_thread.start()
        self.assertTrue(started.wait(timeout=3))
        delete_thread.start()
        time.sleep(0.1)
        self.assertTrue(delete_thread.is_alive())
        release.set()
        review_thread.join(timeout=5)
        delete_thread.join(timeout=5)

        self.assertTrue(review_result[0][0])
        self.assertTrue(delete_result[0][0], delete_result[0][1])
        self.assertFalse(self.service._project_path(project_id).exists())
        del self.service._run_review_once
        self.assertFalse(self.service.execute_review_run(project_id, "run-boundary")[0])
        self.assertFalse(self.service._project_path(project_id).exists())

    def test_delete_cannot_split_review_run_and_runtime_task_creation(self):
        project_id = "project-task-creation-boundary"
        self._add_project(project_id)
        self._assume_ready_for_lock_or_transaction_test()
        app_service = object.__new__(FilingChangeReviewAppService)
        app_service.service = self.service
        task_store = RuntimeTaskStore(connection=self.connection, ensure_schema=False)
        create_entered = threading.Event()
        release_creation = threading.Event()
        worker_start_entered = threading.Event()
        release_worker_start = threading.Event()
        blocking_store = _BlockingRuntimeTaskStore(task_store, create_entered, release_creation)
        start_results = []
        start_errors = []
        delete_results = []

        def _start_request():
            try:
                start_results.append(filing_change_review_controller.start_review(project_id))
            except Exception as exc:
                start_errors.append(exc)

        def _blocking_worker_start(_task_id: str, _project_id: str, _run_id: str) -> bool:
            worker_start_entered.set()
            if not release_worker_start.wait(timeout=5):
                raise TimeoutError("worker start boundary was not released")
            return True

        with (
            mock.patch.object(filing_change_review_controller, "get_service", return_value=app_service),
            mock.patch.object(filing_change_review_controller, "get_runtime_task_store", return_value=blocking_store),
            mock.patch.object(filing_change_review_controller, "_prune_tasks", return_value=None),
            mock.patch.object(filing_change_review_controller, "_run_task_async", side_effect=_blocking_worker_start),
        ):
            start_thread = threading.Thread(target=_start_request)
            delete_thread = threading.Thread(
                target=lambda: delete_results.append(self.service.delete_project(project_id))
            )
            start_thread.start()
            self.assertTrue(create_entered.wait(timeout=3))
            delete_thread.start()
            time.sleep(0.1)
            self.assertTrue(delete_thread.is_alive())
            release_creation.set()
            self.assertTrue(worker_start_entered.wait(timeout=3))
            self.assertTrue(delete_thread.is_alive())
            release_worker_start.set()
            start_thread.join(timeout=5)
            delete_thread.join(timeout=5)
            late_start_result = filing_change_review_controller.start_review(project_id)

        self.assertEqual(start_errors, [])
        self.assertEqual(len(start_results), 1)
        self.assertEqual(json.loads(start_results[0])["code"], 200)
        self.assertEqual(len(delete_results), 1)
        self.assertTrue(delete_results[0][0], delete_results[0][1])
        self.assertIsInstance(late_start_result, tuple)
        self.assertEqual(late_start_result[1], 400)
        session = self.connection.get_session()
        try:
            self.assertEqual(session.query(FilingChangeProject).filter_by(project_id=project_id).count(), 0)
            self.assertEqual(session.query(FilingChangeReviewRun).filter_by(project_id=project_id).count(), 0)
            self.assertEqual(
                session.query(RuntimeTask).filter_by(domain="filing_change_review", project_id=project_id).count(),
                0,
            )
        finally:
            session.close()

    def test_application_form_staging_round_trip_and_cleanup(self):
        project_id = "project-application-staging"
        self._add_project(project_id)
        self.service.form_parser = _FormParser()

        ok, message, staged = self.service.stage_application_form_import(
            project_id,
            _Upload("safe-name.docx", b"word-content"),
        )
        self.assertTrue(ok, message)
        self.assertEqual("safe-name.docx", staged["original_file_name"])
        self.assertEqual(len(b"word-content"), staged["file_size"])
        self.assertNotIn("content", staged)
        staging_path = (
            self.service._project_path(project_id)
            / ".staging"
            / "application_form"
            / staged["staging_file"]
        )
        self.assertTrue(staging_path.is_file())

        ok, message, imported = self.service.import_staged_application_form(
            project_id,
            staged["staging_file"],
            staged["original_file_name"],
        )
        self.assertTrue(ok, message)
        # 此staging夹具仅返回自由文本和非字段字典，保存成功不代表字段识别成功。
        self.assertEqual("partial", imported["parse_status"])
        self.assertEqual("form text", imported["raw_text"])
        self.service.cleanup_application_form_staging(project_id, staged["staging_file"])
        self.assertFalse(staging_path.exists())

    def test_application_form_staging_rejects_invalid_empty_and_oversized_files(self):
        project_id = "project-application-staging-invalid"
        self._add_project(project_id)
        cases = [
            _Upload("../unsafe.exe", b"payload"),
            _Upload("empty.pdf", b""),
        ]
        for upload in cases:
            with self.subTest(file_name=upload.filename):
                ok, _, data = self.service.stage_application_form_import(project_id, upload)
                self.assertFalse(ok)
                if upload.filename == 'empty.pdf':
                    self.assertEqual('empty_file', data['code'])
                    self.assertEqual('failed', self.service.get_application_form(project_id)[2]['latest_attempt']['content_status'])
                else:
                    self.assertIsNone(data)

        with mock.patch.dict(os.environ, {"MAX_CONTENT_LENGTH_MB": "1"}):
            ok, message, data = self.service.stage_application_form_import(
                project_id,
                _Upload("large.pdf", b"x" * (1024 * 1024 + 1)),
            )
        self.assertFalse(ok)
        self.assertIn("上限", message)
        self.assertEqual('failed', data['content_status'])
        staging_dir = self.service._project_path(project_id) / ".staging" / "application_form"
        self.assertFalse(staging_dir.exists())

        external = self.base_dir / "external.pdf"
        external.write_bytes(b"must remain")
        self.service.cleanup_application_form_staging(project_id, str(external))
        self.assertEqual(b"must remain", external.read_bytes())


if __name__ == "__main__":
    unittest.main()
