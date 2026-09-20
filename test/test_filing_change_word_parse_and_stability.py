import json
import sys
import unittest
from datetime import datetime
from pathlib import Path

from docx import Document

ROOT_DIR = Path(__file__).resolve().parents[1]
PACKAGE_PARENT = ROOT_DIR.parent
if str(PACKAGE_PARENT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_PARENT))

from agent.agent_backend.utils.parser.docx_parser import parse_docx
from agent.agent_backend.utils.parser.docx_markdown_parser import parse_docx_to_markdown_json
from agent.agent_backend.services.filing_change_stability_service import FilingChangeStabilityService


class FilingChangeWordParseAndStabilityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        run_tag = datetime.now().strftime("%Y%m%d_%H%M%S")
        cls.output_dir = Path("test") / "filing_change_word_demo" / run_tag
        cls.output_dir.mkdir(parents=True, exist_ok=True)

        cls.docx_path = cls.output_dir / "备案变更_稳定性研究测试文档.docx"
        cls.payload_path = cls.output_dir / "备案变更_稳定性研究测试文档.parsed.json"
        cls.chunk_path = cls.output_dir / "备案变更_稳定性研究测试文档.chunks.json"
        cls.markdown_path = cls.output_dir / "备案变更_稳定性研究测试文档.parsed.md"
        cls.stability_path = cls.output_dir / "备案变更_稳定性研究测试文档.stability.json"

        cls._build_demo_docx(cls.docx_path)

    @staticmethod
    def _build_demo_docx(docx_path: Path) -> None:
        doc = Document()
        doc.add_heading("延长药品有效期稳定性研究资料", level=1)
        doc.add_paragraph("本文件用于测试备案变更类审评中的 Word 标题解析、表格解析和稳定性趋势分析。")

        doc.add_heading("一、变更研究概述", level=2)
        doc.add_paragraph("本次申请拟将药品有效期由24个月延长至36个月，贮藏条件保持阴凉处保存。")

        doc.add_heading("（一）稳定性研究", level=3)
        doc.add_paragraph("对三批样品在长期试验条件下进行考察，记录关键质量指标在不同时间点的变化情况。")

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["考察时间（月）", "含量（%）", "有关物质（%）", "水分（%）"]
        for idx, value in enumerate(headers):
            table.rows[0].cells[idx].text = value

        data_rows = [
            ["限度", "90-110", "≤1.0", "≤5.0"],
            ["0", "99.8", "0.20", "2.1"],
            ["3", "99.5", "0.25", "2.4"],
            ["6", "99.1", "0.31", "2.8"],
            ["9", "98.9", "0.36", "3.0"],
            ["12", "98.5", "0.42", "3.2"],
            ["18", "98.1", "0.55", "3.5"],
            ["24", "97.8", "0.68", "3.8"],
            ["36", "97.2", "0.82", "4.1"],
        ]
        for row_values in data_rows:
            row = table.add_row()
            for idx, value in enumerate(row_values):
                row.cells[idx].text = value

        doc.add_paragraph("结论：在36个月考察期内，各项指标未见明显异常变化。")
        doc.save(docx_path)

    @staticmethod
    def _collect_markdown_tables(payload: dict) -> list:
        tables = []

        def walk_sections(sections: list, path: list) -> None:
            for section in sections or []:
                if not isinstance(section, dict):
                    continue
                title = str(section.get("section_title", "") or "").strip()
                current_path = path + ([title] if title else [])
                for table in section.get("tables", []) or []:
                    if not isinstance(table, dict):
                        continue
                    markdown = str(table.get("markdown", "") or "").strip()
                    columns = table.get("columns", []) or []
                    data = table.get("data", []) or []
                    if markdown and columns and data:
                        tables.append(
                            {
                                "section_title": " > ".join(current_path),
                                "headers": columns,
                                "rows": data,
                                "markdown": markdown,
                            }
                        )
                walk_sections(section.get("children_sections", []) or [], current_path)

        walk_sections(payload.get("sections", []) or [], [])
        return tables

    @staticmethod
    def _render_sections_to_markdown(sections: list) -> str:
        blocks = []

        def walk(items: list, level: int = 1) -> None:
            for section in items or []:
                if not isinstance(section, dict):
                    continue
                title = str(section.get("section_title", "") or "").strip()
                content = str(section.get("content", "") or "").strip()
                if title:
                    blocks.append(f"{'#' * max(1, min(level, 6))} {title}")
                if content:
                    blocks.append(content)
                walk(section.get("children_sections", []) or [], level + 1)

        walk(sections or [], 1)
        return "\n\n".join([item for item in blocks if item]).strip() + "\n"

    def test_generate_demo_docx_and_run_parse(self) -> None:
        payload = parse_docx_to_markdown_json(
            input_path=self.docx_path,
            title=None,
            embed_images=False,
            skip_toc=True,
            filter_header_footer=True,
            merge_continuous_tables=True,
        )
        chunks = parse_docx(str(self.docx_path))
        markdown_text = self._render_sections_to_markdown(payload.get("sections", []) or [])

        self.payload_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        self.chunk_path.write_text(json.dumps(chunks, ensure_ascii=False, indent=2), encoding="utf-8")
        self.markdown_path.write_text(markdown_text, encoding="utf-8")

        self.assertTrue(payload.get("sections"))
        self.assertIn("延长药品有效期稳定性研究资料", json.dumps(payload, ensure_ascii=False))
        self.assertIn("一、变更研究概述", json.dumps(payload, ensure_ascii=False))
        self.assertTrue(
            "（一）稳定性研究" in json.dumps(payload, ensure_ascii=False)
            or "(一)稳定性研究" in json.dumps(payload, ensure_ascii=False)
        )

        markdown_tables = self._collect_markdown_tables(payload)
        self.assertTrue(markdown_tables, "应至少解析出一张 Markdown 表格")
        self.assertTrue(
            "考察时间（月）" in markdown_tables[0].get("markdown", "")
            or "考察时间(月)" in markdown_tables[0].get("markdown", "")
        )
        self.assertTrue(chunks, "应生成结构化 chunk")
        self.assertTrue(
            any(
                ("考察时间（月）" in str(item.get("text", "")) or "考察时间(月)" in str(item.get("text", "")))
                for item in chunks
            )
        )
        self.assertIn("# 延长药品有效期稳定性研究资料", markdown_text)
        self.assertIn("## 一、变更研究概述", markdown_text)
        self.assertTrue("### （一）稳定性研究" in markdown_text or "### (一)稳定性研究" in markdown_text)
        self.assertIn("| 考察时间", markdown_text)

    def test_run_stability_analysis_on_generated_docx(self) -> None:
        payload = parse_docx_to_markdown_json(
            input_path=self.docx_path,
            title=None,
            embed_images=False,
            skip_toc=True,
            filter_header_footer=True,
            merge_continuous_tables=True,
        )
        markdown_tables = self._collect_markdown_tables(payload)
        service = FilingChangeStabilityService()
        result = service.run(
            application_form={
                "form_json": {
                    "item_15_validity_period": {
                        "sub_fields": {
                            "proposed_validity_period": "36个月"
                        }
                    }
                }
            },
            submissions=[
                {
                    "doc_id": "demo_doc",
                    "file_name": self.docx_path.name,
                    "extracted_json": {
                        "markdown_tables": markdown_tables,
                    },
                }
            ],
            project_id="demo_project",
        )

        self.stability_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

        self.assertEqual(result.get("result"), "需人工确认")
        self.assertFalse(result["coverage"]["requirements_known"], "示例表没有逐批方案，不能直接宣称支持延长")
        self.assertEqual(result.get("proposed_validity_period"), "36个月")
        self.assertGreaterEqual(float(result.get("coverage_month", 0) or 0), 36.0)
        self.assertTrue(result.get("key_indicators"))
        self.assertTrue(result.get("charts"))
        self.assertTrue(any(item.get("indicator") == "含量(%)" for item in result.get("key_indicators", [])))


if __name__ == "__main__":
    unittest.main()
