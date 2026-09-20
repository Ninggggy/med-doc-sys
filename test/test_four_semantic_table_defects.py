"""输入事实独立写定，正常路径为真实 Word/PDF/OCR 与隔离数据库。"""
import io
import os
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from agent.test import test_five_parser_defects as previous
from agent.test.test_filing_change_review_service_regression import _Upload
from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
from agent.agent_backend.services.filing_form_revision import validity_cross_checks
from agent.agent_backend.utils.parser.form_field_semantics import selected_option, validity_values
from agent.agent_backend.utils.parser.pdf_page_extractor import extract_pdf_pages
from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import OCRServiceClient
from agent.agent_backend.services.filing_parse_outcome import outcome, ParseFailure


def word_file(path, rows):
    doc = Document()
    for row in rows:
        table = doc.add_table(rows=1, cols=len(row))
        for cell, text in zip(table.rows[0].cells, row): cell.text = text
    doc.save(path)


def pdf_file(path, rows):
    with fitz.open() as doc:
        page = doc.new_page(width=950, height=1100)
        # 原生字符及符号有真实位置，避免用图片模拟正常 PDF 勾选。
        font = os.getenv('FILING_TEST_FONT', '/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
        page.insert_font(fontname='formfont', fontfile=font)
        for i, row in enumerate(rows):
            x = 35
            for text in row:
                page.insert_text((x, 45+i*60), text, fontname='formfont', fontsize=12)
                x += max(200, fitz.get_text_length(text, fontname='china-s', fontsize=12)+15)
        doc.save(path)


def borderless_file(path, *, uncertain=False, ordinary=False):
    with fitz.open() as source, fitz.open() as doc:
        page = source.new_page(width=600, height=400)
        if ordinary:
            for i in range(5):
                prefix = f'{i+1}. ' if ordinary == 'list' else ''
                page.insert_text((35, 55+i*30), prefix+f'Paragraph {i+1} describes the study in words.', fontsize=12)
                if ordinary == 'margin_list':
                    page.insert_text((12, 55+i*30), f'{i+1}.', fontsize=12)
                elif ordinary not in ('paragraph', 'list'):
                    page.insert_text((330, 55+i*30), 'Another column of narrative.', fontsize=12)
        else:
            page.insert_text((40, 40), 'Table 1. Assay results', fontsize=14)
            for x, text in [(40, 'Sample'), (240, 'Assay (mg)'), (430, 'Recovery (%)')]:
                page.insert_text((x, 80), text, fontsize=12)
            for y, name, assay, recovery in [(120, 'Alpha', '-1.25', '95%'),
                                             (170, 'Batch', '', '98.0%'),
                                             (220, 'Gamma', '50.0', '50.0')]:
                page.insert_text((40, y), name, fontsize=12)
                if assay: page.insert_text((240 if not uncertain else 190+y/3, y), assay, fontsize=12)
                page.insert_text((430 if not uncertain else 280+y/4, y), recovery, fontsize=12)
            page.insert_text((40, 186), 'two', fontsize=12)
            page.insert_text((40, 260), '* Note: values in mg; retention 0.5 min.', fontsize=11)
        scan = doc.new_page(width=600, height=400)
        scan.insert_image(scan.rect, stream=page.get_pixmap(matrix=fitz.Matrix(3, 3)).tobytes('png'))
        doc.save(path)


class SemanticTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def parse(self, rows, ext):
        path = self.root / ('form.'+ext)
        (word_file if ext == 'docx' else pdf_file)(path, rows)
        return FilingChangeFormParserService().parse_form_file(str(path))['form_json']

    def test_validity_object_in_real_word_and_pdf(self):
        for ext in ['docx', 'pdf']:
            for text, expected in [
                ('本品有效期保持24个月，稳定性考察由6个月延长至12个月', {}),
                ('稳定性考察由6个月延长至12个月；药品有效期由18个月延长至24个月', {'original_validity_period':'18个月','proposed_validity_period':'24个月'}),
                ('药品有效期由18个月延长至24个月，稳定性考察由6个月延长至12个月', {'original_validity_period':'18个月','proposed_validity_period':'24个月'}),
                ('证书有效期由6个月延长至12个月；检验周期由3个月调整为6个月；截止日期2028年8月', {}),
                ('药品有效期保持24个月，加速试验由6个月延长至12个月', {}),
            ]:
                with self.subTest(ext=ext, text=text):
                    form = self.parse([['药品有效期', '24个月'], ['补充申请的内容', text]], ext)
                    f = form['item_15_validity_period']
                    self.assertEqual(f.get('reported_value'), '24个月')
                    for key in ['original_validity_period','proposed_validity_period']:
                        self.assertEqual(f['sub_fields'][key], expected.get(key,''))

    def test_validity_titles_columns_and_conflicts(self):
        form = self.parse([['药品有效期','原有效期','18个月'], ['药品有效期','拟延长后有效期','24个月']], 'docx')
        self.assertEqual(form['item_15_validity_period']['sub_fields']['original_validity_period'], '18个月')
        self.assertEqual(form['item_15_validity_period']['sub_fields']['proposed_validity_period'], '24个月')
        form = self.parse([['药品有效期', '原批准18个月\n拟变更为24个月']], 'docx')
        self.assertEqual(form['item_15_validity_period']['sub_fields']['proposed_validity_period'], '24个月')
        form = self.parse([['药品有效期', '24个月'], ['补充申请的内容', '药品有效期由18个月延长至24个月；药品有效期由18个月延长至36个月']], 'docx')
        f = form['item_15_validity_period']
        self.assertEqual(f['sub_fields']['proposed_validity_period'], '')
        self.assertTrue(f.get('semantic_issues'))

    def test_validity_cross_checks_keep_object_and_source(self):
        form = FilingChangeFormParserService()._build_full_form('')
        form['item_15_validity_period']['sub_fields']['original_validity_period'] = '18个月'
        sources = [
            ('一致.pdf','药品有效期：18个月'), ('冲突.pdf','药品有效期：24个月'),
            ('证书.pdf','证书有效期：12个月'), ('未知.docx','由6个月延长至12个月'),
        ]
        result = validity_cross_checks(form, [{'text':text,'source_file':name,'is_approval':True,'source_regions':[{'page':2}]} for name,text in sources])
        checks = result['item_15_validity_period']['cross_checks']
        self.assertFalse(any(c.get('value') == '12个月' for c in checks))
        self.assertTrue(any(c['source_file']=='一致.pdf' and c['status']=='一致' for c in checks))
        self.assertTrue(any(c['source_file']=='冲突.pdf' and '冲突' in c['status'] for c in checks))
        self.assertTrue(any(c['source_file']=='未知.docx' and '待核对' in c['status'] for c in checks))
        self.assertTrue(all(c['source_regions']==[{'page':2}] for c in checks))

    def test_option_marker_orientation_and_conflict(self):
        for opts in [['处方药','非处方药'],['是','否'],['使用','不使用'],['已上市','未上市']]:
            a,b=opts
            for text, expected in [(a+'☑ '+b+'□',a), (a+'□ '+b+'☑',b),
                                    ('☑'+a+' □'+b,a),('□'+a+' ☑'+b,b),
                                    (a+'☑\n'+b+'☑',''), (a+'□ '+b+'□',''),
                                    (a+'／'+b,''),(a,a),(a+' ☑ '+b,'')]:
                with self.subTest(text=text): self.assertEqual(selected_option(text,opts),expected)
            self.assertEqual(selected_option(a+'☑\n□'+b, opts), a)

    def test_option_word_pdf_entry_and_cell_boundaries(self):
        for ext in ['docx','pdf']:
            for label,key,options in [('是否为OTC','item_3_otc_type',['处方药','非处方药']),
                                      ('商品名称','item_10_trade_name',['使用','不使用']),
                                      ('受理前药品注册检验','item_19_pre_acceptance_inspection',['是','否'])]:
                for postfix in [False,True]:
                    for chosen in [0,1]:
                        cells = [(v+('☑' if i==chosen else '□')) if postfix else (('☑' if i==chosen else '□')+v) for i,v in enumerate(options)]
                        with self.subTest(ext=ext,cells=cells):
                            form=self.parse([[label,cells[0],'',cells[1]]],ext)
                            self.assertEqual(form[key]['value'],options[chosen])
        form=self.parse([['商品名称','使用☑','不使用□'],['受理前药品注册检验','是☑','否□']], 'docx')
        self.assertEqual(form['item_10_trade_name']['value'],'使用')
        self.assertEqual(form['item_19_pre_acceptance_inspection']['value'],'是')

    def test_unconfirmed_options_keep_evidence_in_both_entries(self):
        for ext in ['docx', 'pdf']:
            for text in ['处方药☑ 非处方药☑', '处方药□ 非处方药□', '处方药 非处方药', '处方药 ☑ 非处方药']:
                with self.subTest(ext=ext, text=text):
                    rows = [['是否为OTC', text]]
                    template = text in ('处方药□ 非处方药□', '处方药 非处方药')
                    if template:
                        with self.assertRaises(ParseFailure) as caught:
                            self.parse(rows, ext)
                        self.assertEqual(caught.exception.code, 'empty_form')
                        rows = [['药品通用名称', '测试药品']] + rows
                    field = self.parse(rows, ext)['item_3_otc_type']
                    self.assertEqual(field['value'], '')
                    if template:
                        self.assertTrue(all(e['selected'] is not True for e in field['selection_evidence']))
                    else:
                        self.assertTrue(field.get('semantic_issues'))
                    self.assertTrue(field['source_regions'])
        field = self.parse([['是否为OTC', '处方药☑'], ['是否为OTC', '非处方药☑']], 'docx')['item_3_otc_type']
        self.assertEqual(field['value'], '')
        self.assertTrue(field['semantic_issues'])

    def test_company_subfields_and_role_boundaries(self):
        rows=[['药品注册申请人','中文名称','甲公司'],['药品注册申请人','法定代表人','张三'],
              ['药品注册申请人','统一社会信用代码及组织机构代码','A123456789'],
              ['制剂生产企业','中文名称','乙公司'],['制剂生产企业','法定代表人','李四'],
              ['制剂生产企业','统一社会信用代码／组织机构代码：','B987654321'],
              ['委托研究机构','研究负责人','王五']]
        form=self.parse(rows,'docx')
        for key, expected in [('item_30_applicant_info',{'applicant_name':'甲公司','legal_representative':'张三','contact':'','credit_or_org_code':'A123456789','license_no':''}),
                              ('item_31_manufacturer_info',{'manufacturer_name':'乙公司','legal_representative':'李四','contact':'','credit_or_org_code':'B987654321','license_no':''}),
                              ('item_32_cro_info',{'research_lead':'王五','contact':''})]:
            for k,v in expected.items(): self.assertEqual(form[key]['sub_fields'][k],v)
        form=self.parse([['药品注册申请人','法定\n代表人：','张三','制剂生产企业','法定代表人','','统一社会信用代码','B123']], 'docx')
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['legal_representative'],'张三')
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['contact'],'')
        self.assertEqual(form['item_31_manufacturer_info']['sub_fields']['contact'],'')
        self.assertEqual(form['item_31_manufacturer_info']['sub_fields']['credit_code'],'B123')
        self.assertEqual(form['item_31_manufacturer_info']['sub_fields']['license_no'],'')

    def test_company_merges_nested_supplier_and_all_roles(self):
        path = self.root/'merged.docx'
        doc = Document()
        table = doc.add_table(rows=3, cols=5)
        for row, values in zip(table.rows, [
                ['药品注册申请人', '中文名称', '负责制剂生产企业技术服务的甲公司', '法定代表人', '张三'],
                ['', '统一社会信用代码', 'A13579', '', ''],
                ['制剂生产企业', '中文名称', '乙公司', '法定代表人', '']]):
            for cell, value in zip(row.cells, values): cell.text = value
        table.cell(0,0).merge(table.cell(1,0))
        table.cell(1,2).merge(table.cell(1,4))
        cro = doc.add_table(rows=1, cols=3)
        for cell,value in zip(cro.rows[0].cells, ['委托研究机构', '中文名称', '丙研究公司']): cell.text=value
        material = doc.add_table(rows=1, cols=2)
        material.cell(0,0).text='原辅包来源'
        nested = material.cell(0,1).add_table(rows=2, cols=3)
        for row,values in zip(nested.rows, [['物料名称','登记号','生产企业'],['测试辅料','Y12345','丁供应商公司']]):
            for cell,value in zip(row.cells,values): cell.text=value
        doc.save(path)
        form = FilingChangeFormParserService().parse_form_file(str(path))['form_json']
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['applicant_name'], '负责制剂生产企业技术服务的甲公司')
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['legal_representative'], '张三')
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['contact'], '')
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['credit_code'], 'A13579')
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['license_no'], '')
        self.assertEqual(form['item_31_manufacturer_info']['sub_fields']['manufacturer_name'], '乙公司')
        self.assertEqual(form['item_31_manufacturer_info']['sub_fields']['contact'], '')
        self.assertEqual(form['item_32_cro_info']['sub_fields']['organization_name'], '丙研究公司')
        self.assertEqual(form['item_17_material_source']['table_rows'][0]['manufacturer'], '丁供应商公司')


class BorderlessTests(unittest.TestCase):
    setUp=SemanticTests.setUp

    @unittest.skipUnless(os.getenv('T2_OCR_URL'),'需要真实 OCR HTTP 服务')
    def test_real_borderless_cells_empty_middle_and_wrap(self):
        path=self.root/'table.pdf';borderless_file(path)
        p=extract_pdf_pages(str(path),ocr_client=OCRServiceClient(os.environ['T2_OCR_URL']),lang='eng')[0]
        self.assertEqual(len(p['tables']),1)
        t=p['tables'][0]
        self.assertEqual(t['rows'],[['Sample','Assay (mg)','Recovery (%)'],['Alpha','-1.25','95%'],['Batch\ntwo','','98.0%'],['Gamma','50.0','50.0']])
        self.assertEqual(len(t['cells']),12)
        self.assertEqual(t['boundary_source'],'inferred_from_text_positions')
        self.assertTrue(t['cells'][4]['word_bboxes_pdf'])
        self.assertEqual(p['text'].count(t['markdown']),1)
        self.assertIn('retention 0.5 min',p['text'])

    @unittest.skipUnless(os.getenv('T2_OCR_URL'),'需要真实 OCR HTTP 服务')
    def test_real_ordinary_columns_are_not_tables(self):
        for layout in [True, 'paragraph', 'list', 'margin_list']:
            with self.subTest(layout=layout):
                path=self.root/'body.pdf';borderless_file(path,ordinary=layout)
                p=extract_pdf_pages(str(path),ocr_client=OCRServiceClient(os.environ['T2_OCR_URL']),lang='eng')[0]
                self.assertEqual(p['tables'],[])
                # 真实 OCR 会漏掉远离正文的边栏序号；覆盖诊断必须保留，不能当作表格缺失。
                self.assertEqual(p['status'], 'partial' if layout == 'margin_list' else 'success')
                if layout == 'margin_list':
                    self.assertTrue(all(e['stage']=='ocr_coverage' for e in p['errors']))
                self.assertFalse(any(x.get('needs_review') for x in p['ocr_regions']))

    @unittest.skipUnless(os.getenv('T2_OCR_URL'),'需要真实 OCR HTTP 服务')
    def test_real_uncertain_table_diagnostic(self):
        path=self.root/'uncertain.pdf';borderless_file(path,uncertain=True)
        p=extract_pdf_pages(str(path),ocr_client=OCRServiceClient(os.environ['T2_OCR_URL']),lang='eng')[0]
        self.assertIn('Alpha',p['text'])
        self.assertEqual(p['status'],'partial')
        result=outcome([p])
        self.assertTrue(any('文字已识别，表格结构未恢复' in e['message'] for e in result['parse_diagnostics']['errors']))

    @unittest.skipUnless(os.getenv('T2_OCR_URL'), '需要真实 OCR HTTP 服务')
    def test_borderless_ctd_and_knowledge_public_read_paths(self):
        from agent.agent_backend.config.settings import settings
        from agent.agent_backend.utils.parser.ctd_paser import parse_bound_section_pdf_to_payload
        from agent.agent_backend.memory.rag.document import DocumentProcessor
        path = self.root/'table.pdf'; borderless_file(path)
        with patch.object(settings, 'ocr_service_url', os.environ['T2_OCR_URL']), patch.object(settings, 'ocr_lang', 'eng'):
            ctd = parse_bound_section_pdf_to_payload(str(path), section_id='3.2.p.5.1')
            self.assertEqual(ctd['sections'][0]['tables'][0]['rows'][2], ['Batch\ntwo', '', '98.0%'])
            processor = DocumentProcessor()
            chunks = processor.chunk_document(processor.parse_to_document(str(path)))
            self.assertEqual(chunks[0].metadata['tables'][0]['rows'][3], ['Gamma', '50.0', '50.0'])
            self.assertEqual(chunks[0].page_start, 1)
            self.assertIn('retention 0.5 min', chunks[0].text)

    def test_borderless_ocr_timeout_empty_and_failure(self):
        path = self.root/'scan.pdf'; borderless_file(path)
        for response in [TimeoutError('injected'), {}, RuntimeError('injected')]:
            with self.subTest(response=response), patch.object(OCRServiceClient, 'image_to_data', **(
                    {'side_effect':response} if isinstance(response, Exception) else {'return_value':response})):
                p = extract_pdf_pages(str(path), ocr_client=OCRServiceClient('http://invalid'), lang='eng')[0]
                self.assertEqual(p['status'], 'failed')
                self.assertTrue(p['errors'])
                self.assertEqual(p['tables'], [])


class PersistenceTests(unittest.TestCase):
    setUp=previous.SourceAndSaveTests.setUp
    tearDown=previous.SourceAndSaveTests.tearDown
    _add_project=previous.SourceAndSaveTests._add_project

    def test_real_approval_and_revision_material_cross_checks(self):
        self._add_project('p'); self.service.form_parser = FilingChangeFormParserService()
        from agent.agent_backend.services.filing_change_material_service import FilingChangeMaterialService
        self.service.material_service = FilingChangeMaterialService()
        path = self.base_dir/'application.docx'
        word_file(path, [['药品有效期', '原有效期18个月']])
        self.assertTrue(self.service.import_application_form('p', _Upload(path.name, path.read_bytes()))[0])
        for name, text in [('批准一致.pdf','药品有效期：18个月'), ('批准冲突.pdf','药品有效期：24个月'),
                           ('证书批准.pdf','证书有效期：12个月'), ('修订说明.docx','由6个月延长至12个月')]:
            material = self.base_dir/name
            (word_file if material.suffix=='.docx' else pdf_file)(material, [[text]])
            ok, msg, uploaded = self.service.upload_submission_files('p', [_Upload(name, material.read_bytes())], '1')
            self.assertTrue(ok, msg)
            doc_id = uploaded['created'][0]['doc_id']
            self.assertTrue(self.service.parse_submission('p', doc_id)[0])
        self.assertTrue(self.service.parse_application_form('p')[0])
        field = self.service.get_application_form('p')[2]['form_json']['item_15_validity_period']
        self.assertEqual(field['sub_fields']['original_validity_period'], '18个月')
        self.assertEqual(field['sub_fields']['proposed_validity_period'], '')
        checks = field['cross_checks']
        self.assertTrue(any(c['source_file']=='批准一致.pdf' and c['status']=='一致' for c in checks))
        self.assertTrue(any(c['source_file']=='批准冲突.pdf' and '冲突' in c['status'] for c in checks))
        self.assertTrue(any(c['source_file']=='修订说明.docx' and '待核对' in c['status'] for c in checks))
        self.assertFalse(any(c.get('value')=='12个月' for c in checks))
        self.assertTrue(all(c['source_regions'] for c in checks))

    def test_manual_choices_parties_and_validity_survive_reparse(self):
        self._add_project('p');self.service.form_parser=FilingChangeFormParserService()
        path=self.base_dir/'form.docx'
        word_file(path,[['药品有效期','24个月'],['补充申请的内容','本品有效期保持24个月，稳定性考察由6个月延长至12个月'],
                        ['是否为OTC','处方药☑','非处方药□'],['药品注册申请人','法定代表人','张三']])
        ok,msg,data=self.service.import_application_form('p',_Upload(path.name,path.read_bytes()));self.assertTrue(ok,msg)
        self.assertEqual(data['form_json']['item_15_validity_period']['sub_fields']['original_validity_period'],'')
        form=deepcopy(data['form_json']);form['item_3_otc_type']['value']='非处方药'
        form['item_30_applicant_info']['sub_fields']['contact']='人工代表'
        form['item_15_validity_period']['sub_fields']['original_validity_period']='20个月'
        self.assertTrue(self.service.save_application_form('p',{'form_json':form})[0])
        self.assertTrue(self.service.parse_application_form('p')[0])
        saved=self.service.get_application_form('p')[2]['form_json']
        self.assertEqual(saved['item_3_otc_type']['value'],'非处方药')
        self.assertEqual(saved['item_30_applicant_info']['sub_fields']['contact'],'人工代表')
        self.assertEqual(saved['item_15_validity_period']['sub_fields']['original_validity_period'],'20个月')

    def test_conflict_refresh_keeps_new_conflict_empty_and_history(self):
        self._add_project('p'); self.service.form_parser = FilingChangeFormParserService()
        path = self.base_dir/'first.docx'
        word_file(path, [['药品有效期', '原有效期18个月，拟延长后有效期24个月'], ['是否为OTC', '处方药☑ 非处方药□']])
        ok, msg, first = self.service.import_application_form('p', _Upload(path.name, path.read_bytes()))
        self.assertTrue(ok, msg)
        word_file(path, [['药品有效期', '原有效期18个月，拟延长后有效期36个月'],
                         ['补充申请的内容', '药品有效期由18个月延长至24个月'], ['是否为OTC', '处方药☑ 非处方药☑']])
        ok, msg, data = self.service.import_application_form('p', _Upload('conflict.docx', path.read_bytes()))
        self.assertTrue(ok, msg)
        form = data['form_json']
        self.assertEqual(form['item_15_validity_period']['sub_fields']['proposed_validity_period'], '')
        self.assertEqual(form['item_3_otc_type']['value'], '')
        self.assertTrue(any(h['value'] == '24个月' and h['historical']
                            for h in form['item_15_validity_period']['value_history']))
        self.assertTrue(form['item_15_validity_period']['semantic_issues'])
        self.assertTrue(form['item_3_otc_type']['semantic_issues'])
        ok, msg, saved = self.service.save_application_form('p', {'form_json':form})
        self.assertTrue(ok, msg)
        self.assertEqual(self.service.get_application_form('p')[2]['form_json'], saved['form_json'])
        self.assertTrue(self.service.parse_application_form('p')[0])
        self.assertTrue(self.service.get_application_form('p')[2]['form_json']['item_3_otc_type']['semantic_issues'])

    @unittest.skipUnless(os.getenv('T2_OCR_URL'), '需要真实 OCR HTTP 服务')
    def test_unresolved_table_survives_service_save_refresh_and_reparse(self):
        self._add_project('p'); self.service.form_parser = FilingChangeFormParserService()
        path = self.base_dir/'uncertain.pdf'
        borderless_file(path, uncertain=True)
        from agent.agent_backend.config.settings import settings
        with patch.object(settings, 'ocr_service_url', os.environ['T2_OCR_URL']), patch.object(settings, 'ocr_lang', 'eng'):
            ok, msg, data = self.service.import_application_form('p', _Upload(path.name, path.read_bytes()))
            self.assertTrue(ok, msg)
            self.assertEqual(data['parse_status'], 'partial')
            self.assertIn('Alpha', data['raw_text'])
            errors = data['parse_diagnostics']['errors']
            self.assertTrue(any(e['code']=='table_structure_unresolved' and e.get('bbox_pdf') for e in errors))
            self.assertTrue(self.service.save_application_form('p', {'form_json':data['form_json']})[0])
            refreshed = self.service.get_application_form('p')[2]
            self.assertEqual(refreshed['parse_status'], 'partial')
            self.assertEqual(refreshed['effective_source']['parse_diagnostics'], data['parse_diagnostics'])
            self.assertTrue(self.service.parse_application_form('p')[0])
            self.assertEqual(self.service.get_application_form('p')[2]['latest_attempt']['content_status'], 'partial')


if __name__=='__main__': unittest.main()
