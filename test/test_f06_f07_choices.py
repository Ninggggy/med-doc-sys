"""F06/F07：独立业务断言；Word/PDF用真实文件和真实解析器，纯规则测试另列。"""
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml

from agent.test.test_four_semantic_table_defects import word_file, pdf_file
from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
from agent.agent_backend.utils.parser.form_field_semantics import selection_evidence, selected_option
from agent.agent_backend.services.filing_parse_outcome import ParseFailure


class RealFormTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='f06-f07-')
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.parser = FilingChangeFormParserService()

    def parse(self, rows, ext, *, inspect_empty=False):
        path = self.root / ('form.' + ext)
        (word_file if ext == 'docx' else pdf_file)(path, rows)
        # 空表只检查解析字段；拒绝替换有效原件由主任务F01集成测试验证。
        result = (self.parser._parse_form_file if inspect_empty else self.parser.parse_form_file)(str(path))
        return result['form_json']

    def assert_source(self, field):
        self.assertTrue(field['source_text'])
        self.assertTrue(field['source_regions'])
        self.assertIn(field['source_file'], ('form.docx', 'form.pdf'))
        for entry in field.get('selection_evidence', []):
            self.assertTrue(entry['source_text'])
            self.assertTrue(entry['source_regions'])
            self.assertEqual(entry['source_file'], field['source_file'])

    def test_unmapped_and_misnumbered_matter_keep_facts(self):
        for ext in ('docx', 'pdf'):
            field = self.parse([['申请事项分类', '☑9.99 新事项', '□1.7 增加规格']], ext)['item_5_application_matter_category']
            self.assertEqual(field['selected_values'], [])
            self.assertEqual([(e['code'], e['name'], e['selected']) for e in field['original_matter']],
                             [('9.99', '新事项', True), ('1.7', '增加规格', False)])
            self.assertTrue(field['semantic_issues'])
            self.assert_source(field)
            field = self.parse([['申请事项分类', '☑1.7 增加规格']], ext)['item_5_application_matter_category']
            self.assertEqual(field['selected_values'], ['1.8'])
            self.assertEqual(field['original_matter'][0]['code'], '1.7')

    def test_all_business_matter_names_map_without_replacing_original_codes(self):
        names = ['变更处方中的辅料', '变更生产工艺', '变更所用原料药的供应商', '变更生产批量',
                 '变更注册标准', '变更包装材料和容器', '变更有效期和贮藏条件', '增加规格',
                 '变更生产场地，境外生产药品', '其他']
        for ext in ('docx', 'pdf'):
            field = self.parse([['申请事项分类', f'☑6.{i} {name}'] for i, name in enumerate(names, 1)], ext)['item_5_application_matter_category']
            self.assertEqual(field['selected_values'], [f'1.{i}' for i in range(1, 11)])
            self.assertEqual([e['code'] for e in field['original_matter']], [f'6.{i}' for i in range(1, 11)])
            self.assertTrue(all(e['selected'] for e in field['original_matter']))

    def test_labels_and_unmarked_candidates_are_not_selections(self):
        for ext in ('docx', 'pdf'):
            form = self.parse([['剂型', ''], ['申请事项分类', ''], ['是否为OTC', '']], ext, inspect_empty=True)
            for key in ('item_11_dosage_form', 'item_5_application_matter_category', 'item_3_otc_type'):
                self.assertEqual(form[key]['value'], '')
                self.assertEqual(form[key]['selected_values'], [])
                self.assertEqual(form[key]['parse_confidence'], 0)
            rows = [['申请事项分类', '6.8 变更有效期和贮藏条件', '6.9 增加规格']]
            with self.assertRaises(ParseFailure) as caught:
                self.parse(rows, ext)
            self.assertEqual(caught.exception.code, 'empty_form')
            field = self.parse([['药品通用名称', '测试药品']] + rows, ext)['item_5_application_matter_category']
            self.assertEqual(field['selected_values'], [])
            self.assertTrue(all(e['selected'] is None for e in field['original_matter']))

    def test_duplicate_conflicting_matter_states(self):
        for ext in ('docx', 'pdf'):
            field = self.parse([['申请事项分类', '☑6.8 变更有效期和贮藏条件', '□6.8 变更有效期和贮藏条件']], ext)['item_5_application_matter_category']
            self.assertEqual(field['selected_values'], [])
            self.assertTrue(field['semantic_issues'])
            self.assertEqual(field['selection_status'], 'conflict')

    def test_plain_dosage_remains_available(self):
        for ext in ('docx', 'pdf'):
            field = self.parse([['剂型', '片剂']], ext)['item_11_dosage_form']
            self.assertEqual(field['selected_values'], ['片剂'])

    def test_single_explicit_matter_with_management_heading(self):
        for ext in ('docx', 'pdf'):
            field = self.parse([['申请事项分类', '6. 已上市化学药品药学变更；6.8 变更有效期和贮藏条件']], ext)['item_5_application_matter_category']
            self.assertEqual(field['selected_values'], ['1.7'])
            leaf = next(e for e in field['original_matter'] if e['code'] == '6.8')
            self.assertEqual(leaf['status'], 'explicit')

    def test_dosage_spaced_title_and_name(self):
        for ext in ('docx', 'pdf'):
            field = self.parse([['剂型', '中 国药 典 剂 型 : 片 剂']], ext)['item_11_dosage_form']
            self.assertEqual(field['selected_values'], ['片剂'])
            self.assertIn('片 剂', field['source_text'])

    def test_unmarked_dosage_list_does_not_become_one_answer(self):
        for ext in ('docx', 'pdf'):
            field = self.parse([['药品通用名称', '测试药品'], ['剂型', '片剂 胶囊剂']], ext)['item_11_dosage_form']
            self.assertEqual(field['selected_values'], [])

    def test_wrapped_checked_matter_name(self):
        for ext in ('docx', 'pdf'):
            for left in ('☑6.8 变更有效期和\n贮藏条件', '6.8 变更有效期和☑\n贮藏条件'):
                field = self.parse([['申请事项分类', left, '□6.9 增加规格']], ext)['item_5_application_matter_category']
                self.assertEqual(field['selected_values'], ['1.7'])
                self.assertIn('贮藏条件', field['original_matter'][0]['name'])

    def test_ambiguous_middle_marker_is_not_a_choice(self):
        for ext in ('docx', 'pdf'):
            field = self.parse([['药品通用名称', '测试药品'], ['剂型', '片剂 ☑ 胶囊剂']], ext)['item_11_dosage_form']
            self.assertEqual(field['selected_values'], [])
            self.assertTrue(field['semantic_issues'])

    def test_wingdings_checked_and_unchecked_names(self):
        doc = Document(); table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '剂型'
        paragraph = table.cell(0, 1).paragraphs[0]
        for code, label in (('0052', '片剂'), ('00A3', '胶囊剂')):
            paragraph._p.append(parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                                         '<w:sym w:font="Wingdings 2" w:char="' + code + '"/></w:r>'))
            paragraph.add_run(label + ' ')
        path = self.root / 'form.docx'; doc.save(path)
        field = self.parser.parse_form_file(str(path))['form_json']['item_11_dosage_form']
        self.assertEqual(field['selected_values'], ['片剂'])
        self.assertEqual(self.parser._extract_checked_labels_from_docx_xml(str(path)), ['片剂'])

    def test_word_paragraph_matter_and_statement(self):
        path = self.root / 'form.docx'
        doc = Document()
        doc.add_paragraph('申请事项分类：6.8 变更有效期和贮藏条件☑\n6.9 增加规格□')
        doc.add_paragraph('专利情况：本申请声明：尚不能确认是否构成他人专利侵权。')
        doc.save(path)
        form = self.parser.parse_form_file(str(path))['form_json']
        self.assertEqual(form['item_5_application_matter_category']['selected_values'], ['1.7'])
        self.assertEqual(form['item_24_patent_info']['value'], '本申请声明：尚不能确认是否构成他人专利侵权。')
        self.assertTrue(form['item_24_patent_info']['semantic_issues'])
        self.assertEqual(form['item_24_patent_info']['source_regions'][0]['paragraph'], 1)

    def test_ooxml_state_and_visible_glyph_are_one_choice(self):
        for field_name, key, labels, expected in (
                ('是否为OTC', 'item_3_otc_type', ['处方药', '非处方药'], '处方药'),
                ('剂型', 'item_11_dosage_form', ['片剂', '胶囊剂'], ['片剂']),
                ('申请事项分类', 'item_5_application_matter_category', ['6.8 变更有效期和贮藏条件', '6.9 增加规格'], ['1.7'])):
            doc = Document(); table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = field_name
            paragraph = table.cell(0, 1).paragraphs[0]
            for state, glyph, label in zip(('1', '0'), ('☒', '☐'), labels):
                paragraph._p.append(parse_xml(
                    '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                    'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
                    '<w:sdtPr><w14:checkbox><w14:checked w14:val="' + state + '"/></w14:checkbox></w:sdtPr>'
                    '<w:sdtContent><w:r><w:t>' + glyph + '</w:t></w:r></w:sdtContent></w:sdt>'))
                paragraph.add_run(label + ' ')
            path = self.root / 'form.docx'; doc.save(path)
            field = self.parser.parse_form_file(str(path))['form_json'][key]
            self.assertEqual(field['value'] if isinstance(expected, str) else field['selected_values'], expected)


# 每种真实输入独立计数，避免几十个子用例只记成“一项通过”。
def choice_case(ext, field_name, key, labels, selected, layout):
    def test(self):
        marks = ['☑' if i in selected else '□' for i in range(2)]
        if layout == 'prefix':
            cells = [mark + label for mark, label in zip(marks, labels)]
        elif layout == 'postfix':
            cells = [label + mark for mark, label in zip(marks, labels)]
        elif layout == 'mixed':
            cells = [labels[0] + marks[0], marks[1] + labels[1]]
        else:
            cells = [' '.join(label + mark for mark, label in zip(marks, labels))]
        rows = [[field_name] + cells]
        if not selected:
            # 保留原全未选输入：公共入口必须拒绝整张空模板；有药名时仍可读取未选证据。
            with self.assertRaises(ParseFailure) as caught:
                self.parse(rows, ext)
            self.assertEqual(caught.exception.code, 'empty_form')
            rows = [['药品通用名称', '测试药品']] + rows
        field = self.parse(rows, ext)[key]
        expected = [('1.7', '1.8')[i] if key.startswith('item_5_') else labels[i] for i in selected]
        self.assertEqual(field['selected_values'], expected)
        self.assert_source(field)
        if key.startswith('item_5_'):
            self.assertEqual([(e['code'], e['name'], e['selected']) for e in field['original_matter']],
                             [('6.8', '变更有效期和贮藏条件', 0 in selected), ('6.9', '增加规格', 1 in selected)])
    return test


for _ext in ('docx', 'pdf'):
    for _name, _key, _labels in (
            ('剂型', 'item_11_dosage_form', ['片剂', '胶囊剂']),
            ('申请事项分类', 'item_5_application_matter_category', ['6.8 变更有效期和贮藏条件', '6.9 增加规格'])):
        for _selected in ([], [0], [1], [0, 1]):
            for _layout in ('prefix', 'postfix', 'mixed', 'one_cell'):
                setattr(RealFormTests, f'test_{_ext}_{_key}_{str(_selected)}_{_layout}',
                        choice_case(_ext, _name, _key, _labels, _selected, _layout))


def radio_case(ext, label, key, labels, postfix):
    def test(self):
        cells = [(name + ('☑' if i == 0 else '□')) if postfix else (('☑' if i == 0 else '□') + name)
                 for i, name in enumerate(labels)]
        field = self.parse([['药品通用名称', '测试药品'], [label] + cells], ext)[key]
        actual = field['sub_fields']['has_change'] if key == 'item_16_prescription' else field['value']
        self.assertEqual(actual, labels[0])
        self.assert_source(field)
    return test


for _ext in ('docx', 'pdf'):
    for _label, _key, _labels in (
            ('本申请属于', 'item_1_application_type', ['境内生产药品补充申请']),
            ('药品注册分类', 'item_2_registration_category', ['化学药品', '中药']),
            ('是否为OTC', 'item_3_otc_type', ['处方药', '非处方药']),
            ('原申请品种状态', 'item_4_original_product_status', ['已上市', '未上市']),
            ('商品名称', 'item_10_trade_name', ['使用', '不使用']),
            ('受理前药品注册检验', 'item_19_pre_acceptance_inspection', ['是', '否'])):
        for _postfix in (False, True):
            setattr(RealFormTests, f'test_radio_{_ext}_{_key}_{_postfix}', radio_case(_ext, _label, _key, _labels, _postfix))


def statement_case(ext, label, key, text, uncertain):
    def test(self):
        field = self.parse([[label, text]], ext)[key]
        self.assertEqual(field['value'], text)
        self.assert_source(field)
        if uncertain:
            self.assertTrue(field['semantic_issues'])
        else:
            self.assertFalse(field.get('semantic_issues'))
    return test


for _ext in ('docx', 'pdf'):
    for _i, (_text, _uncertain) in enumerate((
            ('本申请声明：构成他人专利侵权。', False),
            ('本申请声明：不构成他人专利侵权。', False),
            ('本申请声明：尚不能确认是否构成他人专利侵权。', True),
            ('本申请声明：构成他人去利亿权。', True),
            ('本申请声明：?构成他人专利侵权。', True))):
        setattr(RealFormTests, f'test_{_ext}_patent_{_i}', statement_case(_ext, '专利情况', 'item_24_patent_info', _text, _uncertain))
    for _i, (_text, _uncertain) in enumerate((('首次申请', False), ('非首次申请', False), ('不是首次申请', False),
                                             ('尚不能确认是否首次申请', True), ('?首次申请', True),
                                             ('经核对，本次为首次申请，请按首次申请办理。', False),
                                             ('本品已有申请记录，本次为非首次申请。', False))):
        setattr(RealFormTests, f'test_{_ext}_first_{_i}', statement_case(_ext, '本次申请为', 'item_28_change_related_items', _text, _uncertain))


class ChoiceRuleTests(unittest.TestCase):
    def test_legacy_entry_points_share_checked_rules(self):
        parser = FilingChangeFormParserService()
        for method in (parser._extract_checked_options, parser._extract_pdf_marked_labels):
            self.assertEqual(method('片剂☑ 胶囊剂□'), ['片剂'])
            self.assertEqual(method('☑片剂 □胶囊剂'), ['片剂'])
            self.assertEqual(method('片剂 ☑ 胶囊剂'), [])
            self.assertEqual(method('片剂□ 胶囊剂□'), [])
        self.assertEqual(parser._extract_matter_codes_from_pdf_text('6.8 变更有效期和贮藏条件□'), [])
        self.assertEqual(parser._extract_matter_codes_from_pdf_text('6.8 变更有效期和贮藏条件☑'), ['1.7'])
        self.assertEqual(parser._map_checked_labels_to_matter_codes(['不增加规格']), [])
        self.assertEqual(parser._map_checked_labels_to_matter_codes(['1.7']), [])

    def test_radio_cell_boundaries_and_conflicting_duplicates(self):
        self.assertEqual(selected_option(['处方药☑', '□非处方药'], ['处方药', '非处方药']), '处方药')
        self.assertEqual(selected_option(['处方药☑', '处方药□'], ['处方药', '非处方药']), '')
        self.assertEqual(selected_option(['处方药', '☑', '非处方药'], ['处方药', '非处方药']), '')

    def test_no_sign_can_migrate_across_cells(self):
        result = selection_evidence(['片剂', '☑', '胶囊剂'], multiple=True)
        self.assertEqual(result['selected_values'], [])
        self.assertEqual(result['status'], 'ambiguous')


def mark_case(mark):
    def test(self):
        for text in (mark + '片剂 □胶囊剂', '片剂' + mark + ' 胶囊剂□'):
            self.assertEqual(selection_evidence(text, multiple=True)['selected_values'], ['片剂'])
        self.assertEqual(selected_option('是' + mark + ' 否□', ['是', '否']), '是')
    return test


for _i, _mark in enumerate(('☑', '☒', '√', '✓', '✔', '■', '●', '▣', '(x)', '（X）', '(V)', '(v)', '(√)', '[x]')):
    setattr(ChoiceRuleTests, f'test_marker_{_i}', mark_case(_mark))


if __name__ == '__main__':
    unittest.main()
