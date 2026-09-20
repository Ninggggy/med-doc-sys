"""真实文件与临时数据库回归；正常 OCR 只使用真实服务。"""
import os
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document

from agent.test import test_pdf_page_extraction as pdf_fixture
from agent.test import test_filing_change_review_service_regression as service_fixture
from agent.test.test_pdf_page_extraction import gradient_page
from agent.test.test_filing_change_review_service_regression import _Upload
from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
from agent.agent_backend.utils.parser.pdf_page_extractor import extract_pdf_pages
from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import OCRServiceClient


class LayoutTests(unittest.TestCase):
    setUp = pdf_fixture.PDFPagesTests.setUp
    tearDown = pdf_fixture.PDFPagesTests.tearDown

    def test_word_column_values(self):
        for cells, name, spec in [
            (['药品通用名称', '明确药名', ''], '明确药名', ''),
            (['药品通用名称', '明确药名', '规格', '1%'], '明确药名', '1%'),
            (['药品通用名称', '', '规格', '1%'], '', '1%'),
            (['药品通用名称', '内容', '明确药名'], '明确药名', ''),
        ]:
            with self.subTest(cells=cells):
                doc = Document()
                table = doc.add_table(rows=1, cols=len(cells))
                for cell, text in zip(table.rows[0].cells, cells): cell.text = text
                path = self.path.with_suffix('.docx')
                doc.save(path)
                form = FilingChangeFormParserService().parse_form_file(str(path))['form_json']
                self.assertEqual(form['item_6_generic_name']['value'], name)
                self.assertEqual(form['item_12_specification']['value'], spec)

    def mixed(self, overlay='Archive copy'):
        source = fitz.open()
        p = gradient_page(source)
        doc = fitz.open()
        page = doc.new_page(width=500, height=500)
        page.insert_image(page.rect, stream=p.get_pixmap(matrix=fitz.Matrix(3, 3)).tobytes('png'))
        page.insert_text((40, 290 if overlay == 'Archive copy' else 65), overlay)
        doc.save(self.path)
        doc.close()
        source.close()

    def test_word_multiple_groups_vertical_and_horizontal_merges(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=5)
        table.cell(0, 0).merge(table.cell(1, 0)).text = '药品通用名称'
        table.cell(0, 1).merge(table.cell(0, 2)).text = '明确药名'
        table.cell(1, 1).merge(table.cell(1, 2)).text = '延续说明'
        table.cell(0, 3).merge(table.cell(1, 3)).text = '规格'
        table.cell(0, 4).text = ''
        table.cell(1, 4).text = '1%'
        path = self.path.with_suffix('.docx')
        doc.save(path)
        form = FilingChangeFormParserService().parse_form_file(str(path))['form_json']
        self.assertEqual(form['item_6_generic_name']['value'], '明确药名\n延续说明')
        self.assertEqual(form['item_12_specification']['value'], '1%')

    def test_complete_native_layer_over_scan_needs_no_ocr(self):
        doc = fitz.open()
        page = gradient_page(doc)
        raster = page.get_pixmap(matrix=fitz.Matrix(3, 3)).tobytes('png')
        page.insert_image(page.rect, stream=raster, overlay=False)
        doc.save(self.path)
        doc.close()
        page = extract_pdf_pages(str(self.path))[0]
        self.assertEqual(page['ocr_calls'], 0)
        self.assertEqual(page['status'], 'success')
        pdf_fixture.PDFPagesTests.assert_table(self, page)

    def test_multiheader_chain_and_nonadjacent_tables(self):
        for gap in [False, True]:
            doc = fitz.open()
            for i in range(3):
                gradient_page(doc, 'Gradient elution' + (' (continued)' if i else ''))
                if gap and i == 0: doc.new_page()
            doc.save(self.path)
            doc.close()
            pages = extract_pdf_pages(str(self.path))
            tables = [p['tables'][0] for p in pages if p['tables']]
            if gap: self.assertNotIn('continuation_of', tables[1])
            else: self.assertEqual(tables[1]['continuation_of'], tables[0]['id'])
            self.assertEqual(tables[2]['continuation_of'], tables[1]['id'])
            self.assertEqual(tables[2]['repeated_header_rows'], [0, 1])
            self.assertEqual(tables[2]['rows'][2], ['0.5', '50.0', '50.0'])
            self.assertEqual(tables[2]['page'], 4 if gap else 3)

    def test_mixed_ocr_failures_keep_native_and_missing_region(self):
        self.mixed()
        for options in [dict(side_effect=TimeoutError('timeout')), dict(return_value={'text': []}), dict(side_effect=RuntimeError('offline'))]:
            with self.subTest(options=options), patch.object(OCRServiceClient, 'image_to_data', **options):
                page = extract_pdf_pages(str(self.path))[0]
                self.assertIn('Archive copy', page['text'])
                self.assertEqual(page['status'], 'partial')
                self.assertEqual(page['parse_diagnostics']['failed_pages'], [1])
                self.assertTrue(page['errors'][0]['bbox_pdf'])

    @unittest.skipUnless(os.getenv('T2_OCR_URL'), '真实 OCR 服务未配置')
    def test_real_mixed_overlay(self):
        for overlay in ['Archive copy', 'Overlay']:
            with self.subTest(overlay=overlay):
                self.mixed(overlay)
                page = extract_pdf_pages(str(self.path), ocr_client=OCRServiceClient(os.environ['T2_OCR_URL']), lang='eng')[0]
                self.assertGreater(page['ocr_calls'], 0)
                pdf_fixture.PDFPagesTests.assert_table(self, page)
                self.assertIn(overlay, page['text'])
                self.assertIn('Gradient elution', page['text'])
                self.assertEqual(page['tables'][0]['rows'][3][0], '-1.25')
                self.assertEqual(page['status'], 'success')
                self.assertEqual(page['text'].count(overlay), 1)

    @unittest.skipUnless(os.getenv('T2_OCR_URL'), '真实 OCR 服务未配置')
    def test_real_native_body_and_local_image_table(self):
        pdf_fixture.PDFPagesTests.make(self, local=True)
        page = extract_pdf_pages(str(self.path), ocr_client=OCRServiceClient(os.environ['T2_OCR_URL']), lang='eng')[0]
        pdf_fixture.PDFPagesTests.assert_table(self, page)
        self.assertEqual(page['text'].count('Quality standard native narrative'), 1)
        self.assertEqual(page['text'].count('50.0'), 2)
        self.assertEqual(page['path'], ['native', 'region_ocr'])
        self.assertAlmostEqual(page['tables'][0]['bbox_pdf'][0], 40, delta=1)

    def test_single_header_three_pages_and_different_columns(self):
        for bad_columns in [False, True]:
            doc = fitz.open()
            for i in range(3):
                p = doc.new_page(width=500, height=500)
                p.insert_text((40, 45), 'Assay' + (' (continued)' if i else ''))
                middle = 200 if not bad_columns or i == 0 else 260
                for x in (40, middle, 460): p.draw_line((x, 80), (x, 200))
                for y in (80, 120, 160, 200): p.draw_line((40, y), (460, y))
                for x, y, text in [(50, 105, 'Time'), (middle+10, 105, 'Value'), (50, 145, str(i)), (middle+10, 145, '-1.25'), (50, 185, '50.0'), (middle+10, 185, '50.0')]: p.insert_text((x, y), text)
                p.insert_text((40, 230), 'Note: 95%')
            doc.save(self.path)
            doc.close()
            pages = extract_pdf_pages(str(self.path))
            for i in (1, 2):
                table = pages[i]['tables'][0]
                if bad_columns and i == 1:
                    self.assertNotIn('continuation_of', table)
                else:
                    self.assertEqual(table.get('continuation_of'), pages[i-1]['tables'][0]['id'])
                    self.assertEqual(table['repeated_header_rows'], [0])
                self.assertEqual(table['rows'][2], ['50.0', '50.0'])
                self.assertIn('95%', str(table['notes']))


class SourceAndSaveTests(unittest.TestCase):
    setUp = service_fixture.FilingChangeReviewServiceRegressionTest.setUp
    tearDown = service_fixture.FilingChangeReviewServiceRegressionTest.tearDown
    _add_project = service_fixture.FilingChangeReviewServiceRegressionTest._add_project

    def file(self, suffix):
        path = self.base_dir / ('input' + suffix)
        if suffix == '.docx':
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = '药品通用名称'
            table.cell(0, 1).text = '原始药名'
            doc.save(path)
        else:
            doc = fitz.open()
            doc.new_page().insert_text((40, 50), 'Generic name: Original medicine')
            doc.save(path)
            doc.close()
        return path.read_bytes()

    def test_empty_and_manual_save_are_not_file_parse(self):
        self._add_project('p')
        fresh = self.service.get_application_form('p')[2]
        self.assertIn('item_6_generic_name', fresh['form_json'])
        self.assertEqual(fresh['form_json']['item_6_generic_name']['value'], '')
        for form in [{}, {'item_6_generic_name': {'field_type': 'input', 'value': '人工药名', 'manual_edited': True}}]:
            ok, msg, data = self.service.save_application_form('p', {'form_json': form})
            self.assertTrue(ok, msg)
            self.assertEqual(data['parse_status'], 'not_parsed')
            self.assertFalse(data.get('original_file_id'))
            self.assertFalse(data.get('effective_source'))
            self.assertEqual(self.service.get_application_form('p')[2]['parse_status'], 'not_parsed')

    def test_failed_replacement_then_reparse_source_and_manual(self):
        self.service.form_parser = FilingChangeFormParserService()
        for n, (good, bad) in enumerate([('.docx', '.pdf'), ('.pdf', '.docx'), ('.docx', '.docx')]):
            for legacy in [False, True]:
                pid = f'p{n}{legacy}'
                self._add_project(pid)
                content = self.file(good)
                original_name = '有效原件' + good
                ok, msg, old = self.service.import_application_form(pid, _Upload(original_name, content))
                self.assertTrue(ok, msg)
                form = deepcopy(old['form_json'])
                form['item_6_generic_name']['value'] = '人工保护值'
                form['item_6_generic_name']['manual_edited'] = True
                self.assertTrue(self.service.save_application_form(pid, {'form_json': form})[0])
                if legacy:
                    manifest = self.service._load_submission_manifest(pid)
                    manifest.pop('application_form_source', None)
                    self.service._save_submission_manifest(pid, manifest)
                self.assertFalse(self.service.import_application_form(pid, _Upload('失败替换' + bad, b'damaged'))[0])
                current = self.service.get_application_form(pid)[2]
                self.assertEqual(current['original_file_id'], old['original_file_id'])
                stored = self.service._project_path(pid) / 'application_form' / (old['original_file_id'] + good)
                self.assertEqual(stored.read_bytes(), content)
                self.assertEqual(current['latest_attempt']['content_status'], 'failed')
                ok, msg, new = self.service.parse_application_form(pid)
                self.assertTrue(ok, msg)
                self.assertEqual(new['source_file_name'], original_name)
                self.assertEqual(new['source_file_type'], good[1:])
                refreshed = self.service.get_application_form(pid)[2]
                self.assertEqual(refreshed['form_json']['item_6_generic_name']['value'], '人工保护值')
                self.assertEqual(refreshed['effective_source']['source_file_name'], original_name)
                self.assertEqual(refreshed['latest_attempt']['content_status'], 'success')
                self.assertTrue(any(a['content_status'] == 'failed' for a in self.service._load_submission_manifest(pid)['parse_attempts']))

    def test_history_without_matching_records_uses_safe_file_name(self):
        self.service.form_parser = FilingChangeFormParserService()
        self._add_project('p')
        old = self.service.import_application_form('p', _Upload('原名称.docx', self.file('.docx')))[2]
        manifest = self.service._load_submission_manifest('p')
        for key in ['application_form_source', 'application_form_attempt', 'parse_attempts']: manifest.pop(key, None)
        self.service._save_submission_manifest('p', manifest)
        self.assertFalse(self.service.import_application_form('p', _Upload('坏文件.pdf', b'bad'))[0])
        ok, msg, new = self.service.parse_application_form('p')
        self.assertTrue(ok, msg)
        self.assertEqual(new['source_file_name'], old['original_file_id'] + '.docx')

    def test_commit_failure_preserves_effective_source_and_manual_values(self):
        self.service.form_parser = FilingChangeFormParserService()
        self._add_project('p')
        content = self.file('.docx')
        old = self.service.import_application_form('p', _Upload('有效.docx', content))[2]
        form = deepcopy(old['form_json'])
        form['item_6_generic_name']['value'] = '人工值'
        self.service.save_application_form('p', {'form_json': form})
        previous = self.service.get_application_form('p')[2]
        with patch.object(self.service, '_install_file_payloads', side_effect=PermissionError('injected replacement failure')):
            self.assertFalse(self.service.import_application_form('p', _Upload('新.pdf', self.file('.pdf')))[0])
        unchanged = self.service.get_application_form('p')[2]
        for key in ['original_file_id', 'effective_source', 'form_json']:
            self.assertEqual(unchanged[key], previous[key])
        self.connection.fail_next_commit = True
        self.assertFalse(self.service.import_application_form('p', _Upload('新.pdf', self.file('.pdf')))[0])
        saved = self.service.get_application_form('p')[2]
        self.assertEqual(saved['effective_source'], previous['effective_source'])
        self.assertEqual(saved['form_json'], previous['form_json'])
        self.assertEqual(saved['original_file_id'], previous['original_file_id'])
        stored = self.service._project_path('p') / 'application_form' / (saved['original_file_id'] + '.docx')
        self.assertEqual(stored.read_bytes(), content)
        self.assertTrue(self.service.parse_application_form('p')[0])
        self.assertEqual(self.service.get_application_form('p')[2]['form_json']['item_6_generic_name']['value'], '人工值')

    def test_manual_save_preserves_partial_success_and_latest_failure(self):
        self.service.form_parser = FilingChangeFormParserService()
        for partial in [False, True]:
            pid = 'p' + str(partial)
            self._add_project(pid)
            doc = fitz.open()
            doc.new_page().insert_text((40, 50), 'Generic name: Valid medicine')
            if partial:
                p = doc.new_page()
                p.draw_line((20, 20), (300, 300))
            content = doc.tobytes()
            doc.close()
            with patch.object(OCRServiceClient, 'image_to_data', side_effect=TimeoutError('timeout')):
                ok, msg, data = self.service.import_application_form(pid, _Upload('有效.pdf', content))
            self.assertTrue(ok, msg)
            self.assertEqual(data['parse_status'], 'partial' if partial else 'success')
            self.assertFalse(self.service.import_application_form(pid, _Upload('坏.docx', b'broken'))[0])
            before = self.service.get_application_form(pid)[2]
            form = deepcopy(before['form_json'])
            form['item_6_generic_name']['value'] = '人工值'
            self.assertTrue(self.service.save_application_form(pid, {'form_json': form})[0])
            after = self.service.get_application_form(pid)[2]
            for key in ['parse_status', 'effective_source', 'latest_attempt', 'original_file_id']:
                self.assertEqual(after[key], before[key])
            self.assertTrue(after['form_json']['item_6_generic_name']['manual_modified'])

    def test_autosaved_empty_form_then_bad_first_upload(self):
        self.service.form_parser = FilingChangeFormParserService()
        self._add_project('p')
        self.service.save_application_form('p', {'form_json': {}})
        for content in [b'', b'%PDF-1.7 broken']:
            self.assertFalse(self.service.import_application_form('p', _Upload('坏.pdf', content))[0])
            data = self.service.get_application_form('p')[2]
            self.assertEqual(data['parse_status'], 'not_parsed')
            self.assertFalse(data['original_file_id'])
            self.assertFalse(data['effective_source'])
            self.assertEqual(data['latest_attempt']['content_status'], 'failed')


if __name__ == '__main__':
    unittest.main()
