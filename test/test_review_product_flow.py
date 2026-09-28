"""开发侧功能测试：独立已知文本 + 真实上传/解析/修订API，不是原件人工验收。"""
import io
import os
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from agent.deploy.review_test_app import create_app


class ProductFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory(prefix='review-product-')
        cls.app=create_app(cls.tmp.name,cls.tmp.name)
        cls.service=cls.app.extensions['review_service']

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def setUp(self):
        self.client=self.app.test_client()
        self.base='/api/filing-change-review/projects'
        self.project=self.request('post',self.base,{'project_name':'自动化独立材料'})['project_id']
        self.root=f'{self.base}/{self.project}'

    def request(self,method,url,payload=None,status=200,**kwargs):
        r=getattr(self.client,method)(url,json=payload,**kwargs)
        self.assertEqual(r.status_code,status,r.get_data(as_text=True))
        return json.loads(r.data).get('data')

    def upload(self,text='生产范围：片剂，仅限口服，不得生产注射剂。'):
        data=self.request('post',self.root+'/submissions/upload',data={'files':(io.BytesIO(text.encode()),'独立测试材料.txt'),'material_category':'application_info','material_sub_category':'2'})
        doc=data['created'][0]['doc_id']
        self.request('post',self.root+f'/submissions/{doc}/parse')
        return doc,self.root+f'/parse-review/submission/{doc}'

    def context(self,url):return self.request('get',url)
    def save(self,url,ctx,items,status=200):
        return self.request('post',url,{'source_identity':ctx['source_identity'],'expected_revision':ctx['revision'],'items':items},status=status)
    def item(self,target,text,action='edit_value'):
        return dict(issue_key=target['issue_key'],source_value=target['source_value'],text=text,
            verified_value=text,action=action,reason='根据独立测试原文模拟用户核对',reviewer={'id':'forged'})

    def test_edit_refresh_identity_source_and_no_ocr(self):
        doc,url=self.upload();ctx=self.context(url);original=ctx['original_chunks'];t=ctx['targets'][0]
        value='生产范围：片剂，仅限外用，不得生产注射剂。'
        with patch('agent.agent_backend.services.filing_paddle_backend.parse_pdf',side_effect=AssertionError('保存不得调用OCR')):
            self.save(url,ctx,[self.item(t,value)])
            fresh=self.context(url)
        self.assertEqual(fresh['original_chunks'],original)
        self.assertEqual(fresh['effective_chunks'][0]['text'],value)
        self.assertEqual(fresh['items'][0]['reviewer']['id'],'unattributed')
        self.assertEqual(fresh['targets'][0]['value'],value)
        self.assertNotIn('forged',json.dumps(fresh['effective_chunks']))
        self.save(url,ctx,[self.item(t,value)],409)
        bad=self.item(t,value+'新增');bad['verified_value']=value
        self.save(url,fresh,[bad],400)
        self.request('post',self.root+f'/submissions/{doc}/parse')
        self.save(url,fresh,fresh['items'],409)

    def test_states_empty_unknown_and_foreign_target(self):
        doc,url=self.upload();ctx=self.context(url);t=ctx['targets'][0]
        self.save(url,ctx,[self.item(t,'','confirm_value')],400)
        self.save(url,ctx,[self.item(t,'?')],400)
        self.save(url,ctx,[self.item(t,'','confirm_blank')])
        fresh=self.context(url)
        self.assertEqual(fresh['items'][0]['content_state'],'explicit_blank')
        self.assertFalse(fresh['effective_chunks'][0].get('manual_unreadable'))
        self.save(url,fresh,[self.item(t,'日期?','mark_unreadable')])
        fresh=self.context(url)
        self.assertEqual(fresh['effective_chunks'][0]['text'],'')
        self.assertTrue(fresh['effective_chunks'][0]['manual_unreadable'])
        ready=self.request('get',self.root+'/review/readiness')
        self.assertTrue(any(i['code']=='manual_unreadable' for i in ready['blocking_issues']))
        _,other=self.upload('另一个文档')
        self.save(other,self.context(other),fresh['items'],400)

    def test_confirm_and_partial_save_failure(self):
        _,url=self.upload();ctx=self.context(url);t=ctx['targets'][0]
        self.save(url,ctx,[self.item(t,t['value'],'confirm_value')])
        fresh=self.context(url);self.assertTrue(fresh['targets'][0]['resolved'])
        with patch.object(self.service,'_save_submission_manifest',side_effect=OSError('测试磁盘写入失败')):
            self.save(url,fresh,[self.item(t,'生产范围：仅限外用。')],400)
        after=self.context(url)
        self.assertEqual(after['revision'],fresh['revision'])
        self.assertEqual(after['effective_chunks'],fresh['effective_chunks'])
        self.assertFalse(self.request('get',self.root+'/review/readiness')['ready'])


    def pdf_fixture(self,rows):
        import fitz
        from copy import deepcopy
        pdf=fitz.open();page=pdf.new_page(width=600,height=800)
        cells=[];lines=[]
        for r,row in enumerate(rows):
            for c,value in enumerate(row):
                box=[30+c*250,30+r*45,280+c*250,75+r*45]
                page.draw_rect(fitz.Rect(box));page.insert_text((box[0]+4,box[1]+22),value,fontname='china-s',fontsize=11)
                cells.append(dict(row=r,column=c,rowspan=1,colspan=1,text=value,bbox_pdf=box))
                lines.append(dict(text=value,bbox=[box[0]+4,box[1]+5,box[2]-4,box[3]-5],table_index=0,source='independent_test_fixture'))
        table=dict(id='synthetic',table_index=0,page=1,bbox_pdf=[30,30,530,30+45*len(rows)],cells=cells,
            rows=rows,raw_rows=deepcopy(rows),markdown='\n'.join('| '+' | '.join(r)+' |' for r in rows),needs_review=False)
        chunk=dict(page=1,page_bbox=[0,0,600,800],status='success',errors=[],tables=[table],lines=lines,words=deepcopy(lines),
            text=table['markdown'],raw_text=table['markdown'],ocr_backend='ppocr_v6_medium_experimental')
        return pdf.tobytes(),[chunk]

    def corrupt_cell_fixture(self,chunks,index,value):
        # 独立PDF保留已知正确原文，仅模拟保存下来的OCR错字；不运行第二模型。
        chunk=chunks[0];table=chunk['tables'][0];cell=table['cells'][index]
        cell['text']=value
        for name in ('rows','raw_rows'):table[name][cell['row']][cell['column']]=value
        for name in ('lines','words'):chunk[name][index]['text']=value
        table['markdown']='\n'.join('| '+' | '.join(r)+' |' for r in table['raw_rows'])
        chunk['text']=chunk['raw_text']=table['markdown']

    def test_table_scope_and_form_recompute_no_ocr(self):
        from agent.agent_backend.services.filing_change_extraction_service import extract_document
        pdf,chunks=self.pdf_fixture([['生产地址','生产范围'],['测试一路','片剂（仅限口服）']])
        data=self.request('post',self.root+'/submissions/upload',data={'files':(io.BytesIO(pdf),'独立表格.pdf'),'material_category':'2'})
        doc=data['created'][0]['doc_id'];url=self.root+f'/parse-review/submission/{doc}'
        with patch.object(self.service.material_service,'parse_file',return_value=chunks) as parse:
            self.request('post',self.root+f'/submissions/{doc}/parse')
            ctx=self.context(url);t=ctx['targets'][3];value='片剂（仅限外用，不得口服）'
            self.save(url,ctx,[self.item(t,value)])
            fresh=self.context(url)
            self.assertEqual(parse.call_count,1)
        effective=fresh['effective_chunks'];self.assertEqual(effective[0]['tables'][0]['rows'][1][1],value)
        self.assertIn(value,effective[0]['text']);self.assertEqual(fresh['original_chunks'],chunks)
        facts=extract_document('独立表格.pdf',effective)['facts']
        self.assertTrue(any(f['field']=='production_scope' and value==f['raw_value'] for f in facts))
        self.assertTrue(any(value in str(f) and '测试一路' in str(f) for f in facts))
        pdf,formchunks=self.pdf_fixture([['药品通用名称','合成测试片'],['规格','3 mg']])
        self.corrupt_cell_fixture(formchunks,1,'合成错字片')
        with patch('agent.agent_backend.services.filing_paddle_backend.enabled',return_value=True),patch('agent.agent_backend.services.filing_paddle_backend.parse_pdf',return_value=formchunks) as parse:
            result=self.request('post',self.root+'/application-form/import',data={'file':(io.BytesIO(pdf),'独立申请表.pdf')})
            form=self.request('get',self.root+'/application-form');fid=form['original_file_id']
            url=self.root+f'/parse-review/application_form/{fid}';ctx=self.context(url)
            target=next(t for t in ctx['targets'] if t['value']=='合成错字片')
            self.save(url,ctx,[self.item(target,'合成测试片')])
            refreshed=self.request('get',self.root+'/application-form')
            self.assertEqual(refreshed['form_json']['item_6_generic_name']['value'],'合成测试片')
            self.assertEqual(parse.call_count,1)
            self.assertEqual(self.context(url)['original_chunks'],ctx['original_chunks'])

    def test_positive_review_report_real_prerequisites(self):
        import time
        pdf,chunks=self.pdf_fixture([['药品通用名称','合成测试片'],['规格','3 mg']])
        self.corrupt_cell_fixture(chunks,1,'合成错字片')
        with patch('agent.agent_backend.services.filing_paddle_backend.enabled',return_value=True),patch('agent.agent_backend.services.filing_paddle_backend.parse_pdf',return_value=chunks):
            self.request('post',self.root+'/application-form/import',data={'file':(io.BytesIO(pdf),'独立申请表.pdf')})
        form=self.request('get',self.root+'/application-form');formurl=self.root+f"/parse-review/application_form/{form['original_file_id']}"
        formctx=self.context(formurl);target=next(t for t in formctx['targets'] if t['value']=='合成错字片')
        self.save(formurl,formctx,[self.item(target,'合成测试片')])
        texts={
            'application_info':'申请信息\n药品通用名称：合成测试片\n变更事项：延长有效期\n规格：3 mg',
            '1':'药品注册证书\n药品通用名称：合成测试片\n批准文号：国药准字H20260001\n有效期：24个月',
            '2':'药品生产许可证\n生产地址：测试一路\n生产范围：片剂（仅限口服）',
            '4':'质量标准及说明书\n药品通用名称：合成测试片\n规格：3 mg\n贮藏：密封\n有效期：24个月',
            '5':'药学研究资料\n批号：TEST001\n长期稳定性：25℃，24个月\n含量：99.0%\n结论：符合标准'}
        scope_url=None
        for category,text in texts.items():
            data=self.request('post',self.root+'/submissions/upload',data={'files':(io.BytesIO(text.encode()),f'独立材料{category}.txt'),'material_category':category})
            doc=data['created'][0]['doc_id'];self.request('post',self.root+f'/submissions/{doc}/parse')
            self.request('post',self.root+f'/submissions/{doc}/metadata',{'material_category':category,'auto_classified':False})
            if category=='2':scope_url=self.root+f'/parse-review/submission/{doc}'
        ctx=self.context(scope_url);value='药品生产许可证\n生产地址：测试一路\n生产范围：片剂（仅限外用，不得口服）'
        self.save(scope_url,ctx,[self.item(ctx['targets'][0],value)])
        ready=self.request('get',self.root+'/review/readiness')
        self.assertTrue(ready['ready'],ready['blocking_issues'])
        started=self.request('post',self.root+'/review/start',{})
        for _ in range(200):
            response=self.client.get(self.root+'/report')
            if response.status_code==200 and json.loads(response.data)['data'].get('report_id'):break
            time.sleep(.05)
        self.assertEqual(response.status_code,200,response.get_data(as_text=True))
        report=json.loads(response.data)['data']
        self.assertIn('合成测试片',report['report_content'])
        self.assertNotIn('合成错字片',report['report_content'])
        self.assertIn('仅限外用',json.dumps(report,ensure_ascii=False))
        self.assertIn('不得口服',json.dumps(report,ensure_ascii=False))
        evidence=os.getenv('REVIEW_EVIDENCE_ROOT')
        if evidence:
            directory=Path(evidence);directory.mkdir(parents=True,exist_ok=True)
            (directory/'positive-flow.json').write_text(json.dumps(dict(test_only=True,not_real_material_acceptance=True,report=report,form=self.request('get',self.root+'/application-form'),scope=self.context(scope_url)),ensure_ascii=False,indent=2))
        # 报告保留当时内容；后续修订只使其过期，不覆盖历史报告。
        fresh=self.context(scope_url);self.save(scope_url,fresh,[self.item(fresh['targets'][0],value.replace('外用','研究'))])
        later=self.request('get',self.root+'/report')
        self.assertIn('仅限外用',json.dumps(later,ensure_ascii=False))
        self.assertEqual(later['input_state']['status'],'stale')

    def test_two_issues_only_one_resolved_and_concurrent_actor(self):
        from copy import deepcopy
        pdf,_=self.pdf_fixture([['测试原文','测试原文']])
        lines=[{'text':'旧识别甲','bbox':[30,30,180,50]},{'text':'旧识别乙','bbox':[30,80,180,100]}]
        chunks=[dict(page=1,page_bbox=[0,0,600,800],status='partial',tables=[],words=deepcopy(lines),lines=lines,
            text='旧识别甲\n旧识别乙',raw_text='旧识别甲\n旧识别乙',errors=[{'code':'ocr_quality','bbox_pdf':l['bbox']} for l in lines])]
        data=self.request('post',self.root+'/submissions/upload',data={'files':(io.BytesIO(pdf),'两处问题.pdf'),'material_category':'2'})
        doc=data['created'][0]['doc_id'];url=self.root+f'/parse-review/submission/{doc}'
        with patch.object(self.service.material_service,'parse_file',return_value=chunks):self.request('post',self.root+f'/submissions/{doc}/parse')
        ctx=self.context(url);first=ctx['issues'][0]
        item=dict(issue_key=first['issue_key'],action='correct_text',text='独立测试甲',reason='独立测试模拟核对')
        self.save(url,ctx,[item])
        fresh=self.context(url)
        self.assertEqual(sum(not i['resolved'] for i in fresh['issues']),1)
        self.assertFalse(self.request('get',self.root+'/review/readiness')['ready'])
        other=self.app.test_client()
        changed={**item,'text':'独立测试甲补字'}
        response=other.post(url,json=dict(source_identity=fresh['source_identity'],expected_revision=fresh['revision'],items=[changed]))
        self.assertEqual(response.status_code,200,response.get_data(as_text=True))
        newer=self.context(url);self.assertEqual(newer['items'][0]['reviewer']['id'],'unattributed')
        self.save(url,fresh,[item],409)

    def test_word_logical_table_no_fake_coordinates(self):
        from docx import Document
        from agent.agent_backend.services.filing_change_submission_parser import parse_filing_change_docx
        docx=Document();table=docx.add_table(rows=2,cols=2)
        for ri,row in enumerate([['生产地址','生产范围'],['测试二路','片剂（仅限口服）']]):
            for ci,value in enumerate(row):table.cell(ri,ci).text=value
        file=io.BytesIO();docx.save(file)
        data=self.request('post',self.root+'/submissions/upload',data={'files':(io.BytesIO(file.getvalue()),'逻辑表格.docx'),'material_category':'2'})
        doc=data['created'][0]['doc_id']
        # 强制复用产品已有的结构化Word解析器，输入仍为本次上传的真实测试文件。
        with patch.object(self.service.material_service,'parse_file',side_effect=parse_filing_change_docx):
            self.request('post',self.root+f'/submissions/{doc}/parse')
        url=self.root+f'/parse-review/submission/{doc}';ctx=self.context(url)
        target=next(t for t in ctx['targets'] if t['value']=='片剂（仅限口服）')
        self.assertEqual(target['bbox_pdf'],[])
        self.save(url,ctx,[self.item(target,'片剂（仅限外用，不得口服）')])
        fresh=self.context(url);table=fresh['effective_chunks'][0]['tables'][0]
        self.assertIn('仅限外用',table['raw_rows'][1][1]);self.assertIn('不得口服',str(table['structured_data']))
        self.assertEqual(fresh['original_chunks'],ctx['original_chunks'])

    def test_no_login_required(self):
        client=self.app.test_client()
        self.assertEqual(client.post(self.base,json={'project_name':'无登录测试'}).status_code,200)
        self.assertEqual(client.get('/api/review-session').status_code,404)


if __name__=='__main__':unittest.main()
