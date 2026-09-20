"""真实隔离数据库验证：上下文失败终止、旧结果保留、限度报告同源。"""
import json
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from test_filing_change_review_service_regression import (
    FilingChangeReviewServiceRegressionTest as Fixture, _RuleList, _ReportCapture,
    _ReviewOrchestrator, filing_change_review_controller as controller,
)
from agent.agent_backend.database.mysql.db_model import FilingChangeReviewRun, FilingChangeReviewResult, FilingChangeReviewReport
from agent.agent_backend.llm.errors import LLMContextError
from agent.agent_backend.services.runtime_task_store import RuntimeTaskStore
from agent.agent_backend.services.filing_change_stability_service import FilingChangeStabilityService
from agent.agent_backend.services.filing_change_submission_parser import _FilingChangeDocxParser
from agent.agent_backend.services.filing_change_report_service import FilingChangeReportService


class ReviewPersistenceAudit(unittest.TestCase):
    setUp = Fixture.setUp
    tearDown = Fixture.tearDown
    _add_project = Fixture._add_project

    def test_context_failure_is_persisted_and_old_result_is_unchanged(self):
        pid='synthetic-budget'
        self._add_project(pid)
        self.service.rule_service=_RuleList()
        self.service.orchestrator=_ReviewOrchestrator()
        self.service.report_service=_ReportCapture(self.root_dir)
        ok, msg, first=self.service.start_review(pid)
        self.assertTrue(ok,msg)
        with self.connection.get_session() as s:
            old=s.query(FilingChangeReviewResult).one().conclusion_json
        err=LLMContextError(actual_chars=11,limit_chars=10,stage='filing_overall_summary')
        err.partial_result={'quality_standard_check':{'result':'需人工确认'}}
        self.service.orchestrator=Mock()
        self.service.orchestrator.run.side_effect=err
        ok,msg,new=self.service.prepare_review_run(pid)
        self.assertTrue(ok,msg)
        report=Mock()
        self.service.report_service=report
        store=RuntimeTaskStore(connection=self.connection,ensure_schema=False)
        store.create_task(task_id='budget-task',domain='filing_change_review',project_id=pid,task_type='review',run_id=new['run_id'])
        actual_thread=threading.Thread
        class InlineThread:
            def __init__(self,target,**kwargs):self.target=target
            def start(self):self.target()
        def isolated_thread(*args, **kwargs):
            if str(kwargs.get('name', '')).startswith('filing-change-review-'):
                return InlineThread(*args, **kwargs)
            return actual_thread(*args, **kwargs)
        with patch.object(controller,'get_service',return_value=self.service),patch.object(controller,'get_runtime_task_store',return_value=store),patch.object(controller.threading,'Thread',isolated_thread):
            self.assertTrue(controller._run_task_async('budget-task',pid,new['run_id']))
        report.generate_report.assert_not_called()
        self.service.orchestrator.run.assert_called_once()
        with self.connection.get_session() as s:
            run=s.query(FilingChangeReviewRun).filter_by(run_id=new['run_id']).one()
            self.assertEqual(run.status,'failed')
            self.assertIsNotNone(run.finished_at)
            failure=json.loads(run.result_json)
            self.assertFalse(failure['review_complete'])
            self.assertEqual(failure['error']['code'],'context_limit_exceeded')
            self.assertEqual(failure['partial_result'],err.partial_result)
            self.assertEqual(s.query(FilingChangeReviewResult).one().conclusion_json,old)
            self.assertEqual(s.query(FilingChangeReviewReport).count(),1)
        task=store.get_task('budget-task')
        self.assertEqual(task['status'],'failed')
        self.assertEqual(task['result'],failure)

    def test_column_units_and_export_use_same_decision(self):
        parser=_FilingChangeDocxParser('synthetic.docx')
        parsed=parser._parse_stability_table([
            ['考察项目','可接受标准（μg/mL）','时间点','检测结果（mg/mL）'],
            ['含量','≤300','0月','0.3'],
            ['含量','≤300','1月','0.31'],
        ],'稳定性')
        service=FilingChangeStabilityService()
        rows=[service._enrich_record(r) for r in parsed['records']]
        self.assertEqual([r['within_standard'] for r in rows],[True,False])
        self.assertEqual(rows[0]['result_unit_source'],'result_column_header')
        self.assertEqual(rows[0]['limit_unit_source'],'limit_column_header')
        pending=service._enrich_record({'indicator':'未知单位','result_text':'1widgets','standard_text':'≤2widgets'})
        unmarked=service._enrich_record({'indicator':'pH','result_text':'5','standard_text':'4～6'})
        rows.extend([pending,unmarked])
        limit=service._build_limit_check(rows,[rows[1]],[pending])
        self.assertEqual(limit['status_code'],'out_of_spec')
        self.assertIn('未进行单位换算',limit['reminder'])
        payload={'stability_trend_analysis':{'limit_check':limit,'records':rows}}
        # 持久化序列化往返不能改变判定；真实Word从同一服务端结果生成。
        payload=json.loads(json.dumps(payload))
        generated=FilingChangeReportService(self.root_dir).generate_report('synthetic','units',payload)
        self.assertIn('限度判断：已超限',generated['markdown'])
        self.assertIn('结果列表头/标准列表头',generated['markdown'])
        doc=Document(generated['word_path'])
        text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for expected in ('不符合限度','待确认','未进行单位换算','结果列表头/标准列表头'):
            self.assertIn(expected,text)


def load_tests(loader, tests, pattern):
    return loader.loadTestsFromTestCase(ReviewPersistenceAudit)


if __name__=='__main__':unittest.main()
