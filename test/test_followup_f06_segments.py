"""F06：明确分段独立归属；真实文件、共享规则和旧入口使用相同业务预期。"""
import json
import os
import unittest
from pathlib import Path

import fitz
from docx import Document

from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
from agent.agent_backend.utils.parser.form_field_semantics import (
    option_evidence, selected_option, selection_evidence,
)


ROOT = Path('/tmp/followup-20260906/f06') / os.environ.get('F06_EVIDENCE_PHASE', 'manual')
FAMILIES = (
    ('dosage', '剂型', 'item_11_dosage_form', ('片剂', '胶囊剂'), ('片剂', '胶囊剂'), True),
    ('matter', '申请事项分类', 'item_5_application_matter_category',
     ('6.8 变更有效期和贮藏条件', '6.9 增加规格'), ('1.7', '1.8'), True),
    ('otc', '是否为OTC', 'item_3_otc_type', ('处方药', '非处方药'), ('处方药', '非处方药'), False),
    ('boolean', '受理前药品注册检验', 'item_19_pre_acceptance_inspection', ('是', '否'), ('是', '否'), False),
    ('use', '商品名称', 'item_10_trade_name', ('使用', '不使用'), ('使用', '不使用'), False),
)
# 预期由输入事实写定；多勾仅对单选字段构成冲突。
CASES = (
    ('original', '☑{a}；{b}□', (0,), 'confirmed'),
    ('reverse_selection', '□{a}；{b}☑', (1,), 'confirmed'),
    ('reverse_direction', '{a}☑；□{b}', (0,), 'confirmed'),
    ('unchecked', '□{a}；{b}□', (), 'unconfirmed'),
    ('both', '☑{a}；{b}☑', (0, 1), 'confirmed'),
    ('ambiguous_middle', '{a} ☑ {b}', (), 'ambiguous'),
    ('duplicate_conflict', '☑{a}；{a}□', (), 'conflict'),
    ('ascii_semicolon', '{a}☑;□{b}', (0,), 'confirmed'),
    ('newline', '☑{a}\n{b}□', (0,), 'confirmed'),
    ('parenthesized_marks', '(x){a}；{b}( )', (0,), 'confirmed'),
    ('punctuation', '☑{a}，；{b}□。', (0,), 'confirmed'),
)


def word_file(path, title, text, paragraph=False):
    doc = Document()
    if paragraph:
        doc.add_paragraph('药品通用名称：测试药品')
        doc.add_paragraph(title + '：' + text)
    else:
        for label, value in (('药品通用名称', '测试药品'), (title, text)):
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = label, value
    doc.save(path)


def pdf_file(path, title, text, bordered=False):
    with fitz.open() as doc:
        page = doc.new_page(width=950, height=1100)
        page.insert_font(fontname='formfont', fontfile=os.environ.get('FILING_TEST_FONT', '/System/Library/Fonts/Supplemental/Arial Unicode.ttf'))
        for y, label, value in ((45, '药品通用名称', '测试药品'), (105, title, text)):
            page.insert_text((35, y), label, fontname='formfont', fontsize=12)
            page.insert_text((235, y), value, fontname='formfont', fontsize=12)
            if bordered:
                page.draw_rect(fitz.Rect(25, y - 20, 225, y + 35))
                page.draw_rect(fitz.Rect(225, y - 20, 920, y + 35))
        doc.save(path)


class SegmentTests(unittest.TestCase):
    def setUp(self):
        ROOT.mkdir(parents=True, exist_ok=True)
        self.parser = FilingChangeFormParserService()

    def save_result(self, result):
        (ROOT / (self._testMethodName + '.json')).write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')

    def check_file(self, family, case, ext, *, paragraph=False, bordered=False):
        _, title, key, labels, mapped, multiple = family
        _, template, indices, status = case
        text = template.format(a=labels[0], b=labels[1])
        path = ROOT / (self._testMethodName + '.' + ext)
        if ext == 'docx':
            word_file(path, title, text, paragraph)
        else:
            pdf_file(path, title, text, bordered)
        result = self.parser.parse_form_file(str(path))
        field = result['form_json'][key]
        self.save_result({'input': text, 'field': field, 'content_status': result['content_status']})
        if not multiple and len(indices) > 1:
            indices, status = (), 'conflict'
        expected = [mapped[i] for i in indices]
        self.assertEqual(field['selected_values'] if multiple else field['value'],
                         expected if multiple else (expected[0] if expected else ''))
        self.assertEqual(field['selection_status'], status)
        if status in ('ambiguous', 'conflict'):
            self.assertTrue(field.get('semantic_issues'))
        elif status == 'confirmed':
            self.assertFalse(field.get('semantic_issues'))
        else:
            # 全未勾选的必要事实是两个候选均为 False，不限定整表空白提示策略。
            self.assertEqual([e['selected'] for e in field['selection_evidence']], [False, False])
        self.assertTrue(field['source_text'])
        self.assertTrue(field['source_regions'])
        self.assertEqual(field['source_file'], path.name)
        for entry in field['selection_evidence']:
            self.assertTrue(entry['source_text'])
            self.assertTrue(entry['source_regions'])
            self.assertEqual(entry['source_file'], path.name)
        if status == 'confirmed':
            self.assertEqual(result['content_status'], 'success')
        if family[0] == 'matter' and case[0] == 'original':
            self.assertEqual([(e['code'], e['name'], e['selected']) for e in field['original_matter']],
                             [('6.8', '变更有效期和贮藏条件', True), ('6.9', '增加规格', False)])

    def test_no_marker_can_cross_semicolon_or_cell(self):
        for source in ('片剂；☑；胶囊剂', ['片剂', '☑', '胶囊剂']):
            result = selection_evidence(source, multiple=True)
            self.assertEqual(result['selected_values'], [])
            self.assertEqual(result['status'], 'ambiguous')

    def test_unmarked_segment_stays_unconfirmed(self):
        result = selection_evidence('☑片剂；胶囊剂', multiple=True)
        self.assertEqual(result['selected_values'], ['片剂'])
        self.assertIsNone(result['entries'][1]['selected'])

    def test_real_ambiguity_is_local_and_retained(self):
        result = selection_evidence('☑片剂；胶囊剂 ☑ 颗粒剂', multiple=True)
        self.assertEqual(result['selected_values'], ['片剂'])
        self.assertEqual(result['status'], 'ambiguous')
        self.assertEqual([e['status'] for e in result['entries']], ['selected', 'ambiguous', 'ambiguous'])
        self.assertEqual(selection_evidence('☑片剂；胶囊剂 ☑ 颗粒剂')['selected_values'], [])

    def test_duplicate_contradiction_does_not_erase_other_confirmed_choices(self):
        result = selection_evidence('☑片剂；片剂□；☑胶囊剂', multiple=True)
        self.assertEqual(result['selected_values'], ['胶囊剂'])
        self.assertEqual(result['status'], 'conflict')
        self.assertEqual([e['status'] for e in result['entries']], ['conflict', 'conflict', 'selected'])

    def test_normal_punctuation_and_parentheses_belong_to_names(self):
        for name in ('片剂（薄膜衣）', '片剂(薄膜衣)', '变更生产场地，境外生产药品', '片剂（说明；薄膜衣）'):
            result = selection_evidence('☑' + name + '；胶囊剂□', multiple=True)
            self.assertEqual(result['selected_values'], [name])
            self.assertEqual(result['status'], 'confirmed')

    def test_wrapped_matter_and_segment_boundary(self):
        for source in ('6.8 变更有效期和☑\n贮藏条件；□6.9 增加规格',
                       '☑6.8 变更有效期和\n贮藏条件；6.9 增加规格□'):
            result = selection_evidence(source, multiple=True)
            self.assertEqual(result['selected_values'], ['变更有效期和\n贮藏条件'])
            self.assertEqual(result['status'], 'confirmed')


def file_case(family, case, ext, **kwargs):
    def test(self):
        self.check_file(family, case, ext, **kwargs)
    return test


def rule_case(family, case):
    def test(self):
        _, _, _, labels, _, multiple = family
        _, template, indices, status = case
        source = template.format(a=labels[0], b=labels[1])
        names = [label.split(' ', 1)[1] if family[0] == 'matter' else label for label in labels]
        expected = [names[i] for i in indices]
        if not multiple and len(indices) > 1:
            expected, status = [], 'conflict'
        result = selection_evidence(source, None if multiple else labels, multiple=multiple)
        self.save_result({'input': source, 'evidence': result})
        self.assertEqual(result['selected_values'], expected)
        self.assertEqual(result['status'], status)
        if not multiple:
            self.assertEqual(option_evidence(source, labels), result)
            self.assertEqual(selected_option(source, labels), expected[0] if expected else '')
            if family[0] == 'boolean':
                self.assertEqual(self.parser._normalize_loose_boolean_text(source), expected[0] if expected else '')
        else:
            for method in (self.parser._extract_checked_options, self.parser._extract_pdf_marked_labels):
                self.assertEqual(method(source), expected)
            if family[0] == 'matter':
                self.assertEqual(self.parser._extract_matter_codes_from_pdf_text(source),
                                 [family[4][i] for i in indices])
    return test


def xml_case(family, case):
    def test(self):
        labels = family[3]
        source = case[1].format(a=labels[0], b=labels[1])
        expected = [labels[i].split(' ', 1)[1] if family[0] == 'matter' else labels[i]
                    for i in case[2]]
        path = ROOT / (self._testMethodName + '.docx')
        doc = Document()
        doc.add_paragraph(source)
        doc.save(path)
        self.assertEqual(self.parser._extract_checked_labels_from_wsym(doc._element.xml), expected)
        self.assertEqual(self.parser._extract_checked_labels_from_docx_xml(str(path)), expected)
    return test


for family in FAMILIES:
    for case in CASES:
        setattr(SegmentTests, f'test_rule_{family[0]}_{case[0]}', rule_case(family, case))
        for ext in ('docx', 'pdf'):
            setattr(SegmentTests, f'test_{ext}_{family[0]}_{case[0]}', file_case(family, case, ext))
    setattr(SegmentTests, f'test_paragraph_{family[0]}', file_case(family, CASES[0], 'docx', paragraph=True))
    setattr(SegmentTests, f'test_bordered_pdf_{family[0]}', file_case(family, CASES[0], 'pdf', bordered=True))
    for case in (CASES[0], CASES[2]):
        setattr(SegmentTests, f'test_xml_{family[0]}_{case[0]}', xml_case(family, case))


if __name__ == '__main__':
    unittest.main()
