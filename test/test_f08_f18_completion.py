"""F08–F18：独立输入事实，真实文件/解析/任务/SQLite；同步点仅控制时序。"""
import io
import multiprocessing
import json
import os
import tempfile
import threading
import time
import unittest
from datetime import datetime
from zoneinfo import ZoneInfo
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
from unittest.mock import Mock

import fitz
from docx import Document
from docx.oxml import parse_xml
from flask import Flask
from agent.test.test_four_semantic_table_defects import word_file, pdf_file
from agent.test.test_filing_change_review_service_regression import _TestConnection, _Upload, filing_change_review_controller as controller
from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
from agent.agent_backend.services.filing_change_review_service import FilingChangeReviewService
from agent.agent_backend.services.filing_change_material_service import FilingChangeMaterialService
from agent.agent_backend.services.runtime_task_store import RuntimeTaskStore
from agent.agent_backend.services.filing_form_revision import merge_parse, save_revision, validity_cross_checks
from agent.agent_backend.utils.parser.pdf_page_extractor import extract_pdf_pages
from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import OCRServiceClient


def service_at(root, connection=None):
    service = object.__new__(FilingChangeReviewService)
    service.root_dir = Path(root) / 'review'; service.root_dir.mkdir(exist_ok=True)
    service.db_conn = connection or _TestConnection(Path(root) / 'db.sqlite')
    service.form_parser = FilingChangeFormParserService()
    service.material_service = FilingChangeMaterialService()
    return service


def review_inputs(service,pid):
    # 仅截获送入远程审评生成器前的真实取数，不伪造任何解析字段或表格。
    from agent.test.test_filing_change_review_service_regression import _RuleList, _ReviewOrchestrator, _ReportCapture
    service.rule_service=_RuleList();service.orchestrator=Mock()
    service.orchestrator.run.return_value=_ReviewOrchestrator.run()
    service.report_service=_ReportCapture(service.root_dir)
    ok,message,run=service.prepare_review_run(pid)
    if not ok:raise AssertionError(message)
    ok,message,_=service.execute_review_run(pid,run['run_id'])
    if not ok:raise AssertionError(message)
    return service.orchestrator.run.call_args.kwargs


def concurrent_request(root, pid, barrier, release, responses, calls):
    service=service_at(root); store=RuntimeTaskStore(connection=service.db_conn)
    controller._service=service;controller._runtime_task_store=store
    app=Flask('parallel-request');app.register_blueprint(controller.filing_change_review_bp)
    real=service.form_parser.parse_form_file
    def paused(path):
        calls.put(os.getpid())
        if not release.wait(25):raise RuntimeError('parent did not release parser')
        return real(path)
    service.form_parser.parse_form_file=paused
    barrier.wait(15)
    with app.test_client() as client:
        response=client.post('/filing-change-review/projects/'+pid+'/application-form/import/start',
            data={'file':(io.BytesIO((Path(root)/'form.docx').read_bytes()),'form.docx')})
        responses.put((response.status_code,json.loads(response.get_data(as_text=True))))
    release.wait(25)
    deadline=time.monotonic()+15
    while time.monotonic()<deadline and any(t['status'] in ('pending','running') for t in store.list_project_parse_tasks(pid,'filing_change_review')):time.sleep(.02)
    service.db_conn.engine.dispose()


def checkbox(paragraph, label, checked=True, glyph='☒', postfix=False, sym=False):
    display = ('<w:r><w:sym w:font="Wingdings 2" w:char="F052"/></w:r>' if sym else
               '<w:r><w:t>'+glyph+'</w:t></w:r>')
    text = '<w:r><w:t>'+label+'</w:t></w:r>'
    paragraph._p.append(parse_xml('<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"><w:sdtPr><w14:checkbox>'
        '<w14:checked w14:val="'+('1' if checked else '0')+'"/></w14:checkbox></w:sdtPr><w:sdtContent>'
        +(text+display if postfix else display+text)+'</w:sdtContent></w:sdt>'))


def table_pdf(path, title='Assay results', uncertain=False, scan=False, ordinary=False):
    with fitz.open() as src, fitz.open() as dst:
        p=src.new_page(width=600,height=400)
        if title: p.insert_text((40,40),title,fontname='china-s' if '结果' in title else 'helv',fontsize=14)
        rows = [(80,'Sample','Assay (mg)','Recovery (%)'),(120,'Alpha','-1.25','95%'),
                (170,'Batch','','98.0%'),(220,'Gamma','50.0','50.0')]
        if ordinary:
            rows=[(80+i*35,'Paragraph describes a study.','Another narrative column.','Continued text.') for i in range(4)]
        for y,a,b,c in rows:
            for x,t in [(40,a),(190+y/3 if uncertain and y>80 else 240,b),(280+y/4 if uncertain and y>80 else 430,c)]:
                if t:p.insert_text((x,y),t,fontsize=12)
        if not ordinary:
            p.insert_text((40,186),'two',fontsize=12)
            p.insert_text((40,260),'* Note: values in mg; retention 0.5 min.',fontsize=11)
        if scan:dst.new_page(width=600,height=400).insert_image(p.rect,stream=p.get_pixmap(matrix=fitz.Matrix(3,3)).tobytes('png'))
        (dst if scan else src).save(path)


class FieldsAndPersistence(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name); self.parser=FilingChangeFormParserService()
        from agent.agent_backend.config.settings import settings
        self.addCleanup(patch.stopall)
        patch.object(settings,'ocr_service_url',os.environ.get('T2_OCR_URL','http://127.0.0.1:18765')).start()

    def test_checkbox_controls(self):
        cases=[('是否为OTC','item_3_otc_type','非处方药','处方药'),
               ('受理前药品注册检验','item_19_pre_acceptance_inspection','是','否'),
               ('商品名称','item_10_trade_name','使用','不使用'),
               ('剂型','item_11_dosage_form','片剂','胶囊剂'),
               ('申请事项分类','item_5_application_matter_category','6.8 变更有效期和贮藏条件','6.9 增加规格')]
        for title,key,yes,no in cases:
            for glyph in ('☒','☑','✓','✔','■','▣'):
                for postfix in (True,False):
                    with self.subTest(title=title,glyph=glyph,postfix=postfix):
                        doc=Document(); t=doc.add_table(rows=1,cols=2);t.cell(0,0).text=title
                        p=t.cell(0,1).paragraphs[0]
                        checkbox(p,yes,glyph=glyph,postfix=postfix)
                        checkbox(p,no,False,'□',not postfix)
                        path=self.root/'control.docx';doc.save(path)
                        f=self.parser.parse_form_file(str(path))['form_json'][key]
                        self.assertEqual(f.get('selection_status','confirmed'),'confirmed')
                        self.assertFalse(f.get('semantic_issues'), f)
                        if key.endswith(('otc_type','inspection','trade_name')): self.assertEqual(f['value'],yes)
                        else:self.assertTrue(f['selected_values'])

    def test_checkbox_real_conflict_and_symbol(self):
        for mismatch in (True,False):
            doc=Document();t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='是否为OTC';p=t.cell(0,1).paragraphs[0]
            checkbox(p,'处方药',True,'□' if mismatch else '☒',True)
            checkbox(p,'非处方药',not mismatch,'□' if mismatch else '☒',False)
            path=self.root/'conflict.docx';doc.save(path)
            f=self.parser.parse_form_file(str(path))['form_json']['item_3_otc_type']
            self.assertEqual(f['value'],'');self.assertTrue(f['semantic_issues'])
        doc=Document();t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='剂型'
        checkbox(t.cell(0,1).paragraphs[0],'片剂',sym=True,postfix=True)
        path=self.root/'symbol.docx';doc.save(path)
        self.assertEqual(self.parser.parse_form_file(str(path))['form_json']['item_11_dosage_form']['selected_values'],['片剂'])

    def test_control_name_and_glyph_share_a_text_run(self):
        for postfix in (True,False):
            doc=Document();t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='是否为OTC';p=t.cell(0,1).paragraphs[0]
            for label,state,glyph in [('非处方药','1','☑'),('处方药','0','□')]:
                value=label+glyph if postfix else glyph+label
                p._p.append(parse_xml('<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"><w:sdtPr><w14:checkbox><w14:checked w14:val="'+state+'"/></w14:checkbox></w:sdtPr><w:sdtContent><w:r><w:t>'+value+'</w:t></w:r></w:sdtContent></w:sdt>'))
            path=self.root/'same-run.docx';doc.save(path)
            f=self.parser.parse_form_file(str(path))['form_json']['item_3_otc_type']
            self.assertEqual(f['value'],'非处方药');self.assertFalse(f.get('semantic_issues'))

    def test_prescription_control_conflict_cannot_leave_a_yes_value(self):
        doc=Document();t=doc.add_table(rows=1,cols=3)
        t.cell(0,0).text='处方';t.cell(0,1).text='是否有变更'
        checkbox(t.cell(0,2).paragraphs[0],'是',True,'□')
        doc.add_paragraph('药品通用名称：验证用药品')
        path=self.root/'prescription-conflict.docx';doc.save(path)
        f=self.parser.parse_form_file(str(path))['form_json']['item_16_prescription']
        self.assertEqual(f['sub_fields']['has_change'],'');self.assertTrue(f['semantic_issues'])

    def test_all_material_rows_are_retained(self):
        doc=Document();doc.add_paragraph('原辅包来源');t=doc.add_table(rows=56,cols=4)
        rows=[['物料名称','登记号','受理号','供应商']]+[[f'物料{i}',f'{i:05d}','','供应商甲'] for i in range(55)]
        for r,vs in zip(t.rows,rows):
            for c,v in zip(r.cells,vs):c.text=v
        path=self.root/'materials55.docx';doc.save(path)
        f=self.parser.parse_form_file(str(path))['form_json']['item_17_material_source']
        self.assertEqual(len(f['table_rows']),55)
        self.assertEqual(f['table_rows'][-1]['register_no'],'00054')
        service=service_at(self.root)
        try:
            pid=service.create_project({'project_name':'完整明细持久化'})[2]['project_id']
            self.assertTrue(service.import_application_form(pid,_Upload(path.name,path.read_bytes()))[0])
            fresh=service_at(self.root,service.db_conn)
            self.assertEqual(fresh.get_application_form(pid)[2]['form_json']['item_17_material_source']['table_rows'],f['table_rows'])
            self.assertTrue(fresh.parse_application_form(pid)[0])
            self.assertEqual(review_inputs(fresh,pid)['application_form']['form_json']['item_17_material_source']['table_rows'],f['table_rows'])
        finally:service.db_conn.engine.dispose()

    def test_material_continuations_stop_at_next_item(self):
        doc=Document();doc.add_paragraph('原辅包来源')
        rows=[['乙醇','00123','','甲供应商'],['包装瓶','','00002','乙供应商']]
        for values in rows:
            t=doc.add_table(rows=2,cols=4)
            for r,vals in zip(t.rows,[['物料名称','登记号','受理号','供应商'],values]):
                for c,v in zip(r.cells,vals):c.text=v
        doc.add_paragraph('其他相关情况')
        t=doc.add_table(rows=2,cols=4)
        for r,vals in zip(t.rows,[['物料名称','登记号','受理号','供应商'],['其他事项','999','888','无关供应商']]):
            for c,v in zip(r.cells,vals):c.text=v
        path=self.root/'continuation.docx';doc.save(path)
        f=self.parser.parse_form_file(str(path))['form_json']['item_17_material_source']
        self.assertEqual(f['table_rows'],[dict(zip(['material_name','register_no','accept_no','manufacturer'],r)) for r in rows])
        self.assertTrue(f['source_regions'])

    def test_party_explicit_fields_save_refresh_reparse(self):
        rows=[['药品通用名称','角色测试药品'],['药品注册申请人','中文名称','甲公司'],
              ['药品注册申请人','法定代表人','法人甲'],['药品注册申请人','联系人','联系乙'],
              ['药品注册申请人','注册地址','注册丙址'],['制剂生产企业','中文名称','丁公司'],
              ['制剂生产企业','注册地址','企业注册戊址'],['制剂生产企业','生产地址','生产己址'],
              ['委托研究机构','中文名称','庚研究所'],['委托研究机构','研究负责人','负责人辛'],
              ['委托研究机构','地址','研究壬址'],['委托研究机构','联系电话','01012345678'],
              ['商品名称','使用☑；商品名：明确商品名'],['药品注册分类','化学药品；分类号：2.2'],
              ['受理前药品注册检验','是☑；检品编号：ABC123'],['其他相关情况','已有三次申请记录'],
              ['其他特别申明事项','明确声明正文。第二句保留。'],['处方','活性成分：测试成分；辅料：乙醇；是否有变更：否']]
        expected={30:{'applicant_name':'甲公司','legal_representative':'法人甲','contact':'联系乙','registered_address':'注册丙址'},
                  31:{'manufacturer_name':'丁公司','registered_address':'企业注册戊址','production_address':'生产己址'},
                  32:{'organization_name':'庚研究所','research_lead':'负责人辛','address':'研究壬址','phone':'01012345678'},
                  10:{'trade_name':'明确商品名'},2:{'class_no':'2.2'},19:{'sample_inspection_no':'ABC123'},
                  16:{'active_ingredients':'测试成分','excipients':'乙醇','has_change':'否'}}
        for ext in ('docx','pdf'):
            with self.subTest(ext=ext):
                path=self.root/('roles.'+ext);(word_file if ext=='docx' else pdf_file)(path,rows)
                service=service_at(self.root);pid=service.create_project({'project_name':'隔离字段'})[2]['project_id']
                self.assertTrue(service.import_application_form(pid,_Upload(path.name,path.read_bytes()))[0])
                fresh=service_at(self.root,service.db_conn)
                form=fresh.get_application_form(pid)[2]['form_json']
                for number,fields in expected.items():
                    key=next(k for k in form if k.startswith('item_'+str(number)+'_'))
                    for field,value in fields.items():self.assertEqual(form[key]['sub_fields'][field],value,(ext,key,field))
                self.assertEqual(form['item_29_other_related_info']['value'],'已有三次申请记录')
                self.assertEqual(form['special_statement']['value'],'明确声明正文。第二句保留。')
                edited=deepcopy(form);edited['item_30_applicant_info']['sub_fields']['contact']='人工联系人'
                self.assertTrue(fresh.save_application_form(pid,{'form_json':edited})[0])
                self.assertTrue(fresh.parse_application_form(pid)[0])
                result=service_at(self.root,service.db_conn).get_application_form(pid)[2]['form_json']
                self.assertEqual(result['item_30_applicant_info']['sub_fields']['contact'],'人工联系人')
                self.assertEqual(result['item_30_applicant_info']['sub_fields']['legal_representative'],'法人甲')
                self.assertEqual(result['item_30_applicant_info']['candidates']['sub_fields.contact']['value'],'联系乙')
                self.assertEqual(review_inputs(fresh,pid)['application_form']['form_json'],result)
                service.db_conn.engine.dispose()

    def test_legacy_party_values_are_not_guessed(self):
        service=service_at(self.root)
        try:
            pid=service.create_project({'project_name':'历史字段'})[2]['project_id']
            form=self.parser.normalize_form_json({})
            form['item_30_applicant_info']['sub_fields'].update(contact='旧人员',address='旧地址')
            self.assertTrue(service.save_application_form(pid,{'form_json':form})[0])
            data=service.get_application_form(pid)[2]
            # 真正历史数据没有逐值来源；直接写入隔离旧记录以覆盖兼容读取。
            from agent.agent_backend.database.mysql.db_model import FilingChangeApplicationForm
            session=service.db_conn.get_session();row=session.query(FilingChangeApplicationForm).filter_by(project_id=pid).one()
            old=json.loads(row.form_json);old['item_30_applicant_info'].pop('value_sources',None)
            row.form_json=json.dumps(old,ensure_ascii=False);session.commit();session.close()
            data=service_at(self.root,service.db_conn).get_application_form(pid)[2]
            self.assertEqual(len(data['legacy_role_notes']),2)
            sub=data['form_json']['item_30_applicant_info']['sub_fields']
            self.assertEqual(sub['contact'],'旧人员');self.assertEqual(sub['address'],'旧地址')
            self.assertEqual(sub['legal_representative'],'');self.assertEqual(sub['registered_address'],'')
            self.assertEqual(review_inputs(service,pid)['application_form']['legacy_role_notes'],data['legacy_role_notes'])
            # 旧自动字段在新文件替换后进入历史，仍不得把不明角色改写成确定来源。
            session=service.db_conn.get_session();row=session.query(FilingChangeApplicationForm).filter_by(project_id=pid).one()
            old=json.loads(row.form_json)
            old['item_30_applicant_info'].update(manual_modified=False,manual_paths=[])
            row.form_json=json.dumps(old,ensure_ascii=False);session.commit();session.close()
            replacement=self.root/'new-role.docx'
            word_file(replacement,[['药品注册申请人','中文名称：甲公司；联系人：新联系人；注册地址：明确注册地址']])
            self.assertTrue(service.import_application_form(pid,_Upload(replacement.name,replacement.read_bytes()))[0])
            final=service_at(self.root,service.db_conn).get_application_form(pid)[2]['form_json']['item_30_applicant_info']
            history=[x for x in final['value_history'] if x.get('value') in ('旧人员','旧地址')]
            self.assertEqual(len(history),2)
            self.assertTrue(all(x['recognition_status']=='legacy_unspecified' and x['reason'] for x in history))
            self.assertEqual(final['sub_fields']['contact'],'新联系人')
            self.assertEqual(final['sub_fields']['registered_address'],'明确注册地址')
        finally:service.db_conn.engine.dispose()

    def test_year_shorthand_and_manual_equivalent_candidate(self):
        path=self.root/'years.docx';word_file(path,[['药品有效期','原有效期2年、拟变更为3年']])
        f=self.parser.parse_form_file(str(path))['form_json']
        field=f['item_15_validity_period'];self.assertEqual(field['sub_fields']['original_validity_period'],'2年')
        self.assertEqual(field['sub_fields']['proposed_validity_period'],'3年')
        edited=deepcopy(f);edited['item_15_validity_period']['sub_fields']['original_validity_period']='24个月'
        merged=merge_parse(save_revision(f,edited),f,mode='reparse')
        self.assertNotIn('sub_fields.original_validity_period',merged['item_15_validity_period']['candidates'])
        self.assertEqual(merged['item_15_validity_period']['sub_fields']['original_validity_period'],'24个月')

    def test_equivalent_approval_and_revision_use_distinct_time_roles(self):
        service=service_at(self.root)
        try:
            pid=service.create_project({'project_name':'期限对象与跨文件等价'})[2]['project_id']
            path=self.root/'application.docx'
            word_file(path,[['药品有效期','原有效期24个月，拟延长后有效期36个月']])
            self.assertTrue(service.import_application_form(pid,_Upload(path.name,path.read_bytes()))[0])
            for name,text in [('批准文件.pdf','药品有效期：2年'),('修订说明.docx','药品有效期由2年延长至3年；复验有效期由3个月延长至6个月')]:
                p=self.root/name;(pdf_file if p.suffix=='.pdf' else word_file)(p,[[text]])
                ok,_,data=service.upload_submission_files(pid,[_Upload(name,p.read_bytes())],'1')
                self.assertTrue(ok)
                self.assertTrue(service.parse_submission(pid,data['created'][0]['doc_id'])[0])
            self.assertTrue(service.parse_application_form(pid)[0])
            fresh=service_at(self.root,service.db_conn)
            f=fresh.get_application_form(pid)[2]['form_json']['item_15_validity_period']
            checks=f['cross_checks'];self.assertTrue(checks)
            self.assertTrue(all(c['status']=='一致' for c in checks),checks)
            self.assertEqual({c['source_file'] for c in checks},{'批准文件.pdf','修订说明.docx'})
            self.assertTrue(all(c['source_regions'] and c['source_text'] for c in checks))
            self.assertEqual(f['sub_fields']['original_validity_period'],'24个月')
            self.assertEqual(f['sub_fields']['proposed_validity_period'],'36个月')
            self.assertEqual(review_inputs(fresh,pid)['application_form']['form_json']['item_15_validity_period'],f)
        finally:service.db_conn.engine.dispose()


class NativeTables(unittest.TestCase):
    def setUp(self):
        from agent.agent_backend.config.settings import settings
        self.addCleanup(patch.stopall)
        patch.object(settings,'ocr_service_url',os.environ.get('T2_OCR_URL','http://127.0.0.1:18765')).start()
        patch.object(settings,'ocr_lang','eng').start()

    def test_public_filing_ctd_and_knowledge_entries(self):
        from agent.agent_backend.utils.parser.ctd_paser import parse_pdf_to_markdown_json
        from agent.agent_backend.memory.rag.document import DocumentProcessor
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);service=service_at(root)
            try:
                for uncertain,scan in ((False,False),(True,False),(False,True),(True,True)):
                    path=root/'public.pdf';table_pdf(path,title='3.2.P.5.1 Specification',uncertain=uncertain,scan=scan)
                    pid=service.create_project({'project_name':'公共结构入口'})[2]['project_id']
                    ok,_,upload=service.upload_submission_files(pid,[_Upload(path.name,path.read_bytes())],'5')
                    self.assertTrue(ok);docid=upload['created'][0]['doc_id']
                    self.assertTrue(service.parse_submission(pid,docid)[0])
                    fresh=service_at(root,service.db_conn)
                    data=fresh.get_submission_parsed_markdown(pid,docid)[2]
                    status='partial' if uncertain else 'success'
                    self.assertEqual(data['parse_diagnostics']['status'],status)
                    self.assertIn('Alpha',data['markdown'])
                    if uncertain:self.assertTrue(data['parse_diagnostics']['errors'])
                    else:self.assertIn('-1.25',data['markdown'])
                    ctd=parse_pdf_to_markdown_json(str(path),[{'section_id':'3.2.p.5.1','section_name':'Specification'}])
                    self.assertEqual(ctd['parse_diagnostics']['status'],status)
                    processor=DocumentProcessor();chunks=processor.chunk_document(processor.parse_to_document(str(path)))
                    self.assertTrue(chunks)
                    self.assertEqual(chunks[0].metadata['parse_diagnostics']['status'],status)
                    if not uncertain:
                        self.assertEqual(ctd['source_pages'][0]['tables'][0]['rows'][2],['Batch\ntwo','','98.0%'])
                        self.assertEqual(chunks[0].metadata['tables'][0]['rows'][2],['Batch\ntwo','','98.0%'])
                    else:self.assertTrue(chunks[0].metadata['parse_diagnostics']['errors'])
                    selected=review_inputs(service,pid)['submissions'][0]
                    self.assertEqual(selected['parse_status'],status)
                    if uncertain:self.assertTrue(selected['parse_diagnostics']['errors'])
            finally:service.db_conn.engine.dispose()

    def test_titles_sources_and_unresolved(self):
        with tempfile.TemporaryDirectory() as tmp:
            for scan in (False,True):
                for title in ('Table 1. Assay results','Assay results','检测结果',''):
                    for uncertain in (False,True):
                        with self.subTest(scan=scan,title=title,uncertain=uncertain):
                            path=Path(tmp)/'table.pdf';table_pdf(path,title,uncertain,scan)
                            row=extract_pdf_pages(str(path),ocr_client=OCRServiceClient(os.environ.get('T2_OCR_URL','http://127.0.0.1:18765')),lang='eng')[0]
                            if not scan:self.assertEqual(row['ocr_calls'],0)
                            if uncertain:
                                self.assertEqual(row['status'],'partial');self.assertTrue(row['unresolved_tables']);self.assertTrue(row['errors'])
                            else:
                                self.assertEqual(len(row['tables']),1)
                                expected = [['Alpha','-1.25','95%'],['Batch\ntwo','','98.0%'],['Gamma','50.0','50.0']]
                                if scan and row['tables'][0]['rows'][1][1] == '1.25':
                                    # 原已复现负号丢失：新要求保留主结果并明确隔离，不能静默改字。
                                    cell = next(c for c in row['tables'][0]['cells'] if c['row'] == 1 and c['column'] == 1)
                                    self.assertEqual(cell['numeric_status'], 'numeric_uncertain')
                                    self.assertTrue(any(v['status'] == 'numeric_uncertain' and '-1.25' in v['secondary']
                                                        for v in cell['numeric_verification']))
                                    self.assertEqual(row['status'], 'partial')
                                    self.assertTrue(any(e.get('code') == 'numeric_uncertain' for e in row['errors']))
                                    expected[0][1] = '1.25'
                                self.assertEqual(row['tables'][0]['rows'][1:], expected)
                                self.assertEqual(row['tables'][0]['page'],1)
                                self.assertTrue(all(c['bbox_pdf'] for c in row['tables'][0]['cells']))
                                self.assertEqual(row['tables'][0]['source'],'ocr_alignment' if scan else 'native_alignment')

    def test_narrative_not_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'paragraph.pdf';table_pdf(path,ordinary=True)
            row=extract_pdf_pages(str(path))[0]
            self.assertFalse(row['tables']);self.assertFalse(row.get('unresolved_tables'));self.assertEqual(row['status'],'success')


class Tasks(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        self.service=service_at(self.root);self.addCleanup(self.service.db_conn.engine.dispose)
        self.store=RuntimeTaskStore(connection=self.service.db_conn)
        self.addCleanup(patch.stopall)
        patch.object(controller,'_service',self.service).start();patch.object(controller,'_runtime_task_store',self.store).start()
        self.app=Flask('f08-f18');self.app.register_blueprint(controller.filing_change_review_bp)
        self.pid=self.service.create_project({'project_name':'隔离任务'})[2]['project_id']
        self.path=self.root/'form.docx';word_file(self.path,[['药品通用名称','任务测试药品']]);self.payload=self.path.read_bytes()

    def post(self,payload=None,name='form.docx',pid=None):
        with self.app.test_client() as client:
            r=client.post('/filing-change-review/projects/'+(pid or self.pid)+'/application-form/import/start',data={'file':(io.BytesIO(payload or self.payload),name)})
            return r.status_code,json.loads(r.get_data(as_text=True))

    def test_duplicate_returns_before_parser_finishes(self):
        entered=threading.Event();release=threading.Event();calls=[]
        real=self.service.form_parser.parse_form_file
        def paused(path):
            calls.append(path);entered.set()
            if not release.wait(15):raise RuntimeError('test release timeout')
            return real(path)
        with patch.object(self.service.form_parser,'parse_form_file',side_effect=paused):
            first=self.post();self.assertEqual(first[0],200);self.assertTrue(entered.wait(5))
            try:
                start=time.monotonic();second=self.post(name='same-bytes.docx')
                self.assertLess(time.monotonic()-start,2)
                self.assertEqual(second[0],200);self.assertEqual(second[1]['data']['task_id'],first[1]['data']['task_id'])
                other=self.root/'other.docx';word_file(other,[['药品通用名称','不同文件']])
                self.assertEqual(self.post(other.read_bytes())[0],409)
                with self.app.test_client() as client:
                    self.assertEqual(client.post('/filing-change-review/projects/'+self.pid+'/application-form/parse/start').status_code,409)
                self.assertEqual(len(calls),1)
            finally:
                release.set()
                deadline=time.monotonic()+10
                while time.monotonic()<deadline and self.store.get_task(first[1]['data']['task_id'])['status'] in ('pending','running'):time.sleep(.02)
            self.assertEqual(self.store.get_task(first[1]['data']['task_id'])['status'],'completed')
            self.assertEqual(len(calls),1)

    def test_same_bytes_in_other_project_are_independent(self):
        other=self.service.create_project({'project_name':'独立项目'})[2]['project_id']
        release=threading.Event();calls=[];real=self.service.form_parser.parse_form_file
        def paused(path):
            calls.append(path)
            if not release.wait(15):raise RuntimeError('test release timeout')
            return real(path)
        with patch.object(self.service.form_parser,'parse_form_file',side_effect=paused):
            try:
                one=self.post();two=self.post(pid=other)
                self.assertEqual(one[0],200);self.assertEqual(two[0],200)
                self.assertNotEqual(one[1]['data']['task_id'],two[1]['data']['task_id'])
            finally:
                release.set()
                deadline=time.monotonic()+10
                while time.monotonic()<deadline:
                    if all(self.store.get_task(x[1]['data']['task_id'])['status']=='completed' for x in (one,two)):break
                    time.sleep(.02)
            self.assertEqual(len(calls),2)
            self.assertEqual(self.service.get_application_form(other)[2]['form_json']['item_6_generic_name']['value'],'任务测试药品')

    def test_old_failure_cleanup_and_late_worker_cannot_replace_new_success(self):
        self.assertTrue(self.service.import_application_form(self.pid,_Upload(self.path.name,self.payload))[0])
        now=time.time()+1
        for taskid,created in [('old',now),('new',now+130)]:
            self.store.create_task(task_id=taskid,domain='filing_change_review',project_id=self.pid,task_type='parse_application_form',now=created)
            self.store.claim_task(taskid,now=created)
            if taskid=='old':
                self.store.recover_stale_active_tasks(domain='filing_change_review',stale_after_seconds=120,now=now+121)
            else:
                self.assertTrue(self.service.parse_application_form(self.pid)[0])
                self.store.finish_task(taskid,status='completed',result={'content_status':'success','content_available':True},now=now+131)
        self.assertFalse(self.store.finish_task('old',status='completed',now=now+132))
        self.store.prune_finished(domain='filing_change_review',retention_seconds=7200,now=now+8000)
        final=service_at(self.root,self.service.db_conn).get_application_form(self.pid)[2]
        self.assertEqual(final['latest_attempt']['task_id'],'new')
        self.assertEqual(final['latest_attempt']['content_status'],'success')

    def test_interruption_survives_cleanup_and_new_connection(self):
        self.assertTrue(self.service.import_application_form(self.pid,_Upload(self.path.name,self.payload))[0])
        now=time.time()+1
        self.store.create_task(task_id='interrupted',domain='filing_change_review',project_id=self.pid,task_type='parse_application_form',now=now)
        self.store.claim_task('interrupted',now=now)
        self.store.recover_stale_active_tasks(domain='filing_change_review',stale_after_seconds=120,now=now+121)
        before=self.service.get_application_form(self.pid)[2]
        self.assertEqual(before['latest_attempt']['content_status'],'failed');self.assertEqual(before['parse_status'],'success')
        self.assertEqual(self.store.prune_finished(domain='filing_change_review',retention_seconds=7200,now=now+7322),1)
        self.assertEqual(self.store.prune_finished(domain='filing_change_review',retention_seconds=7200,now=now+7322),0)
        self.assertIsNone(self.store.get_task('interrupted'))
        fresh=service_at(self.root,_TestConnection(self.root/'db.sqlite'));self.addCleanup(fresh.db_conn.engine.dispose)
        after=fresh.get_application_form(self.pid)[2]
        self.assertEqual(after['latest_attempt'],before['latest_attempt']);self.assertEqual(after['original_file_id'],before['original_file_id'])
        self.assertTrue(after['latest_attempt']['message'])
        with patch.object(fresh,'_now',return_value=datetime.fromtimestamp(now+8000,ZoneInfo('Asia/Shanghai'))):
            self.assertTrue(fresh.parse_application_form(self.pid)[0])
        self.assertEqual(fresh.get_application_form(self.pid)[2]['latest_attempt']['content_status'],'success')

    def test_near_simultaneous_processes_only_parse_once(self):
        context=multiprocessing.get_context('spawn')
        barrier=context.Barrier(2);release=context.Event();responses=context.Queue();calls=context.Queue()
        workers=[context.Process(target=concurrent_request,args=(str(self.root),self.pid,barrier,release,responses,calls)) for _ in range(2)]
        for worker in workers:worker.start()
        try:
            results=[responses.get(timeout=20) for _ in workers]
            self.assertIn(200,[r[0] for r in results])
            self.assertTrue(all(r[0] in (200,409) for r in results),results)
            calls.get(timeout=5)
            self.assertEqual(len([t for t in self.store.list_project_parse_tasks(self.pid,'filing_change_review') if t['status'] in ('pending','running')]),1)
        finally:
            release.set()
            for worker in workers:
                worker.join(20)
                if worker.is_alive():worker.terminate();worker.join(5)
        self.assertTrue(all(worker.exitcode==0 for worker in workers))
        self.assertTrue(calls.empty(),'两个进程不得分别执行解析')
        tasks=self.store.list_project_parse_tasks(self.pid,'filing_change_review')
        self.assertEqual(len(tasks),1);self.assertEqual(tasks[0]['status'],'completed')

    def test_prune_failure_retains_only_failure_evidence(self):
        self.store.create_task(task_id='prune-fail',domain='filing_change_review',project_id=self.pid,task_type='parse_application_form',now=1)
        self.store.claim_task('prune-fail',now=1)
        self.store.finish_task('prune-fail',status='failed',message='断连失败',now=2)
        with patch.object(self.store,'_persist_filing_attempt',side_effect=RuntimeError('test database write failure')):
            with self.assertRaises(RuntimeError):self.store.prune_finished(domain='filing_change_review',retention_seconds=7200,now=8000)
        self.assertIsNotNone(self.store.get_task('prune-fail'))

    def test_batch_and_submission_cleanup_keep_per_file_status(self):
        from agent.agent_backend.database.mysql.db_model import FilingChangeSubmissionFile
        from agent.agent_backend.services.filing_parse_outcome import parse_task_id
        p=self.root/'material.pdf'
        with fitz.open() as doc:
            doc.new_page().insert_text((40,50),'Assay study: 95 percent recovery.');doc.save(p)
        ok,_,uploaded=self.service.upload_submission_files(self.pid,[_Upload(p.name,p.read_bytes())],'5')
        self.assertTrue(ok,uploaded)
        docid=uploaded['created'][0]['doc_id']
        self.assertTrue(self.service.parse_submission(self.pid,docid)[0])
        now=time.time()+1
        self.store.create_task(task_id='batch-stop',domain='filing_change_review',project_id=self.pid,
            task_type='parse_submissions_batch',payload={'doc_ids':[docid]},now=now)
        self.store.claim_task('batch-stop',now=now)
        self.store.recover_stale_active_tasks(domain='filing_change_review',stale_after_seconds=120,now=now+121)
        before=self.service.list_submissions(self.pid,{})
        self.store.prune_finished(domain='filing_change_review',retention_seconds=7200,now=now+8000)
        after=service_at(self.root,self.service.db_conn).list_submissions(self.pid,{})
        self.assertEqual(after,before)
        entries=after['list']
        self.assertEqual(entries[0]['latest_attempt']['content_status'],'failed')
        self.assertEqual(entries[0]['parse_status'],'success')
        selected=review_inputs(self.service,self.pid)['submissions'][0]
        self.assertEqual(selected['latest_attempt']['content_status'],'failed')

    def test_real_reparse_admission_persists_original_source(self):
        self.assertTrue(self.service.import_application_form(self.pid,_Upload(self.path.name,self.payload))[0])
        original=self.service.get_application_form(self.pid)[2]
        with patch.object(controller,'_run_parse_task_async',return_value=True):
            with self.app.test_client() as client:
                response=client.post('/filing-change-review/projects/'+self.pid+'/application-form/parse/start')
                taskid=json.loads(response.get_data(as_text=True))['data']['task_id']
        self.store.claim_task(taskid)
        self.store.recover_stale_active_tasks(domain='filing_change_review',stale_after_seconds=120,now=time.time()+121)
        self.store.prune_finished(domain='filing_change_review',retention_seconds=7200,now=time.time()+8000)
        attempt=self.service.get_application_form(self.pid)[2]['latest_attempt']
        self.assertEqual(attempt['source_file_id'],original['original_file_id'])
        self.assertEqual(attempt['source_file_name'],self.path.name)
        self.assertEqual(attempt['task_id'],taskid)


if __name__=='__main__':unittest.main()
