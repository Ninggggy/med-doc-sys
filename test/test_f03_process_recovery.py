"""真实 Word、隔离 SQLite 和 spawn 子进程；仅终止本测试创建的进程。"""
import io
import multiprocessing
import os
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
from agent.agent_backend.services.filing_change_review_service import FilingChangeReviewService
from agent.agent_backend.services.filing_parse_outcome import parse_task_id
from agent.agent_backend.services.runtime_task_store import RuntimeTaskStore
from agent.test.test_filing_change_review_service_regression import _TestConnection, _Upload


def service_at(base):
    service = object.__new__(FilingChangeReviewService)
    service.root_dir = Path(base) / 'review'
    service.root_dir.mkdir(exist_ok=True)
    service.db_conn = _TestConnection(Path(base) / 'db.sqlite')
    service.form_parser = FilingChangeFormParserService()
    return service


def document_bytes(name):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '药品通用名称'
    table.cell(0, 1).text = name
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def pause_worker(base, project_id, window, reached, action='replace'):
    service = service_at(base)

    def pause():
        reached.set()
        # 不在待终止进程中持有 Event.wait 的条件锁，避免父进程清理死锁。
        time.sleep(45)
        raise RuntimeError('parent did not terminate its test process')

    if window in ('install_pending', 'quarantine_move'):
        original = os.replace

        def replace(source, target):
            original(source, target)
            target = Path(target)
            if ((window == 'install_pending' and target.suffix == '.pending') or
                    (window == 'quarantine_move' and target.parent.parent.name == '.trash')):
                pause()

        os.replace = replace
    elif window in ('installed', 'quarantined', 'restored', 'cleaned'):
        method = {'installed': '_install_file_payloads', 'quarantined': '_quarantine_paths',
                  'restored': '_restore_quarantined_paths', 'cleaned': '_finalize_quarantine'}[window]
        original = getattr(service, method)

        def wrapped(*args, **kwargs):
            result = original(*args, **kwargs)
            pause()
            return result

        setattr(service, method, wrapped)
    elif window in ('manifest', 'recovery_manifest'):
        original = service._save_submission_manifest

        def save(project, data):
            if window == 'recovery_manifest' and 'application_form_transaction' not in data:
                pause()
            result = original(project, data)
            transaction = data.get('application_form_transaction', {})
            if window == 'manifest' and transaction.get('source'):
                pause()
            return result

        service._save_submission_manifest = save
    elif window in ('before_commit', 'after_commit'):
        original = service.db_conn.get_session

        def get_session():
            session = original()
            commit = session.commit

            def wrapped_commit():
                if window == 'before_commit':
                    pause()
                commit()
                if window == 'after_commit':
                    pause()

            session.commit = wrapped_commit
            return session

        service.db_conn.get_session = get_session

    if action == 'recover':
        service.get_application_form(project_id)
    elif action == 'reparse':
        service.parse_application_form(project_id)
    else:
        service.import_application_form(project_id, _Upload('同名申请表.docx', document_bytes('新药品')))
    service.db_conn.engine.dispose()


class F03ProcessRecoveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.service = self.fresh()
        self.pid = self.service.create_project({'project_name': 'F03进程恢复'})[2]['project_id']
        self.old_bytes = document_bytes('原有效药品')
        ok, message, self.original = self.service.import_application_form(
            self.pid, _Upload('同名申请表.docx', self.old_bytes))
        self.assertTrue(ok, message)
        self.original_id = self.original['original_file_id']

    def fresh(self):
        service = service_at(self.base)
        self.addCleanup(service.db_conn.engine.dispose)
        return service

    def terminate_at(self, window, action='replace'):
        ctx = multiprocessing.get_context('spawn')
        reached = ctx.Event()
        worker = ctx.Process(target=pause_worker, args=(str(self.base), self.pid, window, reached, action))
        worker.start()
        try:
            self.assertTrue(reached.wait(25), f'{window}: child exited={worker.exitcode}')
            worker.terminate()
            worker.join(10)
            self.assertFalse(worker.is_alive())
            self.assertLess(worker.exitcode, 0)
        finally:
            if worker.is_alive():
                worker.kill()
            worker.join(10)
            worker.close()

    def verify_recovery(self, committed=False, reparse=False):
        fresh = self.fresh()
        ok, message, current = fresh.get_application_form(self.pid)
        self.assertTrue(ok, message)
        file_id = current['original_file_id']
        if committed and not reparse:
            self.assertNotEqual(self.original_id, file_id)
        else:
            self.assertEqual(self.original_id, file_id)
        expected = '新药品' if committed and not reparse else '原有效药品'
        self.assertIn(expected, current['raw_text'])
        self.assertEqual('success', current['parse_status'])
        path = fresh._project_path(self.pid) / 'application_form' / (file_id + '.docx')
        self.assertTrue(path.is_file())
        if not committed or reparse:
            self.assertEqual(self.old_bytes, path.read_bytes())
        self.assertEqual(file_id, current['effective_source']['source_file_id'])
        self.assertEqual('同名申请表.docx', current['effective_source']['source_file_name'])
        self.assertEqual('success' if committed else 'failed', current['latest_attempt']['content_status'])
        if not committed:
            self.assertEqual('task_interrupted', current['latest_attempt']['code'])
        manifest = fresh._load_submission_manifest(self.pid)
        self.assertNotIn('application_form_transaction', manifest)
        self.assertEqual(2, len(manifest['parse_attempts']))
        self.assertEqual(['success', 'success' if committed else 'failed'],
                         [a['content_status'] for a in manifest['parse_attempts']])
        self.assertFalse(list((fresh.root_dir / '.trash').rglob('*.docx')))
        self.assertEqual([path], list(path.parent.glob('*.docx')))
        self.assertFalse(list(path.parent.glob('.*.pending')))
        # 另一个连接再次恢复，历史和有效内容不能变化。
        again = self.fresh()
        self.assertEqual(current, again.get_application_form(self.pid)[2])
        self.assertEqual(manifest, again._load_submission_manifest(self.pid))
        ok, message, parsed = again.parse_application_form(self.pid)
        self.assertTrue(ok, message)
        self.assertEqual(file_id, parsed['original_file_id'])
        self.assertIn(expected, parsed['raw_text'])
        self.assertEqual(file_id, again.get_application_form(self.pid)[2]['original_file_id'])
        for field in parsed['form_json'].values():
            if isinstance(field, dict) and 'field_type' in field:
                self.assertEqual(file_id, field['source_file_id'])
                for source in field.get('value_sources', {}).values():
                    self.assertEqual(file_id, source['source_file_id'])

    def test_exit_after_install(self):
        self.terminate_at('installed')
        self.verify_recovery()

    def test_exit_inside_install(self):
        self.terminate_at('install_pending')
        self.verify_recovery()

    def test_exit_inside_quarantine(self):
        self.terminate_at('quarantine_move')
        self.verify_recovery()

    def test_exit_after_quarantine(self):
        self.terminate_at('quarantined')
        self.verify_recovery()

    def test_exit_after_manifest(self):
        self.terminate_at('manifest')
        self.verify_recovery()

    def test_exit_before_commit(self):
        self.terminate_at('before_commit')
        self.verify_recovery()

    def test_exit_after_commit_before_cleanup(self):
        self.terminate_at('after_commit')
        self.verify_recovery(committed=True)

    def test_recovery_exit_after_restore(self):
        self.terminate_at('quarantined')
        self.terminate_at('restored', 'recover')
        self.verify_recovery()

    def test_recovery_exit_after_cleanup(self):
        self.terminate_at('after_commit')
        self.terminate_at('cleaned', 'recover')
        self.verify_recovery(committed=True)

    def test_recovery_exit_before_manifest_clear(self):
        self.terminate_at('quarantined')
        self.terminate_at('recovery_manifest', 'recover')
        self.verify_recovery()

    def test_reparse_exit_before_commit(self):
        self.terminate_at('before_commit', 'reparse')
        self.verify_recovery(reparse=True)

    def test_reparse_exit_after_commit(self):
        self.terminate_at('after_commit', 'reparse')
        self.verify_recovery(committed=True, reparse=True)

    def test_commit_acknowledgement_failure_keeps_committed_version(self):
        original = self.service.db_conn.get_session

        def get_session():
            session = original()
            commit = session.commit

            def commit_then_fail():
                commit()
                raise ConnectionError('lost commit acknowledgement')

            session.commit = commit_then_fail
            return session

        with mock.patch.object(self.service.db_conn, 'get_session', side_effect=get_session):
            ok, message, _ = self.service.import_application_form(
                self.pid, _Upload('同名申请表.docx', document_bytes('新药品')))
        self.assertTrue(ok, message)
        self.verify_recovery(committed=True)

    def test_invalidated_task_cannot_commit_replacement(self):
        store = RuntimeTaskStore(connection=self.service.db_conn, ensure_schema=False)
        store.create_task(task_id='f03-interrupted', domain='filing_change_review', project_id=self.pid,
                          task_type='import_application_form', now=100)
        # 未处于 running 的任务不能提交；保留现有 RuntimeTask 行锁检查。
        token = parse_task_id.set('f03-interrupted')
        try:
            ok, _, attempt = self.service.import_application_form(
                self.pid, _Upload('同名申请表.docx', document_bytes('新药品')))
        finally:
            parse_task_id.reset(token)
        self.assertFalse(ok)
        self.assertEqual('task_interrupted', attempt['code'])
        current = self.fresh().get_application_form(self.pid)[2]
        self.assertEqual(self.original_id, current['original_file_id'])
        self.assertIn('原有效药品', current['raw_text'])
        self.assertEqual('failed', current['latest_attempt']['content_status'])
        manifest = self.service._load_submission_manifest(self.pid)
        self.assertEqual('f03-interrupted', manifest['parse_attempts'][-1]['task_id'])
        self.assertEqual(attempt, manifest['parse_attempts'][-1])
        self.assertNotEqual(self.original_id, attempt['source_file_id'])
        self.assertNotIn('application_form_transaction', manifest)

    def test_first_import_exit_does_not_claim_effective_content(self):
        self.pid = self.service.create_project({'project_name': 'F03首次导入中断'})[2]['project_id']
        self.terminate_at('installed')
        fresh = self.fresh()
        current = fresh.get_application_form(self.pid)[2]
        self.assertEqual('not_parsed', current['parse_status'])
        self.assertFalse(current.get('original_file_id'))
        self.assertEqual('failed', current['latest_attempt']['content_status'])
        manifest = fresh._load_submission_manifest(self.pid)
        self.assertEqual(1, len(manifest['parse_attempts']))
        self.assertNotIn('application_form_source', manifest)
        self.assertNotIn('application_form_transaction', manifest)
        self.assertFalse(list((fresh._project_path(self.pid) / 'application_form').iterdir()))
        ok, message, _ = fresh.import_application_form(self.pid, _Upload('申请表.docx', self.old_bytes))
        self.assertTrue(ok, message)

    def test_restore_error_preserves_transaction_for_next_connection(self):
        self.terminate_at('quarantined')
        with mock.patch.object(self.service, '_restore_quarantined_paths', side_effect=OSError('restore unavailable')):
            with self.assertRaises(OSError):
                self.service.get_application_form(self.pid)
        self.assertIn('application_form_transaction', self.service._load_submission_manifest(self.pid))
        self.verify_recovery()


if __name__ == '__main__':
    unittest.main()
