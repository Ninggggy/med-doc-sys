"""期望来自手工输入事实，真实生成文件后经现有 PDF/Word 入口解析。"""
import io
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from docx import Document
import fitz
from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService
from agent.agent_backend.services.filing_form_revision import save_revision, merge_parse, validity_cross_checks
from agent.agent_backend.utils.parser.form_field_semantics import selected_option, validity_values
from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import item_start


FACTS = [
    ('3. 是否为OTC', '非处方药'),
    ('5. 申请事项分类', '6. 已上市化学药品药学变更；6.8 变更有效期和贮藏条件'),
    ('6. 药品通用名称', '测试喷雾剂'),
    ('7. 英文名称/拉丁名称', 'Example Hydrochloride Spray'),
    ('8. 汉语拼音', 'Ceshi Penwuji'),
    ('9. 化学名称', '(E)-N-(6,6-二甲基-2-庚烯-4-炔基)-N-甲基盐酸盐'),
    ('10. 商品名称', '不使用'), ('11. 剂型', '中国药典剂型: 喷雾剂'),
    ('12. 规格', '1% (g:g; 每瓶30mL)'), ('13. 同品种其他规格', ''),
    ('14. 包装', '直接接触药品的包装材料和容器: 聚乙烯瓶\n包装规格: 1瓶/盒'),
    ('15. 药品有效期', '原有效期: 18个月\n拟延长后有效期: 24个月'),
    ('16. 处方', '活性成分/中药成分/...: 测试成分\n辅料: 乙醇、纯化水\n是否有变更: 否'),
    ('20. 主要适应症或者功能主治', '用于局部治疗'),
    ('21. 补充申请的内容', '有效期由18个月延长至24个月'),
    ('22. 提出补充申请的理由', '根据长期稳定性研究结果'),
    ('23. 原批准注册内容及相关信息', '原批准18个月；有效期截止日期2028年8月'),
    ('30. 药品注册申请人', '中文名称: 甲方医药有限公司\n注册地址: 甲市一号'),
    ('生产企业', '中文名称: 乙方制剂有限公司\n注册地址: 乙市二号'),
    ('31. 委托研究机构', '中文名称: 丙方研究有限公司\n研究负责人: 王某'),
]


def make_word(path):
    doc = Document()
    doc.add_paragraph('境内生产药品注册备案表')
    for group in (FACTS[:10], FACTS[10:]):
        table = doc.add_table(rows=0, cols=2)
        for label, value in group:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '17. 原/辅料/包材来源'
    nested = table.cell(0, 2).add_table(rows=2, cols=3)
    for c, v in zip(nested.rows[0].cells, ['供应商', '登记号', '原/辅料/包材名称']): c.text = v
    for c, v in zip(nested.rows[1].cells, ['丁方原料有限公司', 'Y20260001', '测试原料']): c.text = v
    doc.save(path)


def make_pdf(path, offset=0):
    doc = fitz.open()
    page = doc.new_page(width=900, height=1200)
    y = 35
    for label, value in [('境内生产药品注册备案表', '')] + FACTS:
        if offset:
            page.insert_text((25, y), label + ':', fontname='china-s', fontsize=11)
            for line in value.splitlines() or ['']:
                page.insert_text((320, y), line, fontname='china-s', fontsize=11)
                y += 20
            continue
        for line in (label + ': ' + value).splitlines():
            page.insert_text((25 + offset, y), line, fontname='china-s', fontsize=11)
            y += 20
    doc.save(path)
    doc.close()


class T3ParserTests(unittest.TestCase):
    def test_options(self):
        for text, options, expected in [
            ('非处方药', ['处方药', '非处方药'], '非处方药'),
            ('处方药／非处方药', ['处方药', '非处方药'], ''),
            ('☐处方药 ☑非处方药', ['处方药', '非处方药'], '非处方药'),
            ('☐处方药 ☐非处方药', ['处方药', '非处方药'], ''),
            ('是否有变更：否', ['是', '否'], '否'),
            ('是否', ['是', '否'], ''), ('是／否', ['是', '否'], ''),
            ('使用／不使用', ['使用', '不使用'], ''), ('不使用', ['使用', '不使用'], '不使用'),
            ('已上市／未上市', ['已上市', '未上市'], ''), ('未上市', ['已上市', '未上市'], '未上市'),
        ]:
            with self.subTest(text=text): self.assertEqual(selected_option(text, options), expected)

    def test_heading_and_validity(self):
        self.assertIsNone(item_start('6. 已上市化学药品药学变更事项'))
        self.assertIsNone(item_start('6.8 变更有效期和贮藏条件'))
        self.assertEqual(item_start('31. 委托研究机构：丙公司')[0], 32)
        for text in ['有效期由18个月延长至24个月', '原批准有效期18个月，拟变更为24个月', '考察24个月；有效期拟从获批的18个月延长至24个月；截止2028年8月']:
            self.assertEqual(validity_values(text), {'original_validity_period': '18个月', 'proposed_validity_period': '24个月'})
        self.assertEqual(validity_values('考察0、6、18、24个月；截止2028年8月'), {})

    def check_form(self, form):
        self.assertEqual(form['original_form_type'], '备案表')
        for key, expected in [('item_3_otc_type', '非处方药'), ('item_6_generic_name', '测试喷雾剂'), ('item_7_english_or_latin_name', 'Example Hydrochloride Spray'), ('item_8_pinyin', 'Ceshi Penwuji'), ('item_9_chemical_name', FACTS[5][1]), ('item_12_specification', FACTS[8][1]), ('item_13_other_specifications', '')]:
            self.assertEqual(form[key]['value'], expected, key)
        self.assertEqual(form['item_5_application_matter_category']['selected_values'], ['1.7'])
        self.assertEqual(form['item_5_application_matter_category']['original_matter'][0]['code'], '6.8')
        self.assertIn('喷雾剂', form['item_11_dosage_form']['selected_values'][0])
        self.assertEqual(form['item_16_prescription']['sub_fields']['has_change'], '否')
        self.assertEqual(form['item_16_prescription']['sub_fields']['excipients'], '乙醇、纯化水')
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['applicant_name'], '甲方医药有限公司')
        self.assertEqual(form['item_31_manufacturer_info']['sub_fields']['manufacturer_name'], '乙方制剂有限公司')
        self.assertEqual(form['item_32_cro_info']['sub_fields']['organization_name'], '丙方研究有限公司')
        self.assertEqual(form['item_15_validity_period']['sub_fields']['original_validity_period'], '18个月')
        self.assertEqual(form['item_15_validity_period']['sub_fields']['proposed_validity_period'], '24个月')

    def test_real_pdf_two_positions(self):
        with tempfile.TemporaryDirectory() as tmp:
            for offset in (0, 85):
                path = Path(tmp) / f'layout{offset}.pdf'
                make_pdf(path, offset)
                result = FilingChangeFormParserService().parse_form_file(str(path))
                self.check_form(result['form_json'])
                self.assertTrue(result['form_json']['item_6_generic_name']['source_regions'][0]['bbox_pdf'])
                self.assertTrue(result['parse_diagnostics'])

    def test_real_word_tables_merged_and_nested(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'form.docx'
            make_word(path)
            doc = Document(path)
            doc.add_paragraph('未注明角色：戊方有限公司')
            doc.save(path)
            result = FilingChangeFormParserService().parse_form_file(str(path))
            self.check_form(result['form_json'])
            self.assertEqual(result['form_json']['item_17_material_source']['table_rows'], [{'material_name': '测试原料', 'register_no': 'Y20260001', 'accept_no': '', 'manufacturer': '丁方原料有限公司'}])
            self.assertTrue(any('戊方' in candidate['source_text'] and '待核对' in candidate['status'] for candidate in result['form_json']['unassigned_party_sources']))

    def test_repository_word_values_and_blank_validity(self):
        path = Path(__file__).resolve().parents[1] / 'task/change_review/1.药品注册补充申请表-变更申请用例.docx'
        form = FilingChangeFormParserService().parse_form_file(str(path))['form_json']
        self.assertEqual(form['item_6_generic_name']['value'], '药品1')
        self.assertEqual(form['item_7_english_or_latin_name']['value'], 'Aadasd')
        self.assertEqual(form['item_14_packaging']['sub_fields']['packaging_specification'], '100ml')
        self.assertEqual(form['item_30_applicant_info']['sub_fields']['applicant_name'], '申请')
        self.assertEqual(form['item_31_manufacturer_info']['sub_fields']['manufacturer_name'], '南京')
        self.assertEqual(form['item_15_validity_period']['sub_fields']['original_validity_period'], '')

    def test_word_checkbox_content_controls_and_vertical_merge(self):
        from docx.oxml import parse_xml
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            table = doc.add_table(rows=2, cols=3)
            table.cell(0, 0).merge(table.cell(1, 0)).text = '16. 处方'
            table.cell(0, 1).text = '活性成分'
            table.cell(0, 2).text = '测试成分'
            table.cell(1, 1).text = '是否有变更'
            paragraph = table.cell(1, 2).paragraphs[0]
            paragraph._p.append(parse_xml('<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"><w:sdtPr><w14:checkbox><w14:checked w14:val="1"/></w14:checkbox></w:sdtPr><w:sdtContent><w:r><w:t>否</w:t></w:r></w:sdtContent></w:sdt>'))
            path = Path(tmp) / 'controls.docx'
            doc.save(path)
            form = FilingChangeFormParserService().parse_form_file(str(path))['form_json']
            self.assertEqual(form['item_16_prescription']['sub_fields']['active_ingredients'], '测试成分')
            self.assertEqual(form['item_16_prescription']['sub_fields']['has_change'], '否')

    def test_cross_checks_keep_provenance(self):
        form = FilingChangeFormParserService()._build_full_form('')
        field = form['item_15_validity_period']
        field['sub_fields']['original_validity_period'] = '18个月'
        field['source_text'] = '申请表原文'
        validity_cross_checks(form, [{'text': '有效期：18个月', 'source_file': '批准文件', 'is_approval': True}, {'text': '原批准有效期12个月，拟变更为24个月', 'source_file': '修订说明'}])
        self.assertEqual([c['status'] for c in field['cross_checks']], ['一致', '存在冲突，需人工确认', '缺失，候选补充值待确认'])
        self.assertEqual(field['source_text'], '申请表原文')
        self.assertEqual(field['sub_fields']['proposed_validity_period'], '')


class T3RevisionTests(unittest.TestCase):
    def test_real_service_save_refresh_reparse_and_failure(self):
        from agent.test.test_filing_change_review_service_regression import _TestConnection, _Upload
        from agent.agent_backend.services.filing_change_review_service import FilingChangeReviewService
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            connection = _TestConnection(root / 't3.sqlite')
            service = object.__new__(FilingChangeReviewService)
            service.root_dir = root / 'review'
            service.root_dir.mkdir()
            service.db_conn = connection
            service.form_parser = FilingChangeFormParserService()
            ok, _, project = service.create_project({'project_name': 'T3测试'})
            self.assertTrue(ok)
            pid = project['project_id']
            path = root / '事实申请表.docx'
            make_word(path)
            ok, message, imported = service.import_application_form(pid, _Upload(path.name, path.read_bytes()))
            self.assertTrue(ok, message)
            key = 'item_15_validity_period'
            edited = deepcopy(imported['form_json'])
            edited[key]['sub_fields']['original_validity_period'] = '20个月'
            ok, message, saved = service.save_application_form(pid, {'form_json': edited, 'confidence': 1})
            self.assertTrue(ok, message)
            self.assertEqual(saved['confidence'], imported['confidence'])
            self.assertEqual(saved['raw_text'], imported['raw_text'])
            self.assertEqual(saved['parse_status'], imported['parse_status'])
            self.assertEqual(saved['form_json'][key]['manual_paths'], ['sub_fields.original_validity_period'])
            # 原源文件在临时项目中变更，验证只锁实际改动的子字段。
            stored = service._project_path(pid) / 'application_form' / (saved['original_file_id'] + '.docx')
            doc = Document(stored)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if '拟延长后有效期: 24个月' in cell.text:
                            cell.text = cell.text.replace('拟延长后有效期: 24个月', '拟延长后有效期: 36个月')
                        if '有效期由18个月延长至24个月' in cell.text:
                            cell.text = cell.text.replace('有效期由18个月延长至24个月', '有效期由18个月延长至36个月')
            doc.save(stored)
            ok, message, parsed = service.parse_application_form(pid)
            self.assertTrue(ok, message)
            refreshed = service.get_application_form(pid)[2]
            self.assertEqual(refreshed['form_json'], parsed['form_json'])
            field = refreshed['form_json'][key]
            self.assertEqual(field['sub_fields']['original_validity_period'], '20个月')
            self.assertEqual(field['sub_fields']['proposed_validity_period'], '36个月')
            self.assertEqual(field['candidates']['sub_fields.original_validity_period']['source_file'], '事实申请表.docx')
            self.assertTrue(field['candidates']['sub_fields.original_validity_period']['source_regions'])
            ok, _, adopted = service.save_application_form(pid, {'form_json': refreshed['form_json'], 'resolutions': {key: {'sub_fields.original_validity_period': 'adopt'}}})
            self.assertTrue(ok)
            self.assertFalse(adopted['form_json'][key]['manual_modified'])
            # T1 失败后仍能读到已采纳值和完整修订历史。
            stored = service._project_path(pid) / 'application_form' / (adopted['original_file_id'] + '.docx')
            stored.write_bytes(b'broken')
            self.assertFalse(service.parse_application_form(pid)[0])
            failed = service.get_application_form(pid)[2]
            self.assertEqual(failed['form_json'], adopted['form_json'])
            self.assertEqual(failed['latest_attempt']['content_status'], 'failed')
            connection.engine.dispose()

    def test_weak_new_result_does_not_replace_reliable_value(self):
        old = {'x': {'field_type': 'input', 'value': '清楚的原文', 'parse_confidence': .9}}
        new = {'x': {'field_type': 'input', 'value': '全文兜底', 'parse_confidence': .5}}
        self.assertEqual(merge_parse(old, new)['x']['value'], '清楚的原文')

    def test_edit_only_one_subfield_reparse_and_resolve(self):
        original = FilingChangeFormParserService()._build_full_form('')
        key = 'item_15_validity_period'
        original[key].update(parse_confidence=0.75, source_text='原有效期18个月，拟变更为24个月')
        original[key]['sub_fields'].update(original_validity_period='18个月', proposed_validity_period='24个月')
        unchanged = save_revision(original, deepcopy(original))
        self.assertFalse(unchanged[key]['manual_modified'])
        edited = deepcopy(original)
        edited[key]['sub_fields']['original_validity_period'] = '20个月'
        saved = save_revision(original, edited)
        self.assertEqual(saved[key]['manual_paths'], ['sub_fields.original_validity_period'])
        fresh = deepcopy(original)
        fresh[key]['sub_fields']['proposed_validity_period'] = '36个月'
        merged = merge_parse(saved, fresh)
        self.assertEqual(merged[key]['sub_fields']['original_validity_period'], '20个月')
        self.assertEqual(merged[key]['sub_fields']['proposed_validity_period'], '36个月')
        self.assertEqual(merged[key]['candidates']['sub_fields.original_validity_period']['source_text'], original[key]['source_text'])
        kept = save_revision(merged, merged, {key: {'sub_fields.original_validity_period': 'keep'}})
        self.assertEqual(kept[key]['candidates'], {})
        self.assertTrue(kept[key]['manual_modified'])
        adopted = save_revision(merged, merged, {key: {'sub_fields.original_validity_period': 'adopt'}})
        self.assertEqual(adopted[key]['sub_fields']['original_validity_period'], '18个月')
        self.assertFalse(adopted[key]['manual_modified'])
        with self.assertRaisesRegex(ValueError, '候选结果已更新'):
            save_revision(merged, merged, {key: {'sub_fields.original_validity_period': {'action': 'adopt', 'candidate': {'value': '旧候选'}}}})
        fresh[key]['sub_fields']['proposed_validity_period'] = ''
        self.assertEqual(merge_parse(merged, fresh)[key]['sub_fields']['proposed_validity_period'], '36个月')

    def test_table_and_legacy_protection(self):
        key = 'item_17_material_source'
        old = {key: {'field_type': 'table', 'table_rows': [{'material_name': '人工值'}], 'manual_modified': True, 'parse_confidence': .7}}
        new = {key: {'field_type': 'table', 'table_rows': [{'material_name': '新识别'}], 'parse_confidence': .8, 'source_text': '原料明细'}}
        merged = merge_parse(old, new)
        self.assertEqual(merged[key]['table_rows'], old[key]['table_rows'])
        self.assertEqual(merged[key]['candidates']['table_rows']['value'], new[key]['table_rows'])


if __name__ == '__main__':
    unittest.main()
