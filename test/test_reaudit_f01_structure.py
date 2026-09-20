"""F01 再复核：结构中的模板残留不能授权替换有效原件。"""
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from agent.test.test_f01_f02_form_replacement import service_at, _Upload, NAME
from agent.test.test_four_semantic_table_defects import word_file, pdf_file


EMPTY = [['药品通用名称', ''], ['是否为OTC', '处方药□', '非处方药□'], ['剂型', '片剂□', '胶囊剂□']]
FOOTERS = ['1', '第1页 共1页', '填写说明：请在相应方框内打勾。']


def document_bytes(root, extension, *, footer='', rows=None, pages=1, header='', detail=False, loose=''):
    path = root / ('input.' + extension)
    rows = EMPTY if rows is None else rows
    if extension == 'docx':
        word_file(path, rows)
        doc = Document(path)
        doc.sections[0].header.paragraphs[0].text = header
        doc.sections[0].footer.paragraphs[0].text = footer
        if loose:
            doc.add_paragraph(loose)
        if detail:
            table = doc.add_table(rows=4, cols=3)
            for cell, text in zip(table.rows[0].cells, ['序号', '物料名称', '登记号']):
                cell.text = text
            for i in range(1, 4):
                table.cell(i, 0).text = str(i)
        stream = io.BytesIO(); doc.save(stream)
        return stream.getvalue()
    pdf_file(path, rows)
    with fitz.open(path) as doc:
        if pages > 1:
            doc.fullcopy_page(0)
        for index, page in enumerate(doc):
            for text, xy in [(footer.replace('第1页', f'第{index+1}页'), (400, 1075)),
                             (header, (35, 16)), (loose, (35, 420))]:
                if text:
                    page.insert_text(xy, text, fontname='china-s', fontsize=10)
            if detail:
                for i, row in enumerate([['序号','物料名称','登记号'], ['1','',''], ['2','',''], ['3','','']]):
                    for j, text in enumerate(row):
                        page.insert_text((35+j*200, 280+i*32), text, fontname='china-s', fontsize=12)
        doc.subset_fonts()
        return doc.tobytes(garbage=4, deflate=True)


class StructureTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='reaudit-f01-')
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        from agent.agent_backend.config.settings import settings
        if os.getenv('T2_OCR_URL'):
            self.enterContext(patch.object(settings, 'ocr_service_url', os.environ['T2_OCR_URL']))
        self.service = service_at(self.root)
        self.addCleanup(self.service.db_conn.engine.dispose)
        self.pid = self.service.create_project({'project_name':'页脚与填写证据'})[2]['project_id']

    def read(self):
        fresh = service_at(self.root)
        try:
            return fresh.get_application_form(self.pid)[2]
        finally:
            fresh.db_conn.engine.dispose()

    def upload(self, data, extension):
        return self.service.import_application_form(self.pid, _Upload('本次.'+extension, data))

    def check_empty(self, extension, *, existing=True, **options):
        if existing:
            original = document_bytes(self.root, 'docx', rows=[['药品通用名称','此前有效药品']])
            self.assertTrue(self.upload(original, 'docx')[0])
            before = self.read()
            path = self.service._project_path(self.pid) / 'application_form' / (before['original_file_id']+'.docx')
        data = document_bytes(self.root, extension, **options)
        answer = self.upload(data, extension)
        self.assertFalse(answer[0], '模板残留不得进入有效替换流程')
        self.assertEqual(answer[2]['code'], 'empty_form')
        for _ in range(2):
            after = self.read()
            self.assertEqual(after['latest_attempt']['content_status'], 'failed')
            self.assertEqual(after['latest_attempt']['code'], 'empty_form')
            if existing:
                self.assertEqual(after['original_file_id'], before['original_file_id'])
                self.assertEqual(after['form_json'], before['form_json'])
                self.assertEqual(after['effective_source'], before['effective_source'])
                self.assertEqual(path.read_bytes(), original)
                self.assertTrue(Document(io.BytesIO(path.read_bytes())).tables)
            else:
                self.assertFalse(after.get('original_file_id'))
        if existing:
            self.assertTrue(self.service.parse_application_form(self.pid)[0])
            after = self.read()
            self.assertEqual(after['original_file_id'], before['original_file_id'])
            self.assertEqual(after['form_json'][NAME]['value'], '此前有效药品')
            self.assertEqual(path.read_bytes(), original)
            manifest = self.service._load_submission_manifest(self.pid)
            self.assertTrue(any(a.get('code') == 'empty_form' for a in manifest['parse_attempts']))

    def test_original_footer_number(self):
        self.check_empty('pdf', footer=FOOTERS[0])

    def test_original_footer_pages(self):
        self.check_empty('pdf', footer=FOOTERS[1])

    def test_original_footer_instruction(self):
        self.check_empty('pdf', footer=FOOTERS[2])

    def check_positive(self, extension, rows, key=None, expected=None, footer=FOOTERS[2]):
        self.assertTrue(self.upload(document_bytes(self.root,'docx', rows=[['药品通用名称','此前有效药品']]),'docx')[0])
        data = document_bytes(self.root, extension, rows=rows, footer=footer)
        answer = self.upload(data, extension)
        self.assertTrue(answer[0], answer[1:])
        state = self.read()
        for _ in range(2):
            self.assertEqual(state['latest_attempt']['content_status'], state['parse_status'])
            self.assertTrue(state['effective_source']['parse_diagnostics']['replacement_eligible'])
            if key:
                field = state['form_json'][key]
                if field['field_type'] == 'group_checkbox':
                    self.assertEqual(field['selected_values'], [expected])
                    self.assertEqual([e['name'] for e in field['selection_evidence'] if e['selected']], [expected])
                else:
                    self.assertEqual(field['value'], expected)
            else:
                self.assertEqual(state['parse_status'],'partial')
                self.assertIn(expected,state['raw_text'])
            path = self.service._project_path(self.pid) / 'application_form' / (state['original_file_id']+'.'+extension)
            self.assertEqual(path.read_bytes(),data)
            self.assertTrue(self.service.parse_application_form(self.pid)[0])
            state = self.read()

    def test_uncertain_loose_text_cannot_authorize_replacement(self):
        original = document_bytes(self.root,'docx',rows=[['药品通用名称','此前有效药品']])
        self.assertTrue(self.upload(original,'docx')[0])
        before = self.read()
        data = document_bytes(self.root,'pdf',footer='尚待核对的附注')
        answer = self.upload(data,'pdf')
        self.assertFalse(answer[0])
        self.assertEqual(answer[2]['code'],'form_content_unconfirmed')
        state = self.read()
        self.assertEqual(state['form_json'],before['form_json'])
        self.assertEqual(state['original_file_id'],before['original_file_id'])
        self.assertEqual(state['effective_source'],before['effective_source'])
        diagnostics = state['latest_attempt']['parse_diagnostics']
        self.assertIn('尚待核对的附注',diagnostics['raw_text'])
        self.assertTrue(any(r['role']=='uncertain' and r['bbox_pdf'] for r in diagnostics['form_content_regions']))

    def test_uncertain_first_import_retains_raw_with_review(self):
        answer = self.upload(document_bytes(self.root,'pdf',footer='尚待核对的附注'),'pdf')
        self.assertTrue(answer[0])
        state = self.read()
        self.assertEqual(state['parse_status'],'partial')
        self.assertFalse(state['effective_source']['parse_diagnostics']['replacement_eligible'])
        self.assertIn('尚待核对的附注',state['raw_text'])

    def test_footer_failure_review_keeps_effective_content_and_failed_attempt(self):
        from agent.agent_backend.services.filing_change_extraction_service import FilingChangeExtractionService
        from agent.agent_backend.services.filing_change_review_orchestrator import FilingChangeReviewOrchestrator
        from agent.agent_backend.services.filing_change_quality_check_service import FilingChangeQualityCheckService
        from agent.agent_backend.services.filing_change_stability_service import FilingChangeStabilityService
        self.assertTrue(self.upload(document_bytes(self.root,'docx',rows=[['药品通用名称','此前有效药品']]),'docx')[0])
        self.assertFalse(self.upload(document_bytes(self.root,'pdf',footer=FOOTERS[0]),'pdf')[0])
        state = self.read()
        extracted = FilingChangeExtractionService().run(state,[],{})
        drug_facts = [fact for fact in extracted['facts'] if fact['field'] == 'drug_name']
        self.assertEqual([fact['raw_value'] for fact in drug_facts], ['此前有效药品'])
        self.assertEqual(extracted['application_form'][NAME]['value_sources']['value']['source_file_id'],state['original_file_id'])
        # 按真实编排入口汇总申请表尝试；只关闭不属于本轮的远程模型调用。
        reviewer = FilingChangeReviewOrchestrator(self.service.form_parser,FilingChangeQualityCheckService(),FilingChangeStabilityService())
        with patch.object(reviewer.llm_gate,'should_call_llm',return_value={'allow':False,'reason':'测试不调用外部模型'}):
            reviewed = reviewer.run(self.pid,state,[],[],[],{},{'missing_items':[]})
        self.assertEqual(reviewed['formal_review']['missing_materials'],[])
        self.assertTrue(any(i['status']=='failed' for i in reviewed['formal_review']['parse_issues']))

    def test_word_separate_field_value_preserves_instruction_meaning(self):
        doc = Document()
        doc.add_paragraph('药品通用名称：')
        doc.add_paragraph(FOOTERS[2])
        stream = io.BytesIO(); doc.save(stream)
        self.assertTrue(self.upload(stream.getvalue(),'docx')[0])
        self.assertEqual(self.read()['form_json'][NAME]['value'],FOOTERS[2])

    def test_business_label_and_value_in_pdf_margin_remain_filling(self):
        with fitz.open() as doc:
            page = doc.new_page(width=900,height=600)
            page.insert_text((35,575),'药品通用名称：1',fontname='china-s',fontsize=12)
            data = doc.tobytes()
        self.assertTrue(self.upload(data,'pdf')[0])
        state = self.read()
        self.assertEqual(state['form_json'][NAME]['value'],'1')
        self.assertTrue(state['effective_source']['parse_diagnostics']['replacement_eligible'])

    def test_numbered_detail_actual_values_not_removed(self):
        for ext in ('docx','pdf'):
            with self.subTest(ext=ext):
                data = document_bytes(self.root,ext,rows=[['药品通用名称','1'],['原辅包来源',''],
                    ['序号','物料名称','登记号'],['1','测试原料','1234']],footer=FOOTERS[0])
                self.assertTrue(self.upload(data,ext)[0])
                state = self.read()
                self.assertEqual(state['form_json'][NAME]['value'],'1')
                self.assertIn('1234',state['raw_text'])
                self.assertFalse(any(r['text']=='1234' and r['role']!='body' for r in state['effective_source']['parse_diagnostics']['form_content_regions']))

    def test_initial_import_with_persisted_empty_form_row(self):
        self.assertTrue(self.service.save_application_form(self.pid,{'form_json':{}})[0])
        self.assertFalse(self.read().get('original_file_id'))
        self.assertTrue(self.upload(document_bytes(self.root,'pdf',footer='尚待核对的附注'),'pdf')[0])
        state = self.read()
        self.assertEqual(state['parse_status'],'partial')
        self.assertFalse(state['effective_source']['parse_diagnostics']['replacement_eligible'])
        self.assertTrue(self.service.parse_application_form(self.pid)[0])

    def check_unconfirmed(self, ext, options):
        data = document_bytes(self.root,'docx',rows=[['药品通用名称','此前有效药品']])
        self.assertTrue(self.upload(data,'docx')[0])
        before = self.read()
        answer = self.upload(document_bytes(self.root,ext,**options),ext)
        self.assertFalse(answer[0])
        self.assertEqual(answer[2]['code'],'form_content_unconfirmed')
        state = self.read()
        self.assertEqual(state['form_json'],before['form_json'])
        self.assertEqual(state['effective_source'],before['effective_source'])
        self.assertEqual(state['original_file_id'],before['original_file_id'])
        path = self.service._project_path(self.pid)/'application_form'/(before['original_file_id']+'.docx')
        self.assertEqual(path.read_bytes(),data)
        self.assertTrue(state['latest_attempt']['parse_diagnostics']['raw_text'])
        self.assertTrue(state['latest_attempt']['parse_diagnostics']['form_content_regions'])


def _empty_case(ext, existing, options):
    def run(self):
        self.check_empty(ext,existing=existing,**options)
    return run


for _ext in ('pdf','docx'):
    for _existing in (False,True):
        for _label,_opts in [
            *[(f'footer_{i}',{'footer':t}) for i,t in enumerate(FOOTERS)],
            ('repeated',{'pages':2,'header':'申请资料（签章处）','footer':'第1页 共2页'}),
            ('detail',{'detail':True,'footer':FOOTERS[1]}),
            ('body_instruction',{'loose':FOOTERS[2]}),
            ('numbered_instruction',{'loose':'填写说明：1.请逐项填写，空项不填。'}),
            ('split_instruction',{'rows':EMPTY+[['填写说明','请在相应方框内打勾。']]}),
            ('table_instruction',{'rows':EMPTY+[[FOOTERS[2]]]}),
        ]:
            setattr(StructureTests,f'test_empty_{_ext}_{_existing}_{_label}',_empty_case(_ext,_existing,_opts))


def _positive_case(ext, rows, key, value):
    def run(self):
        self.check_positive(ext,rows,key,value)
    return run


for _ext in ('pdf','docx'):
    for _label,_rows,_key,_value in [
        ('name_footer',[['药品通用名称','真实药品']]+EMPTY[1:],NAME,'真实药品'),
        ('number_in_name',[['药品通用名称','1']],NAME,'1'),
        ('instruction_in_name',[['药品通用名称',FOOTERS[2]]],NAME,FOOTERS[2]),
        ('fill_instruction_in_name',[['药品通用名称','填表说明：请在相应方框内打勾。']],NAME,'填表说明：请在相应方框内打勾。'),
        ('pagewords_in_name',[['药品通用名称',FOOTERS[1]]],NAME,FOOTERS[1]),
        ('number_spec',[['规格','1']], 'item_12_specification','1'),
        ('selected_only',[['剂型','☑片剂；胶囊剂□']], 'item_11_dosage_form','片剂'),
        ('unmapped',[['其他自填栏目','这是本次真实填写的补充说明']],None,'这是本次真实填写的补充说明'),
        ('unrecognized_field',[['是否为OTC','目前状态待补充确认']],None,'目前状态待补充确认'),
    ]:
        setattr(StructureTests,f'test_positive_{_ext}_{_label}',_positive_case(_ext,_rows,_key,_value))


def _uncertain_case(ext,options):
    def run(self):
        self.check_unconfirmed(ext,options)
    return run


for _ext in ('pdf','docx'):
    for _label,_opts in [
        ('unknown_headers',{'rows':EMPTY+[['其他资料名称','其他资料说明'],['','']]}),
        ('body_version',{'loose':'表格版本：2026年09月'}),
        ('unanchored_prose',{'loose':'本次实际补充了稳定性研究数据'}),
    ]:
        setattr(StructureTests,f'test_unconfirmed_{_ext}_{_label}',_uncertain_case(_ext,_opts))


def _merged_case(top,value):
    def run(self):
        with fitz.open() as doc:
            page=doc.new_page(width=950,height=1100)
            page.insert_font(fontname='formfont',fontfile=os.environ.get('FILING_TEST_FONT', '/System/Library/Fonts/Supplemental/Arial Unicode.ttf'))
            for y in (top,top+190):
                page.draw_line((35,y),(650,y))
            page.draw_line((220,top+95),(650,top+95))
            for x in (35,220,650):
                page.draw_line((x,top),(x,top+190))
            page.insert_text((50,top+30),'药品通用名称',fontname='formfont',fontsize=12)
            page.insert_text((240,top+165),value,fontname='formfont',fontsize=12)
            doc.subset_fonts()
            data=doc.tobytes(garbage=4,deflate=True)
        self.assertTrue(self.upload(document_bytes(self.root,'docx',rows=[['药品通用名称','此前有效药品']]),'docx')[0])
        self.assertTrue(self.upload(data,'pdf')[0])
        for _ in range(2):
            state=self.read()
            self.assertEqual(state['form_json'][NAME]['value'],value)
            regions=state['form_json'][NAME]['value_sources']['value']['source_regions']
            self.assertTrue(any(r['bbox_pdf'][1] > top+95 for r in regions))
            self.assertTrue(self.service.parse_application_form(self.pid)[0])
    return run


for _top in (250,900):
    for _label,_value in [('number','1'),('instruction',FOOTERS[2])]:
        setattr(StructureTests,f'test_merged_real_cell_{_top}_{_label}',_merged_case(_top,_value))


def _header_case(filled,existing):
    def run(self):
        from docx.oxml import OxmlElement
        if existing:
            original=document_bytes(self.root,'docx',rows=[['药品通用名称','此前有效药品']])
            self.assertTrue(self.upload(original,'docx')[0])
            before=self.read()
        doc=Document(io.BytesIO(document_bytes(self.root,'docx')))
        table=doc.add_table(rows=4,cols=2)
        table.cell(0,0).text='补充说明';table.cell(0,1).text='备注'
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        if filled:
            table.cell(1,0).text='本次补充了稳定性研究数据'
        stream=io.BytesIO();doc.save(stream)
        answer=self.upload(stream.getvalue(),'docx')
        self.assertEqual(answer[0],filled)
        state=self.read()
        if filled:
            self.assertEqual(state['parse_status'],'partial')
            self.assertTrue(state['effective_source']['parse_diagnostics']['replacement_eligible'])
            self.assertIn('本次补充了稳定性研究数据',state['raw_text'])
            self.assertTrue(self.service.parse_application_form(self.pid)[0])
        else:
            self.assertEqual(answer[2]['code'],'empty_form')
            if existing:
                self.assertEqual(state['form_json'],before['form_json'])
                self.assertEqual(state['effective_source'],before['effective_source'])
                self.assertEqual(state['original_file_id'],before['original_file_id'])
                original_path=self.service._project_path(self.pid)/'application_form'/(before['original_file_id']+'.docx')
                self.assertEqual(original_path.read_bytes(),original)
    return run


for _filled in (False,True):
    for _existing in (False,True):
        setattr(StructureTests,f'test_word_explicit_header_filled_{_filled}_existing_{_existing}',_header_case(_filled,_existing))


if __name__ == '__main__':
    unittest.main()
