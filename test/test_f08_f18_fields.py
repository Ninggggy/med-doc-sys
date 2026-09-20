"""F12–F18独立业务期望；真实文件，不使用原审查错误输出作为期望。"""
import tempfile
import unittest
from pathlib import Path
from docx import Document
from agent.test.test_four_semantic_table_defects import word_file, pdf_file
from agent.agent_backend.services.filing_change_form_parser_service import FilingChangeFormParserService


class Fields(unittest.TestCase):
    def test_explicit_fields(self):
        cases = [
            ('商品名称', ['使用☑','商品名：明确商品名'], 'item_10_trade_name', {'trade_name':'明确商品名'}),
            ('药品注册分类',['化学药品','分类号：2.2'],'item_2_registration_category',{'class_no':'2.2'}),
            ('受理前药品注册检验',['是☑','检品编号：ABC123'],'item_19_pre_acceptance_inspection',{'sample_inspection_no':'ABC123'}),
            ('处方',['活性成分：测试成分；辅料：乙醇；是否有变更：否'],'item_16_prescription',{'active_ingredients':'测试成分','excipients':'乙醇','has_change':'否'}),
            ('委托研究机构',['中文名称','甲研究所'],'item_32_cro_info',{'organization_name':'甲研究所'}),
            ('委托研究机构',['研究负责人','张三'],'item_32_cro_info',{'research_lead':'张三'}),
            ('委托研究机构',['地址','甲市一路1号'],'item_32_cro_info',{'address':'甲市一路1号'}),
            ('药品注册申请人',['中文名称','甲公司'],'item_30_applicant_info',{'applicant_name':'甲公司'}),
            ('药品注册申请人',['法定代表人','张三'],'item_30_applicant_info',{'legal_representative':'张三'}),
            ('制剂生产企业',['生产地址','乙市二路二号'],'item_31_manufacturer_info',{'production_address':'乙市二路二号'}),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            for ext in ('docx','pdf'):
                for label, tail, key, expected in cases:
                    with self.subTest(ext=ext,label=label,expected=expected):
                        path=Path(tmp)/('input.'+ext)
                        (word_file if ext=='docx' else pdf_file)(path,[[label]+tail])
                        field=FilingChangeFormParserService().parse_form_file(str(path))['form_json'][key]
                        for k,v in expected.items():self.assertEqual(field['sub_fields'].get(k),v)

    def test_periods(self):
        cases=[
            ('原有效期18月，拟延长后有效期24月','药品有效期由18个月延长至24个月','18月','24月',False),
            ('原有效期2年，拟延长后有效期3年','','2年','3年',False),
            ('原有效期2年，拟延长后有效期3年','药品有效期由24个月延长至36个月','2年','3年',False),
            ('药品有效期由18个月延长至24个月，复验有效期由3个月延长至6个月','','18个月','24个月',False),
            ('24个月','本品有效期保持24个月；药品有效期由18个月延长至36个月','','',True),
            ('原有效期18个月，拟延长后有效期24个月','药品有效期由18个月延长至36个月','','',True),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            for ext in ('docx','pdf'):
                for value,body,before,after,conflict in cases:
                    with self.subTest(ext=ext,value=value,body=body):
                        path=Path(tmp)/('input.'+ext)
                        (word_file if ext=='docx' else pdf_file)(path,[['药品有效期',value],['补充申请的内容',body]])
                        f=FilingChangeFormParserService().parse_form_file(str(path))['form_json']['item_15_validity_period']
                        self.assertEqual(f['sub_fields']['original_validity_period'],before)
                        self.assertEqual(f['sub_fields']['proposed_validity_period'],after)
                        self.assertEqual(any('冲突' in i['reason'] for i in f.get('semantic_issues',[])),conflict)

    def test_statements(self):
        with tempfile.TemporaryDirectory() as tmp:
            for ext in ('docx','pdf'):
                for label,key,value in [('其他相关情况','item_29_other_related_info','已有三次申请记录。'),('其他特别申明事项','special_statement','本申请没有其他申明。')]:
                    with self.subTest(ext=ext,label=label):
                        path=Path(tmp)/('input.'+ext)
                        (word_file if ext=='docx' else pdf_file)(path,[[label,value]])
                        self.assertEqual(FilingChangeFormParserService().parse_form_file(str(path))['form_json'][key]['value'],value)

    def test_material_layouts(self):
        for layout in ('separate','same_table'):
            with self.subTest(layout=layout),tempfile.TemporaryDirectory() as tmp:
                doc=Document()
                if layout=='separate':
                    doc.add_table(rows=1,cols=2).cell(0,0).text='原辅包来源'
                    table=doc.add_table(rows=3,cols=4); offset=0
                else:
                    table=doc.add_table(rows=4,cols=4);table.cell(0,0).merge(table.cell(0,3)).text='原辅包来源';offset=1
                rows=[['物料名称','登记号','受理号','供应商'],['乙醇','Y123','A001','甲供应商公司'],['包装瓶','B456','B002','乙供应商公司']]
                for row,vals in zip(list(table.rows)[offset:],rows):
                    for cell,value in zip(row.cells,vals):cell.text=value
                path=Path(tmp)/'input.docx';doc.save(path)
                f=FilingChangeFormParserService().parse_form_file(str(path))['form_json']
                self.assertEqual(f['item_17_material_source']['table_rows'],[dict(zip(['material_name','register_no','accept_no','manufacturer'],r)) for r in rows[1:]])
                self.assertEqual(f['item_31_manufacturer_info']['sub_fields']['manufacturer_name'],'')
