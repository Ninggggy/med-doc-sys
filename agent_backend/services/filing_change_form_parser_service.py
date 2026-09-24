from pathlib import Path
import re
import zipfile
from copy import deepcopy
from agent.agent_backend.utils.parser.form_field_semantics import selected_option, option_evidence, selection_evidence, heading, validity_values, validity_evidence, normalized_validity
from agent.agent_backend.services.filing_form_revision import values, present
from typing import Any, Dict

from agent.agent_backend.utils.parser import ParserManager
from agent.agent_backend.services.filing_parse_outcome import check_file, form_outcome, MATERIAL_SOURCE_ALIASES
from agent.agent_backend.utils.parser.docx_markdown_parser import convert_doc_to_docx
from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import ocr_pdf_region, parse_drug_supplement_pdf, ocr_region_errors


class FilingChangeFormParserService:
    # 入口标签识别与字段赋值共用映射，值仍保留原文，不依据主体名称推断归属。
    PARTY_SUBFIELDS = {
        30: {'中文名称':'applicant_name', '英文名称':None, '申请人名称':'applicant_name', '法定代表人':'legal_representative', '法人代表':'legal_representative',
             '联系人':'contact', '注册地址':'registered_address', '注册地址(住所)':'registered_address', '地址':'address', '联系电话':'phone', '电话':'phone', '手机':'mobile', '通讯地址':'mailing_address', '注册申请负责人':'registration_contact',
             '统一社会信用代码及组织机构代码':'credit_or_org_code', '统一社会信用代码/组织机构代码':'credit_or_org_code',
             '统一社会信用代码':'credit_code', '组织机构代码':'organization_code',
             '《药品生产许可证》编号':'license_no', '药品生产许可证编号':'license_no', '药品生产许可证号':'license_no'},
        31: {'中文名称':'manufacturer_name', '英文名称':None, '企业名称':'manufacturer_name', '法定代表人':'legal_representative', '法人代表':'legal_representative',
             '联系人':'contact', '注册地址':'registered_address', '注册地址(住所)':'registered_address', '生产地址':'production_address', '地址':'address', '联系电话':'phone', '电话':'phone', '手机':'mobile', '通讯地址':'mailing_address', '注册申请负责人':'registration_contact',
             '统一社会信用代码及组织机构代码':'credit_or_org_code', '统一社会信用代码/组织机构代码':'credit_or_org_code',
             '统一社会信用代码':'credit_code', '组织机构代码':'organization_code',
             '《药品生产许可证》编号':'license_no', '药品生产许可证编号':'license_no', '药品生产许可证号':'license_no'},
        32: {'中文名称':'organization_name', '研究机构名称':'organization_name', '研究负责人':'research_lead',
             '联系人':'contact', '注册地址':'registered_address', '地址':'address', '联系电话':'phone', '电话':'phone', '手机':'mobile', '通讯地址':'mailing_address', '注册申请负责人':'registration_contact'},
    }

    @staticmethod
    def _subfield_label(text):
        return re.sub(r'\s+', '', str(text)).replace('／','/').replace('（','(').replace('）',')').strip(':：')

    def parse_form_file(self, file_path: str) -> Dict[str, Any]:
        check_file(file_path)
        parsed = self._parse_form_file(file_path)
        for field in parsed.get('form_json', {}).values():
            if isinstance(field, dict) and 'field_type' in field:
                field['source_file'] = Path(file_path).name
                for evidence in field.get('value_sources', {}).values():
                    evidence['source_file'] = Path(file_path).name
                for evidence in field.get('selection_evidence', []):
                    evidence['source_file'] = Path(file_path).name
                for issue in field.get('semantic_issues', []):
                    issue['source_file'] = Path(file_path).name
        schema = parsed.get('form_json', {})
        known_blocks = [self._compact_text(schema.get(key, {}).get('source_text', '')) for key in (
            'item_17_material_source', 'item_30_applicant_info', 'item_31_manufacturer_info', 'item_32_cro_info')]
        schema['unassigned_party_sources'] = [
            {'source_text': line, 'source_file': Path(file_path).name, 'status': '主体角色不明确，待核对'}
            for line in str(parsed.get('raw_text', '')).splitlines()
            if re.search(r'公\s*司', line) and not any(self._compact_text(line) in block for block in known_blocks)
        ]
        parsed.update(form_outcome(parsed))
        return parsed

    def _parse_form_file(self, file_path: str) -> Dict[str, Any]:
        path = Path(file_path)
        ext = path.suffix.lower().lstrip(".")
        if ext == "doc":
            converted = self._try_convert_doc_to_docx(path)
            if converted is not None:
                return self._parse_docx_form(str(converted))
        if ext == "docx":
            return self._parse_docx_form(str(path))
        if ext == "pdf":
            return self._parse_pdf_form(str(path))
        if not ParserManager.is_supported(ext):
            raw_text = path.read_text(encoding="utf-8", errors="ignore")
            return {"form_json": self._build_full_form(raw_text), "raw_text": raw_text, "confidence": 0.5}
        rows = ParserManager.parse(str(path), ext_hint=ext)
        raw_text = "\n".join(str(item.get("text", "") or "") for item in rows if isinstance(item, dict))
        return {"form_json": self._build_full_form(raw_text), "raw_text": raw_text, "confidence": 0.8}

    def _try_convert_doc_to_docx(self, path: Path) -> Path | None:
        try:
            return convert_doc_to_docx(path)
        except Exception:
            return None

    def _parse_docx_form(self, file_path: str) -> Dict[str, Any]:
        raw_text = ""
        table_rows = []
        structured_records = []
        try:
            from docx import Document

            doc = Document(file_path)
            from agent.agent_backend.utils.parser.form_content_evidence import word_form_structure
            structure = word_form_structure(doc)
            paras = [p.text.strip() for p in doc.paragraphs if str(p.text or "").strip()]
            raw_text = "\n".join(paras)
            for tbl in doc.tables:
                rows = []
                for row in tbl.rows:
                    cells = [str(c.text or "").strip() for c in row.cells]
                    if any(cells):
                        rows.append(cells)
                if rows:
                    table_rows.append(rows)
                    for row in rows:
                        raw_text += "\n" + " | ".join(row)
            excluded_paragraphs = {r['paragraph'] for r in structure['regions']
                                   if r['role'] != 'body' and 'paragraph' in r}
            excluded_rows = {(r['table'],r['row']) for r in structure['regions'] if r['reason']=='Word重复列头结构'}
            structured_records = self._extract_docx_structured_records(doc, excluded_paragraphs=excluded_paragraphs,
                                                                       excluded_table_rows=excluded_rows)
        except Exception:
            raise
        body_text = structure['body_text']
        schema = self._build_full_form(body_text)
        self._apply_structured_records_to_schema(schema, structured_records)
        # 按实际正文顺序关联平铺、独立续表与嵌套表，并在下一事项处结束。
        material = schema['item_17_material_source']
        details, regions = self._docx_material_details(doc)
        if details:
            material['table_rows'] = details
            material['source_regions'] = regions
            material['source_text'] = '\n'.join(r['source_text'] for r in regions)
            material['parse_confidence'] = .85
        self._fill_core_fields_from_text(schema, body_text)
        self._finalize_fields(schema)
        # 状态与显示原文单独留证；规范化勾选文本不替代原始控件事实。
        controls = []
        for index, control in enumerate(n for n in doc.element.iter() if n.tag.rsplit('}',1)[-1]=='sdt'):
            states = [dict(n.attrib) for n in control.iter() if n.tag.rsplit('}',1)[-1]=='checked']
            if not states: continue
            controls.append({'control':index, 'coordinate_unit':'word_control', 'states':states,
                             'display_text': ''.join(str(n.text or '') for n in control.iter() if n.tag.rsplit('}',1)[-1]=='t'),
                             'symbols':[dict(n.attrib) for n in control.iter() if n.tag.rsplit('}',1)[-1]=='sym']})
        if controls: schema['word_control_sources'] = controls
        return {
            "form_json": schema,
            "raw_text": raw_text,
            "structured_records": structured_records,
            "form_structure": structure,
            "confidence": 0.85,
        }

    def _parse_pdf_form(self, file_path: str) -> Dict[str, Any]:
        from agent.agent_backend.services.filing_paddle_backend import enabled,parse_pdf
        if enabled():
            from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import build_pdf_form_from_pages
            parsed=build_pdf_form_from_pages(parse_pdf(file_path),source_file=file_path)
            parsed['ocr_backend']='ppocr_v6_medium_experimental'
            return self.form_from_pdf_result(parsed)
        parsed = parse_drug_supplement_pdf(file_path)
        return self.form_from_pdf_result(parsed, file_path=file_path)

    def form_from_pdf_result(self, parsed, *, file_path=None):
        """原件解析和人工修订共用字段映射；无file_path时不允许再次OCR。"""
        raw_text = str(parsed.get("raw_text", "") or "")
        # 显式实验后端已完成识别；后续字段消费者不得另起旧OCR覆盖证据。
        if parsed.get('ocr_backend') == 'ppocr_v6_medium_experimental':
            file_path = None
        structure = parsed.get('form_structure')
        body_text = structure['body_text'] if structure is not None else raw_text
        extra_blocks = parsed.get("extra_blocks", {}) if isinstance(parsed, dict) else {}
        region_errors = []
        token = ocr_region_errors.set(region_errors)
        try:
            items = self._refine_pdf_items(parsed, file_path) if file_path else deepcopy(parsed.get('items', []))
        finally:
            ocr_region_errors.reset(token)

        schema = self._build_full_form(body_text)
        if isinstance(extra_blocks, dict):
            special_statement = str(extra_blocks.get("其他特别申明事项", "") or "").strip()
            if special_statement:
                schema["special_statement"]["value"] = special_statement
                schema["special_statement"]["parse_confidence"] = 0.7

        self._apply_pdf_items_to_schema(schema, items, raw_text=body_text)
        self._apply_pdf_global_fallbacks(schema, body_text)
        token = ocr_region_errors.set(region_errors)
        try:
            if file_path:
                self._apply_pdf_region_ocr_enhancements(schema, items, file_path)
        finally:
            ocr_region_errors.reset(token)
        for error in region_errors:
            page = next((p for p in parsed.get('pages', []) if p.get('page') == error['page']), None)
            if page is not None:
                page.setdefault('errors', []).append(error)
            else:
                parsed.setdefault('region_errors', []).append(error)
        self._sanitize_pdf_field_values(schema)
        self._fill_core_fields_from_text(schema, body_text)
        self._finalize_fields(schema)
        return {
            "form_json": schema,
            "raw_text": raw_text,
            "pdf_parse_result": parsed,
            "form_structure": structure,
            "confidence": 0.75,
        }

    def _refine_pdf_items(self, parsed, file_path):
        """只补识别低质量申请字段。原页、原 OCR、表格格线和补识别证据各自保留。"""
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import item_start, parse_subfields, compact_for_matching, SECTION_LINES
        items = deepcopy(parsed.get('items', []))
        pages = parsed.get('pages', [])
        for item in items:
            number = item['item_no']
            if number not in (3, 11, 16, 17, 20, 25, 30, 31):
                continue
            evidence = []
            for page in pages:
                if not page.get('ocr_calls') or not page.get('ocr_quality'):
                    continue
                if not any(q.get('low_confidence_word_count') for q in page['ocr_quality']):
                    continue
                heads = [(line, item_start(line['text'])) for line in page['lines'] if item_start(line['text'])]
                anchors = [line for line, found in heads if found[0] == number]
                if not anchors:
                    continue
                for anchor in anchors:
                    # 并列标题沿用原区域分配，不能将整行裁成单个字段。
                    if any(other is not anchor and abs(other['bbox'][1]-anchor['bbox'][1]) < 4 for other, _ in heads):
                        continue
                    boundaries = [line['bbox'][1] for line, _ in heads if line['bbox'][1] > anchor['bbox'][3]-2]
                    boundaries += [line['bbox'][1] for line in page['lines']
                                   if compact_for_matching(line['text']).strip(' |:') in SECTION_LINES
                                   and line['bbox'][1] > anchor['bbox'][3]]
                    end = min(boundaries + [page['page_bbox'][3]])
                    lines = [line for line in page['lines'] if anchor['bbox'][1] <= line['bbox'][1] < end]
                    box = [min(l['bbox'][0] for l in lines)-2, anchor['bbox'][1]-2,
                           max(l['bbox'][2] for l in lines)+2, end]
                    region = {'page': page['page'], 'bbox_pdf': box, 'coordinate_unit': 'pdf_point'}
                    if number == 17:
                        for table in item.get('tables', []):
                            if table['page'] != page['page']:
                                continue
                            table['original_rows'] = deepcopy(table['rows'])
                            for cell in table.get('cells', []):
                                cell_region = {**region, 'bbox_pdf': cell['bbox_pdf']}
                                text = ocr_pdf_region(file_path, cell_region, zoom=6, margin_pt=-1, psm=6, lang='chi_sim')
                                # 真空白仍为空；原文和单元格下标保留，不能凭位置填默认内容。
                                evidence.append({'source_text': text, 'original_text': cell.get('text', ''),
                                                 'source_regions': [cell_region], 'row': cell['row'], 'column': cell['column'], 'method': 'cell_ocr'})
                                if text:
                                    cell['original_text'] = cell.get('text', '')
                                    cell['text'] = re.sub(r'\s+', '', text)
                                    table['rows'][cell['row']][cell['column']] = cell['text']
                            # 补识别后的表格视图统一来自同一行列；原 OCR 另留 original_rows/raw_text。
                            table['columns'], table['data'] = deepcopy(table['rows'][0]), deepcopy(table['rows'][1:])
                            def markdown_row(row):
                                return '| ' + ' | '.join(str(v).replace('|', '\\|').replace('\n', ' ') for v in row) + ' |'
                            table['markdown'] = '\n'.join([markdown_row(table['columns']), markdown_row(['---'] * len(table['columns']))]
                                                         + [markdown_row(row) for row in table['data']])
                        continue
                    zoom = 5 if number in (3, 11, 20) else 4
                    text = ocr_pdf_region(file_path, region, zoom=zoom, margin_pt=0, psm=6, lang='chi_sim+eng' if number in (3, 11, 20) else 'chi_sim')
                    evidence.append({'source_text': text, 'source_regions': [region], 'method': 'region_ocr', 'zoom': zoom})
                    if not text:
                        continue
                    found = heading(text)
                    if found and found[0] != number:
                        continue
                    text = found[2] if found else text
                    # 签字空栏不是法定代表人姓名的第二个候选。
                    if number in (30, 31):
                        text = re.split(r'法定代表人\s*[(（]?\s*签名', text)[0].strip()
                    original_subfields = parse_subfields(item.get('raw_text', ''))
                    subfields = parse_subfields(text)
                    if number in (30, 31):
                        for line in lines:
                            compact = compact_for_matching(line['text'])
                            if compact.startswith('法定代表人:') and '签' not in compact:
                                line_region = {**region, 'bbox_pdf': [line['bbox'][0]-1, line['bbox'][1]-1, line['bbox'][2]+1, line['bbox'][3]+1]}
                                read = ocr_pdf_region(file_path, line_region, zoom=3, margin_pt=0, psm=7, lang='chi_sim')
                                evidence.append({'source_text': read, 'source_regions': [line_region], 'method': 'subfield_ocr'})
                                match = re.match(r'法定代表人\s*[:,;]\s*(.*?)\s*(?:职位|$)', read)
                                if match:
                                    subfields['法定代表人'] = match[1].strip()
                                    item.setdefault('subfield_sources', {})['法定代表人'] = {'source_text': read, 'source_regions': [line_region]}
                            if '统一社会信用代码' in compact:
                                words = [w for w in page.get('words', []) if line['bbox'][1] <= w['bbox'][1] < line['bbox'][3]
                                         and any(c.isdigit() for c in w['text'])]
                                if words:
                                    code_region = {**region, 'bbox_pdf': [min(w['bbox'][0] for w in words), min(w['bbox'][1] for w in words),
                                                                        max(w['bbox'][2] for w in words), max(w['bbox'][3] for w in words)]}
                                    read = ocr_pdf_region(file_path, code_region, zoom=2.5, margin_pt=1, psm=7, lang='eng')
                                    evidence.append({'source_text': read, 'source_regions': [code_region], 'method': 'subfield_ocr'})
                                    code = re.sub(r'\s+', '', read)
                                    label = next((k for k in subfields if k.startswith('统一社会信用代码')), None)
                                    if label and isinstance(subfields[label], str):
                                        previous = re.sub(r'\s+', '', subfields[label])
                                        if len(previous) == 18 or len(code) == 18:
                                            candidates = [
                                                {'value': previous, 'source_text': text, 'source_regions': [region]},
                                                {'value': code, 'source_text': read, 'source_regions': [code_region]},
                                            ]
                                            self._resolve_credit_ocr(item, subfields, label, candidates, number)
                                            evidence[-1].update(accepted=bool(code and code == subfields[label]),
                                                validation='按18位合法字符集及同字段候选一致性选择；不进行字符猜替换')
                        address_lines = [line for line in lines if compact_for_matching(line['text']).startswith('注册地址')]
                        if '注册地址' not in subfields and address_lines:
                            item.setdefault('refinement_issues', []).append({'path': 'sub_fields.registered_address',
                                'reason': '已定位注册地址行，但局部识别未能可靠恢复标签及完整地址；原文和区域保留，本字段未完成',
                                'source_text': '\n'.join(line['text'] for line in address_lines) or text,
                                'source_regions': [{**region, 'bbox_pdf': line['bbox']} for line in address_lines] or [region]})
                    if number == 16:
                        # 低质量整块识别不能直接确认成分；按真实子字段边界复核。
                        if not self._refine_active_ingredients(file_path, item, lines, region, subfields, evidence, text) and original_subfields.get('活性成分/中药成分/...'):
                            subfields['活性成分/中药成分/...'] = original_subfields['活性成分/中药成分/...']
                        for line in lines:
                            words = [w for w in page.get('words', []) if line['bbox'][0] <= w['bbox'][0] < line['bbox'][2]
                                     and line['bbox'][1] <= w['bbox'][1] < line['bbox'][3]]
                            if len(words) != 2:
                                continue
                            reads = [ocr_pdf_region(file_path, {**region, 'bbox_pdf': w['bbox']}, zoom=3,
                                                    margin_pt=1, psm=7, lang='chi_sim') for w in words]
                            if re.sub(r'\W', '', reads[0]) == '是否有变更' and reads[1].strip() in ('是', '否'):
                                subfields['是否有变更'] = reads[1].strip()
                                evidence.append({'source_text': '：'.join(reads), 'source_regions': [{**region, 'bbox_pdf': line['bbox']}],
                                                 'method': 'label_and_value_ocr', 'zoom': 3})
                    item.setdefault('original_raw_text', item.get('raw_text', ''))
                    item.update(raw_text=text, subfields=subfields, option_fragments=[text], source_regions=[region])
                    if number == 25:
                        controls = self._pdf_image_choices(file_path, region, ['是', '否'])
                        item['visual_choices'] = controls
                        if controls:
                            item['option_fragments'] = [('☑' if c['selected'] else '□') + c['name'] for c in controls]
            if evidence:
                item['recognition_evidence'] = evidence
        parsed['refined_items'] = items
        return items

    def _resolve_credit_ocr(self, item, subfields, label, candidates, number):
        # GB 32100 的18位字符集不含 I/O/S/V/Z；只校验格式，不据此声称主体合法。
        # https://scjgj.beijing.gov.cn/zwxx/scjgdt/202407/t20240719_3753300.html
        valid = [c for c in candidates if re.fullmatch(r'[0-9ABCDEFGHJKLMNPQRTUWXY]{18}', c['value'])]
        unique = {c['value'] for c in valid}
        item.setdefault('credit_ocr_candidates', []).extend(deepcopy(candidates))
        if len(unique) == 1:
            selected = valid[0]
            subfields[label] = selected['value']
            item.setdefault('subfield_sources', {})[label] = deepcopy(selected)
            return
        subfields[label] = ''
        target = self.PARTY_SUBFIELDS[number].get(self._subfield_label(label), 'credit_or_org_code')
        item.setdefault('refinement_issues', []).append({
            'path': 'sub_fields.' + target,
            'reason': '统一社会信用代码识别结果冲突，待核对' if unique else '统一社会信用代码未识别为符合18位字符规则的值，待核对',
            'candidates': deepcopy(candidates),
            'source_regions': [r for c in candidates for r in c['source_regions']],
        })

    def _refine_active_ingredients(self, file_path, item, lines, region, subfields, evidence, region_text=''):
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import compact_for_matching, parse_subfields
        label = '活性成分/中药成分/...'
        anchors = [line for line in lines if '活性成' in compact_for_matching(line['text'])]
        if len(anchors) != 1:
            return False
        anchor = anchors[0]
        following = [line for line in lines if line['bbox'][1] > anchor['bbox'][1]
                     and re.match(r'^(?:辅料|是否有变更)[:：]', compact_for_matching(line['text']))]
        if not following:
            return False
        end = min(line['bbox'][1] for line in following)
        body = [line for line in lines if anchor['bbox'][1] <= line['bbox'][1] < end]
        # 连续填写行都属于本子字段；用后续标签截断，不将多行成分缩成第一行。
        box = [min(line['bbox'][0] for line in body), anchor['bbox'][1],
               max(line['bbox'][2] for line in body), min(end, max(line['bbox'][3] for line in body))]
        target = {**region, 'bbox_pdf': box}
        candidates = []
        for zoom in (3, 4):
            read = ocr_pdf_region(file_path, target, zoom=zoom, margin_pt=1,
                                  psm=7 if len(body) == 1 else 6, lang='chi_sim')
            value = parse_subfields(read).get(label, '')
            candidate = {'value': value if isinstance(value, str) else '', 'source_text': read,
                         'source_regions': [target], 'method': 'active_ingredient_ocr', 'zoom': zoom}
            candidates.append(candidate)
            evidence.append(candidate)
        if candidates[0]['value'] and candidates[0]['value'] == candidates[1]['value']:
            subfields[label] = candidates[0]['value']
            item.setdefault('subfield_sources', {})[label] = deepcopy(candidates[0])
        else:
            previous = {'value': subfields.get(label, ''), 'source_text': region_text,
                        'source_regions': [region]}
            subfields[label] = ''
            item.setdefault('refinement_issues', []).append({
                'path': 'sub_fields.active_ingredients', 'reason': '活性成分局部识别不一致或未完整识别，待核对',
                'candidates': [previous, *candidates], 'source_regions': [target],
            })
        return True

    def _pdf_image_choices(self, file_path, region, options):
        """用原 PDF 控件图像边界和中心填充验证单选，标签仍由现有 OCR 识别。"""
        import fitz
        with fitz.open(file_path) as doc:
            page = doc[region['page']-1]
            page.set_rotation(0)
            area = fitz.Rect(region['bbox_pdf'])
            boxes = []
            for image in page.get_image_info():
                rect = fitz.Rect(image['bbox'])
                if rect in area and .85 < rect.width/max(rect.height, .01) < 1.15 and not any(rect.intersects(b) for b in boxes):
                    boxes.append(rect)
            boxes.sort(key=lambda r: r.x0)
            if len(boxes) != len(options) or max(r.y0 for r in boxes)-min(r.y0 for r in boxes) > min(r.height for r in boxes)*.3:
                return []
            entries = []
            for rect in boxes:
                pix = page.get_pixmap(matrix=fitz.Matrix(4, 4), clip=rect, colorspace=fitz.csGRAY, alpha=False)
                samples = pix.samples
                center = [samples[y*pix.width+x] for y in range(int(pix.height*.35), int(pix.height*.65))
                          for x in range(int(pix.width*.35), int(pix.width*.65))]
                brightness = sum(center)/max(len(center), 1)
                if 225 < brightness < 240:
                    return []
                label_region = {**region, 'bbox_pdf': [rect.x1+1, rect.y0-1, rect.x1+rect.width*1.6, rect.y1+2]}
                label = ocr_pdf_region(file_path, label_region, zoom=3, margin_pt=0, psm=7, lang='chi_sim').strip()
                if label not in options:
                    return []
                entries.append({'name': label, 'selected': brightness <= 225, 'status': 'selected' if brightness <= 225 else 'unselected',
                                'source_text': label, 'source_regions': [{**region, 'bbox_pdf': list(rect)}, label_region],
                                'method': 'pdf_control_image_and_label_ocr', 'center_brightness': brightness})
            return entries if {e['name'] for e in entries} == set(options) else []

    def _apply_pdf_items_to_schema(self, schema: Dict[str, Any], items: list[Dict[str, Any]], raw_text: str = "") -> None:
        if not isinstance(schema, dict) or not isinstance(items, list):
            return
        item_map = {}
        for item in items:
            if not isinstance(item, dict):
                continue
            try:
                item_no = int(item.get("item_no"))
            except Exception:
                continue
            item_map[item_no] = item

        self._set_item_source(schema, "item_1_application_type", item_map.get(1))
        self._set_item_source(schema, "item_2_registration_category", item_map.get(2))
        self._set_item_source(schema, "item_3_otc_type", item_map.get(3))
        self._set_item_source(schema, "item_4_original_product_status", item_map.get(4))
        self._set_item_source(schema, "item_5_application_matter_category", item_map.get(5))
        self._set_item_source(schema, "item_6_generic_name", item_map.get(6))
        self._set_item_source(schema, "item_7_english_or_latin_name", item_map.get(7))
        self._set_item_source(schema, "item_8_pinyin", item_map.get(8))
        self._set_item_source(schema, "item_9_chemical_name", item_map.get(9))
        self._set_item_source(schema, "item_10_trade_name", item_map.get(10))
        self._set_item_source(schema, "item_11_dosage_form", item_map.get(11))
        self._set_item_source(schema, "item_12_specification", item_map.get(12))
        self._set_item_source(schema, "item_13_other_specifications", item_map.get(13))
        self._set_item_source(schema, "item_14_packaging", item_map.get(14))
        self._set_item_source(schema, "item_15_validity_period", item_map.get(15))
        self._set_item_source(schema, "item_16_prescription", item_map.get(16))
        self._set_item_source(schema, "item_17_material_source", item_map.get(17))
        self._set_item_source(schema, "item_18_tcm_material_standard", item_map.get(18))
        self._set_item_source(schema, "item_19_pre_acceptance_inspection", item_map.get(19))
        self._set_item_source(schema, "item_20_indication_or_function", item_map.get(20))
        self._set_item_source(schema, "item_21_change_content", item_map.get(21))
        self._set_item_source(schema, "item_22_change_reason", item_map.get(22))
        self._set_item_source(schema, "item_23_original_approval_info", item_map.get(23))
        self._set_item_source(schema, "item_24_patent_info", item_map.get(24))
        self._set_item_source(schema, "item_25_data_protection", item_map.get(25))
        self._set_item_source(schema, "item_26_tcm_protection", item_map.get(26))
        self._set_item_source(schema, "item_27_same_product_period", item_map.get(27))
        self._set_item_source(schema, "item_28_change_related_items", item_map.get(28))
        self._set_item_source(schema, "item_29_other_related_info", item_map.get(29))
        self._set_item_source(schema, "item_30_applicant_info", item_map.get(30))
        self._set_item_source(schema, "item_31_manufacturer_info", item_map.get(31))

        self._set_item_source(schema, "special_statement", item_map.get(33))
        if item_map.get(33):
            schema["special_statement"]["value"] = self._pdf_item_text(item_map[33])

        item_1_text = self._pdf_item_text(item_map.get(1))
        if self._contains_compact(item_1_text, "境内生产药品补充申请") or self._contains_compact(
            self._pdf_item_text({"raw_text": schema.get("item_1_application_type", {}).get("source_text", "")}),
            "境内生产药品补充申请",
        ):
            schema["item_1_application_type"]["value"] = "境内生产药品补充申请"
            schema["item_1_application_type"]["parse_confidence"] = 0.8

        for field_key, item_no in [
            ("item_2_registration_category", 2),
            ("item_3_otc_type", 3),
            ("item_4_original_product_status", 4),
        ]:
            field = schema.get(field_key, {}) or {}
            item_text = self._pdf_item_text(item_map.get(item_no))
            field["value"] = selected_option(item_text, field.get("options", []))
            field["parse_confidence"] = 0.75 if field["value"] else 0.0

        item_5 = item_map.get(5) or {}
        item_5_text = self._pdf_item_text(item_5)
        self._apply_multiple_choice(schema["item_5_application_matter_category"],
                                    item_5.get('option_fragments', item_5_text), matter=True)

        for field_key, item_no in [
            ("item_6_generic_name", 6),
            ("item_7_english_or_latin_name", 7),
            ("item_8_pinyin", 8),
            ("item_9_chemical_name", 9),
            ("item_12_specification", 12),
            ("item_13_other_specifications", 13),
            ("item_18_tcm_material_standard", 18),
            ("item_21_change_content", 21),
            ("item_22_change_reason", 22),
            ("item_23_original_approval_info", 23),
            ("item_24_patent_info", 24),
            ("item_25_data_protection", 25),
            ("item_26_tcm_protection", 26),
            ("item_27_same_product_period", 27),
            ("item_28_change_related_items", 28),
            ("item_29_other_related_info", 29),
        ]:
            item_text = self._pdf_item_text(item_map.get(item_no))
            if item_no == 6:
                item_text = re.sub(r'(?<=[\u4e00-\u9fff])\s*\n\s*(?=[\u4e00-\u9fff])', '', item_text)
            if item_text:
                schema[field_key]["value"] = item_text
                schema[field_key]["parse_confidence"] = 0.6

        item_10_text = self._pdf_item_text(item_map.get(10))
        if item_10_text:
            schema["item_10_trade_name"]["value"] = selected_option(item_10_text, ["使用", "不使用"])
            schema["item_10_trade_name"]["parse_confidence"] = 0.6

        item_11_text = self._pdf_item_text(item_map.get(11))
        if item_11_text:
            self._apply_multiple_choice(schema["item_11_dosage_form"],
                                        item_map[11].get('option_fragments', item_11_text))

        item_14 = item_map.get(14) or {}
        item_14_subfields = item_14.get("subfields", {}) if isinstance(item_14, dict) else {}
        if isinstance(item_14_subfields, dict):
            material = str(item_14_subfields.get("直接接触药品的包装材料和容器", "") or "").strip()
            spec = str(item_14_subfields.get("包装规格", "") or "").strip()
            if material:
                schema["item_14_packaging"]["sub_fields"]["primary_packaging_material"] = material
            if spec:
                schema["item_14_packaging"]["sub_fields"]["packaging_specification"] = spec
            if material or spec:
                schema["item_14_packaging"]["parse_confidence"] = 0.7

        item_15 = item_map.get(15) or {}
        item_15_text = self._pdf_item_text(item_15)
        if item_15_text:
            original = self._first_match(item_15_text, [r"原有效期[:：]?\s*([^\n\r]{1,30})"])
            proposed = self._first_match(item_15_text, [r"(?:拟延长后有效期|变更后有效期)[:：]?\s*([^\n\r]{1,30})"])
            storage = self._first_match(item_15_text, [r"贮藏条件[:：]?\s*([^\n\r]{1,120})"])
            relation = validity_values(item_15_text, context='药品有效期')
            original = relation.get('original_validity_period', '')
            proposed = relation.get('proposed_validity_period', '')
            # 单独“药品有效期”只表示本表填写值，不猜测其变更前后角色。
            if not relation and re.fullmatch(r'\d+\s*个?\s*月', item_15_text):
                schema["item_15_validity_period"]["reported_value"] = item_15_text
            if original:
                schema["item_15_validity_period"]["sub_fields"]["original_validity_period"] = original
            if proposed:
                schema["item_15_validity_period"]["sub_fields"]["proposed_validity_period"] = proposed
            if storage:
                schema["item_15_validity_period"]["sub_fields"]["storage_condition"] = storage
            if original or proposed or storage:
                schema["item_15_validity_period"]["parse_confidence"] = 0.7

        item_16 = item_map.get(16) or {}
        item_16_subfields = item_16.get("subfields", {}) if isinstance(item_16, dict) else {}
        if isinstance(item_16_subfields, dict):
            active = str(item_16_subfields.get("活性成分/中药成分/...", "") or "").strip()
            excipients = str(item_16_subfields.get("辅料", "") or "").strip()
            changed = str(item_16_subfields.get("是否有变更", "") or "").strip()
            if active:
                self._assign_subfield(schema['item_16_prescription'], 'active_ingredients', active,
                                      item_16.get('subfield_sources', {}).get('活性成分/中药成分/...') or
                                      {'source_text': item_16.get('raw_text', ''), 'source_regions': item_16.get('source_regions', [])})
            if excipients:
                schema["item_16_prescription"]["sub_fields"]["excipients"] = excipients
            if changed:
                schema["item_16_prescription"]['option_fragments'] = [changed]
                schema["item_16_prescription"]["sub_fields"]["has_change"] = selected_option(changed, ["是", "否"])
            if active or excipients or changed:
                schema["item_16_prescription"]["parse_confidence"] = 0.65

        item_17 = item_map.get(17) or {}
        if isinstance(item_17, dict):
            table_rows = []
            layout = {}
            for table in item_17.get("tables", []) or []:
                rows = table.get("rows", []) if isinstance(table, dict) else []
                # 人工逻辑单元格没有逐格图像坐标，不伪造网格或沿用上一表列布局。
                edges = sorted({round(c['bbox_pdf'][i], 1) for c in table.get('cells', [])
                                if len(c.get('bbox_pdf') or []) == 4 for i in (0, 2)})
                grid = tuple(round(x-edges[0],1) for x in edges) if edges else ()
                if not grid or layout.get('grid') != grid:
                    layout.clear()
                guessed = self._guess_material_source_rows([rows], layout)
                if layout.get('columns'):
                    layout['grid'] = grid
                table_rows.extend(guessed)
            if table_rows:
                schema["item_17_material_source"]["table_rows"] = table_rows
                schema["item_17_material_source"]["parse_confidence"] = 0.7

        item_19_text = self._pdf_item_text(item_map.get(19))
        if item_19_text:
            schema["item_19_pre_acceptance_inspection"]["value"] = selected_option(item_19_text, ["是", "否"])
            sample_no = self._first_match(item_19_text, [r"检品编号[:：]?\s*([A-Za-z0-9\-_/]+)"])
            if sample_no:
                schema["item_19_pre_acceptance_inspection"]["sub_fields"]["sample_inspection_no"] = sample_no
            if schema["item_19_pre_acceptance_inspection"]["value"] or sample_no:
                schema["item_19_pre_acceptance_inspection"]["parse_confidence"] = 0.6

        item_20 = item_map.get(20) or {}
        item_20_subfields = item_20.get("subfields", {}) if isinstance(item_20, dict) else {}
        item_20_text = self._pdf_item_text(item_20)
        if item_20_text:
            indication = schema['item_20_indication_or_function']['sub_fields']
            category = re.match(r'(?:[（(]\s*分类|适应症分类)\s*[:：]\s*([^\n）)]*)[）)]?\s*(?:\n|$)(.*)', item_20_text, re.S)
            if category:
                indication.update(category=category[1].strip(), description=category[2].strip())
            else:
                indication['description'] = item_20_text
                lines = [line.strip() for line in item_20_text.splitlines() if line.strip()]
                if len(lines) > 1 and re.fullmatch(r'[A-Z\s:：./-]{5,}', lines[0]) and any(re.search(r'[\u4e00-\u9fff]', line) for line in lines[1:]):
                    field = schema['item_20_indication_or_function']
                    field.setdefault('semantic_issues', []).append({'reason': '前置文字无法确认属于分类标签或填写内容，已保留原文，请人工核对',
                        'source_text': lines[0], 'source_regions': item_20.get('source_regions', [])})
        if (
            item_20_text
            and (
                str(schema["item_20_indication_or_function"]["sub_fields"].get("category", "")).strip()
                or str(schema["item_20_indication_or_function"]["sub_fields"].get("description", "")).strip()
            )
        ):
            schema["item_20_indication_or_function"]["parse_confidence"] = 0.55

        applicant_item = item_map.get(30)
        manufacturer_item = item_map.get(31)
        if manufacturer_item and '委托研究' in str(manufacturer_item.get('title', '')):
            item_map[32] = manufacturer_item
            manufacturer_item = None
            schema["item_31_manufacturer_info"]["source_text"] = ""
        cro = item_map.get(32)
        if cro:
            self._set_item_source(schema, "item_32_cro_info", cro)
            self._apply_pdf_party_fields(schema["item_32_cro_info"], cro)
            party = schema["item_32_cro_info"]["sub_fields"]


        applicant_text, manufacturer_text = self._split_pdf_party_blocks(self._pdf_item_text(applicant_item))
        self._apply_pdf_party_fields(
            schema["item_30_applicant_info"],
            applicant_item,
            raw_override=applicant_text or self._pdf_item_text(applicant_item),
            default_name_role="applicant",
        )
        self._apply_pdf_party_fields(
            schema["item_31_manufacturer_info"],
            manufacturer_item,
            raw_override=manufacturer_text or self._pdf_item_text(manufacturer_item),
            default_name_role="manufacturer",
        )
        if not str(schema["item_31_manufacturer_info"]["sub_fields"].get("manufacturer_name", "")).strip():
            global_manufacturer_text = self._extract_global_manufacturer_block(raw_text)
            if global_manufacturer_text:
                self._apply_pdf_party_fields(
                    schema["item_31_manufacturer_info"],
                    manufacturer_item,
                    raw_override=global_manufacturer_text,
                    default_name_role="manufacturer",
                )

    def _apply_pdf_party_fields(
        self,
        field_schema: Dict[str, Any],
        item: Dict[str, Any] | None,
        raw_override: str = "",
        default_name_role: str = "applicant",
    ) -> None:
        if not isinstance(field_schema, dict):
            return
        item_text = str(raw_override or self._pdf_item_text(item) or "")
        subfields = item.get("subfields", {}) if isinstance(item, dict) else {}
        if not isinstance(subfields, dict):
            subfields = {}
        number = 32 if 'organization_name' in field_schema['sub_fields'] else 31 if 'manufacturer_name' in field_schema['sub_fields'] else 30
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import parse_subfields
        parsed = {**subfields, **parse_subfields(item_text)}
        if (item or {}).get('recognition_evidence'):
            parsed.update(subfields)
        for label, value in parsed.items():
            target = self.PARTY_SUBFIELDS[number].get(self._subfield_label(label))
            if target:
                candidates = value if isinstance(value, list) else [value]
                for candidate in candidates:
                    self._assign_subfield(field_schema, target, str(candidate).strip(),
                                          (item or {}).get('subfield_sources', {}).get(label) or
                                          {'source_text':label + '：' + str(candidate), 'source_regions':(item or {}).get('source_regions',[])})
        if item_text:
            field_schema['parse_confidence'] = .75

    def _assign_subfield(self, field, target, value, origin):
        previous = field['sub_fields'].get(target, '')
        source = {'value':value, 'source_text':origin.get('source_text',''),
                  'source_regions':deepcopy(origin.get('source_regions',[])), 'parse_confidence':.88}
        if previous and value and previous != value:
            field.setdefault('semantic_issues', []).append({'reason':'同一主体子字段有多个值，待核对',
                'path':'sub_fields.'+target, 'candidates':[field.get('value_sources',{}).get('sub_fields.'+target, {'value':previous}),source]})
            # 多主体不静默覆盖，保留所有候选和各自位置。
            field['sub_fields'][target] = ''
        else:
            field['sub_fields'][target] = value
        field.setdefault('value_sources', {})['sub_fields.'+target] = source

    def _set_item_source(self, schema: Dict[str, Any], field_key: str, item: Dict[str, Any] | None) -> None:
        if not isinstance(schema, dict):
            return
        field = schema.get(field_key)
        if not isinstance(field, dict):
            return
        if not isinstance(item, dict):
            return
        for key in ("source_regions", "source_pages", "title", "option_fragments", "recognition_evidence", "visual_choices"):
            if key in item:
                field[key] = deepcopy(item[key])
        if item.get('refinement_issues'):
            field.setdefault('semantic_issues', []).extend(deepcopy(item['refinement_issues']))
        source_text = str(item.get("original_raw_text", item.get("raw_text", "")) or "").strip()
        if item.get('recognition_evidence'):
            field['recognized_text'] = item.get('raw_text', '')
        if source_text:
            field["source_text"] = source_text

    def _pdf_item_text(self, item: Dict[str, Any] | None) -> str:
        if not isinstance(item, dict):
            return ""
        return str(item.get("raw_text", "") or "").strip()

    def _first_match(self, text: str, patterns: list[str]) -> str:
        source = str(text or "")
        for pattern in patterns:
            match = re.search(pattern, source)
            if match:
                return str(match.group(1) or "").strip()
        return ""

    def _extract_pdf_marked_labels(self, text: str) -> list[str]:
        return self._extract_checked_options(text)

    def _extract_matter_codes_from_pdf_text(self, text: str) -> list[str]:
        evidence = selection_evidence(text, multiple=True)
        return self._map_checked_labels_to_matter_codes(evidence['selected_values'])

    def _split_pdf_party_blocks(self, text: str) -> tuple[str, str]:
        source = str(text or "")
        if not source.strip():
            return "", ""
        match = re.search(r"(?:\n|^)\s*生产\s*企业\s*[:：]", source)
        marker_index = match.start() if match else -1
        if marker_index < 0:
            return source, ""
        return source[:marker_index].strip(), source[marker_index:].strip()

    def _compact_text(self, text: str) -> str:
        value = self._clean_text(text)
        value = re.sub(r"\s+", "", value)
        value = re.sub(r"[\-—–_:/：;；,.，。()\[\]（）【】]+", "", value)
        return value

    def _contains_compact(self, text: str, keyword: str) -> bool:
        return self._compact_text(keyword) in self._compact_text(text)

    def _extract_company_name(self, text: str) -> str:
        source = str(text or "")
        compact_source = self._compact_text(source)
        patterns = [
            r"(?:中文名称|企业名称|申请人名称|生产企业名称|引文各称|引文名称|名称)[:：]?\s*([^\n\r:：]{2,80}?有\s*限\s*公\s*司)",
        ]
        for pattern in patterns:
            match = re.search(pattern, source)
            if match:
                value = self._compact_text(match.group(1))
                if value:
                    return value
        compact_labeled = re.search(r"(?:中文名称|企业名称|申请人名称|生产企业名称|引文各称|引文名称|名称)(浙江[\u4e00-\u9fa5A-Za-z0-9]{2,40}?有限公司)", compact_source)
        if compact_labeled:
            return self._clean_text(compact_labeled.group(1))
        compact_match = re.search(r"(浙江[\u4e00-\u9fa5A-Za-z0-9]{2,40}?有限公司)", compact_source)
        if compact_match:
            return self._clean_text(compact_match.group(1))
        return ""

    def _extract_unified_social_credit_code(self, text: str) -> str:
        source = str(text or "")
        patterns = [
            r"(?:统一社会信用代码/组织机构代码|统一社会信用代码|社会信用代码)[:：;\s]*([0-9A-Z]{15,25})",
            r"\b([0-9A-Z]{18})\b",
        ]
        for pattern in patterns:
            match = re.search(pattern, source)
            if match:
                return self._clean_text(match.group(1))
        return ""

    def _extract_party_address(self, text: str) -> str:
        source = str(text or "")
        compact_source = self._compact_text(source)
        patterns = [
            r"(?:注册地址\(住所\)|注册地址|生产地址|通讯地址)[:：]?\s*([^\n\r]{6,150})",
        ]
        for pattern in patterns:
            match = re.search(pattern, source)
            if match:
                value = self._clean_text(match.group(1))
                value = re.sub(r"邮编[:：]?[0-9]{0,10}.*$", "", value).strip(" ;；，,")
                if value:
                    return value
        compact_match = re.search(r"(?:注册地址\(住所\)|注册地址|生产地址|通讯地址)[:：]?([\u4e00-\u9fa5A-Za-z0-9\-]{6,120})", compact_source)
        if compact_match:
            return self._clean_text(compact_match.group(1))
        return ""

    def _extract_party_contact_parts(self, text: str) -> list[str]:
        source = str(text or "")
        source = source.replace("@ ", "@").replace(". ", ".").replace(" .", ".")
        out: list[str] = []
        person_match = re.search(r"(?:注册申请负责人|联系人)[:：]?\s*([^\n\r]{1,40})", source)
        if person_match:
            person = self._clean_text(person_match.group(1)).strip(" ;；，,:")
            if person and not self._looks_contact_noise(person):
                out.append(person)
        email_matches = re.findall(r"(?:电子信箱|邮箱|E-?mail)[:：]?\s*([A-Za-z0-9_.+\-]+@[A-Za-z0-9_.\-]+)", source)
        phone_matches = re.findall(r"(?:电话|手机)[:：]?\s*([0-9\-]{7,20})", source)
        for value in email_matches + phone_matches:
            cleaned = self._clean_text(value).strip(" ;；，,")
            if cleaned and cleaned not in out:
                out.append(cleaned)
        return out

    def _looks_contact_noise(self, text: str) -> bool:
        compact = self._compact_text(text)
        if not compact:
            return True
        noise_keywords = ["职位", "电话", "传真", "电子信箱", "手机", "生产许可证", "gmp", "赛默职位"]
        return any(keyword.lower() in compact.lower() for keyword in noise_keywords)

    def _apply_pdf_global_fallbacks(self, schema: Dict[str, Any], raw_text: str) -> None:
        # 全文只能按明确的同一字段标题回填，不能凭某个词在全文中出现。
        for line in str(raw_text or '').splitlines():
            found = heading(line)
            if not found or found[0] not in (1, 2, 3, 4, 10, 19):
                continue
            field = next((v for k, v in schema.items() if k.startswith(f'item_{found[0]}_')), None)
            if not field or field.get('value') or field.get('source_text'):
                continue
            value = selected_option(found[2], field.get('options', []))
            if value:
                field.update(value=value, source_text=line, parse_confidence=0.55)

    def _extract_global_manufacturer_block(self, raw_text: str) -> str:
        lines = str(raw_text or '').splitlines()
        for i, line in enumerate(lines):
            found = heading(line)
            if found and found[0] == 31:
                block = [found[2]]
                for following in lines[i + 1:]:
                    if heading(following):
                        break
                    block.append(following)
                return '\n'.join(block).strip()
        return ""

    def _apply_pdf_region_ocr_enhancements(self, schema: Dict[str, Any], items: list[Dict[str, Any]], file_path: str) -> None:
        # T2 已完成原生/扫描逐页提取。仅对已有定位且低质量字段作区域补识别。
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import parse_subfields
        from agent.agent_backend.services.filing_form_revision import merge_parse
        enhanced = []
        for item in items:
            number = item.get("item_no")
            if item.get('recognition_evidence'):
                continue
            key = next((k for k in schema if k.startswith(f"item_{number}_")), None)
            if not key or not item.get("raw_text") or not item.get("source_regions"):
                continue
            field = schema[key]
            content_values = [v for p, v in values(field).items() if p != "sub_fields.validity_unit"]
            # 选项已识别但未选/冲突不是OCR失败；不能通过区域重识别覆盖这些状态。
            if field.get('selection_evidence'):
                continue
            if any(present(v) for v in content_values) and not (number in (14, 16, 30, 31, 32) and any(not present(v) for v in field.get('sub_fields', {}).values())):
                continue
            text = self._ocr_item_regions(file_path, item["source_regions"], psm=7 if number in (2, 3, 4, 19) else 6, zoom=4, margin_pt=0)
            if not text:
                continue
            found = heading(text)
            if found:
                if found[0] != number:
                    continue
                text = found[2]
            enhanced.append({**item, "raw_text": text, "subfields": parse_subfields(text)})
        if enhanced:
            fresh = self._build_full_form("")
            self._apply_pdf_items_to_schema(fresh, enhanced)
            for field in fresh.values():
                if isinstance(field, dict) and field.get('parse_confidence'):
                    field['parse_confidence'] = 0.55
                    for source in field.get('value_sources', {}).values():
                        source['parse_confidence'] = min(float(source.get('parse_confidence') or .55), .55)
            schema.update(merge_parse(schema, fresh))

    def _ocr_item_regions(
        self,
        file_path: str,
        regions: list[Dict[str, Any]],
        psm: int = 6,
        max_regions: int = 2,
        zoom: float = 2.5,
        margin_pt: float = 6.0,
    ) -> str:
        if not isinstance(regions, list):
            return ""
        texts: list[str] = []
        for region in regions[: max(1, int(max_regions or 1))]:
            if not isinstance(region, dict):
                continue
            text = ocr_pdf_region(file_path, region, psm=psm, zoom=zoom, margin_pt=margin_pt)
            text = self._clean_text(text)
            if text:
                texts.append(text)
        return "\n".join(texts).strip()

    def _looks_like_chinese_drug_name(self, text: str) -> bool:
        value = str(text or "").strip()
        if len(value) < 2:
            return False
        cjk_count = sum(1 for ch in value if "\u4e00" <= ch <= "\u9fff")
        return cjk_count >= 2

    def _extract_value_after_label(self, text: str, label: str) -> str:
        source = str(text or "")
        match = re.search(rf"{re.escape(label)}[:：]?\s*([^\n\r]{{1,80}})", source)
        if match:
            return self._clean_text(match.group(1))
        return self._clean_text(source)

    def _extract_company_name_from_manufacturer_anchor(self, text: str) -> str:
        source = str(text or "")
        compact = self._compact_text(source)
        match = re.search(r"生产企业[:：]?(浙江[\u4e00-\u9fa5A-Za-z0-9]{2,40}?有限公司)", compact)
        if match:
            return self._clean_text(match.group(1))
        return self._extract_company_name(source)

    def _extract_party_address_from_manufacturer_anchor(self, text: str) -> str:
        source = str(text or "")
        compact = self._compact_text(source)
        match = re.search(r"生产地址[:：]?([\u4e00-\u9fa5A-Za-z0-9\-]{6,120})", compact)
        if match:
            return self._clean_text(match.group(1))
        return self._extract_party_address(source)

    def _cleanup_address_text(self, text: str) -> str:
        value = self._clean_text(text)
        value = re.sub(r"(邮编|通讯地址|注册申请负责人|联系人|电话|传真|电子信箱|手机).*$", "", value)
        value = re.sub(r"\(自$", "", value)
        value = value.lstrip(":：;,， ")
        return value.strip(" ;；，,")

    def _extract_registered_address(self, text: str) -> str:
        source = str(text or "")
        match = re.search(r"(?:注册地址\(住所\)|注册地址)[:：;,]?\s*([^\n\r]{6,200})", source)
        if match:
            return self._cleanup_address_text(match.group(1))
        return ""

    def _ocr_packaging_region_from_item15(self, file_path: str, item15_regions: list[Dict[str, Any]]) -> str:
        if not isinstance(item15_regions, list) or not item15_regions:
            return ""
        region = item15_regions[0]
        bbox = region.get("bbox_pdf", []) if isinstance(region, dict) else []
        page = int(region.get("page", 0) or 0) if isinstance(region, dict) else 0
        if page <= 0 or len(bbox) != 4:
            return ""
        x0, y0, x1, _ = [float(item or 0) for item in bbox]
        packaging_region = {
            "page": page,
            "bbox_pdf": [x0 - 8, y0 - 68, x1 + 210, y0 - 2],
        }
        return self._clean_text(ocr_pdf_region(file_path, packaging_region, psm=6, zoom=4.0, margin_pt=0))

    def _looks_meaningful_chinese(self, text: str) -> bool:
        value = self._clean_text(text)
        cjk_count = sum(1 for ch in value if "\u4e00" <= ch <= "\u9fff")
        ascii_count = sum(1 for ch in value if ch.isascii() and ch.isalpha())
        return cjk_count >= 2 and cjk_count >= ascii_count

    def _apply_cross_field_name_normalization(self, schema: Dict[str, Any]) -> None:
        # 英文和拼音只保留为核对来源，不推断或改写明确的中文字段。
        return

    def _sanitize_pdf_field_values(self, schema: Dict[str, Any]) -> None:
        if not isinstance(schema, dict):
            return

        item_13 = schema.get("item_13_other_specifications") or {}
        item_13_value = str(item_13.get("value", "") or "").strip()
        item_13_source = str(item_13.get("source_text", "") or "").strip()
        if self._looks_like_cross_item_pollution(item_13_value) or self._looks_like_cross_item_pollution(item_13_source):
            item_13["value"] = ""
            item_13["parse_confidence"] = 0.0

        item_19 = schema.get("item_19_pre_acceptance_inspection") or {}
        item_19_value = str(item_19.get("value", "") or "").strip()
        item_19_no = str((item_19.get("sub_fields", {}) or {}).get("sample_inspection_no", "") or "").strip()
        item_19_source = str(item_19.get("source_text", "") or "").strip()
        if not item_19_value and not item_19_no and not self._looks_like_meaningful_inspection_text(item_19_source):
            item_19["parse_confidence"] = 0.0

        item_24 = schema.get("item_24_patent_info") or {}
        item_24_value = str(item_24.get("value", "") or "").strip()
        item_24_normalized = self._normalize_patent_statement(item_24_value)
        if item_24_normalized:
            item_24["value"] = item_24_normalized
            item_24["parse_confidence"] = max(float(item_24.get("parse_confidence") or 0), 0.66)

        item_29 = schema.get("item_29_other_related_info") or {}
        item_29_value = str(item_29.get("value", "") or "").strip()
        item_29_source = str(item_29.get("source_text", "") or "").strip()
        if self._looks_like_history_placeholder(item_29_value or item_29_source):
            item_29["value"] = ""
            item_29["parse_confidence"] = 0.0

    def _looks_like_cross_item_pollution(self, text: str) -> bool:
        source = str(text or "").strip()
        if not source:
            return False
        compact = self._compact_text(source)
        return compact.startswith("14包装") or "包装规格" in source or "直接接触药品的包装材料和容器" in source

    def _looks_like_meaningful_inspection_text(self, text: str) -> bool:
        source = str(text or "").strip()
        if not source:
            return False
        if "是" in source or "否" in source or "检品编号" in source:
            return True
        return bool(re.search(r"[A-Za-z0-9]{4,}", source))

    def _extract_meaningful_chinese_line(self, text: str) -> str:
        for line in re.split(r'[\n；;]', str(text or "")):
            cleaned = self._clean_text(line)
            if not cleaned:
                continue
            compact = self._compact_text(cleaned)
            if any(keyword in compact for keyword in ["适应症", "功能主治", "item20", "wiedr", "weagd"]):
                continue
            if self._looks_meaningful_chinese(cleaned):
                return cleaned
        compact_source = self._compact_text(text)
        matches = re.findall(r"[\u4e00-\u9fa5]{2,20}", compact_source)
        meaningful = [item for item in matches if item not in {"主要适应症或者功能主治", "适应症", "功能主治"}]
        return meaningful[-1] if meaningful else ""

    def _normalize_patent_statement(self, text: str) -> str:
        # 只保留实际提取原文，不能凭关键词补出否定词或权属声明。
        return str(text or "").strip()

    def _normalize_loose_boolean_text(self, text: str) -> str:
        return selected_option(text, ["是", "否"])

    def _looks_like_history_placeholder(self, text: str) -> bool:
        source = str(text or "").strip()
        if not source:
            return False
        # 没有特定编号格式不等于空表。Markdown数据行有任何实际内容
        # （包括明确“不适用”或“/”）均保留，不能靠公司名/编号白名单判断。
        after_separator = False
        for line in source.splitlines():
            cells = [cell.strip() for cell in line.strip().strip('|').split('|')]
            if len(cells) > 1 and all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                after_separator = True
                continue
            if after_separator and any(cells):
                return False
        compact = self._compact_text(source)
        headers = ['临床试验暂停/恢复/终止情况/不适用', '临床试验暂停/恢复/终止情况',
                   '受理号', '批件号', '批准内容', '备注']
        if not all(label in compact for label in ('受理号', '批件号', '批准内容')):
            return False
        for header in headers:
            compact = compact.replace(self._compact_text(header), '')
        return not compact.strip('|')

    def _normalize_ascii_wording(self, text: str) -> str:
        value = str(text or "")
        value = value.replace(")", "l").replace("]", "l").replace("|", "l")
        value = re.sub(r"[^A-Za-z0-9 ]+", " ", value)
        value = re.sub(r"\s+", " ", value).strip()
        return value

    def _normalize_english_name(self, text: str) -> str:
        return self._clean_text(text)

    def _normalize_known_drug_name(self, text: str) -> str:
        return self._clean_text(text)

    def _normalize_item_21_content(self, text: str) -> str:
        lines = [self._clean_text(line) for line in str(text or "").splitlines() if self._clean_text(line)]
        values: list[str] = []
        for line in lines:
            content = self._first_match(line, [r"(?:补充申请涉及内容|补充申请的内容)[:：]?\s*(.+)$"])
            if content:
                values.extend([part.strip(" ;；，,") for part in re.split(r"[;；]", content) if part.strip(" ;；，,")])
                continue
            if self._looks_meaningful_chinese(line) and "补充申请" not in line:
                values.append(line)
        normalized_values = [item.replace("有效其", "有效期").strip("“”\"' ") for item in values]
        merged = "；".join(dict.fromkeys([item for item in normalized_values if item]))
        return merged

    def _normalize_packaging_material(self, text: str) -> str:
        value = self._clean_text(text).strip(" ;；，,:：")
        compact = self._compact_text(value)
        if "直接接触药品" in compact or compact in {"zabt", "at", "za"}:
            return ""
        return value

    def _normalize_item_23_content(self, text: str) -> str:
        source = str(text or "")
        fields = [
            ("原申请受理号", self._first_match(source, [r"原申请受理号[;:：,，\s]*([A-Za-z0-9/_-]+)"])),
            ("临床批件编号/临床通知书编号", self._first_match(source, [r"临床批件编号/临床通知书编号[:：]?\s*([^\n\r]+)"])),
            ("原药品批准文号/登记号", self._first_match(source, [r"原药品批准文号/登记号[:：]?\s*([^\n\r]+)"])),
            ("药品标准编号", self._first_match(source, [r"药品标准编号[:：]?\s*([^\n\r]+)"])),
        ]
        lines = []
        for label, value in fields:
            cleaned = self._clean_text(value).strip(" ;；，,")
            if cleaned:
                lines.append(f"{label}: {cleaned}")
        return "\n".join(lines).strip()

    def _extract_docx_structured_records(self, doc, excluded_paragraphs=(), excluded_table_rows=()) -> list[Dict[str, Any]]:
        out = []
        seen_cells = set()
        def walk(table, location):
            previous = None
            for row_index, row in enumerate(table.rows):
                if (location,row_index) in excluded_table_rows:
                    previous = None
                    continue
                cells = []
                repeated = set()
                row_cells = set()
                for cell in row.cells:
                    if cell._tc in row_cells:
                        continue
                    row_cells.add(cell._tc)
                    if cell._tc in seen_cells:
                        repeated.add(len(cells))
                    seen_cells.add(cell._tc)
                    cells.append(cell)
                texts = [self._cell_text_with_checkbox(c) for c in cells]
                # 纵向合并的标题仍界定各列归属，合并的值只读取一次。
                texts = [text if i not in repeated or heading(text) or self._is_docx_subfield(0, text) else ''
                         for i, text in enumerate(texts)]
                found = []
                for index, text in enumerate(texts):
                    candidate = heading(text)
                    if candidate and not (found and self._is_docx_subfield(found[-1][1][0], text)):
                        found.append((index, candidate))
                # 同行字段只消费自己的单元格范围；普通值与尾部空格不充当子标题。
                if found:
                    groups = [(index, found[pos+1][0] if pos+1 < len(found) else len(texts), candidate)
                              for pos, (index, candidate) in enumerate(found)]
                    previous = found[0][1][:2] if len(found) == 1 else None
                elif previous and texts and (len(cells) < len(row.cells) or not texts[0].strip()) and not any(re.fullmatch(r'\d+', t.strip()) for t in texts if t):
                    number, label = previous
                    groups = [(-1, len(texts), (number, label, ''))]
                else:
                    previous = None
                    for index, cell in enumerate(cells):
                        for ti, nested in enumerate(cell.tables):
                            walk(nested, f'{location}/r{row_index}/c{index}/t{ti}')
                    continue
                expanded = []
                for start, end, (number, label, inline) in groups:
                    subs = [i for i in range(start+1, end) if self._is_docx_subfield(number, texts[i])]
                    if subs:
                        for j, pos in enumerate(subs):
                            expanded.append((start, subs[j+1] if j+1<len(subs) else end, number, label,
                                             inline if j==0 else '', texts[pos], pos+1))
                    else:
                        expanded.append((start,end,number,label,inline,'',start+1))
                for start, end, number, label, inline, sub, value_start in expanded:
                    remaining = texts[value_start:end]
                    content = '\n'.join([inline] + remaining).strip()
                    base = {"seq": str(number), "item": label, "source_text": '\n'.join(texts[max(0, start):end]),
                            "option_fragments": [inline] + remaining,
                            "continuation": start in repeated or start == -1,
                            "source_regions": [{"table": location, "row": row_index, "column_start": max(0, start),
                                                "column_end": end, "coordinate_unit": "word_table"}]}
                    pairs = self._extract_key_value_pairs(content) if number in (14, 15, 16, 20, 23, 30, 31, 32) else []
                    if number in (1, 2, 3, 4, 5, 10, 11, 19) or (number == 16 and self._subfield_label(sub) == '是否有变更'):
                        out.append({**base, "sub_field": sub, "value": content})
                    elif pairs:
                        for key, value in pairs:
                            out.append({**base, "sub_field": key, "value": value})
                    else:
                        out.append({**base, "sub_field": sub or "内容", "value": content})
                    for index in range(max(0, start), end):
                        cell = cells[index]
                        if number in (17, 29, 32):
                            for obj in self._extract_nested_table_rows(cell):
                                out.append({**base, "sub_field": "明细表", "value": obj})
                        else:
                            for ti, nested in enumerate(cell.tables):
                                walk(nested, f'{location}/r{row_index}/c{index}/t{ti}')
        for index, table in enumerate(doc.tables):
            walk(table, str(index))
        # 无表格或表外字段仍按标题分段，复用相同语义映射。
        current = None
        for paragraph_index, paragraph in enumerate(doc.paragraphs):
            if paragraph_index in excluded_paragraphs:
                current = None
                continue
            text = self._paragraph_text_with_checkbox(paragraph).strip()
            found = heading(text)
            region = {'paragraph': paragraph_index, 'coordinate_unit': 'word_paragraph'}
            if found:
                current = {"seq": str(found[0]), "item": found[1], "sub_field": "内容", "value": found[2], "source_text": text,
                           "source_regions": [region]}
                out.append(current)
            elif current and text:
                current["value"] += '\n' + text
                current["source_text"] += '\n' + text
                current['source_regions'].append(region)
        return out

    def _is_docx_subfield(self, number, text):
        labels = {
            14: {'直接接触药品的包装材料和容器', '包装规格'},
            15: {'原有效期', '拟延长后有效期', '变更后有效期', '贮藏条件'},
            16: {'活性成分/中药成分/...', '活性成分', '中药成分', '辅料', '是否有变更'},
            20: {'适应症分类', '补充申请涉及内容'},
            23: {'原申请受理号', '原药品批准文号/登记号', '药品标准编号'},
        }
        return self._subfield_label(text) in ({'内容'} | labels.get(number, set()) | set(self.PARTY_SUBFIELDS.get(number, {})))

    def _paragraph_text_with_checkbox(self, paragraph) -> str:
        checked, unchecked = '☑☒√✓✔■●▣', '☐□○'
        def local(node): return node.tag.rsplit('}',1)[-1]
        def attr(node, name, default=''):
            return next((v for k,v in node.attrib.items() if k.endswith('}'+name)), default)
        def render(node):
            tag = local(node)
            if tag == 'sdt':
                properties = next((n for n in node if local(n)=='sdtPr'),None)
                content = next((n for n in node if local(n)=='sdtContent'),None)
                state_node = next((n for n in properties.iter() if local(n)=='checked'),None) if properties is not None else None
                text = ''.join(render(n) for n in content) if content is not None else ''
                if state_node is None: return text
                state = attr(state_node,'val','1') in ('1','true','on')
                marks = list(re.finditer('['+checked+unchecked+']',text))
                conflict = len(marks)>1 or any((m.group() in checked) != state for m in marks)
                # 用状态校验显示，在字形原位置输出一次，保留前置/后置及控件边界。
                if not marks:
                    text = ('☑' if state else '☐') + text
                elif len(marks)==1:
                    text = text[:marks[0].start()] + ('☑' if state else '☐') + text[marks[0].end():]
                if conflict: text = '【控件状态与显示矛盾】'+text
                has_label = bool(re.sub('['+checked+unchecked+r'\s]', '',text))
                return ';'+text+';' if has_label else text
            if tag == 't': return str(node.text or '')
            if tag == 'tab': return ' '
            if tag in ('br','cr'): return '\n'
            if tag == 'sym':
                code = attr(node,'char').upper()
                if code in ('0052','F052'): return '☑'
                if code in ('00A3','F0A3'): return '☐'
                try:
                    glyph = chr(int(code,16))
                    return glyph if glyph in checked+unchecked else ''
                except ValueError: return ''
            return ''.join(render(n) for n in node)
        return render(paragraph._p)

    def _cell_text_with_checkbox(self, cell) -> str:
        lines: list[str] = []
        for paragraph in getattr(cell, "paragraphs", []):
            text = self._clean_text(self._paragraph_text_with_checkbox(paragraph))
            if text:
                lines.append(text)
        return "\n".join(lines)

    def _clean_text(self, text: str) -> str:
        value = str(text or "").replace("\xa0", " ").replace("\u3000", " ")
        value = re.sub(r"[ \t]+", " ", value)
        return value.strip()

    def _extract_checked_options(self, text: str) -> list[str]:
        return selection_evidence(text, multiple=True, require_mark=True)['selected_values']

    def _extract_key_value_pairs(self, text: str) -> list[tuple[str, str]]:
        out: list[tuple[str, str]] = []
        for line in re.split(r'[\n；;]', str(text or "")):
            line = self._clean_text(line)
            if not line:
                continue
            pairs = re.findall(r"([^：:\s]{1,40})[：:][ \t]*([^：:]+?)(?=\s+[^：:\s]{1,40}[：:]|$)", line)
            for key, val in pairs:
                k = self._clean_text(key)
                v = self._clean_text(val).strip("_ ：:")
                if k and v:
                    out.append((k, v))
        return out

    def _extract_nested_table_rows(self, cell) -> list[Dict[str, str]]:
        out: list[Dict[str, str]] = []
        for table in getattr(cell, "tables", []):
            rows: list[list[str]] = []
            for row in table.rows:
                values = [self._clean_text(self._cell_text_with_checkbox(c)) for c in row.cells]
                if any(values):
                    rows.append(values)
            if len(rows) < 2:
                continue
            headers = rows[0]
            for values in rows[1:]:
                if any(values):
                    out.append(dict(zip(headers, values)))
        return out

    def _apply_structured_records_to_schema(self, schema: Dict[str, Any], records: list[Dict[str, Any]]) -> None:
        if not isinstance(records, list) or not isinstance(schema, dict):
            return
        for rec in records:
            seq = str((rec or {}).get("seq", ""))
            value = (rec or {}).get("value", "")
            key = "special_statement" if seq == "33" else next((k for k in schema if k.startswith(f"item_{seq}_")), None)
            if key and rec.get('continuation') and seq in ('6', '7', '8', '9', '12', '13') and value:
                previous_value = schema[key].get('value', '')
                value = '\n'.join(v for v in (previous_value, str(value)) if v)
            if key:
                field = schema[key]
                source = str(rec.get("source_text", value) or "")
                if source and source not in field.get("source_text", ""):
                    field["source_text"] = (field.get("source_text", "") + "\n" + source).strip()
                field.setdefault("source_regions", []).extend(rec.get("source_regions", []))
            if seq in ("1", "2", "3", "4", "10", "19"):
                schema[key].setdefault('option_fragments', []).extend(rec.get('option_fragments', [value]))
                schema[key]["value"] = selected_option(schema[key]['option_fragments'], schema[key]["options"])
            if seq == "33":
                schema[key]["value"] = str(value)
            elif seq == "1":
                if "境内生产药品补充申请" in str(value):
                    schema["item_1_application_type"]["value"] = "境内生产药品补充申请"
                if str(value).startswith("是"):
                    schema["item_1_application_type"]["sub_fields"]["shared_application_form_no"] = str(value).replace("是", "").strip("，,：:_ ")
            elif seq == "5":
                schema[key].setdefault('option_fragments', []).extend(rec.get('option_fragments', [value]))
                self._apply_multiple_choice(schema[key], schema[key]['option_fragments'], matter=True)
            elif seq == "6":
                if str(value):
                    schema["item_6_generic_name"]["value"] = str(value)
            elif seq == "7":
                if str(value):
                    schema["item_7_english_or_latin_name"]["value"] = str(value)
            elif seq == "8":
                if str(value):
                    schema["item_8_pinyin"]["value"] = str(value)
            elif seq == "9":
                if str(value):
                    schema["item_9_chemical_name"]["value"] = str(value)
            elif seq == "11":
                schema[key].setdefault('option_fragments', []).extend(rec.get('option_fragments', [value]))
                self._apply_multiple_choice(schema[key], schema[key]['option_fragments'])
            elif seq == "12":
                if str(value):
                    schema["item_12_specification"]["value"] = str(value)
            elif seq == "13":
                if str(value):
                    schema["item_13_other_specifications"]["value"] = str(value)
            elif seq == "14":
                sub = str((rec or {}).get("sub_field", ""))
                if "包装规格" in sub:
                    schema["item_14_packaging"]["sub_fields"]["packaging_specification"] = str(value)
                elif "包装材料" in sub or "容器" in sub:
                    schema["item_14_packaging"]["sub_fields"]["primary_packaging_material"] = str(value)
            elif seq == "15":
                txt = str(value)
                if re.fullmatch(r'\d+\s*个?\s*月', txt) and str(rec.get('sub_field', '')) in ('', '内容'):
                    schema['item_15_validity_period']['reported_value'] = txt
                if "原有效期" in str((rec or {}).get("sub_field", "")):
                    schema["item_15_validity_period"]["sub_fields"]["original_validity_period"] = txt
                if "拟延长后有效期" in str((rec or {}).get("sub_field", "")) or "变更后有效期" in str((rec or {}).get("sub_field", "")):
                    schema["item_15_validity_period"]["sub_fields"]["proposed_validity_period"] = txt
                if "贮藏条件" in str((rec or {}).get("sub_field", "")):
                    schema["item_15_validity_period"]["sub_fields"]["storage_condition"] = txt
            elif seq == "16":
                sub = str((rec or {}).get("sub_field", ""))
                if "活性成分" in sub or "中药成分" in sub:
                    schema["item_16_prescription"]["sub_fields"]["active_ingredients"] = str(value)
                elif "辅料" in sub:
                    schema["item_16_prescription"]["sub_fields"]["excipients"] = str(value)
                elif "变更" in sub:
                    field = schema["item_16_prescription"]
                    field.setdefault('option_fragments', []).extend(rec.get('option_fragments', [value]))
                    field["sub_fields"]["has_change"] = selected_option(field['option_fragments'], ["是", "否"])
            elif seq == "17" and isinstance(value, dict):
                row = {
                    "material_name": str(value.get("原/辅料/包材名称", value.get("material_name", ""))),
                    "register_no": str(value.get("登记号", value.get("register_no", ""))),
                    "accept_no": str(value.get("受理号", value.get("accept_no", ""))),
                    "manufacturer": str(value.get("生产企业名称", value.get("供应商", value.get("manufacturer", "")))),
                }
                guessed = self._guess_material_source_rows([[list(value), list(value.values())]]) or self._guess_material_source_rows([[
                    ["原/辅料/包材名称", "登记号", "受理号", "供应商"],
                    [row["material_name"], row["register_no"], row["accept_no"], row["manufacturer"]],
                ]])
                if guessed:
                    schema["item_17_material_source"].setdefault("table_rows", []).extend(guessed)
            elif seq == "18":
                if str(value):
                    schema["item_18_tcm_material_standard"]["value"] = str(value)
            elif seq == "20":
                sub = str((rec or {}).get("sub_field", ""))
                if "分类" in sub:
                    schema["item_20_indication_or_function"]["sub_fields"]["category"] = str(value)
                else:
                    schema["item_20_indication_or_function"]["sub_fields"]["description"] = str(value)
            elif seq == "21":
                if str(value):
                    current = str(schema["item_21_change_content"]["value"] or "")
                    schema["item_21_change_content"]["value"] = f"{current}；{value}".strip("；")
            elif seq == "22":
                if str(value):
                    schema["item_22_change_reason"]["value"] = str(value)
            elif seq == "23":
                if str(value):
                    current = str(schema["item_23_original_approval_info"]["value"] or "")
                    sub = str((rec or {}).get("sub_field", ""))
                    line = f"{sub}：{value}" if sub else str(value)
                    schema["item_23_original_approval_info"]["value"] = f"{current}\n{line}".strip()
            elif seq == "24":
                if str(value):
                    current = str(schema["item_24_patent_info"]["value"] or "")
                    sub = str((rec or {}).get("sub_field", ""))
                    line = f"{sub}：{value}" if sub and sub != '内容' else str(value)
                    schema["item_24_patent_info"]["value"] = f"{current}\n{line}".strip()
            elif seq == "25":
                if str(value):
                    schema["item_25_data_protection"]["value"] = str(value)
            elif seq == "26":
                if str(value):
                    schema["item_26_tcm_protection"]["value"] = str(value)
            elif seq == "27":
                if str(value):
                    schema["item_27_same_product_period"]["value"] = str(value)
            elif seq == "28":
                if str(value):
                    current = str(schema["item_28_change_related_items"]["value"] or "")
                    schema["item_28_change_related_items"]["value"] = f"{current}\n{value}".strip()
            elif seq == "29":
                if value:
                    current = str(schema["item_29_other_related_info"]["value"] or "")
                    schema["item_29_other_related_info"]["value"] = f"{current}\n{value}".strip()
            elif seq in ('30','31'):
                target = self.PARTY_SUBFIELDS[int(seq)].get(self._subfield_label(rec.get('sub_field','')))
                if target:
                    self._assign_subfield(schema[key], target, str(value), rec)
            elif seq == "32":
                entries = value.items() if isinstance(value, dict) else [(rec.get('sub_field',''),value)]
                for sub, content in entries:
                    target = self.PARTY_SUBFIELDS[32].get(self._subfield_label(sub))
                    if target:
                        self._assign_subfield(schema[key], target, str(content), rec)

        for key in [
            "item_1_application_type",
            "item_2_registration_category",
            "item_3_otc_type",
            "item_4_original_product_status",
            "item_5_application_matter_category",
            "item_6_generic_name",
            "item_7_english_or_latin_name",
            "item_8_pinyin",
            "item_9_chemical_name",
            "item_10_trade_name",
            "item_11_dosage_form",
            "item_12_specification",
            "item_13_other_specifications",
            "item_14_packaging",
            "item_15_validity_period",
            "item_16_prescription",
            "item_17_material_source",
            "item_18_tcm_material_standard",
            "item_19_pre_acceptance_inspection",
            "item_20_indication_or_function",
            "item_21_change_content",
            "item_22_change_reason",
            "item_23_original_approval_info",
            "item_24_patent_info",
            "item_25_data_protection",
            "item_26_tcm_protection",
            "item_27_same_product_period",
            "item_28_change_related_items",
            "item_29_other_related_info",
            "item_30_applicant_info",
            "item_31_manufacturer_info",
            "item_32_cro_info",
        ]:
            if key in schema:
                field = schema[key]
                field["parse_confidence"] = 0.88 if any(present(v) for p, v in values(field).items() if p != "sub_fields.validity_unit") else 0.0
        self._finalize_fields(schema)

    def _extract_checked_labels_from_docx_xml(self, file_path: str) -> list[str]:
        try:
            with zipfile.ZipFile(file_path, "r") as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
        except (OSError, KeyError, zipfile.BadZipFile):
            return []
        return self._extract_checked_labels_from_wsym(xml)

    def _extract_checked_labels_from_wsym(self, xml: str) -> list[str]:
        # 所有XML入口使用同一段落解码及符号归属规则，不能按整段“含已选”取全部名称。
        from docx.oxml import parse_xml
        from types import SimpleNamespace
        root = parse_xml(xml.encode('utf-8') if isinstance(xml, str) else xml)
        labels = []
        for node in root.iter():
            if str(node.tag).split('}')[-1] == 'p':
                text = self._paragraph_text_with_checkbox(SimpleNamespace(_p=node))
                labels.extend(self._extract_checked_options(text))
        return list(dict.fromkeys(labels))

    def _map_checked_labels_to_matter_codes(self, labels: list[str]) -> list[str]:
        mapping = {
            "变更有效期和贮藏条件": "1.7",
            "变更处方中的辅料": "1.1",
            "变更生产工艺": "1.2",
            "变更所用原料药的供应商": "1.3",
            "变更生产批量": "1.4",
            "变更注册标准": "1.5",
            "变更包装材料和容器": "1.6",
            "增加规格": "1.8",
            "变更生产场地": "1.9",
            "变更生产场地，境外生产药品": "1.9",
            "其他": "1.10",
        }
        out = []
        for label in labels:
            # 调用者只传已确认的名称，仍拒绝带有未选标记或否定前缀的条目。
            evidence = selection_evidence(re.sub(r'\s*\n\s*', '', label), multiple=True)
            if evidence['status'] != 'confirmed':
                continue
            name = evidence['selected_values'][0] if len(evidence['selected_values']) == 1 else ''
            for key, code in mapping.items():
                if self._compact_text(key).replace(',', '，') == self._compact_text(name).replace(',', '，') and code not in out:
                    out.append(code)
        return out

    def _apply_multiple_choice(self, field, fragments, *, matter=False):
        fragments = list(fragments) if isinstance(fragments, (list, tuple)) else [fragments]
        cleaned = []
        for fragment in fragments:
            text = str(fragment or '')
            found = heading(text)
            if found and found[0] in (5, 11):
                text = found[2]
            text = re.sub(r'^中\s*国\s*药\s*典\s*剂\s*型\s*[:：]?\s*', '', text)
            if text.strip():
                if not matter:
                    # 两个完整剂型名称的空格是候选边界，单个名称内的字距另行清理。
                    cleaned.extend(re.split(r'(?<=剂)[ \t]+(?=\S+剂(?:\s|$))', text))
                else:
                    cleaned.append(text)
        explicit_dosage = bool(re.search(r'中国\s*药典\s*剂型\s*[:：]', field.get('recognized_text', '')))
        evidence = selection_evidence(cleaned, multiple=True,
            require_mark=bool(not matter and field.get('recognition_evidence') and not explicit_dosage))
        field['selection_status'] = evidence['status']
        field['selection_evidence'] = evidence['entries']
        selected, issues = [], []
        for entry in evidence['entries']:
            entry['source_regions'] = deepcopy(field.get('source_regions', []))
            if matter:
                codes = self._map_checked_labels_to_matter_codes([entry['name']]) if entry['selected'] is True else []
                entry['mapped_codes'] = codes
                if entry['selected'] is True and not codes:
                    issues.append({'reason': '原表事项已选中，但没有明确的内部类别映射，待核对', **entry})
                selected.extend(codes)
            elif entry['selected'] is True:
                name = re.sub(r'\s+', '', entry['name'])
                if field.get('recognition_evidence') and not re.fullmatch(r'[\u4e00-\u9fff]{1,24}(?:剂|片|丸|膏|散|胶囊|颗粒|饮片)', name):
                    entry.update(selected=None, status='unconfirmed')
                    issues.append({'reason': '剂型识别文字不能确认为剂型名称，保留候选及原文', **entry})
                else:
                    selected.append(re.sub(r'(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])', '', entry['name']))
        if evidence['status'] != 'confirmed' and cleaned:
            issues.append({'reason': '选项归属冲突，待核对' if evidence['status'] in ('ambiguous', 'conflict') else '选项未确认，待核对',
                           'source_text': field.get('source_text', ''), 'source_regions': field.get('source_regions', [])})
        field['selected_values'] = list(dict.fromkeys(selected))
        if not matter and issues and not selected and evidence['status'] == 'confirmed':
            field['selection_status'] = 'unconfirmed'
        field['parse_confidence'] = .8 if selected else 0.0
        if issues:
            field['semantic_issues'] = issues
        else:
            field.pop('semantic_issues', None)
        if matter:
            field['original_matter'] = [entry for entry in evidence['entries'] if entry['status'] != 'category']
            field['mapping_note'] = '内部业务类别仅按确认选中的事项名称映射；原表编号、名称、状态及来源独立保留。'

    def _fill_core_fields_from_text(self, schema: Dict[str, Any], raw_text: str) -> None:
        # 禁止通过跨换行的空白匹配借用下一个字段。
        for line in str(raw_text or '').splitlines():
            found = heading(line)
            if not found or found[0] not in (6, 21, 22):
                continue
            key = next(k for k in schema if k.startswith(f'item_{found[0]}_'))
            field = schema[key]
            if not field.get('value') and not field.get('source_text') and found[2] and not heading(found[2]):
                field.update(value=found[2], source_text=line, parse_confidence=0.55)
        from agent.agent_backend.utils.parser.drug_supplement_pdf_parser import parse_subfields
        for key, label, target in [('item_10_trade_name','商品名','trade_name'),
                                    ('item_2_registration_category','分类号','class_no'),
                                    ('item_19_pre_acceptance_inspection','检品编号','sample_inspection_no')]:
            block = schema[key]
            # 保留片段边界再排除独立选项；不可先压平换行，让选项变成附属值。
            pieces = re.split(r'[\n;；|]', block.get('source_text',''))
            option_pattern = r'[☑☒☐□√✓✔■●▣○ ]*(?:' + '|'.join(re.escape(o) for o in block.get('options',[]) if isinstance(o,str)) + r')[☑☒☐□√✓✔■●▣○ ]*'
            content = '\n'.join(p for p in pieces if not re.fullmatch(option_pattern,p.strip()))
            found = parse_subfields(content).get(label)
            if isinstance(found,str):
                self._assign_subfield(block,target,found.strip(),block)
        field = schema["item_15_validity_period"]
        evidence, issues = [], []
        # 字段标题只为本字段提供默认话题，不能扩散到其他材料或考察时长。
        for key in ("item_15_validity_period", "item_21_change_content", "item_23_original_approval_info"):
            block = schema[key]
            source = block.get('source_text') or block.get('value', '')
            detail = validity_evidence(source, context='药品有效期' if key == 'item_15_validity_period' else '')
            origin = {'source_text':source, 'source_field':key, 'source_regions':block.get('source_regions',[])}
            issues.extend({**issue, **origin} for issue in detail['issues'])
            if detail['reported_value'] and key == 'item_15_validity_period':
                field['reported_value'] = detail['reported_value']
            evidence.extend({'path':path, 'value':value, 'normalized_months':normalized_validity(value), **origin,
                             **{k:candidate[k] for k in ('historical','stage','stage_marker','object','segment_text','relation_text','relation_values') if k in candidate}}
                            for candidate in detail['candidates'] for path,value in candidate.items() if path in ('original_validity_period','proposed_validity_period','reported_value','unchanged_value'))
            evidence.extend({**origin, **candidate, 'path':'unchanged_declaration'}
                            for candidate in detail['candidates'] if candidate.get('unchanged_declaration'))
        for path in ('original_validity_period','proposed_validity_period'):
            candidates = [e for e in evidence if e['path']==path and e.get('stage','current')=='current']
            choices = {normalized_validity(e['value']) for e in candidates}
            if len(choices)>1:
                issues.append({'reason':'申请表中药品有效期关系冲突，待核对','candidates':candidates})
            field['sub_fields'][path] = candidates[0]['value'] if len(choices)==1 and not any('冲突' in issue['reason'] for issue in issues) else ''
            if field['sub_fields'][path]:
                field.setdefault('value_sources',{})['sub_fields.'+path] = {**candidates[0], 'parse_confidence':.75}
                field['parse_confidence'] = max(field['parse_confidence'],.75)
        field['validity_evidence'] = evidence
        unchanged = [e for e in evidence if e['path']=='unchanged_value' and e.get('stage','current')=='current']
        relations = [e for e in evidence if e['path'] in ('original_validity_period','proposed_validity_period') and e.get('stage','current')=='current']
        if unchanged and any(normalized_validity(e['value']) != normalized_validity(u['value']) for e in relations for u in unchanged):
            issues.append({'reason':'药品有效期保持不变与变更关系冲突，待核对', 'candidates':unchanged+relations})
        declarations = [e for e in evidence if e['path']=='unchanged_declaration' and e.get('stage','current')=='current'
                        and e['source_field'] != 'item_23_original_approval_info']
        if declarations and len({e['normalized_months'] for e in relations}) > 1:
            issues.append({'reason':'药品有效期保持不变与变更关系冲突，待核对',
                           'candidates':declarations+relations+[e for e in evidence if e['path']=='reported_value']})
        if issues:
            if any('冲突' in issue['reason'] for issue in issues):
                for path in ('original_validity_period','proposed_validity_period'):
                    field['sub_fields'][path] = ''
                    field.get('value_sources',{}).pop('sub_fields.'+path,None)
            field['semantic_issues'] = issues

    def _docx_material_details(self, doc):
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        active = None
        details, regions = [], []
        layout = {}

        def walk(table, location):
            nonlocal active
            grid = tuple(int(col.w or 0) for col in table._tbl.tblGrid.gridCol_lst)
            if layout.get('columns') and layout.get('grid') != grid:
                layout.clear()
            pending = []
            def flush():
                if active in (None, 17):
                    found = self._guess_material_source_rows([pending], layout)
                    if layout.get('columns'):
                        layout['grid'] = grid
                    if found:
                        details.extend(found)
                        regions.append({'table':location, 'coordinate_unit':'word_table',
                                        'rows':deepcopy(pending), 'source_text':'\n'.join(' | '.join(r) for r in pending)})
                pending.clear()
            for ri, row in enumerate(table.rows):
                cells = list(dict.fromkeys(c._tc for c in row.cells))
                texts = [self._cell_text_with_checkbox(next(c for c in row.cells if c._tc is tc)) for tc in cells]
                header = any(any(a in str(t) for a in MATERIAL_SOURCE_ALIASES['material_name']) for t in texts)
                found = next((heading(t) for t in texts if heading(t)), None) if not header else None
                if found:
                    flush()
                    if found[0] != active:
                        layout.clear()
                    active = found[0]
                pending.append(texts)
                for ci, tc in enumerate(cells):
                    cell = next(c for c in row.cells if c._tc is tc)
                    for ni, nested in enumerate(cell.tables):
                        flush()
                        walk(nested, f'{location}/r{ri}/c{ci}/t{ni}')
            flush()

        for index, element in enumerate(doc.element.body):
            tag = element.tag.rsplit('}',1)[-1]
            if tag == 'p':
                found = heading(self._paragraph_text_with_checkbox(Paragraph(element,doc)))
                if found:
                    if found[0] != active:
                        layout.clear()
                    active = found[0]
                elif self._paragraph_text_with_checkbox(Paragraph(element,doc)).strip():
                    # 无表头的续表只跨相邻表或明确续表标记，普通正文终止继承。
                    if not re.fullmatch(r'\s*(?:续表|续上表|Continued)\s*', Paragraph(element,doc).text, re.I):
                        layout.clear()
            elif tag == 'tbl':
                walk(Table(element,doc),str(index))
        return details, regions

    def _guess_material_source_rows(self, tables: list[list[list[str]]], layout=None) -> list[Dict[str, str]]:
        aliases = MATERIAL_SOURCE_ALIASES
        out = []
        layout = {} if layout is None else layout
        for table in tables:
            columns = layout.get('columns', {})
            for row in table:
                headers = {}
                for index, cell in enumerate(row):
                    text = re.sub(r"\s+", "", str(cell or ""))
                    for key, labels in aliases.items():
                        if any(label in text for label in labels):
                            headers.setdefault(key, index)
                if "manufacturer" in headers:
                    # 企业名称不能同时作为物料名称列。
                    if headers.get("material_name") == headers["manufacturer"]:
                        headers.pop("material_name", None)
                if "material_name" in headers and ("register_no" in headers or "manufacturer" in headers):
                    columns = headers
                    layout.update(columns=headers, width=len(row))
                    continue
                if columns and (len(row) != layout.get('width') or any(heading(str(c)) for c in row)):
                    columns = {}
                    layout.clear()
                if not columns:
                    continue
                obj = {key: self._clean_text(row[index]) if index < len(row) else "" for key, index in columns.items()}
                name = obj.get("material_name", "")
                if not name or name in ("生产", "名称", "企业名称"):
                    continue
                out.append({key: obj.get(key, "") for key in aliases})
        return out


    def _finalize_fields(self, schema):
        for key, field in schema.items():
            if not isinstance(field, dict) or "field_type" not in field:
                continue
            if key == 'item_25_data_protection' and field.get('source_text'):
                evidence = selection_evidence(field.get('option_fragments', field['source_text']), ['是', '否'],
                                              require_mark=bool(field.get('recognition_evidence')))
                field['selection_status'] = evidence['status']
                field['selection_evidence'] = field.get('visual_choices') or [
                    {**entry, 'source_regions': deepcopy(field.get('source_regions', []))} for entry in evidence['entries']]
                field['value'] = '涉及数据保护：' + evidence['value'] if evidence['value'] else ''
                if not evidence['value']:
                    field.setdefault('semantic_issues', []).append({'reason': '数据保护缺少可靠选择依据，原始识别保留待核对',
                        'source_text': field['source_text'], 'source_regions': field.get('source_regions', [])})
            prescription_choice = key == 'item_16_prescription' and bool(field.get('option_fragments'))
            other_prescription_issues = [issue for issue in field.get('semantic_issues', [])
                                        if prescription_choice and issue.get('path') not in (None, 'sub_fields.has_change')]
            if (field.get('field_type') in ('radio', 'radio_with_input') or prescription_choice) and field.get('source_text'):
                evidence = option_evidence(field.get('option_fragments', field['source_text']), ['是', '否'] if prescription_choice else field.get('options', []))
                field['selection_status'] = evidence['status']
                field['selection_evidence'] = [{**entry, 'source_regions': deepcopy(field.get('source_regions', []))}
                                               for entry in evidence['entries']]
                if prescription_choice:
                    field['sub_fields']['has_change'] = evidence['value']
                else:
                    field['value'] = evidence['value']
                if not evidence['value']:
                    field['semantic_issues'] = [{'reason': '选项同时勾选或归属冲突，待核对' if evidence['status'] in ('conflict', 'ambiguous') else '选项未确认，待核对',
                                                  'source_text': field['source_text'], 'source_regions': field.get('source_regions', [])}]
                else:
                    field.pop('semantic_issues', None)
                if other_prescription_issues:
                    field['semantic_issues'] = other_prescription_issues + field.get('semantic_issues', [])
            if key == "item_16_prescription":
                field["sub_fields"]["has_change"] = selected_option(field["sub_fields"].get("has_change", ""), ["是", "否"])
            if key in ("item_9_chemical_name", "item_12_specification"):
                field["value"] = re.sub(r"\s*\n\s*", " ", field.get("value", ""))
            if key in ("item_6_generic_name", "item_9_chemical_name"):
                field['value'] = self._clean_text(field.get('value', ''))
            if key in ('item_24_patent_info', 'item_28_change_related_items'):
                text = field.get('value', '')
                if re.search(r'尚不能确认|无法确认|不能确定|不确定|待确认|待核对|去利|起利|亿权|[�□?？]', self._compact_text(text)):
                    field['semantic_issues'] = [{'reason': '声明含不确定表述或疑似识别模糊，保留原文待核对',
                                                'source_text': field.get('source_text', text),
                                                'source_regions': field.get('source_regions', [])}]
            nonempty = bool(field.get('reported_value')) or any(present(v) for p, v in values(field).items() if p != "sub_fields.validity_unit")
            if '【控件状态与显示矛盾】' in field.get('source_text',''):
                field.setdefault('semantic_issues', []).append({'reason':'Word控件状态与显示字形矛盾，待核对',
                    'source_text':field['source_text'], 'source_regions':field.get('source_regions',[])})
                field['selection_status'] = 'conflict'
                field['value'] = ''
                field['selected_values'] = []
                if prescription_choice:
                    field['sub_fields']['has_change'] = ''
                    field.get('value_sources',{}).pop('sub_fields.has_change',None)
            if not nonempty:
                field["parse_confidence"] = 0.0
            field["recognition_status"] = "extracted" if nonempty else "unrecognized"

    def normalize_form_json(self, form_json: Dict[str, Any]) -> Dict[str, Any]:
        base = self._build_full_form("")
        if not isinstance(form_json, dict):
            return base
        return self._deep_merge(base, form_json)

    def _deep_merge(self, base: Dict[str, Any], custom: Dict[str, Any]) -> Dict[str, Any]:
        out = dict(base)
        for key, val in (custom or {}).items():
            if key in out and isinstance(out[key], dict) and isinstance(val, dict):
                out[key] = self._deep_merge(out[key], val)
            else:
                out[key] = val
        return out

    def _field(
        self,
        field_name: str,
        field_type: str,
        required: bool = False,
        options: list | None = None,
        sub_fields: Dict[str, Any] | None = None,
    ) -> Dict[str, Any]:
        return {
            "field_name": field_name,
            "field_type": field_type,
            "required": bool(required),
            "options": options or [],
            "value": "",
            "selected_values": [],
            "sub_fields": sub_fields or {},
            "source_text": "",
            "parse_confidence": 0.0,
            "manual_modified": False,
        }

    def _build_full_form(self, raw_text: str) -> Dict[str, Any]:
        text = re.sub(r"\s+", "", str(raw_text or ""))
        category_hit = ""
        schema = {
            "task_type": "extend_validity_period",
            "original_form_type": "备案表" if "备案表" in text else ("补充申请表" if "补充申请表" in text else ""),
            "special_statement": self._field("其他特别申明事项", "textarea"),
            "item_1_application_type": self._field(
                "本申请属于",
                "radio",
                True,
                ["境内生产药品补充申请"],
                {"shared_application_form_no": ""},
            ),
            "item_2_registration_category": self._field(
                "药品注册分类",
                "radio",
                True,
                ["中药", "预防用生物制品", "化学药品", "治疗用生物制品", "按生物制品管制的诊断试剂"],
                {"class_no": ""},
            ),
            "item_3_otc_type": self._field("是否为OTC", "radio", True, ["处方药", "非处方药"]),
            "item_4_original_product_status": self._field(
                "原申请品种状态", "radio", True, ["已上市", "未上市", "已批准临床", "在审评或审批中"]
            ),
            "item_5_application_matter_category": self._field(
                "申请事项分类",
                "tree_checkbox",
                True,
                [
                    {"code": "1.1", "label": "变更处方中的辅料"},
                    {"code": "1.2", "label": "变更生产工艺"},
                    {"code": "1.3", "label": "变更所用原料药的供应商"},
                    {"code": "1.4", "label": "变更生产批量"},
                    {"code": "1.5", "label": "变更注册标准"},
                    {"code": "1.6", "label": "变更包装材料和容器"},
                    {"code": "1.7", "label": "变更有效期和贮藏条件"},
                    {"code": "1.8", "label": "增加规格"},
                    {"code": "1.9", "label": "变更生产场地，境外生产药品"},
                    {"code": "1.10", "label": "其他"},
                ],
                {"other_description": ""},
            ),
            "item_6_generic_name": self._field("药品通用名称", "input", True),
            "item_7_english_or_latin_name": self._field("英文名称/拉丁名称", "input"),
            "item_8_pinyin": self._field("汉语拼音", "input"),
            "item_9_chemical_name": self._field("化学名称", "textarea"),
            "item_10_trade_name": self._field("商品名称", "radio_with_input", False, ["不使用", "使用"], {"trade_name": ""}),
            "item_11_dosage_form": self._field("剂型", "group_checkbox", True, []),
            "item_12_specification": self._field("规格", "textarea", True),
            "item_13_other_specifications": self._field("同品种其他规格", "textarea"),
            "item_14_packaging": self._field(
                "包装",
                "group",
                True,
                [],
                {"primary_packaging_material": "", "packaging_specification": ""},
            ),
            "item_15_validity_period": self._field(
                "药品有效期",
                "group",
                True,
                [],
                {
                    "original_validity_period": "",
                    "proposed_validity_period": "",
                    "validity_unit": "月",
                    "storage_condition": "",
                },
            ),
            "item_16_prescription": self._field(
                "处方",
                "group",
                False,
                [],
                {"active_ingredients": "", "excipients": "", "has_change": ""},
            ),
            "item_17_material_source": self._field(
                "原/辅料/包材来源",
                "table",
                False,
                [],
                {"columns": ["material_name", "register_no", "accept_no", "manufacturer"]},
            ),
            "item_18_tcm_material_standard": self._field("中药材标准", "textarea"),
            "item_19_pre_acceptance_inspection": self._field(
                "受理前药品注册检验", "radio_with_input", False, ["是", "否"], {"sample_inspection_no": ""}
            ),
            "item_20_indication_or_function": self._field(
                "主要适应症或者功能主治", "group", False, [], {"category": "", "description": ""}
            ),
            "item_21_change_content": self._field("补充申请的内容", "textarea", True),
            "item_22_change_reason": self._field("提出补充申请理由", "textarea", True),
            "item_23_original_approval_info": self._field("原批准注册内容及相关信息", "textarea", True),
            "item_24_patent_info": self._field("专利情况", "textarea"),
            "item_25_data_protection": self._field("数据保护", "textarea"),
            "item_26_tcm_protection": self._field("中药品种保护", "textarea"),
            "item_27_same_product_period": self._field("同品种新药检测期", "textarea"),
            "item_28_change_related_items": self._field("本次申请涉及事项", "textarea"),
            "item_29_other_related_info": self._field("其他相关情况", "textarea"),
            "item_30_applicant_info": self._field(
                "申请人信息",
                "group",
                True,
                [],
                {"applicant_name": "", "license_no": "", "credit_code":"", "organization_code":"", "credit_or_org_code":"", "contact": "", "address": "", "legal_representative":"", "registered_address":"", "mailing_address":"", "phone":"", "mobile":"", "registration_contact":""},
            ),
            "item_31_manufacturer_info": self._field(
                "生产企业信息",
                "group",
                True,
                [],
                {"manufacturer_name": "", "license_no": "", "credit_code":"", "organization_code":"", "credit_or_org_code":"", "contact": "", "address": "", "legal_representative":"", "registered_address":"", "production_address":"", "mailing_address":"", "phone":"", "mobile":"", "registration_contact":""},
            ),
            "item_32_cro_info": self._field(
                "委托研究机构信息",
                "group",
                False,
                [],
                {"organization_name": "", "contact": "", "address": "", "research_lead":"", "registered_address":"", "mailing_address":"", "phone":"", "mobile":"", "registration_contact":""},
            ),
        }
        if category_hit:
            schema["item_5_application_matter_category"]["selected_values"] = [category_hit]
            schema["item_5_application_matter_category"]["parse_confidence"] = 0.8
        return schema
