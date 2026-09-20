from datetime import datetime
from zoneinfo import ZoneInfo
import re
from io import BytesIO
from pathlib import Path
from typing import Any, Dict

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None


class FilingChangeReportService:
    def __init__(self, root_dir: Path) -> None:
        self.root_dir = Path(root_dir)

    @staticmethod
    def _sanitize_xml_text(value: Any) -> str:
        text = "" if value is None else str(value)
        cleaned_chars = []
        for ch in text:
            code = ord(ch)
            if ch in "\t\n\r" or code >= 32:
                cleaned_chars.append(ch)
        return "".join(cleaned_chars).replace("\x00", "")

    def generate_report(self, project_id: str, run_id: str, result_payload: Dict[str, Any]) -> Dict[str, Any]:
        report_dir = self.root_dir / "projects" / project_id / "reports"
        report_dir.mkdir(parents=True, exist_ok=True)
        now_str = result_payload.get("report_generated_at") or datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M:%S")
        markdown = self._build_markdown(project_id, run_id, result_payload, now_str)
        markdown = self._sanitize_xml_text(markdown)
        md_path = report_dir / f"{run_id}.md"
        md_path.write_text(markdown, encoding="utf-8")
        word_path = report_dir / f"{run_id}.docx"
        if Document is not None:
            doc = Document()
            from docx.shared import Pt, Inches, RGBColor
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Bullet 2'):
                style = doc.styles[style_name]; style.font.name = 'Arial Unicode MS'; style.font.color.rgb = RGBColor(0, 0, 0)
                style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
            doc.styles['Normal'].font.size = Pt(9)
            doc.sections[0].left_margin = Inches(.65); doc.sections[0].right_margin = Inches(.65)
            self._write_markdown_to_docx(doc, markdown)
            for table in doc.tables:
                if table.rows:
                    header = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(header)
            self._append_charts(doc, result_payload)
            doc.save(str(word_path))
        else:
            raise RuntimeError('Word导出依赖python-docx不可用，不能以文本文件冒充Word')
        return {"markdown": markdown, "word_path": str(word_path), "markdown_path": str(md_path), "generated_at": now_str}

    def _append_charts(self, doc, payload):
        from docx.shared import Inches
        # 从同轮SVG生成可直接显示的图像；原始数值与来源仍保留在完整表格中。
        for index, chart in enumerate((payload.get('stability_trend_analysis') or {}).get('charts') or []):
            svg = chart.get('svg_content')
            if not svg: continue
            doc.add_heading('趋势图 ' + str(chart.get('indicator') or index + 1), level=2)
            import fitz
            with fitz.open(stream=svg.encode('utf-8'), filetype='svg') as chart_document:
                png = chart_document[0].get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes('png')
            doc.add_picture(BytesIO(png), width=Inches(6))
            doc.add_paragraph('各点的数值、批次、条件与来源见本轮完整稳定性数据表。')

    def _write_markdown_to_docx(self, doc: Any, markdown_text: str) -> None:
        """将 Markdown 粗粒度映射为 Word 标题/正文/列表，确保导出报告可直接编辑。"""
        text = self._sanitize_xml_text(markdown_text or "")
        table = None
        for raw in text.splitlines():
            line = str(raw or "").rstrip()
            if line.startswith('|') and line.endswith('|'):
                cells = [x.strip().replace('<br>', '\n').replace('&#124;', '|') for x in line.strip('|').split('|')]
                if all(re.fullmatch(r':?-+:?', x) for x in cells): continue
                if table is None or len(table.columns) != len(cells):
                    table = doc.add_table(rows=0, cols=len(cells)); table.style = 'Table Grid'
                row = table.add_row()
                for cell, value in zip(row.cells, cells): cell.text = value
                continue
            table = None
            if not line:
                doc.add_paragraph("")
                continue
            if line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
                continue
            if line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
                continue
            if line.startswith("- "):
                doc.add_paragraph(line[2:].strip(), style="List Bullet")
                continue
            if line.startswith("  - "):
                doc.add_paragraph(line[4:].strip(), style="List Bullet 2")
                continue
            doc.add_paragraph(line)

    def _build_markdown(self, project_id: str, run_id: str, result_payload: Dict[str, Any], now_str: str) -> str:
        def value(v):
            return self._sanitize_xml_text(v if v is not None else '历史结果未记录').replace('|', '&#124;').replace('\n', '<br>')
        p = result_payload
        lines = ['# 药品备案变更审评报告', '', f'- 审评轮次：{run_id}',
                 f"- 审评完成时间（北京时间）：{p.get('review_completed_at') or '历史结果未记录'}",
                 f'- 报告生成时间（北京时间）：{now_str}', '', '## AI总体结论']
        overall = p.get('overall_conclusion') or {}
        if any(isinstance(call, dict) and call.get('scene') == 'overall_summary'
               and call.get('status') == 'fallback' for call in p.get('llm_calls') or []):
            lines += ['解释排序未完成，使用原始事实顺序；不改变已完成的规则检查结论。']
        lines += [str(overall.get('result') or '历史结果未记录'), str(overall.get('summary') or '历史结果未记录'),
                  '## 变更事项类别', str((p.get('change_category_suggestion') or {}).get('suggested_category') or '历史结果未记录'),
                  str((p.get('change_category_suggestion') or {}).get('reason') or '历史结果未记录'),
                  '## 技术支持情况', str((p.get('technical_support') or {}).get('status') or '历史结果未记录'),
                  '## 建议处理方式', str(p.get('recommended_action') or '历史结果未记录'),
                  '## 人工复核意见（本轮独立保存）']
        manual = p.get('manual_confirmation') or {}
        lines += [str(manual.get('comment') or '本轮尚未记录人工意见'),
                  '复核人：' + str(manual.get('reviewer') or '未记录'), '复核时间：' + str(manual.get('confirmed_at') or '未记录'),
                  '## 问题数量与处理明细', '口径：按本轮检查发现计数；同一材料不同字段、批次、条件及规则发现分别保留。未知不计为零。']
        for key, label in [('file_missing', '文件缺失'), ('content_missing', '字段或内容缺项'), ('parse_error', '解析异常'), ('conflict', '信息冲突')]:
            lines.append(label + '：' + value((p.get('issue_counts') or {}).get(key)))
        for issue in p.get('issues') or []:
            lines += ['### ' + str(issue.get('module', '检查')), str(issue.get('status')) + '：' + str(issue.get('reason')),
                      '处理要求：' + str(issue.get('required_action'))]
        for key, label in [('formal_review', '形式审查'), ('consistency_check', '一致性'), ('quality_standard_check', '质量'), ('stability_trend_analysis', '稳定性')]:
            data = p.get(key) or {}
            lines += ['## ' + label, str(data.get('result') or data.get('summary') or ('见逐项检查' if data else '历史结果未记录'))]
            for field in ('summary', 'missing_materials', 'content_missing', 'parse_issues', 'not_applicable'):
                if field in data:
                    lines += self._fmt_value(data[field])
            if key == 'quality_standard_check':
                bridge = data.get('standard_bridge') or {}
                if bridge: lines += self._fmt_value(bridge)
            if key == 'stability_trend_analysis':
                lines += ['方案覆盖：' + str((data.get('coverage') or {}).get('reason') or '未记录'),
                          '显著变化：' + str((data.get('significant_change_assessment') or {}).get('status') or '未记录')]
                limit_check = data.get('limit_check') or {}
                lines += ['限度判断：' + str(limit_check.get('status') or '历史结果未记录'),
                          str(limit_check.get('reminder') or '')]
                for kind, title in [('out_of_spec', '超限明细'), ('undecidable', '待确认明细')]:
                    details = limit_check.get(kind + '_details') or []
                    count = limit_check.get(kind + '_count', len(details))
                    lines.append('### ' + title + '（记录总数：' + str(count) + '，已保存明细：' + str(len(details)) + '）')
                    if isinstance(count, (int, float)) and count > len(details):
                        lines.append('历史明细不完整：仅展示已保存记录，缺失明细未自动重算。')
                    for detail in details:
                        lines.append('- ' + '；'.join(str(detail.get(field) or '未记录') for field in
                            ('indicator', 'context', 'result_text', 'limit_text', 'reason_text', 'source_file_name')))
                        lines.append('  结果单位：' + str(detail.get('result_unit') or '未标注') + '；标准单位：' + str(detail.get('limit_unit') or '未标注'))
                        if detail.get('source_position'):
                            lines += self._fmt_value({'来源位置': detail['source_position']})
                        if detail.get('numeric_verification'):
                            lines += self._fmt_value({'OCR数值候选与核验位置': detail['numeric_verification']})
                        for component in detail.get('components') or []:
                            state = {True: '未超限', False: '已超限', None: '待确认'}.get(component.get('within_standard'), '待确认')
                            lines.append('  - 组成项 ' + str(component.get('item') or '未识别') + '：' + state +
                                         '；结果 ' + str(component.get('result') or '缺失') + '；标准 ' + str(component.get('limit') or '缺失') +
                                         '；' + str(component.get('reason_text') or component.get('reason') or '原因未记录'))
        lines += ['## 完整规则列表']
        if 'rule_results' not in p: lines.append('历史结果未记录')
        for rule in p.get('rule_results') or []:
            lines += ['### ' + str(rule.get('rule_code')) + ' ' + str(rule.get('rule_name', '')),
                      '执行：' + str(rule.get('execution_status') or '未记录') + '；业务状态：' + str(rule.get('status') or '无业务结论'),
                      '依据：' + str(rule.get('basis_source') or '未记录'), '证据编号：' + '、'.join(rule.get('evidence_ids') or [])]
            for sub in rule.get('subchecks') or []:
                lines.append(str(sub.get('name')) + '：' + str(sub.get('status')) + '；' + str(sub.get('reason')))
        for key, label in [('records', '自制稳定性数据'), ('reference_records', '参比稳定性数据')]:
            records = (p.get('stability_trend_analysis') or {}).get(key)
            lines += ['## ' + label]
            if records is None: lines.append('历史结果未记录'); continue
            lines += ['| 批次及条件 | 项目及时间 | 原值及标准 | 判断 | 来源 |', '| --- | --- | --- | --- | --- |']
            unit_sources = {'result': '检测结果原文', 'standard': '标准原文', 'indicator': '指标',
                            'record.unit': '关联单位字段', 'result_column_header': '结果列表头',
                            'limit_column_header': '标准列表头'}
            for r in records:
                cells = [' / '.join(str(r.get(k) or '') for k in ('batch_no', 'condition', 'orientation', 'specification', 'packaging')),
                         str(r.get('indicator', '')) + ' / ' + str(r.get('time_point', r.get('month', ''))),
                         str(r.get('result_text', '')) + ' / ' + str(r.get('limit_text', '')),
                         {True: '符合限度', False: '不符合限度', None: '待确认'}.get(r.get('within_standard'), '待确认')
                         + ('；' + str(r['unit_note']) if r.get('unit_note') else '')
                         + ('；单位来源：' + unit_sources.get(r.get('result_unit_source'), '未标注') + '/' + unit_sources.get(r.get('limit_unit_source'), '未标注')
                            if r.get('result_unit_source') or r.get('limit_unit_source') else ''),
                         self._record_position(r)]
                lines.append('| ' + ' | '.join(value(c) for c in cells) + ' |')
            manual_records = [(r, [v for v in r.get('numeric_verification', [])
                                  if isinstance(v, dict) and v.get('status') == 'manual_confirmed'])
                              for r in records]
            manual_records = [(r, checks) for r, checks in manual_records if checks]
            if manual_records:
                lines += ['### 人工数值确认（' + label + '）',
                          '本轮采用人工确认值；不表示原OCR识别正确，不改写历史轮次。']
                for record, checks in manual_records:
                    lines.append(str(record.get('indicator') or '未记录指标') + '；' + self._record_position(record))
                    for check in checks:
                        lines += self._fmt_value({
                            '确认状态': 'manual_confirmed',
                            '原OCR候选': check.get('primary', ''),
                            '局部复核候选': check.get('secondary', ''),
                            '确认值': check.get('manual_value', ''),
                            '核对依据': check.get('manual_reason', ''),
                            '原数值状态': check.get('ocr_status', '历史未记录'),
                            '原页位置': {'page': check.get('page'), 'bbox_pdf': check.get('bbox_pdf', [])},
                        })
        lines += ['## 本轮证据内容']
        if (p.get('reference_usage') or {}).get('status') == 'not_used':
            lines.append('本轮未自动采用参考资料正文；参考库启用不代表本轮已引用。')
        lines += ['| 证据及材料 | 位置与使用状态 | 原文 |', '| --- | --- | --- |']
        for ref in p.get('evidence_refs') or []:
            cells = [str(ref.get('evidence_id', '')) + ' ' + str(ref.get('title') or '来源未记录'),
                     str(ref.get('position') or '位置未记录') + ' / ' + str(ref.get('availability') or '状态未记录'),
                     str(ref.get('snippet') or '原文未记录')]
            lines.append('| ' + ' | '.join(value(c) for c in cells) + ' |')
        return '\n\n'.join(lines).replace(' |\n\n| ', ' |\n| ')

    @staticmethod
    def _record_position(record):
        source = record.get('source_position') or {}
        name = source.get('file_name') or record.get('source_file_name') or record.get('source_doc_id') or '来源未记录'
        parts = [str(name)]
        for key, label in [('page', '页'), ('table', '表'), ('row', '行'), ('cell', '列')]:
            if key == 'page' and str(name).lower().endswith(('.doc', '.docx')):
                continue
            value = source.get(key)
            if value is not None and value != '': parts.append(label + ' ' + str(value))
        if len(parts) == 1: parts.append('具体位置未记录')
        return ' / '.join(parts)

    @staticmethod
    def _looks_like_raw_dict_markdown(markdown_text: str) -> bool:
        text = str(markdown_text or "")
        if "{'" in text or "'}" in text:
            return True
        if "## 一、总体审评建议" in text and "{'result':" in text:
            return True
        return False

    @staticmethod
    def _fmt_scalar(value: Any) -> str:
        return ("" if value is None else str(value)).replace("\r\n", "\n").replace("\r", "\n").strip()

    @classmethod
    def _fmt_value(cls, value: Any, indent: int = 0) -> list:
        prefix = "  " * indent
        lines = []
        if isinstance(value, dict):
            for key, val in value.items():
                if isinstance(val, (dict, list)):
                    lines.append(f"{prefix}- {key}:")
                    lines.extend(cls._fmt_value(val, indent + 1))
                else:
                    lines.append(f"{prefix}- {key}: {cls._fmt_scalar(val)}")
            return lines
        if isinstance(value, list):
            if not value:
                return [f"{prefix}- （无）"]
            for item in value:
                if isinstance(item, (dict, list)):
                    lines.append(f"{prefix}-")
                    lines.extend(cls._fmt_value(item, indent + 1))
                else:
                    lines.append(f"{prefix}- {cls._fmt_scalar(item)}")
            return lines
        return [f"{prefix}{cls._fmt_scalar(value)}"]

    def _build_markdown_from_structured(
        self, project_id: str, run_id: str, result_payload: Dict[str, Any], now_str: str
    ) -> str:
        lines = [
            "# 药品备案变更类审评报告草稿",
            "",
            f"- 生成时间：{now_str}",
            "",
        ]
        report_draft = result_payload.get("review_report_draft", {}) if isinstance(result_payload, dict) else {}
        sections = report_draft.get("sections", []) if isinstance(report_draft, dict) else []
        has_overall = any(
            isinstance(sec, dict) and "总体审评建议" in str(sec.get("name", "")) for sec in (sections or [])
        )
        if not has_overall:
            lines.append("## 一、总体审评建议")
            overall = result_payload.get("overall_conclusion", {}) if isinstance(result_payload, dict) else {}
            lines.extend(self._fmt_value(overall))
            lines.append("")

        if isinstance(sections, list) and sections:
            for sec in sections:
                if not isinstance(sec, dict):
                    continue
                lines.append(f"## {sec.get('name', '未命名章节')}")
                lines.extend(self._fmt_value(sec.get("conclusion", "")))
                evs = sec.get("evidence_refs", [])
                if isinstance(evs, list) and evs:
                    lines.append("### 证据引用")
                    for e in evs[:8]:
                        if not isinstance(e, dict):
                            continue
                        title = e.get("doc_name", "") or e.get("title", "")
                        lines.append(f"- 文档：{title}（doc_id={e.get('doc_id','')}，位置={e.get('position','')}）")
                        snippet = str(e.get("snippet", "") or "").replace("\n", " ").strip()
                        if snippet:
                            lines.append(f"  - 内容片段：{snippet[:220]}")
                        rule_code = e.get("rule_code", "")
                        rule_name = e.get("rule_name", "")
                        if rule_code or rule_name:
                            lines.append(f"  - 引用规则：{rule_code} {rule_name}".strip())
                lines.append("")
        else:
            lines.append("## 二、资料形式审查")
            lines.extend(self._fmt_value(result_payload.get("formal_review", {})))
            lines.append("")
            lines.append("## 三、变更管理类别建议")
            lines.extend(self._fmt_value(result_payload.get("change_category_suggestion", {})))
            lines.append("")
            lines.append("## 四、质量标准比对")
            lines.extend(self._fmt_value(result_payload.get("quality_standard_check", {})))
            lines.append("")
            lines.append("## 五、稳定性趋势分析")
            lines.extend(self._fmt_value(result_payload.get("stability_trend_analysis", {})))
            lines.append("")

        lines.append("## 九、补正通知书草稿")
        notice = result_payload.get("correction_notice_draft", {}) if isinstance(result_payload, dict) else {}
        if isinstance(notice, dict):
            for idx, item in enumerate((notice.get("items", []) or []), start=1):
                if not isinstance(item, dict):
                    continue
                lines.append(f"- {idx}. 问题：{self._fmt_scalar(item.get('issue', ''))}")
                lines.append(f"  - 依据：{self._fmt_scalar(item.get('basis', ''))}")
                lines.append(f"  - 补充要求：{self._fmt_scalar(item.get('required_supplement', ''))}")
        else:
            lines.extend(self._fmt_value(notice))
        lines.append("")
        lines.append("## 十、AI初步结论建议")
        conclusion = result_payload.get("conclusion", {}) if isinstance(result_payload, dict) else {}
        if isinstance(conclusion, dict):
            lines.append(f"- AI初步建议：{self._fmt_scalar(conclusion.get('ai_judgement', conclusion.get('result', '')))}")
            lines.append(f"- 处理建议：{self._fmt_scalar(conclusion.get('suggestion', 'AI结果仅供审评员参考。'))}")
        else:
            lines.extend(self._fmt_value(conclusion))
        lines.append("")
        return "\n".join(lines)
